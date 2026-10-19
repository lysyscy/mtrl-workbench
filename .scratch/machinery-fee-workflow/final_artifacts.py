"""Recheck only the owned local UI fixture and final native report copies."""
from pathlib import Path
import hashlib
import json
import shutil
import subprocess
from urllib.parse import urlencode, unquote
from urllib.request import Request, urlopen

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / '.artifacts/machinery-fee-workflow/implementation-20261019'
BASE = 'http://127.0.0.1:8777'


def fetch(path, payload=None):
    data = None if payload is None else json.dumps(payload).encode('utf-8')
    request = Request(BASE + path, data=data, headers={'Content-Type': 'application/json'})
    with urlopen(request, timeout=60) as response:
        assert response.headers.get('X-Isolated-Fixture') == 'true'
        return response.read(), dict(response.headers)


state = json.loads(fetch('/api/m06/workbench?ym=2026-09')[0])
assert sum(len(item['rows']) for item in state['details']) == 1966
assert any('隔离验证：仅保存核查备注' in item.get('text', '') for item in state['matters'])
preview = json.loads(fetch('/api/m06/workbench/i8-fetch', {'request_id': 'final-offline-logging-check'})[0])
assert preview['counts'] == {'usage': 307, 'project': 69, 'team': 129}
assert preview['run_record_id'] > 0
results = []
for ym in ('2027-01', '2026-09'):
    query = urlencode({'ym': ym, 'projects': json.dumps(state['settings']['shown_projects'], ensure_ascii=False)})
    report = json.loads(fetch('/api/m06/workbench/preview?' + query)[0])
    content, headers = fetch('/api/m06/workbench/export?' + query + '&snapshot_token=' + report['snapshot_token'])
    filename = unquote(next(v for k, v in headers.items() if k.lower() == 'x-export-filename'))
    assert Path(filename).name == filename
    path = OUT / 'exports' / filename
    assert hashlib.sha256(content).digest() == hashlib.sha256(path.read_bytes()).digest()
    book = load_workbook(path)
    assert book.sheetnames == ['工程公司汇总表', '零租设备月报表', '包月外租设备月报', '年度汇总']
    annual = book.worksheets[3]
    assert annual['D1'].value == f'{int(ym[:4]) - 1}年1月'
    assert annual['P1'].value == f'{ym[:4]}年1月'
    assert all(not sheet.row_dimensions[row].hidden for sheet in book for row in range(2, sheet.max_row + 1))
    book.close()
    (OUT / 'export-verification-input.json').write_text(json.dumps({'path': str(path), 'ym': ym}, ensure_ascii=False), encoding='utf-8')
    subprocess.run([str(ROOT / 'backend/.venv/Scripts/python.exe'), '-X', 'utf8',
                    str(Path(__file__).with_name('verify_excel.py'))], cwd=ROOT, check=True)
    native = json.loads((OUT / 'native-spreadsheet-verification.json').read_text(encoding='utf-8'))
    assert native['opened'] and native['original_export_unchanged']
    shutil.copyfile(OUT / 'mechanical-print-preview.pdf', OUT / f'mechanical-print-preview-{ym}.pdf')
    shutil.copyfile(OUT / 'native-spreadsheet-verification.json', OUT / f'native-spreadsheet-verification-{ym}.json')
    results.append({'ym': ym, 'path': str(path), 'bytes': len(content), 'warnings': report['warnings'],
                    'native_open': True, 'hash_unchanged': True})
(OUT / 'final-artifacts.json').write_text(json.dumps({'scope': 'isolated-only', 'details': 1966,
    'offline_fetch_run_record_id': preview['run_record_id'], 'exports': results}, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'details': 1966, 'exports': len(results), 'offline_fetch_logged': True}, ensure_ascii=False))
