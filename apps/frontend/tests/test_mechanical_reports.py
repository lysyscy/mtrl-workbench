"""Real four-sheet workbook, independent export choices and company summary."""
import json
import re
from decimal import Decimal
from urllib.parse import parse_qs,urlparse
import openpyxl
from playwright.sync_api import expect
from test_preview import server,browser_page
from live_data import get,open_page,grid_row


def test_actual_four_sheet_download_and_independent_project_preferences(server,browser_page,api,projects,disposable,tmp_path):
    page,_=browser_page;state=get(api,'/api/m06/workbench');assert state['projects']
    eligible=[p for p in projects if p['project_type']=='normal' and p['status']=='active' and p['warehouse_name'] in [v['warehouse'] for v in state['projects']]]
    assert len(eligible)>=2,'真实副本缺少两个报表项目，导出顺序未覆盖'
    project,other=eligible[:2]
    open_page(page,server,'mechanical','机械费')
    page.get_by_role('button',name='显示项目',exact=True).click();dialog=page.get_by_role('dialog',name='选择显示项目及顺序',exact=True)
    dialog.get_by_role('button',name='全不选',exact=True).click();dialog.get_by_role('checkbox',name=project['disp'],exact=True).check();dialog.get_by_role('button',name='应用项目选择',exact=True).click()
    expect(page.get_by_role('region',name='公司汇总',exact=True)).to_be_visible()
    page.get_by_role('button',name='导出台账…',exact=True).click();dialog=page.get_by_role('dialog',name='导出台账',exact=True)
    expect(dialog.get_by_role('checkbox',name=project['disp'],exact=True)).to_be_checked()
    dialog.get_by_role('button',name='全不选',exact=True).click()
    for item in [project,other]:dialog.get_by_role('checkbox',name=item['disp'],exact=True).check()
    # The shared list moves every candidate; put the other chosen project first.
    up=dialog.get_by_role('button',name='上移'+other['disp'],exact=True)
    while up.is_enabled():up.click()
    ym=page.get_by_label('月份',exact=True).input_value()
    warehouses=[other['warehouse_name'],project['warehouse_name']]
    report=get(api,'/api/m06/workbench/preview',ym=ym,projects=json.dumps(warehouses,ensure_ascii=False))
    download=dialog.get_by_role('button',name='保存到本地',exact=True);expect(download).to_be_enabled(timeout=30000)
    requests=[];page.on('request',lambda request:requests.append(request.url) if '/workbench/export?' in request.url else None)
    with page.expect_download() as event:download.click()
    path=tmp_path/'mechanical.xlsx';event.value.save_as(path)
    params=parse_qs(urlparse(requests[-1]).query);assert json.loads(params['projects'][0])==warehouses
    workbook=openpyxl.load_workbook(path,data_only=True);assert workbook.sheetnames==[sheet['name'] for sheet in report['sheets']];assert len(workbook.sheetnames)==4
    assert all(sheet.print_area and sheet.print_title_rows for sheet in workbook)
    sheet=report['sheets'][0]
    checked=0
    for item in [other,project]:
        position=next(row['row'] for row in sheet['project_rows'] if row['warehouse']==item['warehouse_name'])
        assert workbook.worksheets[0].cell(position,2).value==item['disp']
        for cell in sheet['rows'][position-1]:
            if cell['column']>=4 and re.fullmatch(r'-?\d+(?:\.\d+)?',cell['value']):
                assert Decimal(str(workbook.worksheets[0].cell(position,cell['column']).value))==Decimal(cell['value']);checked+=1
    assert checked,'真实报送项目没有金额格，API到文件金额未覆盖'
    workbook.close()
    # Export choices never change the main display selection, and survive reload.
    assert get(api,'/api/m06/workbench',ym=ym)['settings']['shown_projects']==[project['warehouse_name']]
    page.reload();page.get_by_role('button',name='导出台账…',exact=True).click()
    expect(dialog.locator('.export-project-row').first.get_by_role('checkbox',name=other['disp'],exact=True)).to_be_checked()
    for item in [other,project]:expect(dialog.get_by_role('checkbox',name=item['disp'],exact=True)).to_be_checked()
    dialog.get_by_role('button',name='全不选',exact=True).click();dialog.get_by_role('button',name='关闭',exact=True).click()
    page.reload();page.get_by_role('button',name='导出台账…',exact=True).click()
    expect(dialog.get_by_role('checkbox',checked=True)).to_have_count(0)
    expect(dialog.get_by_role('button',name='保存到本地',exact=True)).to_be_disabled()
    dialog.get_by_role('button',name='关闭',exact=True).click()


def test_actual_company_summary_layout_and_amounts_readonly(server,browser_page,api,projects):
    page,_=browser_page;page.set_viewport_size({'width':2400,'height':1200})
    open_page(page,server,'mechanical','机械费');grid=page.get_by_role('region',name='公司汇总',exact=True)
    expect(grid).to_be_visible();ym=page.get_by_label('月份',exact=True).input_value()
    state=get(api,'/api/m06/workbench',ym=ym);warehouses=state['settings']['shown_projects'];assert warehouses
    preview=get(api,'/api/m06/workbench/preview',ym=ym,projects=json.dumps(warehouses,ensure_ascii=False))
    count=len(preview['sheets'][0]['project_rows']);assert count
    page.wait_for_function("""count=>{const g=document.querySelector('.mechanical-grid');const last=g?.querySelector('[row-index="'+(count-1)+'"]');const total=g?.querySelector('.ag-row-pinned');return last&&total&&Math.abs(total.getBoundingClientRect().top-last.getBoundingClientRect().bottom)<2;}""",arg=count)
    scale=page.evaluate("Number(getComputedStyle(document.documentElement).getPropertyValue('--fscale'))")
    assert grid.locator('.ag-header-cell[col-id="c1"]').bounding_box()['width']<=200*scale+1
    assert grid.locator('.ag-header-cell[col-id="c14"]').bounding_box()['width']<=144*scale+1
    project=next(p for p in projects if p['warehouse_name']==warehouses[0])
    position=next(item['row'] for item in preview['sheets'][0]['project_rows'] if item['warehouse']==warehouses[0])
    row=grid_row(page,'.mechanical-grid',project['id'])
    for dom,cells in [(row,preview['sheets'][0]['rows'][position-1]),(grid.locator('.ag-row-pinned'),preview['sheets'][0]['rows'][-1])]:
        values={cell['column']-1:cell['value'] for cell in cells}
        for column in [4,5,7,8,9,10,12,13]:
            expected=values[column].replace(',','');cell=dom.locator('[col-id="c'+str(column)+'"]')
            actual=(cell.locator('input').input_value() if cell.locator('input').count() else cell.inner_text()).replace(',','')
            if re.fullmatch(r'-?\d+(?:\.\d+)?',expected):assert Decimal(actual)==Decimal(expected)
            else:assert actual=='—'
    page.set_viewport_size({'width':2400,'height':500})
    page.wait_for_function("document.querySelector('.mechanical-grid').getBoundingClientRect().height<340")
    header=grid.locator('.ag-header').bounding_box();grid.locator('[row-index="'+str(count-1)+'"]').scroll_into_view_if_needed()
    assert abs(grid.locator('.ag-header').bounding_box()['y']-header['y'])<1
    # Same-scope reads retain the real internal editor, including its unsaved draft.
    field=grid.locator('[col-id=c12] input').first
    field.scroll_into_view_if_needed();field.fill('321.09');field.focus()
    label=field.get_attribute('aria-label');held=[]
    page.route('**/api/m06/workbench/preview?*',lambda route:held.append(route))
    with page.expect_request('**/api/m06/workbench/preview?*'):
        page.evaluate("window.dispatchEvent(new Event('offline'));window.dispatchEvent(new Event('online'))")
    try:
        expect(grid).to_be_visible();expect(field).to_be_focused();expect(field).to_have_value('321.09')
        assert held
    finally:
        for route in held:route.continue_()
        page.unroute('**/api/m06/workbench/preview?*')
    page.wait_for_load_state('networkidle')
    expect(page.get_by_label(label,exact=True)).to_have_value('321.09')
    expect(page.get_by_label(label,exact=True)).to_be_focused()
