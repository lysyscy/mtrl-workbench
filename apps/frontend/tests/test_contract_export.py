"""Actual contract filter/sort IDs and downloaded workbook values."""
from decimal import Decimal
import re
import openpyxl
import pytest
from playwright.sync_api import expect
from test_preview import server,browser_page
from live_data import get,open_page
from preview_select import pick_option


def test_real_contract_order_export_and_explicit_empty_range(server,browser_page,api,tmp_path):
    page,_=browser_page;rows=get(api,'/api/m11/list')['rows'];assert rows
    open_page(page,server,'contracts','集采台账')
    pick_option(page,'纳入状态','纳入台账')
    page.get_by_role('columnheader',name=re.compile('^合同金额$')).click()
    expected=[r['id'] for r in sorted((r for r in rows if r['tag']=='纳入台账'),key=lambda r:Decimal(str(r['contract_amt'] or 0)))]
    assert expected,'真实副本没有纳入合同，导出代表未覆盖'
    with page.expect_request(lambda r:r.url.endswith('/api/m11/export')) as request, page.expect_download() as event:
        page.get_by_role('button',name='导出 Excel',exact=True).click()
    assert request.value.post_data_json['ids']==expected
    path=tmp_path/'contracts.xlsx';event.value.save_as(path)
    workbook=openpyxl.load_workbook(path,data_only=True)
    assert workbook.sheetnames==['集采台账'];values=list(workbook.active.values);assert len(values)==len(expected)+2
    mapping={r['id']:r for r in rows}
    for cells,cid in zip(values[1:-1],expected):
        row=mapping[cid]
        assert cells[2]==row['material'] and cells[3]==(row['contract_no'] or '—')
        assert cells[5]==pytest.approx(float(row['contract_amt']),abs=.005)
    assert values[-1][5]==pytest.approx(sum(float(mapping[cid]['contract_amt']) for cid in expected),abs=.01)
    workbook.close()
    page.get_by_label('搜索',exact=True).fill('no-match-contract-7efb35')
    expect(page.get_by_text('没有符合条件的合同，可调整筛选。',exact=True)).to_be_visible()
    with page.expect_request(lambda r:r.url.endswith('/api/m11/export')) as request, page.expect_download() as event:
        page.get_by_role('button',name='导出 Excel',exact=True).click()
    assert request.value.post_data_json['ids']==[]
    path=tmp_path/'empty.xlsx';event.value.save_as(path);workbook=openpyxl.load_workbook(path,data_only=True);assert workbook.active.max_row==1;workbook.close()


def test_contract_download_is_disabled_anonymously(server,browser_page):
    page,_=browser_page;open_page(page,server,'contracts','集采台账',authenticated=False)
    expect(page.get_by_role('button',name='导出 Excel',exact=True)).to_be_disabled()
    assert page.request.post(server+'/api/m11/export',data={'ids':[]}).status==401
