"""Fee item identity, monthly value, history and browser integration on a disposable real copy."""
from decimal import Decimal
from uuid import uuid4
from playwright.sync_api import expect
from test_preview import server,browser_page
from test_fees import sample,open_sample
from live_data import get


def post(api,path,payload):
    response=api.post(path,json=payload);response.raise_for_status();data=response.json();assert data['ok'],data;return data


def test_item_months_export_and_history(server,api,projects,disposable):
    project=next(p['warehouse_name'] for p in projects if p['warehouse_name'] and p['project_type']!='shell')
    name='核验费用'+uuid4().hex[:10];renamed=name+'改';year='2026'
    scope={'project':project};item=post(api,'/api/m09/fee/item/create',{**scope,'name':name})['row_id']
    def year_data():return get(api,'/api/m09/fee/year',year=year,pj=project)
    def row():return next(r for r in year_data()['rows'] if r.get('item_id')==item)
    assert row()['months']==[None]*12
    post(api,'/api/m09/fee/item/amount',{**scope,'item_id':item,'ym':year+'-03','value':'12.34'})
    post(api,'/api/m09/fee/item/amount',{**scope,'item_id':item,'ym':year+'-09','value':'0'})
    post(api,'/api/m09/fee/item/rename',{**scope,'item_id':item,'value':renamed})
    assert row()['item']==renamed and Decimal(row()['months'][2])==Decimal('12.34') and Decimal(row()['months'][8])==0
    month=get(api,'/api/m09/fee',ym=year+'-03',pj=project)
    assert [(x['item'],Decimal(x['amount'])) for x in month['manuals'] if x['item']==renamed]==[(renamed,Decimal('12.34'))]
    from io import BytesIO
    from openpyxl import load_workbook
    export=api.get('/api/m05/ledger/export',params={'projects':project,'ym':year+'-03'});export.raise_for_status()
    workbook=load_workbook(BytesIO(export.content),data_only=True)
    try:
        sheet=workbook['特殊费用月报表']
        found=[list(row) for row in sheet.iter_rows(values_only=True) if renamed in row]
        assert len(found)==1 and any(str(value) in ('12.34','12.340') for value in found[0])
    finally:workbook.close()
    post(api,'/api/m09/fee/item/delete',{**scope,'item_id':item})
    assert not any(r.get('item_id')==item for r in year_data()['rows'])
    history=api.get('/api/history/status').json()['undo'];assert history
    post(api,'/api/history/undo',{'operation_id':history['id']})
    assert row()['item']==renamed and Decimal(row()['months'][2])==Decimal('12.34') and Decimal(row()['months'][8])==0
    redo=api.get('/api/history/status').json()['redo'];assert redo
    post(api,'/api/history/redo',{'operation_id':redo['id']})
    assert not any(r.get('item_id')==item for r in year_data()['rows'])


def test_fee_item_browser_entry_and_reload(server,browser_page,api,projects,disposable):
    page,_=browser_page;project,data=sample(api,projects);open_sample(page,server,project,data)
    year=data['year'];key=project['warehouse_name'];before=get(api,'/api/m09/fee/year',year=year,pj=key)
    ids={r.get('item_id') for r in before['rows']};page.get_by_role('button',name='增加手填项').click()
    expect(page.locator('.fees-page').get_by_role('status').filter(has_text='已增加')).to_be_visible(timeout=15000)
    after=get(api,'/api/m09/fee/year',year=year,pj=key);new=[r for r in after['rows'] if r.get('item_id') and r['item_id'] not in ids];assert len(new)==1
    name=new[0]['item'];edit=page.get_by_label(f'{name}费用项名称',exact=True);expect(edit).to_be_visible();renamed='界面核验'+uuid4().hex[:8];edit.fill(renamed);edit.press('Enter')
    expect(page.locator('.fees-page').get_by_role('status').filter(has_text='已保存')).to_be_visible(timeout=15000)
    amount=page.get_by_label(f'{year}-03 {renamed}金额',exact=True);expect(amount).to_be_visible(timeout=15000);amount.fill('42.50');amount.press('Enter')
    expect(page.locator('.fees-page').get_by_role('status').filter(has_text='已保存')).to_be_visible(timeout=15000)
    def saved():return next(r for r in get(api,'/api/m09/fee/year',year=year,pj=key)['rows'] if r.get('item_id')==new[0]['item_id'])
    assert Decimal(saved()['months'][2])==Decimal('42.50')
    page.reload();open_sample(page,server,project,data)
    expect(page.get_by_label(f'{year}-03 {renamed}金额',exact=True)).to_have_value('42.50',timeout=15000)


def test_legacy_monthly_add_undo_and_batch(server,api,projects,disposable):
    project=next(p['warehouse_name'] for p in projects if p['warehouse_name'] and p['project_type']!='shell')
    name='旧入口核验'+uuid4().hex[:10];scope={'project':project}
    post(api,'/api/m09/fee/add',{**scope,'ym':'2026-03','name':name,'raw':'7.25'})
    history=api.get('/api/history/status').json()['undo'];assert history
    post(api,'/api/history/undo',{'operation_id':history['id']})
    assert not any(r['item']==name for r in get(api,'/api/m09/fee',ym='2026-03',pj=project)['manuals'])
    redo=api.get('/api/history/status').json()['redo'];assert redo
    post(api,'/api/history/redo',{'operation_id':redo['id']})
    post(api,'/api/m09/fee/batch',{**scope,'items':[name],'months':[{'ym':'2026-03','values':{name:'9.50'}},{'ym':'2026-04','values':{name:'0'}}]})
    for ym,want in [('2026-03','9.50'),('2026-04','0')]:
        matches=[r for r in get(api,'/api/m09/fee',ym=ym,pj=project)['manuals'] if r['item']==name]
        assert len(matches)==1 and Decimal(matches[0]['amount'])==Decimal(want)


def test_next_month_draft_and_focus_survive_previous_save(server,browser_page,api,projects,disposable):
    page,_=browser_page;project,data=sample(api,projects)
    name='连续录入'+uuid4().hex[:8]
    item=post(api,'/api/m09/fee/item/create',{'project':project['warehouse_name'],'name':name})['row_id']
    open_sample(page,server,project,data)
    first=page.get_by_label(f'{data["year"]}-03 {name}金额',exact=True)
    second=page.get_by_label(f'{data["year"]}-04 {name}金额',exact=True)
    pending=[]
    def hold(route):
        response=route.fetch();assert response.ok
        pending.append((route,response))
    page.route('**/api/m09/fee/item/amount',hold)
    try:
        first.fill('17.25');first.press('Tab')
        second.fill('23.50')
        expect(second).to_be_focused()
        assert len(pending)==1
        route,response=pending.pop();route.fulfill(response=response)
        expect(page.locator('.fees-page').get_by_role('status').filter(has_text='已保存')).to_be_visible(timeout=15000)
        expect(second).to_have_value('23.50');expect(second).to_be_focused()
        page.unroute('**/api/m09/fee/item/amount',hold)
        second.press('Enter')
        expect(page.locator('.fees-page').get_by_role('status').filter(has_text='已保存')).to_be_visible(timeout=15000)
        page.wait_for_function("""() => !document.querySelector('.fee-draft-feedback')""")
        row=next(r for r in get(api,'/api/m09/fee/year',year=data['year'],pj=project['warehouse_name'])['rows'] if r.get('item_id')==item)
        assert Decimal(row['months'][2])==Decimal('17.25') and Decimal(row['months'][3])==Decimal('23.50')
    finally:
        for route,response in pending:route.fulfill(response=response)
