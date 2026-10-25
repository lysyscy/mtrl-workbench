"""Inline annual editing keeps original month, project and record identity."""
from decimal import Decimal
import re
import time

import pytest
from playwright.sync_api import expect
from sqlalchemy import select
from test_preview import server, browser_page, login
from tools.audit.workspace_paths import output


@pytest.fixture(scope='module')
def preview_seed(make_preview_seed, preview_session):
    def seed(env):
        from app.mod01_projects.models import Project
        from app.mod08_inv.models import InvDetail
        from app.mod09_fee.models import FeeManualItem
        with preview_session(env) as s:
            pj = s.scalar(select(Project).where(Project.warehouse_name == '隔离验收仓库'))
            other = s.scalar(select(Project).where(Project.warehouse_name == '长名称验收仓库'))
            s.add_all([FeeManualItem(project_id=pid,ym=ym,item_name=name,amount=Decimal(amount)) for pid,ym,name,amount in [
                (pj.id,'2024-03','运费','10.10'),(pj.id,'2025-12','运费','-2.10'),
                (pj.id,'2026-01','运费','10'),(pj.id,'2026-01','运费','20'),(pj.id,'2026-12','运费','-5'),
                (pj.id,'2027-01','未来项','999'),(other.id,'2026-01','运费','999'),
                (pj.id,'2028-01','安全文明施工费（按 CBS 分类）','20'),
                (pj.id,'2026-09','油卡','0'),(pj.id,'2026-09',' 油卡','1'),
                (pj.id,'2025-12','大额','999999.99'),(pj.id,'2026-01','大额','1234567.89'),
            ]])
            s.add_all([InvDetail(project_id=pj.id,warehouse_name='隔离验收仓库',ym=ym,direction=direction,cbs_name=cbs,amount=Decimal(amount),doc_date='2099-01-01') for ym,direction,cbs,amount in [
                ('2025-01','出库','CBS-安文费','100'),('2026-02','出库','CBS-安文费','20'),
                ('2028-01','出库','CBS-安文费','100'),('2028-02','出库','CBS-安文费','100'),
                ('2026-02','入库','CBS-安文费','1000'),('2026-02','出库','其他','1000'),('2027-02','出库','CBS-安文费','1000'),
            ]])
            s.commit()
    return make_preview_seed(seed)


def annual(page, base, year='2026'):
    response=page.request.get(base+'/api/m09/fee/year',params={'year':year,'pj':'隔离验收仓库'})
    assert response.status==200
    return response.json()


def monthly(page,base,ym):
    response=page.request.get(base+'/api/m09/fee',params={'ym':ym,'pj':'隔离验收仓库'})
    assert response.status==200
    return response.json()['manuals']


def editor(page,ym,item,field='金额'):
    return page.get_by_role('textbox',name=re.compile('^'+re.escape(f'{ym} {item}{field}')+r' #\d+$'),include_hidden=field=='项名')


def wait_pair(page,base,ym,item,value):
    deadline=time.monotonic()+15
    while time.monotonic()<deadline:
        if any(r['item']==item and Decimal(r['amount'])==Decimal(value) for r in monthly(page,base,ym)):
            return
        page.wait_for_timeout(100)
    raise AssertionError(monthly(page,base,ym))


def open_fees(page,base,authenticated=True):
    if authenticated:
        login(page,base)
        page.get_by_role('link',name='MOD-09 特殊费用',exact=True).click()
    else:
        page.goto(base+'/preview/fees')
    expect(page.get_by_label('特殊费用年度表',exact=True)).to_be_visible()


def test_annual_original_history_duplicates_and_inline_record_edits(server,browser_page):
    page,_=browser_page
    month_reads=[]
    page.on('request',lambda request:month_reads.append(request.url) if '/api/m09/fee?' in request.url else None)
    open_fees(page,server)
    initial=annual(page,server)
    assert month_reads==[]
    assert len(initial['records'])==6
    assert [r['index'] for r in initial['records'] if r['ym']=='2026-01']==[0,1,2]
    rows={r['item']:r for r in initial['rows']}
    assert (rows['运费']['prior'],rows['运费']['months'][0],rows['运费']['total'])==('8.00','30.00','33.00')
    assert rows['运费']['months'][1] is None
    auto=rows['安全文明施工费（按 CBS 分类）']
    assert (auto['prior'],auto['months'][1],auto['total'])==('100.00','20.00','120.00')
    assert rows['油卡']['months'][8]=='0.00' and ' 油卡' in rows and '未来项' not in rows
    grid=page.get_by_label('特殊费用年度表',exact=True)
    assert grid.get_by_role('columnheader').count()==15
    assert grid.locator('[col-id=prior] input,[col-id=total] input').count()==0
    assert grid.locator('.ag-row').filter(has_text='安全文明施工费（按 CBS 分类）').get_by_role('textbox').count()==0
    assert page.get_by_role('dialog').count()==0
    amounts=editor(page,'2026-01','运费')
    expect(amounts).to_have_count(2)
    amounts.first.fill('15');amounts.first.press('Tab')
    wait_pair(page,server,'2026-01','运费','15')
    expect(amounts.first).to_have_value('15.00')
    assert [Decimal(r['amount']) for r in monthly(page,server,'2026-01') if r['item']=='运费']==[Decimal('15'),Decimal('20')]
    assert {r['item']:r for r in annual(page,server)['rows']}['运费']['months'][0]=='35.00'
    assert len(month_reads)==1
    amounts.first.fill('abc');amounts.first.press('Tab')
    expect(page.get_by_role('alert').first).to_contain_text('金额不合法')
    assert [Decimal(r['amount']) for r in monthly(page,server,'2026-01') if r['item']=='运费']==[Decimal('15'),Decimal('20')]
    page.get_by_role('button',name='取消这项修改',exact=True).click()
    name=editor(page,'2026-01','运费','项名').nth(1)
    name.locator('xpath=ancestor::details').locator('summary').click()
    name.fill('专用运费');name.press('Tab')
    wait_pair(page,server,'2026-01','专用运费','20')
    expect(editor(page,'2026-01','专用运费')).to_have_value('20.00')
    assert {r['item']:r for r in annual(page,server)['rows']}['运费']['prior']=='8.00'
    editor(page,'2026-01','专用运费','项名').locator('xpath=ancestor::details').locator('summary').click()
    page.get_by_role('button',name=re.compile(r'^2026-01 专用运费删除 #')).click()
    page.get_by_role('button',name='删除',exact=True).last.click()
    expect(editor(page,'2026-01','专用运费')).to_have_count(0)
    assert any(r['item']=='运费' and Decimal(r['amount'])==15 for r in monthly(page,server,'2026-01'))
    page.screenshot(path=str(output('fee-inline-records-synthetic.png',kind='screenshots')),full_page=True)


def test_add_immediately_creates_inline_row_and_empty_zero_negative(server,browser_page):
    page,_=browser_page
    open_fees(page,server)
    ym=page.get_by_label('新增费用核算月',exact=True).input_value()
    before=monthly(page,server,ym)
    page.get_by_role('button',name='增加手填项',exact=True).click()
    name=editor(page,ym,'新项','项名')
    expect(name).to_have_count(1)
    assert page.get_by_role('dialog').count()==0
    assert len(monthly(page,server,ym))==len(before)+1
    name.fill('清零样例');name.press('Tab')
    amount=editor(page,ym,'清零样例')
    expect(amount).to_have_value(re.compile(r'^0(\.00)?$'))
    for raw,expected in [('5','5.00'),('','0.00'),('0','0.00'),('-12.34','-12.34')]:
        amount.fill(raw);amount.press('Tab')
        wait_pair(page,server,ym,'清零样例',expected)
        expect(amount).to_have_value(re.compile(r'^0(\.00)?$') if expected=='0.00' else expected)
    blank=page.get_by_role('textbox',name='2026-03 运费金额',exact=True)
    expect(blank).to_have_value('')
    blank.fill('0');blank.press('Tab')
    wait_pair(page,server,'2026-03','运费','0')
    expect(editor(page,'2026-03','运费')).to_have_value(re.compile(r'^0(\.00)?$'))
    blank=page.locator('input[aria-label="2026-04  油卡金额"]')
    blank.fill('-3.5');blank.press('Tab')
    wait_pair(page,server,'2026-04',' 油卡','-3.5')
    assert {r['item']:r for r in annual(page,server)['rows']}['运费']['months'][4] is None


def test_existing_add_contract_validates_before_create_and_returns_identity(server,browser_page):
    page,_=browser_page
    login(page,server)
    scope={'ym':'2026-06','project':'隔离验收仓库'}
    before=monthly(page,server,'2026-06')
    invalid=page.request.post(server+'/api/m09/fee/add',data={**scope,'name':'非法新项','raw':'abc'})
    assert invalid.status==200 and invalid.json()['ok'] is False
    assert monthly(page,server,'2026-06')==before
    for name,raw in [(' 原名','0'),(' 原名','-2.25')]:
        saved=page.request.post(server+'/api/m09/fee/add',data={**scope,'name':name,'raw':raw}).json()
        assert saved['ok'] and isinstance(saved['row_id'],int)
        assert any(r['id']==saved['row_id'] and r['item']==name and Decimal(r['amount'])==Decimal(raw) for r in monthly(page,server,'2026-06'))
    assert len(monthly(page,server,'2026-06'))==len(before)+2


def test_late_amount_receipt_preserves_next_draft_and_focus(server,browser_page):
    page,_=browser_page
    open_fees(page,server)
    held=[]
    def hold(route):
        response=route.fetch();held.append((route,response))
    page.route('**/api/m09/fee/amount',hold)
    first=editor(page,'2026-01','运费').first
    other=editor(page,'2026-01','大额').first
    first.fill('17');first.press('Tab')
    expect(page.get_by_role('status')).to_contain_text('正在保存')
    other.fill('-22.25')
    expect(other).to_be_focused()
    assert len(held)==1
    route,response=held.pop();route.fulfill(response=response)
    page.unroute('**/api/m09/fee/amount')
    expect(other).to_have_value('-22.25')
    expect(other).to_be_focused()
    other.press('Tab')
    wait_pair(page,server,'2026-01','大额','-22.25')
    wait_pair(page,server,'2026-01','运费','17')


def test_damaged_add_receipt_rereads_without_readding(server,browser_page):
    page,_=browser_page
    open_fees(page,server)
    writes=[]
    def damaged(route):
        writes.append(route.request.post_data)
        response=route.fetch();assert response.ok
        route.fulfill(response=response,body='{')
    page.route('**/api/m09/fee/add',damaged)
    blank=page.get_by_role('textbox',name='2026-05 运费金额',exact=True)
    blank.fill('-4.25');blank.press('Tab')
    expect(page.get_by_role('alert').first).to_contain_text('保存结果待核实')
    assert len(writes)==1
    expect(page.get_by_role('button',name='增加手填项',exact=True)).to_be_disabled()
    wait_pair(page,server,'2026-05','运费','-4.25')
    page.unroute('**/api/m09/fee/add')
    page.get_by_role('button',name='重读核对',exact=True).click()
    expect(page.get_by_role('status')).to_contain_text('已重读')
    page.get_by_role('button',name='已核对，保留输入继续',exact=True).click()
    field=editor(page,'2026-05','运费')
    expect(field).to_have_value('-4.25')
    field.fill('-4.5');field.press('Tab')
    wait_pair(page,server,'2026-05','运费','-4.5')
    assert len([r for r in monthly(page,server,'2026-05') if r['item']=='运费'])==1


def test_fee_writes_require_login_and_change_nothing(server,browser_page):
    import httpx
    page,_=browser_page
    open_fees(page,server,authenticated=False)
    expect(page.get_by_role('button',name='增加手填项',exact=True)).to_be_disabled()
    expect(editor(page,'2026-01','运费').first).to_be_disabled()
    before=monthly(page,server,'2026-01')
    with httpx.Client(base_url=server) as client:
        for path,payload in [
            ('add',{'name':'越权项','raw':'999'}),('name',{'index':0,'value':'越权'}),
            ('amount',{'index':0,'value':'999'}),('delete',{'index':0}),
        ]:
            response=client.post('/api/m09/fee/'+path,json={'ym':'2026-01','project':'隔离验收仓库',**payload})
            assert response.status_code==401
    assert monthly(page,server,'2026-01')==before


def test_fee_sheet_exports_in_ledger_workbook(server,browser_page):
    """费用随逐月欠款台账三表工作表导出：真实隔离数据经真实模板构建器落本轮 .artifacts。

    ⚠️ 端点全链（`/api/m05/ledger/export`）会把成品写进正式 `EXPORT_ROOT`（`PROJECT_ROOT/exports`，
    代码内固定、无配置覆盖），按本票硬边界不触发；此处取同一构建器 `ledger_export.build_xlsx` ＋
    同一页面数据（`/api/m09/fee`、`/api/m05/ledger`）在隔离数据上产出，落点限定在本轮 .artifacts。
    """
    import re
    from decimal import Decimal
    from pathlib import Path
    from app.report_export import ledger_export as LE
    from tools.audit.workspace_paths import output_dir
    page,_ = browser_page
    login(page,server)
    page.get_by_role('link',name='MOD-09 特殊费用',exact=True).click()
    expect(page.get_by_text('随逐月欠款台账导出 · 特殊费用月报表',exact=True)).to_be_visible()
    assert page.get_by_role('button',name=re.compile('导出')).count() == 0    # 费用页无独立导出
    fee = page.request.get(server+'/api/m09/fee',params={'ym':'2026-09','pj':'隔离验收仓库'}).json()
    assert fee['project'] == '隔离验收仓库'
    projects = page.request.get(server+'/api/m01/projects?status=all&keyword=').json()
    pid = next(p['id'] for p in projects if p['warehouse_name'] == '隔离验收仓库')
    ledger = page.request.get(server+'/api/m05/ledger',params={'project_id':pid,'end_ym':'2026-09'}).json()
    sections = {cat: next(s for s in ledger['sides'] if s['category'] == cat) for cat in ('material','mechanical')}
    sections[LE.SHEET_FEE] = fee
    out_dir = output_dir('exports') / 'tail415-ledger'
    path = Path(LE.build_xlsx(sections,out_dir,display=fee['project_display'],ym='2026-09'))
    assert str(path).startswith(str(output_dir('exports')))              # 只落本轮 .artifacts
    assert not str(path).startswith(str(Path(__file__).resolve().parents[3] / 'exports'))
    assert path.name == f"{fee['project_display']}欠款台账_2026年9月.xlsx"
    import openpyxl
    wb = openpyxl.load_workbook(path)
    assert set(wb.sheetnames) == {'辅材','机械','特殊费用月报表'}
    ws = wb['特殊费用月报表']
    assert ws['A1'].value == f"{fee['project_display']}费用汇总表"
    assert '9月份' in {ws.cell(row=2,column=c).value for c in range(1,ws.max_column+1)}
    items = [fee['auto'],*fee['manuals']]
    for offset,item in enumerate(items):
        row = 3+offset
        assert ws.cell(row=row,column=1).value == (item['item'] or '').strip()
        assert float(ws.cell(row=row,column=2).value) == float(item['amount'])
        assert float(ws.cell(row=row,column=3).value) == float(item['amount'])   # 单月列 ⇒ 合计 = 该月值
    assert ws.cell(row=3+len(items),column=1).value is None    # 行数随数据项派生，无残留
    assert sum(Decimal(it['amount']) for it in items) == Decimal(fee['total'])
    wb.close()
    page.screenshot(path=str(output('fee-export-workbook.png',kind='screenshots')),full_page=True)


def test_repeated_add_is_visible_as_separate_annual_rows(server,browser_page):
    page,_=browser_page
    login(page,server)
    scope={'project':'隔离验收仓库'}
    originals=[]
    for ym,name,raw in [('2026-07','新项','2'),('2026-07','新项','3'),('2026-08','新项2','4')]:
        response=page.request.post(server+'/api/m09/fee/add',data={**scope,'ym':ym,'name':name,'raw':raw})
        assert response.ok and response.json()['ok']
        originals.append(response.json()['row_id'])
    page.get_by_role('link',name='MOD-09 特殊费用',exact=True).click()
    expect(page.get_by_label('特殊费用年度表',exact=True)).to_be_visible()
    before=annual(page,server)
    add=page.get_by_role('button',name='增加手填项',exact=True)
    ym=page.get_by_label('新增费用核算月',exact=True).input_value()
    for name in ['新项3','新项4']:
        expect(add).to_be_enabled()
        add.click()
        expect(editor(page,ym,name,'项名')).to_be_visible()
        assert any(row['item']==name for row in annual(page,server)['rows'])
    after=annual(page,server)
    assert len(after['rows'])==len(before['rows'])+2
    assert [row['item'] for row in after['records'] if row['id'] in originals]==['新项','新项','新项2']
    assert [Decimal(row['amount']) for row in monthly(page,server,'2026-07') if row['id'] in originals]==[Decimal('2'),Decimal('3')]


def test_automatic_amount_remains_visible_beside_same_name_manual_records(server,browser_page):
    page,_=browser_page
    open_fees(page,server)
    year=page.get_by_label('费用年份',exact=True)
    year.fill('2028');year.press('Enter')
    name='安全文明施工费（按 CBS 分类）'
    manual=editor(page,'2028-01',name)
    expect(manual).to_have_value('20.00')
    grid=page.get_by_label('特殊费用年度表',exact=True)
    row=grid.locator('.ag-row').filter(has=manual)
    january=row.locator('[col-id="month-1"]')
    february=row.locator('[col-id="month-2"]')
    expect(january.locator('.fee-auto-amount')).to_have_text('自动 100.00')
    expect(january.get_by_text('合计 120.00',exact=True)).to_be_visible()
    expect(february.locator('.fee-auto-amount')).to_have_text('自动 100.00')
    expect(february.get_by_role('textbox')).to_have_value('')
    assert january.locator('.fee-auto-amount input').count()==0
    assert february.locator('.fee-auto-amount input').count()==0
    for raw,total in [('-5','95.00'),('0','100.00')]:
        manual.fill(raw);manual.press('Tab')
        wait_pair(page,server,'2028-01',name,raw)
        expect(january.get_by_text('合计 '+total,exact=True)).to_be_visible()
        expect(january.locator('.fee-auto-amount')).to_have_text('自动 100.00')
        assert {r['item']:r for r in annual(page,server,'2028')['rows']}[name]['months'][0]==total
        monthly_data=page.request.get(server+'/api/m09/fee',params={'ym':'2028-01','pj':'隔离验收仓库'}).json()
        assert Decimal(monthly_data['auto']['amount'])==Decimal('100')
