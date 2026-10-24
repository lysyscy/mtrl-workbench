"""TAIL-414: inventory query, manual fetch boundary, adjustments and real export file.

Runs against the real isolated backend (own schema) with the shared preview build.
The i8 source-cache directory is redirected into this run's artifacts by a
runtime-only hook (sitecustomize injected through the seed env), so neither the
worktree `var/` state nor the formal export root is touched: exports land in
`.artifacts/active/<effort>/<run>/exports` via the preview server's own guard.
"""
import io
import os
import re
import subprocess
import sys
import zipfile
from pathlib import Path

import httpx
import openpyxl
import pytest
from playwright.sync_api import expect
from test_preview import server, browser_page, login
from preview_select import open_popup, pick_option
from tools.audit.workspace_paths import output, output_dir


@pytest.fixture(scope='module')
def preview_seed(make_preview_seed):
    def seed(env):
        runtime = Path(output_dir('work')) / 'tail414-runtime'
        runtime.mkdir(parents=True, exist_ok=True)
        cache_dir = runtime / 'i8_source_cache'
        (runtime / 'sitecustomize.py').write_text(
            "import os\nfrom pathlib import Path\n"
            "_dir = os.environ.get('MTRL_INV_TEST_SOURCE_CACHE')\n"
            "if _dir:\n"
            "    try:\n"
            "        from app.mod02_i8 import source_cache\n"
            "        source_cache.DIR = Path(_dir)\n"
            "    except Exception:\n"
            "        pass\n",
            encoding='utf-8')
        env['PYTHONPATH'] = str(runtime) + os.pathsep + env.get('PYTHONPATH', '')
        env['MTRL_INV_TEST_SOURCE_CACHE'] = str(cache_dir)
        code = """
import io, os
from pathlib import Path
from openpyxl import Workbook
from app.report_export.schema_cols import KCDJ_COLS
from app.core.db import session_scope
from app.mod01_projects.models import Project  # noqa: F401 - FK target of inv_detail
from app.mod08_inv.importer import parse_rows
from app.mod08_inv.repository import upsert_details


def grow(seq, year, month, doc, direction, material, prop, spec, unit, qty, amount, tax,
         wh, cbs, vendor='', biz='采购入库', project='隔离验收项目', doc_date='', audit_date='', wbs='9号楼'):
    row = [''] * 28
    row[0], row[1], row[2] = str(seq), year, month
    row[3], row[4], row[5] = biz, project, doc
    name, team, dept = (vendor.split('/') + ['', ''])[:3]
    row[6], row[19], row[20] = name, team, dept
    row[8], row[9], row[10], row[11] = material, unit, spec, prop
    row[12], row[15], row[16] = qty, amount, tax
    row[13] = row[14] = str(round(float(amount) / float(qty), 2))
    row[17], row[18] = direction, wh
    row[21], row[22], row[25], row[27] = doc_date, audit_date, wbs, cbs
    return row


rows_09 = [
    grow(1, '2026', '9', 'TAIL414-D1', '入库', '螺纹钢', '主材', 'HRB400E', '吨', '10', '12345.67', '13950.61',
         '隔离验收仓库', '主体工程', '合成供应商/钢筋班/材料部', doc_date='2026-09-05', audit_date='2026-09-06'),
    grow(2, '2026', '9', 'TAIL414-D1', '出库', '焊条', '辅材', 'J422', '包', '4', '240.00', '271.20',
         '隔离验收仓库', '主体工程', '合成供应商/焊接班/材料部', biz='领用出库', doc_date='2026-09-12', audit_date='2026-09-13'),
    grow(1, '2026', '9', 'TAIL414-D2', '入库', '脚手架', '主材', '6m', '套', '30', '9000.00', '10170.00',
         '隔离归档仓库', '主体工程', '合成供应商/架设班/安全部', doc_date='2026-09-15', audit_date='2026-09-16'),
    grow(2, '2026', '9', 'TAIL414-D2', '出库', '安全网', '辅材', '3x6', '张', '50', '1500.00', '1695.00',
         '隔离归档仓库', '安全文明施工费', '合成供应商/架设班/安全部', biz='领用出库', doc_date='2026-09-16', audit_date='2026-09-17'),
    grow(1, '2026', '9', 'TAIL414-D3', '入库', '螺纹钢', '主材', 'HRB400E', '吨', '2', '2000.00', '2260.00',
         '隔离验收仓库', '主体工程', '合成供应商/钢筋班/材料部', doc_date='2026-09-18', audit_date='2026-09-19'),
]
rows_08 = [
    grow(1, '2026', '8', 'TAIL414-D0', '入库', '水泥', '主材', 'P.O42.5', '吨', '100', '35000.00', '39550.00',
         '隔离验收仓库', '主体工程', '合成供应商/搅拌站/材料部', doc_date='2026-08-20', audit_date='2026-08-21'),
]


def build(rows):
    book = Workbook()
    sheet = book.active
    sheet.title = '报表1'
    sheet.append(['库存单据明细表'] + [''] * 27)
    sheet.append([KCDJ_COLS[i] for i in range(28)])
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


content_09 = build(rows_09)
with session_scope() as session:
    groups = {}
    for row in parse_rows(content_09) + parse_rows(build(rows_08)):
        groups.setdefault((row['ym'], row['warehouse_name']), []).append(row)
    for (ym, warehouse), group in groups.items():
        upsert_details(session, ym=ym, warehouse=warehouse, rows=group)
cache = Path(os.environ['MTRL_INV_TEST_SOURCE_CACHE']) / 'inv' / '2026-09.xlsx'
cache.parent.mkdir(parents=True, exist_ok=True)
cache.write_bytes(content_09)
print('SEED inventory rows parsed and source cache written', flush=True)
"""
        subprocess.run([sys.executable, '-X', 'utf8', '-c', code], env=env, check=True)
    return make_preview_seed(seed)


def rows_of(page, text=''):
    """Live grid body rows only (header contains 合计 via 价税合计; AG Grid keeps an
    opacity-zero ghost during row swaps)."""
    base = page.locator('[role="row"]:not(.ag-opacity-zero):not(.ag-header-row)')
    return base.filter(has_text=text) if text else base


def open_inventory(page, server):
    login(page, server)
    page.get_by_role('link', name='MOD-08 库存单据与调整单').click()
    expect(rows_of(page, 'TAIL414-D1').first).to_be_visible()
    expect(page.get_by_text('5 / 6 行', exact=True)).to_be_visible()


def api_inventory(page, server):
    return page.request.get(server + '/api/m08/inventory').json()


def row_id_of(data, doc, material):
    return next(row['id'] for row in data['rows'] if row['doc_no'] == doc and row['material'] == material)


def set_column_filter(page, column, *, text=None, min_value=None, max_value=None, values=()):
    if not page.locator('.inventory-filter').is_visible():
        page.get_by_role('button', name=re.compile('^列筛选')).click()
    pick_option(page, '筛选列', column)
    if text is not None:
        page.get_by_label('列包含文字', exact=True).fill(text)
    if min_value is not None:
        page.get_by_label('列最小值', exact=True).fill(min_value)
    if max_value is not None:
        page.get_by_label('列最大值', exact=True).fill(max_value)
    for value in values:
        open_popup(page, '列值清单').filter(has_text=re.compile(f'^{re.escape(value)}$')).click()
        page.keyboard.press('Escape')


def back_to_default_month(page):
    """From 全部 back to the default single month through the picker, or reload."""
    pick_option(page, '库存核算期间模式', '月')
    page.get_by_label('库存核算期间', exact=True).click()
    cell = page.locator('.ant-picker-panel td[title="2026-09"]')
    if cell.count():
        cell.click()
    else:
        page.keyboard.press('Escape')
        page.reload()
    expect(page.get_by_text('5 / 6 行', exact=True)).to_be_visible()


def test_writes_need_login_and_fetch_stays_inside_external_boundary(server, browser_page):
    page, _ = browser_page
    page.goto(server + '/preview/inventory')
    expect(page.get_by_text('未登录 · 只读', exact=True)).to_be_visible()
    expect(rows_of(page, 'TAIL414-D1').first).to_be_visible()
    expect(page.get_by_role('button', name='按月抓取…', exact=True)).to_be_disabled()
    # Adjustments open read-only: every write control is disabled without an account.
    rows_of(page, 'TAIL414-D3').get_by_role('button', name='办理调整').click()
    dialog = page.locator('.inventory-dialog')
    expect(dialog.get_by_label('调整后金额')).to_be_disabled()
    expect(dialog.get_by_label('调整备注')).to_be_disabled()
    expect(dialog.get_by_role('button', name='保存调整单')).to_be_disabled()
    expect(dialog.get_by_role('button', name='撤销调整单…')).to_have_count(0)
    dialog.get_by_role('button', name='取消', exact=True).click()
    before = httpx.get(server + '/api/m08/inventory').json()
    with httpx.Client(base_url=server) as client:
        for path, payload in [
            ('/api/m08/adjust', {'row_id': 1, 'amount': '1', 'remark': '未登录写入'}),
            ('/api/m08/adjust/clear', {'row_id': 1}),
            ('/api/m08/fetch', {'month': '2026-09', 'confirm': True}),
        ]:
            response = client.post(path, json=payload)
            assert response.status_code == 401, (path, response.status_code)
            assert '未登录' in response.json()['message']
    assert httpx.get(server + '/api/m08/inventory').json() == before
    # Logged in: the dialog checks /api/m02/config first, refuses without a binding,
    # and merely opening the page or the dialog never sends a fetch request.
    fetch_calls = []
    page.on('request', lambda request: fetch_calls.append(request.url)
            if request.method == 'POST' and request.url.endswith('/api/m08/fetch') else None)
    open_inventory(page, server)
    assert not fetch_calls
    page.get_by_role('button', name='按月抓取…', exact=True).click()
    dialog = page.locator('.inventory-dialog')
    expect(dialog).to_contain_text('当前账号：预览管理员 · i8：未绑定')
    expect(dialog.get_by_role('alert')).to_contain_text('i8 服务地址未配置或格式不正确')
    expect(dialog.get_by_role('button', name='确认按 2026-09 抓取', exact=True)).to_be_disabled()
    assert not fetch_calls
    dialog.get_by_role('button', name='取消', exact=True).click()
    expect(dialog).not_to_be_visible()
    assert not fetch_calls
    # A forced fetch is refused server-side and changes nothing. With no project phid
    # the failure happens before any network action: no run record, no degrade, by design.
    rows_before = api_inventory(page, server)
    response = page.request.post(server + '/api/m08/fetch', data={'month': '2026-09', 'confirm': True})
    assert response.status < 500
    body = response.json()
    assert body.get('ok') is not True and body.get('message')
    rows_after = api_inventory(page, server)
    assert rows_before['rows'] == rows_after['rows']
    records = page.request.get(server + '/api/m02/fetch/records?limit=50').json()
    assert not [record for record in records if record['run_type'] == 'fetch'
                and record['target'].startswith('inv 2026-09')]


def test_period_project_direction_and_column_filters_match_api(server, browser_page):
    page, _ = browser_page
    open_inventory(page, server)
    data = api_inventory(page, server)
    assert [row['doc_no'] for row in data['rows'] if row['ym'] == '2026-09'] == \
        ['TAIL414-D1', 'TAIL414-D1', 'TAIL414-D3', 'TAIL414-D2', 'TAIL414-D2']
    # Default period is last month: the 2026-08 cement row stays out until 全部.
    expect(rows_of(page, '水泥')).to_have_count(0)
    pick_option(page, '库存核算期间模式', '全部')
    expect(page.get_by_text('6 / 6 行', exact=True)).to_be_visible()
    expect(rows_of(page, 'TAIL414-D0')).to_be_visible()
    back_to_default_month(page)
    # Project multi-select is keyed by warehouse; direction narrows further.
    page.get_by_label('库存项目', exact=True).click()
    open_popup(page, '库存项目').filter(has_text='隔离项目').click()
    page.keyboard.press('Escape')
    expect(page.get_by_text('3 / 6 行', exact=True)).to_be_visible()
    expect(rows_of(page, '隔离归档仓库')).to_have_count(0)
    pick_option(page, '库存方向', '入库')
    expect(page.get_by_text('2 / 6 行', exact=True)).to_be_visible()
    totals = rows_of(page, '合计')
    rows = [row for row in data['rows']
            if row['ym'] == '2026-09' and row['warehouse'] == '隔离验收仓库' and row['direction'] == '入库']
    amount = sum(round(float(row['amount']) * 100) for row in rows) / 100
    assert amount == 14345.67
    expect(totals).to_contain_text('14,345.67')
    expect(totals).to_contain_text('12.00')
    expect(totals).not_to_contain_text('1,234.57')
    # Column filters intersect: text, numeric range and a value list.
    set_column_filter(page, '物料', text='螺纹钢')
    expect(page.get_by_text('2 / 6 行', exact=True)).to_be_visible()
    expect(page.get_by_role('button', name='列筛选 · 1 列')).to_be_visible()
    expect(rows_of(page, 'TAIL414-D3')).to_contain_text('2,000.00')
    page.get_by_role('button', name='清除此列').click()
    set_column_filter(page, '金额', min_value='2000', max_value='10000')
    expect(page.get_by_text('1 / 6 行', exact=True)).to_be_visible()
    expect(rows_of(page, 'TAIL414-D3')).to_contain_text('2,000.00')
    expect(rows_of(page, 'TAIL414-D1')).to_have_count(0)
    page.get_by_role('button', name='清除全部列').click()
    expect(page.get_by_text('2 / 6 行', exact=True)).to_be_visible()
    pick_option(page, '库存方向', '全部')
    expect(page.get_by_text('3 / 6 行', exact=True)).to_be_visible()
    set_column_filter(page, '方向', values=['入库'])
    expect(page.get_by_text('2 / 6 行', exact=True)).to_be_visible()
    expect(page.get_by_role('button', name='列筛选 · 1 列')).to_be_visible()
    page.get_by_role('button', name='清除全部列').click()
    expect(page.get_by_text('3 / 6 行', exact=True)).to_be_visible()
    # Nothing is adjusted yet, so the only-adjusted switch yields the empty overlay.
    page.get_by_role('checkbox', name='只看被调整行').check()
    expect(page.get_by_text('当前筛选条件下没有库存单据')).to_be_visible()
    page.get_by_role('checkbox', name='只看被调整行').uncheck()
    expect(page.get_by_text('3 / 6 行', exact=True)).to_be_visible()
    page.screenshot(path=str(output('tail414-filters.png', kind='screenshots')), full_page=True)


def test_adjust_create_reject_uncertain_receipt_and_clear(server, browser_page):
    page, _ = browser_page
    open_inventory(page, server)
    data = api_inventory(page, server)
    target_id = row_id_of(data, 'TAIL414-D3', '螺纹钢')
    row = rows_of(page, 'TAIL414-D3')
    row.get_by_role('button', name='办理调整').click()
    dialog = page.locator('.inventory-dialog')
    expect(dialog.get_by_role('heading')).to_contain_text('调整单 · TAIL414-D3 #1')
    expect(dialog).to_contain_text('i8 原始金额')
    expect(dialog).to_contain_text('2,000.00')
    expect(dialog.get_by_label('调整后金额')).to_have_value('2000.00')
    # Empty amount or remark is refused before any request is sent.
    dialog.get_by_label('调整后金额').fill('')
    save = dialog.get_by_role('button', name='保存调整单')
    expect(save).to_be_disabled()
    expect(dialog).to_contain_text('调整后金额必填。')
    dialog.get_by_label('调整后金额').fill('100')
    expect(dialog).to_contain_text('调整备注必填。')
    expect(save).to_be_disabled()
    # A negative amount passes the form but is rejected by the backend; draft stays.
    dialog.get_by_label('调整备注').fill('TAIL414 核对修正')
    dialog.get_by_label('调整后金额').fill('-5')
    save.click()
    expect(dialog.get_by_role('alert')).to_contain_text('调整后金额不得为负数')
    assert api_inventory(page, server)['adjustments'] == []
    expect(dialog.get_by_label('调整后金额')).to_have_value('-5')
    # A damaged 200 receipt after a real write must lock the dialog, never resubmit.
    writes = []

    def damaged(route):
        writes.append(route.request.post_data)
        response = route.fetch()
        route.fulfill(response=response, body='{')
    page.route('**/api/m08/adjust', damaged)
    dialog.get_by_label('调整后金额').fill('88')
    save.click()
    expect(dialog).to_contain_text('结果待核实，勿重复提交。')
    expect(dialog).to_contain_text('读取结果不完整，请重试；仍失败请联系维护人员。')
    expect(dialog.get_by_label('调整后金额')).to_be_disabled()
    expect(dialog.get_by_role('button', name='保存调整单')).to_be_disabled()
    expect(page.get_by_role('alert').filter(has_text='的调整结果待核实')).to_be_visible()
    assert len(writes) == 1
    page.unroute('**/api/m08/adjust')
    written = api_inventory(page, server)['adjustments']
    assert len(written) == 1 and written[0]['row_id'] == target_id
    assert round(float(written[0]['amount']), 2) == 88.0
    assert round(float(written[0]['before_amount']), 2) == 2000.0
    assert written[0]['operator'] == '预览管理员'
    page.screenshot(path=str(output('tail414-adjust-unknown.png', kind='screenshots')), full_page=True)
    # Only an explicit reread plus confirmation unlocks further editing.
    dialog.get_by_role('button', name='重读核对', exact=True).click()
    expect(dialog).to_contain_text('已重读，请核对当前调整单；输入仍保留。')
    dialog.get_by_role('button', name='已核对，继续编辑').click()
    expect(dialog.get_by_label('调整后金额')).to_be_enabled()
    expect(dialog).to_contain_text('留痕：2,000.00 → 88.00')
    # Updating keeps the original before_amount and stores the new trail.
    dialog.get_by_label('调整后金额').fill('100')
    dialog.get_by_label('调整备注').fill('第二次修正')
    save.click()
    expect(dialog.locator('p[role="status"]')).to_contain_text('调整单已保存；原始数据未改，相关报表已标记重读。')
    expect(dialog).to_contain_text('留痕：2,000.00 → 100.00')
    expect(dialog).to_contain_text('预览管理员')
    dialog.get_by_role('button', name='取消', exact=True).click()
    expect(page.locator('.inventory-dialog')).not_to_be_visible()
    expect(row).to_contain_text('100.00')
    expect(row).to_contain_text('2,000.00')
    saved = api_inventory(page, server)
    saved_row = next(item for item in saved['rows'] if item['id'] == target_id)
    assert round(float(saved_row['adj_amount']), 2) == 100.0
    assert round(float(saved_row['amount']), 2) == 2000.0
    # Only-adjusted filter shows exactly the adjusted row.
    page.get_by_role('checkbox', name='只看被调整行').check()
    expect(page.get_by_text('1 / 6 行', exact=True)).to_be_visible()
    page.get_by_role('checkbox', name='只看被调整行').uncheck()
    expect(page.get_by_text('5 / 6 行', exact=True)).to_be_visible()
    # Reopen through the adjusted-row entry and revoke with the impact confirmation.
    row.get_by_role('button', name='查看 / 调整').click()
    dialog = page.locator('.inventory-dialog')
    expect(dialog).to_contain_text('留痕：2,000.00 → 100.00')
    dialog.get_by_role('button', name='撤销调整单…').click()
    confirm = dialog.get_by_role('alert')
    expect(confirm).to_contain_text('撤销本条调整单：100.00 → 原单 2,000.00。相关统计将重算。')
    dialog.get_by_role('button', name='保留调整').click()
    expect(confirm).not_to_be_visible()
    assert len(api_inventory(page, server)['adjustments']) == 1
    dialog.get_by_role('button', name='撤销调整单…').click()
    dialog.get_by_role('button', name='确认撤销调整单').click()
    expect(page.locator('.inventory-dialog')).not_to_be_visible()
    expect(row).not_to_contain_text('100.00')
    assert api_inventory(page, server)['adjustments'] == []
    expect(page.get_by_text('5 / 6 行', exact=True)).to_be_visible()
    page.screenshot(path=str(output('tail414-adjust-cleared.png', kind='screenshots')), full_page=True)


def test_row_deeplink_overrides_filters_and_clear_restores_them(server, browser_page):
    page, _ = browser_page
    open_inventory(page, server)
    data = api_inventory(page, server)
    rod_id = row_id_of(data, 'TAIL414-D3', '螺纹钢')
    set_column_filter(page, '物料', text='焊条')
    expect(page.get_by_text('1 / 6 行', exact=True)).to_be_visible()
    # In-app navigation keeps the page mounted: filters survive, locate overrides them.
    page.evaluate("id => { history.pushState({}, '', '/preview/inventory?row=' + id);"
                  " window.dispatchEvent(new PopStateEvent('popstate')); }", rod_id)
    located = page.locator('.inventory-location')
    expect(located).to_contain_text('已定位 TAIL414-D3 #1')
    expect(rows_of(page, 'TAIL414-D3')).to_be_visible()
    expect(page.get_by_text('1 / 6 行', exact=True)).to_be_visible()
    expect(page.get_by_label('列包含文字', exact=True)).to_have_value('焊条')
    page.screenshot(path=str(output('tail414-deeplink.png', kind='screenshots')), full_page=True)
    # Clearing removes only the row parameter and the original filter applies again.
    located.get_by_role('button', name='清除定位').click()
    expect(page.locator('.inventory-location')).to_have_count(0)
    assert 'row=' not in page.url
    expect(rows_of(page, 'TAIL414-D1').filter(has_text='焊条')).to_be_visible()
    expect(rows_of(page, 'TAIL414-D3')).to_have_count(0)
    expect(page.get_by_text('1 / 6 行', exact=True)).to_be_visible()
    # A missing or invalid row parameter says so instead of guessing a substitute.
    page.goto(server + '/preview/inventory?row=999999')
    expect(page.locator('.inventory-location')).to_contain_text('原行已不在当前库存中，请重读核对。')
    page.locator('.inventory-location').get_by_role('button', name='清除定位').click()
    expect(page.locator('.inventory-location')).to_have_count(0)
    page.goto(server + '/preview/inventory?row=abc')
    expect(page.locator('.inventory-location')).to_contain_text('定位参数无效。')
    expect(page.get_by_text('5 / 6 行', exact=True)).to_be_visible()


def test_export_real_file_matches_source_and_ignores_page_scope(server, browser_page):
    page, _ = browser_page
    open_inventory(page, server)
    data = api_inventory(page, server)
    target_id = row_id_of(data, 'TAIL414-D3', '螺纹钢')
    created = page.request.post(server + '/api/m08/adjust',
                                data={'row_id': target_id, 'amount': '100', 'remark': '导出隔离核对'}).json()
    assert created.get('ok') is True
    page.reload()
    expect(rows_of(page, 'TAIL414-D3')).to_contain_text('100.00')
    # Page scope (filter + adjustment) must not reach the export.
    set_column_filter(page, '物料', text='焊条')
    expect(page.get_by_text('1 / 6 行', exact=True)).to_be_visible()
    page.get_by_role('button', name='导出文件…', exact=True).click()
    dialog = page.get_by_role('dialog')
    expect(dialog.get_by_role('heading')).to_contain_text('导出 · 2026-09 库存（含数据透视表）')
    expect(dialog).to_contain_text('不受本页项目、方向和列筛选限制')
    expect(dialog).to_contain_text('不叠加调整单')
    with page.expect_download() as event:
        dialog.get_by_role('button', name='保存到本地', exact=True).click()
    download = event.value
    assert download.suggested_filename == '库存单据明细表_2026年9月.xlsx'
    path = output(download.suggested_filename, kind='exports')
    download.save_as(path)

    def number(value):
        return round(float(value), 2)

    def text(value):
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value)
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        assert sorted(workbook.sheetnames) == ['安全文明施工费', '报表1', '过程资料']
        grid = [list(row) for row in workbook['报表1'].iter_rows(values_only=True)]
        header = next(index for index, row in enumerate(grid) if '序号' in row and '金额' in row)
        details = [row for row in grid[header + 1:] if any(cell not in (None, '') for cell in row)]
        assert [(row[5], text(row[0]), row[8], row[17], row[18], row[27]) for row in details] == [
            ('TAIL414-D1', '1', '螺纹钢', '入库', '隔离验收仓库', '主体工程'),
            ('TAIL414-D1', '2', '焊条', '出库', '隔离验收仓库', '主体工程'),
            ('TAIL414-D2', '1', '脚手架', '入库', '隔离归档仓库', '主体工程'),
            ('TAIL414-D2', '2', '安全网', '出库', '隔离归档仓库', '安全文明施工费'),
            ('TAIL414-D3', '1', '螺纹钢', '入库', '隔离验收仓库', '主体工程'),
        ]
        assert [number(row[15]) for row in details] == [12345.67, 240.0, 9000.0, 1500.0, 2000.0]
        assert [number(row[12]) for row in details] == [10.0, 4.0, 30.0, 50.0, 2.0]
        # The two pivot display sheets aggregate the same untouched source.
        display = [list(row) for row in workbook['过程资料'].iter_rows(values_only=True)]
        cells = {number(cell) for row in display for cell in row
                 if isinstance(cell, (int, float)) and not isinstance(cell, bool)}
        assert {14345.67, 240.0, 9000.0, 1500.0} <= cells
        assert 25085.67 in cells or {14585.67, 10500.0} <= cells
        labels = {str(cell) for row in display for cell in row if cell}
        assert {'隔离验收仓库', '隔离归档仓库'} <= labels
        safe = [list(row) for row in workbook['安全文明施工费'].iter_rows(values_only=True)]
        safe_cells = [cell for row in safe for cell in row if cell not in (None, '')]
        assert '隔离归档仓库' in {str(cell) for cell in safe_cells}
        assert '隔离验收仓库' not in {str(cell) for cell in safe_cells}
        assert number(next(cell for cell in safe_cells
                           if isinstance(cell, (int, float)) and not isinstance(cell, bool))) == 1500.0
    finally:
        workbook.close()
    with zipfile.ZipFile(path) as archive:
        parts = archive.namelist()
        assert 'xl/pivotTables/pivotTable1.xml' in parts
        assert 'xl/pivotTables/pivotTable2.xml' in parts
        assert 'refreshOnLoad' in archive.read('xl/pivotCache/pivotCacheDefinition1.xml').decode('utf-8')
    made = list(output_dir('exports').glob('库存单据明细表/2026年9月/*.xlsx'))
    assert made and made[0].name == '库存单据明细表_2026年9月.xlsx'
    expect(dialog).not_to_be_visible()
    page.screenshot(path=str(output('tail414-export.png', kind='screenshots')), full_page=True)
    # A month without a cache is refused with the fetch-first hint, not an empty file.
    for period in ('2026-08', '2026-10'):
        refused = page.request.get(server + f'/api/m08/inventory/export?period={period}').json()
        assert refused.get('ok') is False and ('请先抓取' in refused['message'] or '源文件不在本地' in refused['message'])
    # Cross-month / whole-year ranges keep the single-month hint and block the export.
    pick_option(page, '库存核算期间模式', '全部')
    expect(page.get_by_text('导出请选择单月', exact=True)).to_be_visible()
    expect(page.get_by_role('button', name='导出文件…', exact=True)).to_be_disabled()
    page.screenshot(path=str(output('tail414-export-hint.png', kind='screenshots')), full_page=True)
