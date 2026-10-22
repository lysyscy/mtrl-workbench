"""398 correction: remarks, representative undo recovery and real batch export."""
from __future__ import annotations
import io
import re
import zipfile
import openpyxl
from playwright.sync_api import expect
from test_preview import server, browser_page
from test_ledger import open_ledger, project_id, data, previous_report_month
from tools.audit.workspace_paths import output


def test_normal_month_remark_clear_undo_and_stale_token(server,browser_page):
    page,_=browser_page;open_ledger(page,server);project=project_id(page,server)
    row=page.locator('[row-id="material:month:2025-12"]');note=row.locator('[col-id="remark"] input')
    expect(row.locator('[col-id="payment"] input')).to_have_count(0)
    note.fill('正常月份中文备注');note.press('Tab')
    expect(page.get_by_role('status',name='保存回执')).to_contain_text('已保存')
    def normal(): return next(r for side in data(page,server,project)['sides'] if side['category']=='material' for r in side['rows'] if r['month']=='2025-12' and r['row_kind']=='month')
    assert normal()['remark']=='正常月份中文备注'
    page.get_by_role('button',name='撤销',exact=True).click();expect(page.get_by_role('status',name='保存回执')).to_contain_text('已撤销');assert normal()['remark']==''
    page.get_by_role('button',name='重做',exact=True).click();expect(page.get_by_role('status',name='保存回执')).to_contain_text('已重做');assert normal()['remark']=='正常月份中文备注'
    with page.expect_response(lambda response:'/m05/ledger/remark' in response.url and response.request.method=='POST') as saved:
        note.fill('');note.press('Tab')
    assert saved.value.json()['ok'];expect(note).to_have_value('')
    assert normal()['remark']==''
    pending=[];page.route('**/m05/ledger/remark',lambda route:pending.append(route))
    note.fill('本地未覆盖草稿');note.press('Tab');expect(page.get_by_role('status',name='保存回执')).to_contain_text('正在保存');assert pending
    assert 'token=' in pending[0].request.post_data
    changed=page.request.post(server+'/m05/ledger/remark',form={'project_id':project,'category':'material','ym':'2025-12','remark':'其他编辑已保存'})
    assert changed.json()['ok']
    pending[0].fulfill(response=pending[0].fetch())
    expect(page.get_by_role('status',name='保存回执')).to_contain_text('版本冲突');expect(note).to_have_value('本地未覆盖草稿')
    assert normal()['remark']=='其他编辑已保存'


def test_independent_multi_project_preview_cancel_and_actual_zip(server,browser_page):
    page,_=browser_page;open_ledger(page,server)
    page.get_by_role('button',name='逐月台账导出…',exact=True).click()
    picker=page.get_by_role('group',name='导出项目候选');expect(picker).to_be_visible()
    picker.get_by_role('button',name=re.compile('^☐ 长项目名验收')).click()
    expect(page.get_by_role('button',name='选择项目',exact=True)).to_contain_text('隔离项目')
    page.get_by_role('button',name='确认清单',exact=True).click()
    dialog=page.get_by_role('dialog');expect(dialog).to_contain_text('选了2个项目')
    expect(dialog.get_by_role('row')).to_have_count(3)
    dialog.get_by_role('button',name='项目排序',exact=True).click()
    dialog.get_by_role('button',name='取消',exact=True).click();expect(dialog).to_have_count(0)
    assert not list(page.context.pages)[0].is_closed()
    page.get_by_role('button',name='逐月台账导出…',exact=True).click()
    page.get_by_role('button',name='确认清单',exact=True).click()
    order=dialog.locator('tbody tr td:nth-child(2)').all_text_contents()
    with page.expect_download() as event: page.get_by_role('button',name='保存到本地',exact=True).click()
    download=event.value;path=output(download.suggested_filename,kind='exports');download.save_as(path)
    assert path.suffix=='.zip'
    with zipfile.ZipFile(path) as archive:
        assert len(archive.namelist())==2
        assert [next(disp for disp in order if disp in name) for name in archive.namelist()] == order
        assert any('隔离项目' in n for n in archive.namelist()) and any('长项目名验收' in n for n in archive.namelist())
        for name in archive.namelist():
            workbook=openpyxl.load_workbook(io.BytesIO(archive.read(name)),data_only=True)
            assert workbook.sheetnames==['辅材','机械','特殊费用月报表']
            disp='长项目名验收' if '长项目名验收' in name else '隔离项目'
            expected=data(page,server,project_id(page,server,disp),previous_report_month(page))
            for side in expected['sides']:
                exported=[(str(cells[1]),str(cells[5] or '')) for cells in workbook[side['label']].iter_rows(min_row=3,max_col=6,values_only=True) if re.fullmatch(r'\d{4}-\d{2}',str(cells[1] or ''))]
                assert exported==[(r['month'],r['remark']) for r in side['rows']]
            workbook.close()
    expect(page.get_by_role('button',name='选择项目',exact=True)).to_contain_text('隔离项目')


def test_add_focuses_new_record_at_fixed_historical_month_and_project_survives_navigation(server,browser_page):
    page,_=browser_page;open_ledger(page,server)
    project=project_id(page,server)
    page.get_by_role('region',name='辅材欠款台账',exact=True).get_by_role('columnheader',name='月份（按单据日期）',exact=True).click()
    page.get_by_role('button',name='新增历史行',exact=True).click()
    expect(page.get_by_role('status',name='保存回执')).to_contain_text('已新增历史行')
    page.wait_for_function("document.activeElement?.matches('[col-id=month] input')")
    rid=page.locator('[col-id=month] input:focus').evaluate("e=>Number(e.closest('[row-id]').getAttribute('row-id').split(':').at(-1))")
    actual=next(r for side in data(page,server,project)['sides'] for r in side['rows'] if r['record_id']==rid)
    assert actual['month']=='2025-11'
    draft=page.locator(f'[row-id="material:manual:{rid}"] [col-id=payment] input');draft.fill('bad');draft.press('Tab')
    expect(page.get_by_role('status',name='保存回执')).to_contain_text('保存被拒')
    page.get_by_role('button',name='选择项目',exact=True).click();page.get_by_role('group',name='查看项目候选').get_by_role('button',name=re.compile('^归档样例')).click()
    expect(page.get_by_role('button',name='选择项目',exact=True)).to_contain_text('归档')
    page.get_by_role('link',name='MOD-01 项目档案',exact=True).click();page.get_by_role('link',name='MOD-05 付款流水与欠款统计',exact=True).click()
    expect(page.get_by_role('button',name='选择项目',exact=True)).to_contain_text('归档样例')

    page.get_by_role('button',name='选择项目',exact=True).click();page.get_by_role('group',name='查看项目候选').get_by_role('button',name=re.compile('^隔离项目')).click()
    expect(draft).to_have_value('bad')


def test_restored_numeric_sort_does_not_use_lexical_amount_order(server,browser_page):
    page,_=browser_page;open_ledger(page,server)
    region=page.get_by_role('region',name='辅材欠款台账',exact=True)
    region.get_by_role('columnheader',name='付款',exact=True).click()
    expect(region.get_by_role('columnheader',name='付款',exact=True)).to_have_attribute('aria-sort','ascending')
    values=region.locator('[row-id]:not([row-id=total])').evaluate_all("""rows=>rows.sort((a,b)=>Number(a.getAttribute('row-index'))-Number(b.getAttribute('row-index'))).map(row=>{const cell=row.querySelector('[col-id=payment]');return Number((cell.querySelector('input')?.value||cell.textContent).replaceAll(',',''))})""")
    assert values==sorted(values),values


def test_uncertain_undo_requires_readback_before_another_history_action(server,browser_page):
    from test_ledger import add_row,cell,sync_ledger,wait_value
    page,_=browser_page;open_ledger(page,server);project=project_id(page,server)
    rid=add_row(page,server,project,'撤销响应不确定测试');sync_ledger(page)
    paid=cell(page,rid,'payment');paid.fill('78.90');paid.press('Tab');page.get_by_role('heading',name='逐月欠款台账',exact=True).click();wait_value(page,project,rid,'payment','78.90')
    requests=[]
    def uncertain(route):
        requests.append(route.request.post_data_json);response=route.fetch();assert response.json()['ok'];route.abort('failed')
    page.route('**/api/history/undo',uncertain)
    page.get_by_role('button',name='撤销',exact=True).click()
    expect(page.get_by_role('button',name='重新加载',exact=True)).to_be_visible()
    expect(page.get_by_role('button',name='撤销',exact=True)).to_be_disabled()
    page.keyboard.press('Control+z');assert len(requests)==1
    page.get_by_role('button',name='重新加载',exact=True).click()
    expect(page.get_by_role('button',name='重新加载',exact=True)).to_have_count(0)
    wait_value(page,project,rid,'payment','12.00');assert len(requests)==1
