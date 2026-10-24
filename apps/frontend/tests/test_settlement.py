"""TAIL-412: settlement month report, extract and balances through the real isolated backend.

Seed (fresh schema per run):
- 2025-12 收发结存明细：3 仓 4 行（隔离归档仓库一行故意不对平：10+1-5 != 7）。
- 库存原单入库：焊材（焊丝/焊条）、气体（二氧化碳/氧气/液氩）各月少量命中行，
  外加必须被规则排除的 焊条保温桶 / 氧气瓶；仓库取自项目档案（在建 2 仓 + 归档 1 仓）。
"""
import io
import json
import os
import subprocess
import sys
import zipfile
from datetime import datetime, timedelta, timezone
from decimal import Decimal

import httpx
import openpyxl
import pytest
from playwright.sync_api import expect
from test_preview import server, browser_page, login, output


def default_month() -> str:
    now = datetime.now(timezone(timedelta(hours=8)))
    return (now.replace(day=1) - timedelta(days=1)).strftime('%Y-%m')


DEFAULT_MONTH = default_month()
DEFAULT_YEAR = DEFAULT_MONTH[:4]
UNBALANCED_TOTALS = {'begin': '1510.11', 'in': '12596.69', 'out': '13150.72', 'end': '957.08'}


@pytest.fixture(scope='module')
def preview_seed():
    def seed(env):
        if os.environ.get('MTRL_SETL_TEST_BUILD'):
            env['FRONTEND_PREVIEW_DIR'] = os.environ['MTRL_SETL_TEST_BUILD']
        code = """
from app.core.db import session_scope
from app.mod01_projects.models import Project
from app.mod07_setl.models import SetlDetail
from app.mod08_inv.models import InvDetail
MONTH = '2025-12'
setl = [
    ('隔离归档仓库', 'WL-009', 10.00, 1.00, 5.00, 7.00),
    ('隔离验收仓库', 'WL-001', 1000.01, 250.02, 300.00, 950.03),
    ('隔离验收仓库', 'WL-002', 0.10, 0.00, 0.05, 0.05),
    ('长名称验收仓库', 'WL-001', 500.00, 12345.67, 12845.67, 0.00),
]
inv = [
    ('隔离验收仓库', '2024-06', '焊丝', 100.50),
    ('隔离验收仓库', '2025-12', '二氧化碳', 60.25),
    ('隔离验收仓库', '2025-12', '焊条', 200.75),
    ('隔离验收仓库', '2026-03', '焊丝', 11.11),
    ('隔离验收仓库', '2026-04', '氧气', 22.22),
    ('隔离验收仓库', '2025-12', '焊条保温桶', 999.99),
    ('隔离验收仓库', '2025-12', '氧气瓶', 888.88),
    ('隔离归档仓库', '2025-12', '焊丝', 33.33),
    ('长名称验收仓库', '2026-05', '液氩', 44.44),
]
with session_scope() as s:
    s.add_all([SetlDetail(ym=MONTH, warehouse_name=w, resource_code=c, resource_name='物料' + c,
                          begin_amt=b, in_amt=i, out_amt=o, end_amt=e)
               for w, c, b, i, o, e in setl])
    s.add_all([InvDetail(ym=m, warehouse_name=w, material_name=n, direction='入库', amount=a, qty=1)
               for w, m, n, a in inv])
"""
        subprocess.run([sys.executable, '-X', 'utf8', '-c', code], env=env, check=True)
    return seed


# ============================================================ helpers

def open_settlement(page, server):
    login(page, server)
    page.get_by_role('link', name='MOD-07 收发结存与提取').click()
    expect(page.get_by_role('heading', name='收发结存与提取')).to_be_visible()


def month_toolbar(page):
    return page.locator('.settlement-page > div:not([hidden]) .settlement-toolbar')


def pick_period(page, label, value):
    box = page.locator('label', has_text=label).filter(has=page.locator('.ant-picker'))
    field = box.locator('.ant-picker input')
    field.click()
    field.fill(value)
    field.press('Enter')


def pick_status(page, value):
    box = page.locator('label', has_text='项目状态').filter(has=page.locator('.ant-select'))
    box.locator('.ant-select').click()
    page.locator(f'.ant-select-item-option[title="{value}"]').click()


def switch_tab(page, label):
    page.locator('.settlement-tabs').get_by_text(label, exact=True).click()


def open_extract(page, server):
    open_settlement(page, server)
    switch_tab(page, '焊材 / 气体提取')
    expect(annual(page, '焊材')).to_be_visible()
    expect(year_field(page)).to_have_value(DEFAULT_YEAR)


def year_field(page):
    return page.get_by_role('textbox', name='提取年份')


def annual(page, kind):
    return page.locator('div:not([hidden]) section.settlement-annual[aria-label="%s年度汇总"]' % kind)


def grid_rows(scope):
    return scope.locator('.ag-grid-scrolling-container div[role="row"]')


def pinned_total_row(scope):
    return scope.locator('.ag-grid-pinned-bottom-rows-container div[role="row"]').first


def annual_rows(page, kind):
    return [row.locator('[col-id]').all_text_contents() for row in grid_rows(annual(page, kind)).all()]


def row_cell(scope, row_text, col_id):
    row = grid_rows(scope).filter(has_text=row_text).first
    return row.locator(f'.ag-cell[col-id="{col_id}"]')


def money_text(value: str) -> str:
    sign = '-' if value.startswith('-') else ''
    whole, _, part = value.lstrip('-').partition('.')
    grouped = f'{int(whole):,}' if whole else '0'
    return f'{sign}{grouped}.{part or "0"}'


def total_of(rows, field):
    return sum((Decimal(r[field]) for r in rows), Decimal('0'))


def extract_api(page, server, year='2025', status='在建'):
    return page.request.get(server + '/api/m07/extract', params={'year': year, 'status': status}).json()


def post_balance(page, server, payload):
    return page.request.post(server + '/api/m07/extract/balance',
                             data=json.dumps(payload), headers={'Content-Type': 'application/json'})


def set_balance(page, server, kind, warehouse, raw):
    body = post_balance(page, server, {'kind': kind, 'warehouse': warehouse, 'raw': raw}).json()
    assert body['ok'] is True, body
    return body


def clear_balances(page, server, warehouses=('隔离验收仓库', '隔离归档仓库', '长名称验收仓库')):
    for kind in ('焊材', '气体'):
        for warehouse in warehouses:
            body = post_balance(page, server, {'kind': kind, 'warehouse': warehouse, 'raw': ''}).json()
            assert body['ok'] is True, body


def refresh_extract_view(page):
    pick_period(page, '年份', '2026')
    pick_period(page, '年份', '2025')
    expect(year_field(page)).to_have_value('2025')
    expect(annual(page, '焊材')).to_be_visible()


def balance_input(page, kind, warehouse):
    return page.get_by_label(f'{kind} · {warehouse} · 2025-11及之前', exact=True)


def save_via_input(page, kind, warehouse, value):
    field = balance_input(page, kind, warehouse)
    field.fill(value)
    field.press('Enter')


def feedback(page, text=''):
    locator = page.locator('div:not([hidden]) p[role="status"]')
    return locator.filter(has_text=text) if text else locator


def page_alert(page, text=''):
    locator = page.locator('div:not([hidden]) p.notice[role="alert"]')
    return locator.filter(has_text=text) if text else locator.first


def late_extract_payload(year, kind, warehouse, prior, month_amount, month_index=11):
    """A valid-but-synthetic extract body; used to expose late/stale renders after a range switch."""
    months = [f'{year}-{m:02d}' for m in range(1, 13)]
    cells = [None] * 12
    cells[month_index] = month_amount
    total = str(Decimal(prior) + Decimal(month_amount))
    grid = {'kind': kind, 'year': year, 'months': months,
            'rows': [{'warehouse': warehouse, 'prior': prior, 'balance': None,
                      'cells': cells, 'total': total}],
            'sum_row': [prior] + [None] * 12 + [total]}
    other = {'kind': '气体' if kind == '焊材' else '焊材', 'year': year, 'months': months,
             'rows': [], 'sum_row': [None] * 14}
    grids = {'焊材': grid, '气体': other} if kind == '焊材' else {'焊材': other, '气体': grid}
    return {'annual': grids, 'balances': {k: {} for k in ('焊材', '气体')},
            'amounts': {}, 'annual_totals': {}, 'kinds': ['焊材', '气体']}


# ============================================================ 验收① 月报：行集合 / 顺序 / 分币合计 / 只读 / 切页保活

def test_month_report_rows_order_totals_and_readonly(server, browser_page):
    page, _ = browser_page
    open_settlement(page, server)
    # Default month has no data: empty state, export disabled.
    expect(page.get_by_text('该核算月没有收发结存数据', exact=True)).to_be_visible()
    expect(month_toolbar(page).get_by_role('button', name='导出文件…')).to_be_disabled()
    pick_period(page, '核算月', '2025-12')
    expect(page.locator('.settlement-summary')).to_contain_text('2025-12 · 3 仓 · 4 行 · 4 个文件')
    expect(page.locator('.settlement-summary')).to_contain_text('存在未对平')
    grid = page.locator('.settlement-page > div:not([hidden]) .settlement-grid')
    rows = grid_rows(grid)
    expect(rows).to_have_count(3)
    # Row order and values straight from the backend warehouse aggregation
    # (warehouse name ascending by codepoint: 长 U+957F < 隔 U+9694).
    api = page.request.get(server + '/api/m07/settlement', params={'period': '2025-12'}).json()
    assert api['balance_ok'] is False
    assert [r['name'] for r in api['rows']] == ['长名称验收仓库', '隔离归档仓库', '隔离验收仓库']
    assert rows.locator('[col-id="name"]').all_text_contents() == \
        ['长名称验收仓库', '隔离归档仓库', '隔离验收仓库']
    for index, row in enumerate(api['rows']):
        cells = rows.nth(index).locator('[col-id]')
        assert cells.nth(1).text_content() == str(row['rows'])
        for j, field in enumerate(['begin', 'in', 'out', 'end'], start=2):
            assert cells.nth(j).text_content() == money_text(row[field]), field
    assert rows.nth(1).locator('.status-tag').text_content() == '未对平'
    assert rows.nth(0).locator('.status-tag').text_content() == '已对平'
    assert rows.nth(2).locator('.status-tag').text_content() == '已对平'
    # Pinned total row: sums recomputed independently from API rows must equal the grid.
    total_row = pinned_total_row(grid)
    assert total_row.locator('[col-id="name"]').text_content() == '合计'
    assert total_row.locator('[col-id="rows"]').text_content() == str(sum(r['rows'] for r in api['rows']))
    for field in ['begin', 'in', 'out', 'end']:
        want = money_text(str(total_of(api['rows'], field)))
        assert total_row.locator(f'[col-id="{field}"]').text_content() == want, field
    page.screenshot(path=str(output('settlement-month-2025-12.png', kind='screenshots')), full_page=True)
    # The month report is fully read-only: no antd editors inside its grid.
    assert grid.locator('.ant-input, .ant-select').count() == 0
    # Tab keep-alive: the extract tab and back keep the month view mounted with its data.
    switch_tab(page, '焊材 / 气体提取')
    expect(annual(page, '焊材')).to_be_visible()
    switch_tab(page, '收发结存月报')
    expect(page.locator('.settlement-summary')).to_contain_text('2025-12 · 3 仓 · 4 行 · 4 个文件')
    expect(rows.nth(0).locator('[col-id="name"]')).to_have_text('长名称验收仓库')


# ============================================================ 验收① 提取：年度网格 / 行序 / 状态范围 / 历史列 / 规则排除

def test_extract_grid_order_status_scope_and_history_column(server, browser_page):
    page, _ = browser_page
    open_extract(page, server)
    clear_balances(page, server)
    pick_period(page, '年份', '2026')
    expect(year_field(page)).to_have_value(DEFAULT_YEAR)
    expect(grid_rows(annual(page, '焊材'))).to_have_count(1)
    expect(grid_rows(annual(page, '气体'))).to_have_count(2)
    # Default year/status: archive order, rule exclusions applied, no balance column in 2026.
    weld = annual_rows(page, '焊材')
    gas = annual_rows(page, '气体')
    assert [r[0] for r in weld] == ['隔离验收仓库'], weld
    assert [r[0] for r in gas] == ['隔离验收仓库', '长名称验收仓库'], gas
    # 2026 prior = everything before 2026-01 (2024-06 + 2025-12); empty months show 0.00.
    assert weld[0][1] == '301.25' and weld[0][4] == '11.11' and weld[0][14] == '11.11'
    assert gas[0][1] == '60.25' and gas[0][5] == '22.22' and gas[0][14] == '22.22'
    assert gas[1][1] == '0.00' and gas[1][6] == '44.44' and gas[1][14] == '44.44'
    assert annual(page, '焊材').locator('.ant-input').count() == 0
    assert annual(page, '气体').locator('.ant-input').count() == 0
    api = extract_api(page, server, DEFAULT_YEAR, '在建')
    assert api['balances'] == {'焊材': {}, '气体': {}}
    weld_sum = api['annual']['焊材']['sum_row']
    assert weld_sum[0] == '301.25' and weld_sum[3] == '11.11' and weld_sum[13] == '11.11'
    pick_period(page, '年份', '2025')
    expect(year_field(page)).to_have_value('2025')
    expect(grid_rows(annual(page, '焊材'))).to_have_count(1)
    # Rows are warehouses with data in ANY year (status-filtered), not per selected year.
    expect(grid_rows(annual(page, '气体'))).to_have_count(2)
    # 2025: history balance column appears; earlier months are blank, totals keep backend values.
    weld = annual_rows(page, '焊材')
    gas = annual_rows(page, '气体')
    blank11 = [''] * 11
    assert weld == [['隔离验收仓库', '100.50', *blank11, '200.75', '200.75', '']], weld
    assert gas == [['隔离验收仓库', '0.00', *blank11, '60.25', '60.25', ''],
                   ['长名称验收仓库', '0.00', *blank11, '0.00', '0.00', '']], gas
    expect(balance_input(page, '焊材', '隔离验收仓库')).to_have_value('')
    expect(balance_input(page, '气体', '隔离验收仓库')).to_have_value('')
    # Archived warehouses only appear under 归档 / 全部.
    pick_status(page, '归档')
    expect(grid_rows(annual(page, '焊材'))).to_have_count(1)
    assert [r[0] for r in annual_rows(page, '焊材')] == ['隔离归档仓库']
    assert annual_rows(page, '焊材')[0][13] == '33.33'
    expect(page.get_by_text('当前筛选条件下没有气体数据', exact=True)).to_be_visible()
    pick_status(page, '全部')
    expect(grid_rows(annual(page, '焊材'))).to_have_count(2)
    assert [r[0] for r in annual_rows(page, '焊材')] == ['隔离验收仓库', '隔离归档仓库']
    assert [r[0] for r in annual_rows(page, '气体')] == ['隔离验收仓库', '长名称验收仓库']
    api = extract_api(page, server, '2025', '全部')
    assert [r['warehouse'] for r in api['annual']['焊材']['rows']] == ['隔离验收仓库', '隔离归档仓库']
    page.screenshot(path=str(output('settlement-extract-2025-all.png', kind='screenshots')), full_page=True)


# ============================================================ 验收①②④ 结余保存重读 / 拒绝不部分写入 / API 一致

def test_balance_save_reread_reject_and_api_agreement(server, browser_page):
    page, _ = browser_page
    open_extract(page, server)
    clear_balances(page, server)
    pick_period(page, '年份', '2025')
    refresh_extract_view(page)
    field = balance_input(page, '焊材', '隔离验收仓库')
    expect(field).to_have_value('')
    # A real save goes through the backend and the re-read shows the stored value everywhere.
    save_via_input(page, '焊材', '隔离验收仓库', '1234.56')
    expect(feedback(page, '焊材 · 隔离验收仓库 · 已保存')).to_be_visible()
    api = extract_api(page, server, '2025', '在建')
    assert api['balances']['焊材']['隔离验收仓库'] == '1234.56'
    weld_row = api['annual']['焊材']['rows'][0]
    assert weld_row['balance'] == '1234.56' and weld_row['total'] == '1435.31'
    assert api['annual']['焊材']['sum_row'][13] == '1435.31'
    expect(field).to_have_value('1234.56')
    expect(row_cell(annual(page, '焊材'), '隔离验收仓库', 'total')).to_have_text('1,435.31')
    expect(pinned_total_row(annual(page, '焊材')).locator('[col-id="total"]')).to_have_text('1,435.31')
    # Saving again overwrites the same kind x warehouse cell.
    save_via_input(page, '焊材', '隔离验收仓库', '2')
    expect(feedback(page, '焊材 · 隔离验收仓库 · 已保存')).to_be_visible()
    assert extract_api(page, server)['balances']['焊材']['隔离验收仓库'] == '2.00'
    expect(field).to_have_value('2.00')
    expect(row_cell(annual(page, '焊材'), '隔离验收仓库', 'total')).to_have_text('202.75')
    # 0 is a legal entry: stored and shown, never confused with 未录入.
    save_via_input(page, '焊材', '隔离验收仓库', '0')
    expect(feedback(page, '焊材 · 隔离验收仓库 · 已保存')).to_be_visible()
    assert extract_api(page, server)['balances']['焊材']['隔离验收仓库'] == '0.00'
    expect(field).to_have_value('0.00')
    # Clearing removes the stored row: the API key disappears (未录入 != 0).
    save_via_input(page, '焊材', '隔离验收仓库', '')
    expect(feedback(page, '焊材 · 隔离验收仓库 · 已清空')).to_be_visible()
    balances = extract_api(page, server)['balances']
    assert '隔离验收仓库' not in balances['焊材']
    assert extract_api(page, server)['annual']['焊材']['rows'][0]['balance'] is None
    expect(field).to_have_value('')
    # Invalid input is refused server-side: alert, input kept, nothing written.
    save_via_input(page, '焊材', '隔离验收仓库', 'abc')
    alert = page_alert(page)
    expect(alert).to_contain_text('2025-11及之前金额要填数字（2 位小数 · 不得为负）；本次未保存，原值不变')
    expect(field).to_have_value('abc')
    draft = page.locator('.settlement-draft')
    expect(draft).to_contain_text('焊材 · 隔离验收仓库 · 2025-11及之前：abc')
    expect(draft.get_by_role('alert')).to_contain_text('要填数字')
    assert '隔离验收仓库' not in extract_api(page, server)['balances']['焊材']
    # The failing draft survives switching to the month tab and back (切页保活).
    switch_tab(page, '收发结存月报')
    expect(page.get_by_text('该核算月没有收发结存数据', exact=True)).to_be_visible()
    switch_tab(page, '焊材 / 气体提取')
    expect(page.locator('.settlement-draft')).to_contain_text('2025-11及之前：abc')
    expect(field).to_have_value('abc')
    # Cancelling drops the draft and restores the stored view without any write.
    draft.get_by_role('button', name='取消修改', exact=True).click()
    expect(page.locator('.settlement-draft')).to_have_count(0)
    expect(field).to_have_value('')
    save_via_input(page, '焊材', '隔离验收仓库', '-1.00')
    expect(page_alert(page, '不得为负')).to_be_visible()
    assert '隔离验收仓库' not in extract_api(page, server)['balances']['焊材']
    page.locator('.settlement-draft').get_by_role('button', name='取消修改', exact=True).click()
    # Direct API misuse is rejected with the documented reason and writes nothing.
    rejected = post_balance(page, server, {'kind': '电缆', 'warehouse': '隔离验收仓库', 'raw': '5'}).json()
    assert rejected['ok'] is False and rejected['message'] == '认不出的类别（只认 焊材 / 气体）'
    rejected = post_balance(page, server, {'kind': '焊材', 'warehouse': '   ', 'raw': '5'}).json()
    assert rejected['ok'] is False and rejected['message'] == '没找到这个仓库（键 = 仓库名）'
    assert extract_api(page, server)['balances'] == {'焊材': {}, '气体': {}}
    page.screenshot(path=str(output('settlement-balance-reject.png', kind='screenshots')), full_page=True)


# ============================================================ 验收③ 保存中切范围 + 迟到回执不覆盖 / 未核实锁定

def test_unknown_receipt_scope_switch_and_verification(server, browser_page):
    page, _ = browser_page
    open_extract(page, server)
    clear_balances(page, server)
    pick_period(page, '年份', '2025')
    refresh_extract_view(page)
    # A save left in flight while the range switches: the late OK receipt must not
    # push the old 2025 grid back over the current 2026 view, and the write still lands.
    holder = []
    page.route('**/api/m07/extract/balance', lambda route: holder.append(route))
    save_via_input(page, '焊材', '隔离验收仓库', '66.60')
    expect(feedback(page, '正在保存历史结余…')).to_be_visible()
    assert len(holder) == 1
    pick_period(page, '年份', '2026')
    expect(annual(page, '焊材')).to_be_visible()
    route = holder.pop()
    response = route.fetch()                 # release the real POST after the switch
    assert response.ok
    route.fulfill(response=response)
    expect(feedback(page, '焊材 · 隔离验收仓库 · 已保存')).to_be_visible()
    api = extract_api(page, server, '2026', '在建')
    assert api['balances']['焊材']['隔离验收仓库'] == '66.60'
    assert api['annual']['焊材']['rows'][0]['total'] == '77.71'
    expect(row_cell(annual(page, '焊材'), '隔离验收仓库', 'total')).to_have_text('77.71')
    assert annual_rows(page, '焊材')[0][0] == '隔离验收仓库'
    pick_period(page, '年份', '2025')
    expect(balance_input(page, '焊材', '隔离验收仓库')).to_have_value('66.60')
    # Reset for the damaged-receipt scenario.
    post_balance(page, server, {'kind': '焊材', 'warehouse': '隔离验收仓库', 'raw': ''})
    refresh_extract_view(page)
    expect(balance_input(page, '焊材', '隔离验收仓库')).to_have_value('')
    # A damaged HTTP-200 receipt must never offer a second submission while the real
    # write already landed; the panel locks until an explicit re-read and confirmation.
    writes = []

    def damaged(route):
        writes.append(route.request.post_data)
        response = route.fetch()
        assert response.ok and response.json()['ok'] is True
        route.fulfill(response=response, body='{')
    page.route('**/api/m07/extract/balance', damaged)
    save_via_input(page, '焊材', '隔离验收仓库', '77.70')
    alert = page_alert(page)
    expect(alert).to_contain_text('保存结果待核实')
    expect(alert).to_contain_text('勿重复提交')
    assert len(writes) == 1
    expect(balance_input(page, '焊材', '隔离验收仓库')).to_be_disabled()
    expect(month_toolbar(page).get_by_role('button', name='导出文件…')).to_be_disabled()
    assert extract_api(page, server)['balances']['焊材']['隔离验收仓库'] == '77.70'
    page.unroute('**/api/m07/extract/balance')
    page.get_by_role('button', name='重读核对', exact=True).click()
    expect(feedback(page, '已重读，未保存输入仍保留')).to_be_visible()
    expect(balance_input(page, '焊材', '隔离验收仓库')).to_have_value('77.70')
    page.get_by_role('button', name='已核对，保留输入继续', exact=True).click()
    expect(page.locator('.settlement-draft')).to_contain_text('焊材 · 隔离验收仓库 · 2025-11及之前：77.70')
    page.locator('.settlement-draft').get_by_role('button', name='保存这项', exact=True).click()
    expect(feedback(page, '焊材 · 隔离验收仓库 · 已保存')).to_be_visible()
    expect(page.locator('.settlement-draft')).to_have_count(0)
    assert extract_api(page, server)['balances']['焊材']['隔离验收仓库'] == '77.70'
    page.screenshot(path=str(output('settlement-balance-verified.png', kind='screenshots')), full_page=True)


# ============================================================ 验收③ 快速切范围：迟到响应不覆盖新范围

def test_fast_scope_switch_late_responses_do_not_overwrite(server, browser_page):
    page, _ = browser_page
    open_extract(page, server)
    clear_balances(page, server)
    # Extract: hang the 2025 request, land on 2026, then fulfill the stale response.
    pending = []

    def hang_2025(route):
        if 'year=2025' in route.request.url:
            pending.append(route)
        else:
            route.continue_()
    page.route('**/api/m07/extract?**', hang_2025)
    pick_period(page, '年份', '2025')
    expect(feedback(page, '正在读取库存年度汇总…')).to_be_visible()
    assert len(pending) == 1
    pick_period(page, '年份', '2026')
    expect(year_field(page)).to_have_value(DEFAULT_YEAR)
    expect(annual(page, '焊材')).to_be_visible()
    try:
        pending.pop().fulfill(json=late_extract_payload('2025', '焊材', '迟到2025仓', '1.00', '5.00'))
    except Exception:
        pass  # aborted request: the late payload must render nowhere either way
    assert grid_rows(annual(page, '焊材')).count() == 1
    expect(annual(page, '焊材')).not_to_contain_text('迟到2025仓')
    assert row_cell(annual(page, '焊材'), '隔离验收仓库', '2026-03').text_content() == '11.11'
    page.unroute('**/api/m07/extract?**')
    # Month report: same protection for a hung month response after switching twice.
    switch_tab(page, '收发结存月报')
    pending = []
    page.route('**/api/m07/settlement?**', lambda route: pending.append(route)
               if 'period=2025-12' in route.request.url else route.continue_())
    pick_period(page, '核算月', '2025-12')
    assert len(pending) == 1
    pick_period(page, '核算月', '2026-01')
    expect(page.get_by_text('该核算月没有收发结存数据', exact=True)).to_be_visible()
    synthetic = {'ym': '2025-12',
                 'rows': [{'id': '迟到月仓', 'name': '迟到月仓', 'rows': 9,
                           'begin': '999.99', 'in': '999.99', 'out': '999.99', 'end': '999.99'}],
                 'counts': {'warehouses': 1, 'rows': 9, 'files': 2}, 'balance_ok': True}
    try:
        pending.pop().fulfill(json=synthetic)
    except Exception:
        pass  # aborted request: the late payload must render nowhere either way
    expect(page.get_by_text('该核算月没有收发结存数据', exact=True)).to_be_visible()
    expect(page.locator('.settlement-summary')).to_contain_text('2026-01 · 0 仓 · 0 行 · 1 个文件')
    expect(page.locator('.settlement-page')).not_to_contain_text('迟到月仓')
    page.unroute('**/api/m07/settlement?**')
    page.screenshot(path=str(output('settlement-late-response-guard.png', kind='screenshots')), full_page=True)


# ============================================================ 验收③ 导出窗口：范围变化取消迟到下载

def test_export_dialog_cancels_late_download_on_scope_change(server, browser_page):
    page, _ = browser_page
    open_settlement(page, server)
    pick_period(page, '核算月', '2025-12')
    expect(page.locator('.settlement-summary')).to_contain_text('2025-12 · 3 仓')
    pending = []
    page.route('**/api/m07/settlement/export?**', lambda route: pending.append(route))
    downloads = []

    def on_download(download):
        downloads.append(download)
    page.on('download', on_download)
    month_toolbar(page).get_by_role('button', name='导出文件…').click()
    dialog = page.get_by_role('dialog')
    expect(dialog).to_contain_text('3 个分仓库工作簿 ＋ 1 个总表，打包 ZIP。')
    dialog.get_by_role('button', name='保存到本地', exact=True).click()
    assert len(pending) == 1
    # Leaving the page pauses the account tree: the export aborts and the dialog closes.
    # (The busy dialog's modal overlay swallows real pointer events; a JS click still
    # goes through React Router's client-side navigation.)
    page.evaluate("document.querySelector('a[href=\"/preview/projects\"]').click()")
    expect(page.get_by_role('heading', name='项目档案')).to_be_visible()
    try:
        pending.pop().fulfill(body=b'PK\x05\x06' + bytes(18), content_type='application/zip')
    except Exception:
        pass  # aborted: the late file must not open a download either way
    page.wait_for_timeout(600)
    assert downloads == []
    page.unroute('**/api/m07/settlement/export?**')
    # Returning re-enables the page with its scope intact.
    page.get_by_role('link', name='MOD-07 收发结存与提取').click()
    expect(page.locator('.settlement-summary')).to_contain_text('2025-12 · 3 仓 · 4 行 · 4 个文件')
    expect(month_toolbar(page).get_by_role('button', name='导出文件…')).to_be_enabled()


# ============================================================ 验收⑤ 未登录：页面只读 + API 直发 401 且数据不变

def test_unauthenticated_page_readonly_and_api_401(server, browser_page):
    page, _ = browser_page
    page.goto(server + '/preview/settlement')
    expect(page.get_by_text('未登录 · 只读', exact=True)).to_be_visible()
    expect(page.get_by_text('该核算月没有收发结存数据', exact=True)).to_be_visible()
    switch_tab(page, '焊材 / 气体提取')
    pick_period(page, '年份', '2025')
    expect(year_field(page)).to_have_value('2025')
    expect(grid_rows(annual(page, '焊材'))).to_have_count(1)
    expect(balance_input(page, '焊材', '隔离验收仓库')).to_be_disabled()
    expect(balance_input(page, '气体', '隔离验收仓库')).to_be_disabled()
    before = httpx.get(server + '/api/m07/extract', params={'year': '2025', 'status': '在建'}).json()['balances']
    with httpx.Client(base_url=server) as client:
        response = client.post('/api/m07/extract/balance',
                               json={'kind': '焊材', 'warehouse': '隔离验收仓库', 'raw': '1.00'})
        assert response.status_code == 401
        assert '未登录' in response.json()['message']
        client.cookies.set('mw_preview_397', 'invalid')
        assert client.post('/api/m07/extract/balance',
                           json={'kind': '焊材', 'warehouse': '隔离验收仓库', 'raw': '1.00'}).status_code == 401
    after = httpx.get(server + '/api/m07/extract', params={'year': '2025', 'status': '在建'}).json()['balances']
    assert before == after
    page.screenshot(path=str(output('settlement-unauthenticated.png', kind='screenshots')), full_page=True)


# ============================================================ 验收④ 导出实际文件：月报 zip

SETL_HEADER = ['仓库', '资源分类', '资源编码', '资源名称', '规格型号', '计量单位', '物料性质',
               '期初数量', '期初金额', '收数量', '收单价', '收金额',
               '发数量', '发单价', '发金额', '结存数量', '结存金额']


def test_export_month_zip_actual_file_bytes(server, browser_page):
    page, _ = browser_page
    open_settlement(page, server)
    pick_period(page, '核算月', '2025-12')
    expect(page.locator('.settlement-summary')).to_contain_text('2025-12 · 3 仓 · 4 行 · 4 个文件')
    month_toolbar(page).get_by_role('button', name='导出文件…').click()
    dialog = page.get_by_role('dialog')
    expect(dialog).to_contain_text('3 个分仓库工作簿 ＋ 1 个总表，打包 ZIP。')
    expect(dialog.locator('li')).to_have_count(3)
    with page.expect_download() as event:
        dialog.get_by_role('button', name='保存到本地', exact=True).click()
    download = event.value
    assert download.suggested_filename == '第四工程公司2025年12月收发结存表.zip'
    path = output(download.suggested_filename, kind='exports')
    download.save_as(path)
    with zipfile.ZipFile(path) as archive:
        assert set(archive.namelist()) == {'隔离归档仓库.xlsx', '隔离验收仓库.xlsx',
                                           '长名称验收仓库.xlsx', '第四工程公司2025年12月收发结存表.xlsx'}

        def sheet(name, member):
            wb = openpyxl.load_workbook(io.BytesIO(archive.read(member)), data_only=True)
            try:
                return list(wb[name].values)
            finally:
                wb.close()
        per_warehouse = sheet('Sheet1', '隔离验收仓库.xlsx')
        assert list(per_warehouse[0]) == SETL_HEADER
        assert per_warehouse[1][2] == 'WL-001' and per_warehouse[2][2] == 'WL-002'
        assert per_warehouse[1][8] == pytest.approx(1000.01)
        assert per_warehouse[1][11] == pytest.approx(250.02)
        assert per_warehouse[1][14] == pytest.approx(300.00)
        assert per_warehouse[1][16] == pytest.approx(950.03)
        assert per_warehouse[2][8] == pytest.approx(0.10)
        assert per_warehouse[2][16] == pytest.approx(0.05)
        archived = sheet('Sheet1', '隔离归档仓库.xlsx')
        assert archived[1][2] == 'WL-009' and archived[1][16] == pytest.approx(7.00)
        report = sheet('报表1', '第四工程公司2025年12月收发结存表.xlsx')
        assert list(report[0]) == SETL_HEADER
        assert len(report) == 5
        assert [(r[0], r[2]) for r in report[1:]] == [
            ('长名称验收仓库', 'WL-001'), ('隔离归档仓库', 'WL-009'),
            ('隔离验收仓库', 'WL-001'), ('隔离验收仓库', 'WL-002')]
        assert report[2][16] == pytest.approx(7.00)
        assert report[1][11] == pytest.approx(12345.67)
        pivot = sheet('Sheet1', '第四工程公司2025年12月收发结存表.xlsx')
        assert pivot[3][0] == '长名称验收仓库'
        assert [pivot[3][i] for i in (1, 2, 3, 4)] == pytest.approx([500.00, 12345.67, 12845.67, 0.00])
        assert pivot[4][0] == '隔离归档仓库'
        assert [pivot[4][i] for i in (1, 2, 3, 4)] == pytest.approx([10.00, 1.00, 5.00, 7.00])
        assert pivot[5][0] == '隔离验收仓库'
        assert [pivot[5][i] for i in (1, 2, 3, 4)] == pytest.approx([1000.11, 250.02, 300.05, 950.08])
        assert pivot[6][0] == '总计'
        assert [pivot[6][i] for i in (1, 2, 3, 4)] == pytest.approx(
            [float(UNBALANCED_TOTALS[k]) for k in ('begin', 'in', 'out', 'end')])
        assert len(pivot) == 7


# ============================================================ 验收④ 导出实际文件：提取 xlsx（原表 + 年度表）

def test_export_extract_xlsx_actual_file_bytes(server, browser_page):
    page, _ = browser_page
    open_extract(page, server)
    clear_balances(page, server)
    pick_period(page, '年份', '2025')
    set_balance(page, server, '焊材', '隔离验收仓库', '1234.56')
    set_balance(page, server, '气体', '隔离验收仓库', '0')
    set_balance(page, server, '焊材', '隔离归档仓库', '55.55')
    refresh_extract_view(page)
    expect(balance_input(page, '焊材', '隔离验收仓库')).to_have_value('1234.56')
    toolbar = month_toolbar(page)
    expect(grid_rows(annual(page, '焊材'))).to_have_count(1)
    toolbar.get_by_role('button', name='导出文件…').click()
    dialog = page.get_by_role('dialog')
    expect(dialog).to_contain_text('2025 · 在建 · 焊材 1 仓 / 气体 2 仓')
    with page.expect_download() as event:
        dialog.get_by_role('button', name='保存到本地', exact=True).click()
    download = event.value
    assert download.suggested_filename == '焊材气体提取表_2025年12月.xlsx'
    path = output(download.suggested_filename, kind='exports')
    download.save_as(path)
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        assert wb.sheetnames == ['焊材', '气', '焊材年度汇总', '气体年度汇总']
        weld = list(wb['焊材'].values)
        assert len(weld) == 3
        assert weld[1][0] == '隔离验收仓库'
        assert weld[1][1] == pytest.approx(1234.56)
        assert all(weld[1][c] in (None, '') for c in range(2, 13))
        assert weld[1][13] == pytest.approx(200.75)
        assert weld[1][14] == pytest.approx(1435.31)
        assert weld[2][13] == pytest.approx(200.75)
        assert weld[2][14] == pytest.approx(1435.31)
        assert weld[2][1] in (None, '')
        gas = list(wb['气'].values)
        assert gas[1][0] == '隔离验收仓库'
        assert isinstance(gas[1][1], (int, float)) and gas[1][1] == pytest.approx(0.00)
        assert gas[1][13] == pytest.approx(60.25)
        assert gas[1][14] == pytest.approx(60.25)
        weld_annual = list(wb['焊材年度汇总'].values)
        assert weld_annual[0] == ('项目', '往年金额', *[f'2025-{m:02d}' for m in range(1, 13)],
                                  '合计', '2025-11及之前')
        assert weld_annual[1][0] == '隔离验收仓库'
        assert weld_annual[1][1] == pytest.approx(100.50)
        assert weld_annual[1][13] == pytest.approx(200.75)
        assert weld_annual[1][14] == pytest.approx(1435.31)
        assert weld_annual[1][15] == pytest.approx(1234.56)
        assert weld_annual[2][0] == '合计'
        assert weld_annual[2][14] == pytest.approx(1435.31)
        assert weld_annual[2][15] in (None, '')
        gas_annual = list(wb['气体年度汇总'].values)
        assert gas_annual[1][1] == pytest.approx(0.00)
        assert gas_annual[1][13] == pytest.approx(60.25)
        assert gas_annual[1][15] == pytest.approx(0.00)
    finally:
        wb.close()
    page.screenshot(path=str(output('settlement-export-extract.png', kind='screenshots')), full_page=True)
    # Widening to 全部 pulls the archived warehouse (and its balance) into the same template.
    pick_status(page, '全部')
    expect(grid_rows(annual(page, '焊材'))).to_have_count(2)
    toolbar.get_by_role('button', name='导出文件…').click()
    dialog = page.get_by_role('dialog')
    expect(dialog).to_contain_text('2025 · 全部 · 焊材 2 仓 / 气体 2 仓')
    with page.expect_download() as event:
        dialog.get_by_role('button', name='保存到本地', exact=True).click()
    download = event.value
    assert download.suggested_filename == '焊材气体提取表_2025年12月.xlsx'
    path = output('全部-' + download.suggested_filename, kind='exports')
    download.save_as(path)
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        weld = list(wb['焊材'].values)
        assert [weld[1][0], weld[2][0]] == ['隔离验收仓库', '隔离归档仓库']
        assert weld[2][1] == pytest.approx(55.55)
        assert weld[2][13] == pytest.approx(33.33)
        assert weld[2][14] == pytest.approx(88.88)
        weld_annual = list(wb['焊材年度汇总'].values)
        assert [weld_annual[1][0], weld_annual[2][0]] == ['隔离验收仓库', '隔离归档仓库']
        assert weld_annual[2][1] == pytest.approx(0.00)
        assert weld_annual[2][15] == pytest.approx(55.55)
    finally:
        wb.close()
    # A month without settlement details must answer the documented JSON-200 refusal.
    response = page.request.get(server + '/api/m07/settlement/export', params={'period': '2026-01'})
    assert response.status == 200
    assert response.headers['content-type'].startswith('application/json')
    body = response.json()
    assert body['ok'] is False and '没有任何收发结存明细' in body['message']
    # Cleanup so the module leaves no balance rows behind.
    clear_balances(page, server)
