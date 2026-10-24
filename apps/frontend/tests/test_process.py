"""TAIL-410: process reports maintenance and export through the real isolated backend."""
import io
import re
import shutil
import subprocess
import sys
import urllib.parse
import zipfile
from datetime import datetime
from zoneinfo import ZoneInfo

import httpx
import openpyxl
import pytest
from playwright.sync_api import expect
from preview_select import pick_option
from test_preview import server, browser_page, login
from tools.audit.workspace_paths import output

WH_A = '隔离验收仓库'
WH_B = '长名称验收仓库'
DISP_A, DISP_B = '隔离项目', '长项目名验收'
CATEGORIES = ['主辅材', '工具', '劳保', '周转材料', '临时设施']
MATERIAL_KEYS = [f'{direction}_{cat}' for direction in ('in', 'out') for cat in CATEGORIES]
DATA_KEYS = MATERIAL_KEYS + [f'mech_{i}' for i in range(6)]
MECH_LABELS = ['租分公司设备', '内租设备', '外租零租含税', '外租包月含税', '外租零租不含税', '外租包月不含税']
# 模板列序（唯一出口 = process_export.COL_OF_KEY：临时设施 在 周转材料 前，与页面列序刻意不同）
TEMPLATE_LETTERS = {'in_主辅材': 'C', 'in_工具': 'D', 'in_劳保': 'E', 'in_临时设施': 'F', 'in_周转材料': 'G',
                    'out_主辅材': 'H', 'out_工具': 'I', 'out_劳保': 'J', 'out_临时设施': 'K', 'out_周转材料': 'L',
                    'mech_0': 'M', 'mech_1': 'N', 'mech_2': 'O', 'mech_3': 'P', 'mech_4': 'Q', 'mech_5': 'R'}
SEGMENT_ANCHORS = (('C', 3, 7), ('H', 8, 12), ('M', 13, 16), ('Q', 17, 18))

SEED = """
from datetime import datetime
from decimal import Decimal
from sqlalchemy import select
from app.core.db import session_scope
from app.mod01_projects.models import Project
from app.mod06_proc.models import ProcMechManual
from app.mod08_inv.models import InvAdjustment, InvDetail
with session_scope() as s:
    projects = {p.warehouse_name: p for p in s.execute(select(Project)).scalars()}
    a, b = projects['隔离验收仓库'], projects['长名称验收仓库']
    rows = [
        InvDetail(ym='2026-09', project_id=a.id, warehouse_name='隔离验收仓库', bill_name_full=a.name_full,
                  doc_no='INV-A1', seq='1', direction='入库', material_name='螺纹钢', material_property='主材',
                  unit='吨', qty=10, amount=Decimal('1000.50')),
        InvDetail(ym='2026-09', project_id=a.id, warehouse_name='隔离验收仓库', doc_no='INV-A2', seq='2',
                  direction='入库', material_name='电钻', material_property='工具', unit='台', qty=2,
                  amount=Decimal('200.25')),
        InvDetail(ym='2026-09', project_id=a.id, warehouse_name='隔离验收仓库', doc_no='INV-A4', seq='4',
                  direction='入库', material_name='铁丝', material_property='辅材', unit='卷', qty=5,
                  amount=Decimal('50.25')),
        InvDetail(ym='2026-09', project_id=a.id, warehouse_name='隔离验收仓库', doc_no='INV-A3', seq='3',
                  direction='出库', material_name='水泥', material_property='辅材', unit='袋', qty=30,
                  amount=Decimal('300.75')),
        InvDetail(ym='2026-10', project_id=a.id, warehouse_name='隔离验收仓库', doc_no='INV-B1', seq='1',
                  direction='入库', material_name='手套', material_property='劳保', unit='副', qty=40,
                  amount=Decimal('400.00')),
        InvDetail(ym='2026-10', project_id=a.id, warehouse_name='隔离验收仓库', doc_no='INV-B2', seq='2',
                  direction='出库', material_name='模板', material_property='周转材料', unit='块', qty=5,
                  amount=Decimal('500.50')),
        InvDetail(ym='2026-10', project_id=a.id, warehouse_name='隔离验收仓库', doc_no='INV-B3', seq='3',
                  direction='出库', material_name='活动房', material_property='临时设施', unit='间', qty=1,
                  amount=Decimal('60.00')),
    ]
    s.add_all(rows)
    s.flush()
    s.add(InvAdjustment(row_id=rows[0].id, before_amount=Decimal('1000.50'), amount=Decimal('1100.00'),
                        remark='隔离调整样例', operator='预览管理员', changed_at=datetime(2026, 10, 15, 8, 30)))
    s.add(ProcMechManual(ym='2026-09', project_id=b.id, warehouse_name='长名称验收仓库',
                         mech_0=Decimal('111.00'), mech_1=Decimal('222.00'), mech_2=Decimal('333.00'),
                         mech_3=Decimal('44.00'), mech_4=Decimal('55.00'), mech_5=Decimal('66.00'),
                         source_batches={'__detail_sources': {'mech_2': {'mode': 'detail', 'basis': {}, 'history': []}}}))
"""


@pytest.fixture(scope='module')
def preview_seed(make_preview_seed):
    def seed(env):
        subprocess.run([sys.executable, '-X', 'utf8', '-c', SEED], env=env, check=True)
    return make_preview_seed(seed)


def fmt(value):
    """后端矩阵值（数字 / None）→ 页面 money() 显示文本（空态留空、千分位）。"""
    if value is None or value == '':
        return ''
    return f'{float(value):,.2f}'


def js(value):
    """后端数字 → JS 输入框显示的字符串（111.0 → '111'，77.7 → '77.7'）。"""
    text = repr(float(value))
    return text[:-2] if text.endswith('.0') else text


def matrix_api(page, server):
    data = page.request.get(server + '/api/m06/matrix').json()
    assert data['ok'] is True
    return data['matrices']


def pick_project(page, disp):
    pick_option(page, '过程资料项目', disp)


def col(letter):
    return openpyxl.utils.column_index_from_string(letter)


def cell_value(ws, row, letter):
    return ws.cell(row=row, column=col(letter)).value


def expect_rows(page, matrix, kind):
    """逐数据行核对核算月顺序与材料 10 列（机械格是输入框，另核）。"""
    rows = matrix[kind]
    tbody = page.locator('.process-matrix').locator('tbody')
    for i, api_row in enumerate(rows):
        tr = tbody.locator('tr').nth(i)
        expect(tr.locator('th')).to_have_text(api_row['ym'])
        for j, key in enumerate(MATERIAL_KEYS):
            expect(tr.locator('td').nth(j)).to_have_text(fmt(api_row[key]))


def expect_totals(page, matrix, kind):
    """合计行（16 列逐列）与总求和行（四段）对照后端矩阵。"""
    totals = matrix['totals'][kind]
    row = page.locator('.process-matrix').locator('tfoot tr').nth(0)
    for j, key in enumerate(DATA_KEYS):
        expect(row.locator('td').nth(j)).to_have_text(fmt(totals.get(key)))
    names = ['入库合计', '出库合计', '机械合计', '不含税合计']
    seg_row = page.locator('.process-matrix').locator('tfoot tr').nth(1)
    for i, (name, seg) in enumerate(zip(names, matrix['segments'][kind])):
        expect(seg_row.locator('td').nth(i)).to_have_text(f'{name} {fmt(seg["total"])}')


def check_workbook(path, *, display, rows):
    """下载的真实字节 → openpyxl：标题 / 顺序 / 金额 / 零值留空 / 小计 / 四段合计 / 模板尾块。"""
    wb = openpyxl.load_workbook(str(path), data_only=True)
    try:
        ws = wb['汇总表']
        assert ws['A1'].value == f'{display}材料机械消耗汇总表', ws['A1'].value
        n = len(rows)
        for i, row in enumerate(rows):
            r = 4 + i
            assert cell_value(ws, r, 'A') == i + 1, (r, '序号')
            assert cell_value(ws, r, 'B') == row['ym'], (r, '核算月')
            for key, letter in TEMPLATE_LETTERS.items():
                got, want = cell_value(ws, r, letter), row.get(key)
                if not want:
                    assert got in (None, ''), (r, letter, got, want)
                else:
                    assert round(float(got), 2) == round(float(want), 2), (r, letter, got, want)
        sub_r, tot_r = 4 + n, 5 + n
        assert cell_value(ws, sub_r, 'B') == '小计'
        assert cell_value(ws, tot_r, 'B') == '合计：'
        for key, letter in TEMPLATE_LETTERS.items():
            values = [row.get(key) for row in rows if row.get(key) is not None]
            want = round(sum(map(float, values)), 2) if values else None
            got = cell_value(ws, sub_r, letter)
            if want:
                assert round(float(got), 2) == want, (letter, got, want)
            else:
                assert got in (None, ''), (letter, got)
        for anchor, c1, c2 in SEGMENT_ANCHORS:
            seg = round(sum(float(cell_value(ws, sub_r, openpyxl.utils.get_column_letter(c)) or 0)
                            for c in range(c1, c2 + 1)), 2)
            got = cell_value(ws, tot_r, anchor)
            if seg:
                assert round(float(got), 2) == seg, (anchor, got, seg)
            else:
                assert got in (None, ''), (anchor, got)
    finally:
        wb.close()
    wb = openpyxl.load_workbook(str(path))
    try:
        ws = wb['汇总表']
        n = len(rows)
        assert ws.max_row == n + 7, ws.max_row
        assert str(ws.cell(row=n + 6, column=1).value or '').startswith('1、材料费用')
        for anchor, _c1, _c2 in SEGMENT_ANCHORS:
            formula = ws.cell(row=n + 5, column=col(anchor)).value
            assert isinstance(formula, str) and formula.startswith('=IF('), (anchor, formula)
        assert ws.cell(row=4, column=2).number_format == '@'
    finally:
        wb.close()


def open_process(page, server):
    login(page, server)
    page.goto(server + '/preview/process')
    expect(page.get_by_role('heading', name='过程资料')).to_be_visible()
    expect(page.locator('.process-matrix')).to_be_visible()


def test_tabs_rows_and_totals_match_backend(server, browser_page):
    page, _ = browser_page
    open_process(page, server)
    matrices = matrix_api(page, server)
    assert sorted(matrices) == [WH_B, WH_A]
    # 默认项目 = 候选首位（长名称验收仓库）：只有机械费、无库存的月份，材料 10 列全部保持空态。
    default = matrices[WH_B]
    assert [r['ym'] for r in default['adjusted']] == ['2026-09']
    first_row = page.locator('.process-matrix').locator('tbody tr').nth(0)
    for j, key in enumerate(MATERIAL_KEYS):
        expect(first_row.locator('td').nth(j)).to_have_text('')
    for i, label in enumerate(MECH_LABELS):
        if i == 2:
            continue  # 明细确认格是只读展示，不属于输入框
        expect(page.get_by_label(f'2026-09 {label}')).to_have_value(js(default['adjusted'][0][f'mech_{i}']))
    confirmed = page.locator('.process-cell').filter(has_text='明细确认')
    expect(confirmed).to_have_count(1)
    expect(confirmed).to_contain_text('333.00')
    expect_totals(page, default, 'adjusted')
    # 切到隔离项目：调整后表逐行逐列、逐列合计与四段总求和全部对上后端矩阵。
    pick_project(page, DISP_A)
    current = matrices[WH_A]
    # 物料性质归集：主材+辅材同向合并进「入库主辅材」，其余四类各落其列。
    assert current['raw'][0]['in_主辅材'] == 1050.75 and current['adjusted'][0]['in_主辅材'] == 1150.25
    expect_rows(page, current, 'adjusted')
    expect_totals(page, current, 'adjusted')
    adjusted_cell = page.locator('.process-matrix tbody tr').nth(0).locator('td.process-adjusted')
    expect(adjusted_cell).to_have_count(1)
    expect(adjusted_cell).to_have_text(fmt(current['adjusted'][0]['in_主辅材']))
    # 表① 原始表：同一格回到 i8 原始金额。
    page.locator('.ant-segmented-item').filter(has_text='表① 原始表').click()
    expect_rows(page, current, 'raw')
    expect_totals(page, current, 'raw')
    expect(page.locator('.process-matrix .process-adjusted')).to_have_count(0)
    # 表③ 差异表：单据明细 + 三列合计 + 稳定库存行 ID 定位入口。
    page.locator('.ant-segmented-item').filter(has_text='差异表').click()
    grid = page.get_by_role('grid')
    expect(grid.get_by_role('gridcell', name='INV-A1')).to_be_visible()
    diff_row = grid.locator('[role="row"]').filter(has_text='螺纹钢')
    expect(diff_row).to_contain_text('隔离验收仓库')
    expect(diff_row).to_contain_text('1,000.50')
    expect(diff_row).to_contain_text('1,100.00')
    expect(diff_row).to_contain_text('99.50')
    expect(diff_row).to_contain_text('隔离调整样例')
    expect(diff_row).to_contain_text('预览管理员')
    expect(diff_row).to_contain_text('2026-10-15 08:30')
    expect(grid.get_by_role('link', name='定位原行')).to_have_attribute('href', re.compile(r'/inventory\?row=\d+$'))
    expect(grid.locator('[role="row"]').filter(has_text='合计')).to_contain_text('99.50')
    page.screenshot(path=str(output('tail410-diff-tab.png', kind='screenshots')), full_page=True)


def test_mech_cell_save_reread_keeps_later_input(server, browser_page):
    page, _ = browser_page
    open_process(page, server)
    pick_project(page, DISP_A)
    pending = []

    def hold_post(route):
        if route.request.method == 'POST':
            pending.append(route)
        else:
            route.continue_()

    page.route('**/api/m06/mech', hold_post)
    mech0 = page.get_by_label('2026-09 租分公司设备')
    mech0.fill('123.45')
    mech0.blur()
    # 第一格在途时继续输入下一格：保存后的重读不得覆盖后续输入。
    mech3 = page.get_by_label('2026-09 外租包月含税')
    mech3.fill('300')
    assert len(pending) == 1
    post_body = pending[0].request.post_data or ''
    assert 'mech_0' in post_body and '123.45' in post_body and WH_A in post_body
    pending.pop().continue_()
    expect(page.get_by_role('status').filter(has_text='租分公司设备已保存')).to_be_visible()
    expect(mech3).to_have_value('300')
    mech = page.request.get(server + '/api/m06/mech').json()
    assert mech['mech'][WH_A]['2026-09'][:4] == [123.45, 0, 0, 0]
    page.unroute('**/api/m06/mech')
    # 第二格真实提交落库，同一月两格互不覆盖。
    mech3.blur()
    expect(page.get_by_role('status').filter(has_text='外租包月含税已保存')).to_be_visible()
    mech = page.request.get(server + '/api/m06/mech').json()
    assert mech['mech'][WH_A]['2026-09'][:4] == [123.45, 0, 0, 300]
    row = matrix_api(page, server)[WH_A]['adjusted'][0]
    assert (row['mech_0'], row['mech_3']) == (123.45, 300)
    # 清空 = 显式存 0：机械 0 显示 0.00（材料空态才是空白）。
    mech3.fill('')
    mech3.blur()
    expect(page.get_by_role('status').filter(has_text='外租包月含税已保存')).to_be_visible()
    expect(mech3).to_have_value('0')
    mech = page.request.get(server + '/api/m06/mech').json()
    assert mech['mech'][WH_A]['2026-09'] == [123.45, 0, 0, 0, 0, 0]


def test_invalid_and_negative_amounts_rejected_input_kept(server, browser_page):
    page, _ = browser_page
    open_process(page, server)
    mech0 = page.get_by_label('2026-09 租分公司设备')
    expect(mech0).to_have_value('111')
    before = page.request.get(server + '/api/m06/mech').json()
    mech0.fill('abc')
    mech0.blur()
    alert = page.get_by_role('alert').first
    expect(alert).to_contain_text('「租分公司设备」不是数字')
    expect(mech0).to_have_value('abc')
    expect(page.get_by_text('1 项未保存')).to_be_visible()
    mech0.fill('-5')
    mech0.blur()
    expect(alert).to_contain_text('「租分公司设备」不得为负数（当前 -5.00）')
    expect(mech0).to_have_value('-5')
    after = page.request.get(server + '/api/m06/mech').json()
    assert after['mech'] == before['mech'] and after['count'] == before['count']
    # 修正后重试：只写这一格，其余列不动。
    mech0.fill('150')
    mech0.press('Enter')
    expect(page.get_by_role('status').filter(has_text='租分公司设备已保存')).to_be_visible()
    expect(page.get_by_text('1 项未保存')).to_have_count(0)
    mech = page.request.get(server + '/api/m06/mech').json()
    assert mech['mech'][WH_B]['2026-09'] == [150, 222, 333, 44, 55, 66]


def test_draft_survives_navigation_and_stays_with_project(server, browser_page):
    page, _ = browser_page
    open_process(page, server)
    mech1 = page.get_by_label('2026-09 内租设备')
    mech1.fill('abc')
    mech1.blur()
    expect(page.get_by_role('alert').first).to_contain_text('「内租设备」不是数字')
    expect(page.get_by_text('1 项未保存')).to_be_visible()
    # 切到项目档案再回来：被拒草稿（含输入值）留在账号内存。
    page.get_by_role('link', name='MOD-01 项目档案').click()
    expect(page.get_by_role('heading', name='项目档案')).to_be_visible()
    page.get_by_role('link', name='MOD-06 过程资料').click()
    expect(mech1).to_have_value('abc')
    expect(page.get_by_role('alert').first).to_contain_text('「内租设备」不是数字')
    # 拒绝留在所属项目：切走不串显，切回仍在。
    pick_project(page, DISP_A)
    expect(page.locator('.process-matrix')).to_be_visible()
    expect(page.get_by_role('alert')).to_have_count(0)
    expect(page.get_by_text('1 项未保存')).to_be_visible()
    pick_project(page, DISP_B)
    expect(page.get_by_role('alert').first).to_contain_text('「内租设备」不是数字')
    expect(mech1).to_have_value('abc')
    page.get_by_role('button', name='取消这项修改').click()
    expect(mech1).to_have_value('222')
    expect(page.get_by_text('1 项未保存')).to_have_count(0)


def test_damaged_receipt_blocks_resubmit_until_reread(server, browser_page):
    page, _ = browser_page
    open_process(page, server)
    pick_project(page, DISP_A)
    writes = []

    def damaged(route):
        if route.request.method != 'POST':
            route.continue_()
            return
        writes.append(route.request.post_data)
        response = route.fetch()
        assert response.ok and response.json()['ok'] is True
        route.fulfill(response=response, body='{')

    page.route('**/api/m06/mech', damaged)
    mech1 = page.get_by_label('2026-09 内租设备')
    mech1.fill('88.88')
    mech1.blur()
    alert = page.get_by_role('alert').first
    expect(alert).to_contain_text('保存结果待核实')
    expect(alert).to_contain_text('勿重复提交')
    expect(mech1).to_be_disabled()
    assert len(writes) == 1
    page.screenshot(path=str(output('tail410-uncertain.png', kind='screenshots')), full_page=True)
    # 200 正文损坏前服务端其实已写成：禁止重复提交，直到重读核对。
    mech = page.request.get(server + '/api/m06/mech').json()
    assert mech['mech'][WH_A]['2026-09'][1] == 88.88
    page.unroute('**/api/m06/mech')
    page.get_by_role('button', name='重读资料').click()
    expect(page.get_by_role('status').filter(has_text='已重读，请核对刚才的保存结果。')).to_be_visible()
    page.get_by_role('button', name='已核对，保留输入继续').click()
    expect(mech1).to_be_enabled()
    expect(mech1).to_have_value('88.88')
    expect(page.get_by_text('1 项未保存')).to_be_visible()
    mech1.press('Enter')
    expect(page.get_by_role('status').filter(has_text='内租设备已保存')).to_be_visible()
    expect(page.get_by_text('1 项未保存')).to_have_count(0)
    # 已写成但重读失败：同一把锁，重读核对后输入不丢、值已落库。
    def fail_matrix(route):
        if route.request.method == 'GET' and route.request.url.endswith('/api/m06/matrix'):
            route.fulfill(status=500, json={'message': '矩阵重读失败样例'})
        else:
            route.continue_()

    page.route('**/api/m06/matrix', fail_matrix)
    mech4 = page.get_by_label('2026-09 外租零租不含税')
    mech4.fill('77.7')
    mech4.blur()
    alert = page.get_by_role('alert').first
    expect(alert).to_contain_text('已保存，重读失败。')
    expect(mech4).to_be_disabled()
    page.unroute('**/api/m06/matrix')
    page.get_by_role('button', name='重读资料').click()
    expect(page.get_by_role('status').filter(has_text='已重读，请核对刚才的保存结果。')).to_be_visible()
    page.get_by_role('button', name='已核对，保留输入继续').click()
    expect(mech4).to_have_value('77.7')
    expect(page.get_by_text('1 项未保存')).to_have_count(0)
    mech = page.request.get(server + '/api/m06/mech').json()
    assert mech['mech'][WH_A]['2026-09'] == [123.45, 88.88, 0, 0, 77.7, 0]


def test_detail_confirmed_cell_switch_to_manual_keeps_amount(server, browser_page):
    page, _ = browser_page
    open_process(page, server)
    confirmed = page.locator('.process-cell').filter(has_text='明细确认')
    expect(confirmed).to_have_count(1)
    expect(confirmed).to_contain_text('333.00')
    expect(confirmed.locator('input')).to_have_count(0)
    # 明细确认格拒绝直接改值，且不写入。
    refused = page.request.post(server + '/api/m06/mech',
                                data={'warehouse': WH_B, 'ym': '2026-09', 'values': {'mech_2': '999'}})
    assert refused.status == 400
    assert '明细确认' in refused.json()['message']
    mech = page.request.get(server + '/api/m06/mech').json()
    assert mech['mech'][WH_B]['2026-09'][2] == 333
    assert mech['sources'][WH_B]['2026-09']['mech_2']['mode'] == 'detail'
    # 取消切换保留明细来源。
    page.get_by_role('button', name='切为手填').click()
    dialog = page.locator('.process-source-dialog')
    expect(dialog.get_by_role('heading', name='切为手填 · 外租零租含税')).to_be_visible()
    expect(dialog).to_contain_text('当前金额 333.00 元')
    dialog.get_by_role('button', name='保留明细来源').click()
    expect(dialog).not_to_be_visible()
    expect(page.locator('.process-cell').filter(has_text='明细确认')).to_have_count(1)
    # 确认切换：金额保留、来源变手填、可继续编辑。
    page.get_by_role('button', name='切为手填').click()
    dialog.get_by_role('button', name='确认切为手填').click()
    expect(page.get_by_role('status').filter(has_text='已切为手填，金额保留。')).to_be_visible()
    mech2 = page.get_by_label('2026-09 外租零租含税')
    expect(mech2).to_have_value('333')
    expect(page.locator('.process-cell').filter(has_text='明细确认')).to_have_count(0)
    mech = page.request.get(server + '/api/m06/mech').json()
    assert mech['sources'][WH_B]['2026-09']['mech_2']['mode'] == 'manual'
    assert mech['mech'][WH_B]['2026-09'][2] == 333
    mech2.fill('340')
    mech2.press('Enter')
    expect(page.get_by_role('status').filter(has_text='外租零租含税已保存')).to_be_visible()
    mech = page.request.get(server + '/api/m06/mech').json()
    assert mech['mech'][WH_B]['2026-09'][2] == 340


def test_anonymous_page_is_read_only_and_api_refuses_writes(server, browser_page):
    page, _ = browser_page
    page.goto(server + '/preview/process')
    expect(page.get_by_text('未登录 · 只读', exact=True)).to_be_visible()
    expect(page.locator('.process-matrix')).to_be_visible()
    for label in MECH_LABELS:
        cell_input = page.get_by_label(f'2026-09 {label}')
        if cell_input.count():
            expect(cell_input).to_be_disabled()
    switch = page.get_by_role('button', name='切为手填')
    if switch.count():
        expect(switch).to_be_disabled()
    before = page.request.get(server + '/api/m06/mech').json()
    with httpx.Client(base_url=server) as client:
        for path, payload in [
            ('/api/m06/mech', {'warehouse': WH_B, 'ym': '2026-09', 'values': {'mech_0': '1'}}),
            ('/api/m06/workbench/source-manual', {'warehouse': WH_B, 'ym': '2026-09', 'columns': ['mech_2'],
                                                  'effective_signature': 'x', 'confirmed': True}),
        ]:
            response = client.post(path, json=payload)
            assert response.status_code == 401, (path, response.status_code)
            assert '未登录' in response.json()['message']
    after = page.request.get(server + '/api/m06/mech').json()
    assert (after['mech'], after['sources'], after['count']) == (before['mech'], before['sources'], before['count'])


def test_export_adjusted_download_matches_matrix(server, browser_page):
    page, _ = browser_page
    open_process(page, server)
    pick_project(page, DISP_A)
    page.get_by_role('button', name='批量导出（先选项目）…').click()
    dialog = page.locator('.process-export-dialog')
    expect(dialog.get_by_role('heading', name='批量导出 · 表② 调整后表')).to_be_visible()
    expect(dialog).to_contain_text('已选 1 个项目')
    expect(dialog.get_by_role('checkbox', name=DISP_A)).to_be_checked()
    page.screenshot(path=str(output('tail410-export-dialog.png', kind='screenshots')), full_page=True)
    with page.expect_download() as download_info:
        dialog.get_by_role('button', name='保存到本地').click()
    download = download_info.value
    ym = datetime.now(ZoneInfo('Asia/Shanghai')).strftime('%Y-%m')
    assert download.suggested_filename == f'{DISP_A}_过程资料_表② 调整后表_{ym}.xlsx', download.suggested_filename
    target = output('tail410-调整后表-隔离项目.xlsx', kind='exports')
    shutil.copyfile(download.path(), target)
    expect(dialog).not_to_be_visible()
    check_workbook(target, display=DISP_A, rows=matrix_api(page, server)[WH_A]['adjusted'])


def test_export_raw_diff_and_zip_follow_selection_order(server, browser_page):
    page, _ = browser_page
    open_process(page, server)
    ym = datetime.now(ZoneInfo('Asia/Shanghai')).strftime('%Y-%m')
    matrices = matrix_api(page, server)
    # 表① 原始表单项目：文件名用展示名，金额 = 原始矩阵。
    raw = page.request.get(server + '/api/m06/export', params={'projects': WH_A, 'tab': '0'})
    assert raw.status == 200
    assert urllib.parse.unquote(raw.headers['x-export-filename']) == f'{DISP_A}_过程资料_表① 原始表_{ym}.xlsx'
    raw_path = output('tail410-原始表-隔离项目.xlsx', kind='exports')
    raw_path.write_bytes(raw.body())
    check_workbook(raw_path, display=DISP_A, rows=matrices[WH_A]['raw'])
    # 差异表两项目：打成 ZIP，文件与包内顺序都跟所选项目一致。
    multi = page.request.get(server + '/api/m06/export', params={'projects': f'{WH_A},{WH_B}', 'tab': '2'})
    assert multi.status == 200
    assert multi.headers['content-type'].startswith('application/zip')
    zip_path = output('tail410-差异表-两项目.zip', kind='exports')
    zip_path.write_bytes(multi.body())
    with zipfile.ZipFile(zip_path) as archive:
        assert archive.namelist() == [f'{DISP_A}_过程资料_差异表_{ym}.xlsx', f'{DISP_B}_过程资料_差异表_{ym}.xlsx']
        adjusted, raw_rows = matrices[WH_A]['adjusted'], matrices[WH_A]['raw']
        wb = openpyxl.load_workbook(io.BytesIO(archive.read(f'{DISP_A}_过程资料_差异表_{ym}.xlsx')), data_only=True)
        try:
            ws = wb['汇总表']
            assert ws['A1'].value == f'{DISP_A}材料机械消耗汇总表'
            for i, (adj_row, raw_row) in enumerate(zip(adjusted, raw_rows)):
                r = 4 + i
                assert cell_value(ws, r, 'B') == adj_row['ym']
                for key, letter in TEMPLATE_LETTERS.items():
                    left, right = raw_row.get(key), adj_row.get(key)
                    got = cell_value(ws, r, letter)
                    if left is None and right is None:
                        assert got in (None, ''), (r, letter, got)
                    else:
                        want = round(float(right or 0) - float(left or 0), 2)
                        if want:
                            assert round(float(got), 2) == want, (r, letter, got, want)
                        else:
                            assert got in (None, ''), (r, letter, got)
            delta = round(float(adjusted[0]['in_主辅材']) - float(raw_rows[0]['in_主辅材']), 2)
            assert round(float(cell_value(ws, 4 + len(adjusted), 'C')), 2) == delta
            assert round(float(cell_value(ws, 5 + len(adjusted), 'C')), 2) == delta
        finally:
            wb.close()
        wb = openpyxl.load_workbook(io.BytesIO(archive.read(f'{DISP_B}_过程资料_差异表_{ym}.xlsx')), data_only=True)
        try:
            ws = wb['汇总表']
            assert ws['A1'].value == f'{DISP_B}材料机械消耗汇总表'
            assert cell_value(ws, 4, 'B') == '2026-09'
            for letter in TEMPLATE_LETTERS.values():
                assert cell_value(ws, 4, letter) in (None, ''), letter
            assert cell_value(ws, 5, 'C') in (None, '')
            assert cell_value(ws, 6, 'C') in (None, '')
        finally:
            wb.close()
    # 没有数据的项目被拒做，不产出空表。
    refused = page.request.get(server + '/api/m06/export', params={'projects': '隔离归档仓库', 'tab': '0'})
    assert refused.status == 200
    body = refused.json()
    assert body['ok'] is False and '没有数据' in body['message']
