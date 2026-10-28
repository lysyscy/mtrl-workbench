"""Small offline checks; no workbench database or source workbook writes."""
from hashlib import sha256
import json
from pathlib import Path
from xml.dom.minidom import parseString
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill
import pytest
from sqlalchemy import create_engine, text

from snapshot import configure, validate

configure(Path(__file__).resolve().parents[3])
from audit import compare, read_plan
from annotate import META, RULE, INFO_RULE, annotate, instructions
from projects import context
from methods import load_approved
from app.mod03_notice.plan_audit_methods import apply_changes


@pytest.fixture
def inputs(tmp_path):
    source = tmp_path/'source.xlsx'
    wb = Workbook()
    ws = wb.active
    ws.title = '材机付款计划明细'
    ws.append(['测试甲项目付款计划'])
    ws.append(['单据号', '该流程金额', '本月预付款金额', '收款单位', '付款方式'])
    for row in [('A', 100, 10, '单位甲', '电汇'), ('B', 100, 90, '单位乙', '承兑'),
                ('C', 100, 10, '单位丙', '供应链'), ('C', 100, 20, '单位丙', '供应链'),
                ('', 100, 5, '缺号', '现汇'), ('D', None, 5, '缺金额', '现汇'),
                ('E', 100, 5, '待定', '现金'), ('X', 100, 5, '未知', '现汇')]:
        ws.append(row)
    ws.append(['合计', 700, 150])
    ws['L4'] = '人工付款附注'
    ws['D4'].fill = PatternFill('solid', fgColor='FFCC00')
    ws.conditional_formatting.add('D4', CellIsRule(operator='greaterThan', formula=['500'],
                                                 fill=PatternFill('solid', fgColor='00AA00')))
    wb.create_sheet('其他')['A1'] = '=1+2'
    wb.save(source)
    notices = [dict(id=i, bill_no=doc, approved='100.00', paid='20.00',
                    ownership_id=1, project_name='系统甲', payee='系统单位',
                    planned_pay_method='', excluded=False, i8_revoked=False, jc_contract=None)
               for i, doc in enumerate('ABCDE', 1)]
    snapshot = dict(schema=1, complete=True, captured_at='2026-10-15T10:00:00+08:00',
                    scope='测试完整范围', database_fingerprint='offline', method_column_available=True,
                    notices=notices, projects=[dict(id=1, disp='测试甲', warehouse_name='测试甲仓库', name_full='')])
    return source, snapshot


def test_original_read_audit_and_method_candidates(inputs, tmp_path):
    source, snapshot = inputs
    before = source.read_bytes()
    plan = read_plan(source, '2026-10')
    preview = compare(plan, snapshot)
    assert len(plan['rows']) == 8 and len(plan['structures']) == 11
    assert [r['audit_result'] for r in plan['rows']] == [
        '完全匹配', '超过未付', '重复上报', '重复上报', '无单据号', '输入待核对', '完全匹配', '单据号查不到']
    assert [c['bill_no'] for c in preview['changes']] == ['A', 'E']
    assert preview['changes'][0]['new'] == '现汇'
    assert not plan['source_warnings']
    assert source.read_bytes() == before


def test_in_place_notes_preserve_business_cells_clear_resolved_and_keep_manual(inputs, tmp_path):
    source, snapshot = inputs
    before = source.read_bytes()
    with ZipFile(source) as z:
        original = {name: z.read(name) for name in z.namelist()}
    original_cells = {c.getAttribute('r'): c.toxml() for c in parseString(original['xl/worksheets/sheet1.xml']).getElementsByTagName('c')}
    plan = read_plan(source, '2026-10')
    compare(plan, snapshot)
    receipt = annotate(plan, tmp_path/'backups')
    assert Path(receipt['backup']).read_bytes() == before
    wb = load_workbook(source)
    ws = wb['材机付款计划明细']
    assert META in wb.sheetnames and wb[META].sheet_state == 'veryHidden'
    assert ws['M3'].value is None
    assert '系统未付80' in ws['M4'].value
    assert ws['M2'].value == '核对说明'
    assert any(RULE in rule.formula for cf in ws.conditional_formatting for rule in ws.conditional_formatting[cf])
    wb.close()
    with ZipFile(source) as z:
        for name, content in original.items():
            if name not in {'xl/worksheets/sheet1.xml', 'xl/styles.xml', 'xl/workbook.xml', 'xl/_rels/workbook.xml.rels', '[Content_Types].xml'}:
                assert z.read(name) == content
        cells = {c.getAttribute('r'): c.toxml() for c in parseString(z.read('xl/worksheets/sheet1.xml')).getElementsByTagName('c')}
        assert all(cells[k] == v for k, v in original_cells.items())
    # 模拟用户改单据数据并在说明列追加人工文字，再复核。
    wb = load_workbook(source)
    ws = wb['材机付款计划明细']
    ws['C4'] = 70
    ws['M4'] = '已与项目联系'
    ws['C5'] = 5
    ws['A6'] = 'Z'
    wb.save(source)
    wb.close()
    plan = read_plan(source, '2026-10')
    compare(plan, snapshot)
    annotate(plan, tmp_path/'backups')
    wb = load_workbook(source)
    ws = wb['材机付款计划明细']
    assert ws['M4'].value == '已与项目联系'
    assert ws['M5'].value in (None, '')
    rules = [rule for cf in ws.conditional_formatting for rule in ws.conditional_formatting[cf] if RULE in rule.formula]
    assert len(rules) == 1
    assert 'C4' not in str(list(ws.conditional_formatting))
    wb.close()


@pytest.mark.parametrize('conflict', ['occupied', 'changed', 'open'])
def test_annotation_refuses_to_overwrite_unknown_notes_changed_or_open_source(inputs, tmp_path, conflict):
    source, snapshot = inputs
    if conflict == 'occupied':
        w = load_workbook(source)
        w.active['M3'] = '人工备注'
        w.save(source)
        w.close()
    plan = read_plan(source, '2026-10')
    compare(plan, snapshot)
    if conflict == 'changed':
        source.write_bytes(source.read_bytes()+b'changed')
    if conflict == 'open':
        source.with_name('~$'+source.name).write_bytes(b'locked')
    before = source.read_bytes()
    with pytest.raises(ValueError):
        annotate(plan, tmp_path/'backups')
    assert source.read_bytes() == before


def test_nonunique_system_bill_and_wrong_project_never_write(inputs):
    source, snapshot = inputs
    snapshot['notices'].append(dict(snapshot['notices'][0], id=99))
    preview = compare(read_plan(source, '2026-10'), snapshot)
    assert [c['bill_no'] for c in preview['changes']] == ['E']
    snapshot['notices'].pop()
    snapshot['notices'][0]['ownership_id'] = 99
    assert [c['bill_no'] for c in compare(read_plan(source, '2026-10'), snapshot)['changes']] == ['E']


def test_ai_project_map_uses_fresh_catalog_and_does_not_override_wrong_notice_owner(inputs):
    source, snapshot = inputs
    plan = read_plan(source, '2026-10')
    for row in plan['rows']:
        row['project'] = '甲方简称'
    decision = dict(schema=1, project_catalog_sha256=context(plan,snapshot)['project_catalog_sha256'],
                    decisions=[dict(source_project='甲方简称',project_id=1,reason='本次AI结合项目全名和原表上下文确定')])
    compare(plan,snapshot,decision)
    assert plan['rows'][0]['final_status'] == '核对通过'
    plan = read_plan(source,'2026-10')
    decision['decisions'][0]['source_project'] = '测试甲'
    snapshot['notices'][0]['ownership_id'] = 99
    compare(plan,snapshot,decision)
    assert '项目未唯一对应系统通知书归属' in plan['rows'][0]['input_issues']
    snapshot['projects'].append(dict(id=2,disp='新项目',warehouse_name='新仓库',name_full='新项目'))
    with pytest.raises(ValueError,match='清单已变化'):
        compare(read_plan(source,'2026-10'),snapshot,decision)


@pytest.mark.parametrize('day,paid_label', [('2026-10-13',True),('2026-09-30',False),('unknown',False)])
def test_paid_reminder_uses_plan_month_while_all_historical_money_counts(inputs, day, paid_label, tmp_path):
    source, snapshot = inputs
    notice = snapshot['notices'][1]
    notice.update(paid='100',payments=[dict(amount='10',pay_date='2026-08-01'),dict(amount='90',pay_date=day)])
    plan = read_plan(source,'2026-10')
    preview = compare(plan,snapshot)
    row = plan['rows'][1]
    assert row['audit_result'] == '超过未付'
    assert row['plan_already_paid'] is paid_label
    assert row['system']['paid'] == '100'
    notes, _, blue = instructions(plan)
    assert ('10/13已付同金额' in notes[4]) is paid_label
    assert not any(c['id']==2 for c in preview['changes'])
    annotate(plan,tmp_path/'backup')
    wb = load_workbook(source)
    rules = [r for cf in wb.active.conditional_formatting for r in wb.active.conditional_formatting[cf]]
    assert any(INFO_RULE in r.formula for r in rules) is paid_label
    wb.close()


@pytest.mark.parametrize('payments, label', [
    ([('2026-10-12', '40'), ('2026-10-13', '50')], '10/12、10/13累计已付同金额'),
    ([('2026-10-12', '40')], '10/12已付40元'),
])
def test_paid_same_amount_is_shortened_only_when_equal(inputs, payments, label):
    source, snapshot = inputs
    snapshot['notices'][1].update(paid=str(sum(int(v) for _, v in payments)),
                                payments=[dict(pay_date=d, amount=v) for d, v in payments])
    plan = read_plan(source, '2026-10')
    compare(plan, snapshot)
    notes, _, _ = instructions(plan)
    assert label in notes[4]


def test_blank_amount_is_zero_only_for_totals_not_for_row_match(inputs):
    source,snapshot = inputs
    wb=load_workbook(source)
    wb.active['B11']=700
    wb.active['C11']=150
    wb.save(source)
    wb.close()
    plan=read_plan(source,'2026-10')
    assert not plan['source_warnings']
    compare(plan,snapshot)
    assert plan['rows'][5]['audit_result']=='输入待核对'
    wb=load_workbook(source)
    wb.active['B8']='=1+2'
    wb.save(source)
    wb.close()
    assert any('缺失金额' in w['issue'] for w in read_plan(source,'2026-10')['source_warnings'])


@pytest.mark.parametrize('change', ['incomplete', 'duplicate_id', 'unknown_status', 'paid_mismatch'])
def test_snapshot_rejects_incomplete_or_duplicate_ids(inputs, change):
    _, snapshot = inputs
    if change == 'incomplete':
        snapshot['complete'] = False
    elif change == 'duplicate_id':
        snapshot['notices'].append(snapshot['notices'][0])
    elif change == 'unknown_status':
        del snapshot['notices'][0]['i8_revoked']
    else:
        snapshot['notices'][0]['payments'] = [dict(amount='99',pay_date='2026-10-12')]
    with pytest.raises(ValueError):
        validate(snapshot)


def test_approval_is_bound_to_preview_and_source(inputs, tmp_path):
    source, snapshot = inputs
    preview = compare(read_plan(source, '2026-10'), snapshot)
    path = tmp_path/'preview.json'
    path.write_text(json.dumps(preview), encoding='utf-8')
    digest = sha256(path.read_bytes()).hexdigest()
    assert load_approved(path, digest)['changes']
    with pytest.raises(ValueError, match='SHA'):
        load_approved(path, '0'*64)
    source.write_bytes(source.read_bytes()+b'changed')
    with pytest.raises(ValueError, match='原表已改变'):
        load_approved(path, digest)


@pytest.mark.parametrize('conflict', ['', 'old', 'project', 'revoked', 'duplicate', 'missing', 'approved', 'paid'])
def test_apply_is_atomic_and_updates_only_planned_method(conflict):
    engine = create_engine('sqlite:///:memory:')
    changes = [dict(id=i, bill_no=f'D{i}', ownership_id=1, old='', new='承兑', approved=str(i*100), paid='0') for i in (1, 2)]
    with engine.begin() as conn:
        conn.exec_driver_sql('CREATE TABLE payment_notice (id INTEGER PRIMARY KEY, bill_no TEXT, '
                             'ownership_id INTEGER, planned_pay_method TEXT, excluded BOOL, '
                             'i8_revoked BOOL, amount_approved NUMERIC)')
        conn.exec_driver_sql("INSERT INTO payment_notice VALUES (1,'D1',1,'',0,0,100),(2,'D2',1,'',0,0,200)")
        conn.exec_driver_sql('CREATE TABLE payment_flow (notice_id INTEGER, amount NUMERIC, status TEXT)')
        if conflict == 'old':
            conn.exec_driver_sql("UPDATE payment_notice SET planned_pay_method='现汇' WHERE id=2")
        if conflict == 'project':
            conn.exec_driver_sql('UPDATE payment_notice SET ownership_id=2 WHERE id=2')
        if conflict == 'revoked':
            conn.exec_driver_sql('UPDATE payment_notice SET i8_revoked=1 WHERE id=2')
        if conflict == 'duplicate':
            conn.exec_driver_sql("INSERT INTO payment_notice VALUES (3,'D2',1,'',0,0,200)")
        if conflict == 'missing':
            conn.exec_driver_sql('DELETE FROM payment_notice WHERE id=2')
        if conflict == 'approved':
            conn.exec_driver_sql('UPDATE payment_notice SET amount_approved=199 WHERE id=2')
        if conflict == 'paid':
            conn.exec_driver_sql("INSERT INTO payment_flow VALUES (2, 190, 'effective')")
    if conflict:
        with pytest.raises(ValueError):
            with engine.begin() as conn:
                apply_changes(conn, {'changes': changes})
    else:
        with engine.begin() as conn:
            apply_changes(conn, {'changes': changes})
    with engine.connect() as conn:
        first = conn.execute(text('SELECT planned_pay_method, amount_approved FROM payment_notice WHERE id=1')).one()
        assert first == ('' if conflict else '承兑', 100)
    engine.dispose()



def test_contract_method_conflict_skips_only_that_row_and_marks_original(inputs, tmp_path):
    source, snapshot = inputs
    snapshot['notices'][0]['jc_contract'] = {'id': 61, 'contract_no': 'CGHT-TEST', 'payment_method': '承兑'}
    original = source.read_bytes()
    plan = read_plan(source, '2026-10')
    preview = compare(plan, snapshot)
    assert [c['bill_no'] for c in preview['changes']] == ['E']
    assert '合同付款方式' in '；'.join(plan['rows'][0]['input_issues'])
    backup = annotate(plan, tmp_path/'backup')
    wb = load_workbook(source)
    assert wb['材机付款计划明细']['E3'].value == '电汇'
    assert '承兑' in wb['材机付款计划明细']['M3'].value
    assert wb['材机付款计划明细']['L4'].value == '人工付款附注'
    assert wb['其他']['A1'].value == '=1+2'
    assert Path(backup['backup']).read_bytes() == original


def test_confirmed_method_exception_is_in_preview_and_missing_contract_method_skips(inputs):
    source, snapshot = inputs
    snapshot['notices'][0]['jc_contract'] = {'id': 61, 'contract_no': 'CGHT-TEST', 'payment_method': '承兑'}
    plan=read_plan(source,'2026-10')
    preview=compare(plan,snapshot,confirmed_method_rows=[3])
    a=next(c for c in preview['changes'] if c['bill_no']=='A')
    assert a['method_exception_confirmed'] is True
    assert a['jc_contract']==snapshot['notices'][0]['jc_contract']
    assert a['new']=='现汇'
    snapshot['notices'][0]['jc_contract']['payment_method']=''
    plan=read_plan(source,'2026-10')
    preview=compare(plan,snapshot)
    assert [c['bill_no'] for c in preview['changes']]==['E']
    assert '未填写' in '；'.join(plan['rows'][0]['input_issues'])



def test_old_snapshot_missing_contract_context_cannot_write_methods(inputs):
    source, snapshot = inputs
    snapshot['notices'][0].pop('jc_contract')
    plan=read_plan(source,'2026-10')
    preview=compare(plan,snapshot,confirmed_method_rows=[3])
    assert [c['bill_no'] for c in preview['changes']]==['E']
    assert '上下文缺失' in '；'.join(plan['rows'][0]['input_issues'])


def test_paid_plan_keeps_blue_state_even_if_contract_method_changed(inputs):
    source,snapshot=inputs
    snapshot['notices'][0].update(jc_contract={'id':61,'contract_no':'CGHT-TEST','payment_method':'承兑'},
        paid='20',payments=[{'amount':'20','pay_date':'2026-10-01'}])
    plan=read_plan(source,'2026-10')
    preview=compare(plan,snapshot)
    assert plan['rows'][0]['final_status']=='本月已付'
    assert [c['bill_no'] for c in preview['changes']]==['E']
