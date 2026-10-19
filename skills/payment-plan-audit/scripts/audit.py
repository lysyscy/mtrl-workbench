"""Read a plan, reuse the workbench rules, annotate issues in the original file."""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import date
from decimal import Decimal
from hashlib import sha256
import json
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from snapshot import capture, configure, validate
from annotate import META, annotate
from projects import context, decisions


def norm(value):
    return re.sub(r'[\s：:（）()/、]', '', str(value or ''))


ALIASES = {
    'project': ('项目', '项目名称', '工程名称'),
    'doc_no': ('单据号', '单据编码', '付款单号', '材料：其他支出付资产处/贸易公司（编号）',
               '外租机械：付款通知书（编号）'),
    'flow_amount': ('该流程金额', '流程金额'),
    'planned_amount': ('本月预付款金额', '本月计划付款金额'),
    'payee': ('收款单位',), 'pay_method': ('付款方式', '支付方式'),
    'remark': ('备注',), 'pay_date': ('付款时间',),
}
FIELD_OF = {norm(label): key for key, labels in ALIASES.items() for label in labels}
METHODS = {'现汇': '现汇', '承兑': '承兑', '供应链': '供应链', '其他': '其他',
           '银行承兑': '承兑', '电汇': '现汇', '现金': '现汇'}


def text(value):
    return '' if value is None else str(value).strip()


def read_plan(source, month, sheet_name=None):
    from app.core.money import amount_text, to_money
    digest = sha256(source.read_bytes()).hexdigest()
    formulas = load_workbook(source, data_only=False)
    values = load_workbook(source, data_only=True)
    try:
        candidates = []
        for ws in values:
            if any({'doc_no', 'flow_amount', 'planned_amount'} <= {
                    FIELD_OF.get(norm(c.value)) for c in row}
                   for row in ws.iter_rows(max_row=min(12, ws.max_row))):
                candidates.append(ws.title)
        chosen = sheet_name or (candidates[0] if len(candidates) == 1 else None)
        if chosen not in values.sheetnames:
            raise ValueError(f'请用--sheet明确核对表；可识别候选：{candidates}')
        ws, fs = values[chosen], formulas[chosen]
        rows, structures, warnings = [], [], []
        header, project, title_cell = {}, '', ''
        project_rows, subtotal_rows = [], []

        def total_check(n, detail, kind):
            for key in ('flow_amount', 'planned_amount'):
                col = header.get(key)
                if col is None or not detail:
                    continue
                observed = ws.cell(n, col + 1).value
                if observed is None and fs.cell(n, col + 1).value is None:
                    continue
                known = [r[key] for r in detail]
                if any(r[key] is None and key not in r['blank_amounts'] for r in detail):
                    warnings.append({'cell': f'{chosen}!{get_column_letter(col+1)}{n}',
                                     'amount_field': key,
                                     'issue': f'{kind}包含缺失金额，无法独立验证', 'expected': None,
                                     'observed': text(observed)})
                    continue
                try:
                    got = to_money(amount_text(observed))
                except ValueError:
                    warnings.append({'cell': f'{chosen}!{get_column_letter(col+1)}{n}',
                                     'amount_field': key,
                                     'issue': f'{kind}金额无有效值', 'expected': None,
                                     'observed': text(observed)})
                    continue
                expected = sum((Decimal(v or '0') for v in known), Decimal('0.00'))
                if got != expected:
                    warnings.append({'cell': f'{chosen}!{get_column_letter(col+1)}{n}',
                                     'amount_field': key,
                                     'issue': f'{kind}与独立明细合计不符', 'expected': str(expected),
                                     'observed': str(got), 'difference': str(expected-got),
                                     'formula': text(fs.cell(n, col+1).value)})

        for n in range(1, fs.max_row + 1):
            raw = [c.value for c in fs[n]]
            if not any(v is not None for v in raw):
                continue
            vals = [c.value for c in ws[n]]
            fields = {FIELD_OF[norm(v)]: i for i, v in enumerate(vals) if norm(v) in FIELD_OF}
            first = text(vals[0])
            if '付款计划' in first and '项目' in first:
                project = first.split('项目', 1)[0].strip()
                title_cell = f'{chosen}!A{n}'
                header, project_rows, subtotal_rows = {}, [], []
                kind = '项目标题'
            elif {'flow_amount', 'planned_amount', 'doc_no'} <= fields.keys():
                header = fields
                kind = '表头'
            elif fields and all(v is None or norm(v) in FIELD_OF for v in vals):
                kind = '表头续行'
            elif header:
                cell = lambda key: vals[header[key]] if key in header else None
                doc = text(cell('doc_no'))
                label = re.sub(r'[\s：:]', '', doc)
                if label in ('小计', '合计', '总计', '大合计'):
                    kind = '小计' if label == '小计' else '项目合计'
                    total_check(n, subtotal_rows if kind == '小计' else project_rows, kind)
                    if kind == '小计':
                        subtotal_rows = []
                elif n == fs.max_row and not any(vals[:4]) and any(
                        fs.cell(n, header[k]+1).data_type == 'f'
                        for k in ('flow_amount', 'planned_amount')):
                    kind = '全表总计'
                    total_check(n, rows, kind)
                elif any(cell(k) is not None for k in ('doc_no', 'payee', 'flow_amount', 'planned_amount')):
                    kind = '明细'
                    issues = []
                    item = {'sheet': chosen, 'row': n, 'project': text(cell('project')) or project,
                            'project_title_cell': title_cell, 'doc_no': doc,
                            'cells': {key: f'{get_column_letter(col+1)}{n}' for key, col in header.items()},
                            'payee': text(cell('payee')), 'pay_method_raw': text(cell('pay_method')),
                            'pay_method': METHODS.get(text(cell('pay_method')), ''),
                            'remark': text(cell('remark')), 'hidden': bool(fs.row_dimensions[n].hidden),
                            'original': {get_column_letter(i+1): text(v) for i, v in enumerate(vals) if v is not None},
                            'formulas': {c.coordinate: c.value for c in fs[n] if c.data_type == 'f'}}
                    item['blank_amounts'] = []
                    for key in ('flow_amount', 'planned_amount'):
                        if cell(key) is None and fs.cell(n, header[key]+1).data_type != 'f':
                            item['blank_amounts'].append(key)
                        try:
                            amount = to_money(amount_text(cell(key)))
                            if amount < 0:
                                raise ValueError('负金额')
                            item[key] = str(amount)
                        except ValueError:
                            item[key] = None
                            issues.append(f'{ALIASES[key][0]}缺失或无效')
                    if len(re.findall(r'[A-Za-z]{2,}\d{6,}', doc)) > 1:
                        issues.append('同格多个单据号，未自动拆分')
                    if not item['project']:
                        issues.append('项目未识别')
                    item['cells'].setdefault('project', title_cell.rsplit('!', 1)[-1] or f'A{n}')
                    if not item['pay_method']:
                        issues.append('付款方式为空或需人工确认')
                    item['input_issues'] = issues
                    rows.append(item)
                    project_rows.append(item)
                    subtotal_rows.append(item)
                else:
                    kind = '待确认行'
            else:
                kind = '待确认行'
            structures.append({'sheet': chosen, 'row': n, 'kind': kind,
                               'content': ' | '.join(text(v) for v in vals if v is not None)})
        if not rows:
            raise ValueError('未识别出明细，请确认表与列结构，原件未改动')
        return {'source': str(source), 'source_sha256': digest, 'month': month, 'sheet': chosen,
                'not_audited_sheets': [s for s in values.sheetnames if s not in (chosen, META)],
                'rows': rows, 'structures': structures, 'source_warnings': warnings}
    finally:
        formulas.close()
        values.close()
        if sha256(source.read_bytes()).hexdigest() != digest:
            raise ValueError('核对期间原文件发生变化，请重新核对')


def compare(plan, snapshot, project_map=None):
    from app.mod04_plan.domain import NoticeAmounts, ProjectBrief, audit_rows, match_project
    from app.core.money import to_money
    validate(snapshot)
    selected = decisions(project_map, snapshot, plan)
    plan['project_decisions'] = list(selected.values())
    by_doc = defaultdict(list)
    for notice in snapshot['notices']:
        by_doc[text(notice['bill_no'])].append(notice)
    projects = [ProjectBrief(**{k: p.get(k) for k in ('id', 'disp', 'warehouse_name', 'name_full')},
                             keywords=tuple(p.get('match_keywords') or ())) for p in snapshot['projects']]
    views = {doc: NoticeAmounts(doc, to_money(ns[0]['approved']), to_money(ns[0]['paid']))
             for doc, ns in by_doc.items() if len(ns) == 1 and doc}
    inputs = [dict(r, year_month=plan['month']) for r in plan['rows']]
    results = audit_rows(inputs, views)
    changes = []
    for row, result in zip(plan['rows'], results):
        hits = by_doc.get(row['doc_no'], [])
        match = match_project(row['project'], projects)
        choice = selected.get(row['project'])
        expected_id = choice['project_id'] if choice is not None else match.project_id
        row['project_match'] = next((p.disp or p.warehouse_name for p in projects if p.id == expected_id), '待确认')
        row['project_match_reason'] = choice['reason'] if choice else '现有规则匹配'
        row['audit_result'] = result
        if len(hits) > 1:
            row['input_issues'].append('系统单据号不唯一，未选择其中一条')
        if any('多个单据号' in s or '金额缺失' in s or '金额缺失或无效' in s for s in row['input_issues']):
            row['audit_result'] = '输入待核对'
        n = hits[0] if len(hits) == 1 else None
        row['system'] = n
        row['system_owner'] = next((p.disp or p.warehouse_name or f'#{p.id}' for p in projects
                                    if n and p.id == n.get('ownership_id')), '未归属')
        if expected_id is None and '项目未识别' not in row['input_issues']:
            row['input_issues'].append('项目未识别')
        elif n and expected_id is not None and n.get('ownership_id') != expected_id:
            row['input_issues'].append('项目未唯一对应系统通知书归属')
        if n and (n.get('excluded') or n.get('i8_revoked')):
            row['input_issues'].append('系统通知书已剔除或撤单')
        row['final_status'] = ('核对通过' if row['audit_result'] == '完全匹配' and not row['input_issues']
                               else '待人工核对')
        row['month_payments'] = []
        monthly = defaultdict(Decimal)
        for payment in (n or {}).get('payments', []):
            try:
                paid_date = date.fromisoformat(payment['pay_date'])
            except (ValueError, TypeError):
                continue
            if paid_date.strftime('%Y-%m') == plan['month']:
                monthly[paid_date.isoformat()] += to_money(payment['amount'])
        row['month_payments'] = [{'pay_date': day, 'amount': str(total)} for day, total in sorted(monthly.items())]
        row['plan_already_paid'] = (bool(monthly) and not row['input_issues']
                                   and result in ('完全匹配', '超过未付')
                                   and Decimal(row['planned_amount'] or '0') > 0
                                   and sum(monthly.values()) >= Decimal(row['planned_amount']))
        if row['plan_already_paid']:
            row['final_status'] = '本月已付'
        eligible = (n and row['audit_result'] == '完全匹配' and not row['input_issues']
                    and not row['plan_already_paid'] and Decimal(row['planned_amount'] or 0) > 0)
        row['method_action'] = '待人工处理'
        if eligible:
            old = n.get('planned_pay_method') or ''
            row['method_action'] = '无需变更' if old == row['pay_method'] else '可预览回写'
            if old != row['pay_method']:
                changes.append({'id': n['id'], 'bill_no': n['bill_no'],
                                'ownership_id': n.get('ownership_id'), 'old': old,
                                'new': row['pay_method'], 'source_row': row['row'],
                                'source_method': row['pay_method_raw'],
                                'approved': n['approved'], 'paid': n['paid'],
                                'flow_amount': row['flow_amount'], 'planned_amount': row['planned_amount']})
    if len({c['id'] for c in changes}) != len(changes):
        raise ValueError('回写候选重复定位同一通知书')
    plan['snapshot'] = {k: snapshot[k] for k in ('captured_at', 'scope', 'database_fingerprint', 'method_column_available')}
    plan['summary'] = {'rows': len(plan['rows']), 'projects': len({r['project'] for r in plan['rows']}),
                       'planned_total': str(sum((Decimal(r['planned_amount']) for r in plan['rows']
                                                if r['planned_amount'] is not None), Decimal('0.00'))),
                       'results': dict(Counter(r['audit_result'] for r in plan['rows'])),
                       'row_status': dict(Counter(r['final_status'] for r in plan['rows'])),
                       'with_input_issues': sum(bool(r['input_issues']) for r in plan['rows']),
                       'method_changes': len(changes)}
    return {'schema': 1, 'source': plan['source'], 'source_sha256': plan['source_sha256'],
            'month': plan['month'], 'sheet': plan['sheet'], 'snapshot': plan['snapshot'], 'changes': changes}


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--repo', help='仅开发时使用本地仓库规则；普通使用无需此参数')
    ap.add_argument('--server')
    ap.add_argument('--session-file')
    ap.add_argument('--source', required=True)
    ap.add_argument('--month', required=True)
    ap.add_argument('--sheet')
    ap.add_argument('--snapshot')
    ap.add_argument('--prepare', action='store_true', help='只生成项目匹配上下文，原表不标注')
    ap.add_argument('--project-map', help='Agent根据当前项目清单给出的匹配决策JSON')
    ap.add_argument('--output-dir', required=True)
    args = ap.parse_args()
    configure(args.repo)
    if not args.prepare and not args.project_map:
        raise ValueError('请先用--prepare取得当前项目清单，由AI给出--project-map后再标注原表')
    if not re.fullmatch(r'\d{4}-(0[1-9]|1[0-2])', args.month):
        raise ValueError('月份须为YYYY-MM')
    source = Path(args.source).resolve()
    output = Path(args.output_dir).resolve()
    output.mkdir(parents=True, exist_ok=True)
    for name in ('audit.json', 'snapshot.json', 'method-preview.json', 'context.json'):
        if (output/name).exists() or output/name == source:
            raise ValueError('本次输出已存在，请使用新的独立输出目录')
    snapshot = (json.loads(Path(args.snapshot).read_text(encoding='utf-8')) if args.snapshot
                else capture(server=args.server, session_file=args.session_file))
    validate(snapshot)
    plan = read_plan(source, args.month, args.sheet)
    if args.prepare:
        for name, data in [('snapshot.json', snapshot), ('context.json', context(plan, snapshot))]:
            (output/name).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
        print(json.dumps({'context': str(output/'context.json'), 'source_unchanged': True}, ensure_ascii=False))
        return
    plan['snapshot_origin'] = '离线复核，沿用旧快照' if args.snapshot else '本次系统接口只读快照'
    mapping = json.loads(Path(args.project_map).read_text(encoding='utf-8')) if args.project_map else None
    preview = compare(plan, snapshot, mapping)
    plan['input_source_sha256'] = plan['source_sha256']
    backups = (Path(args.repo)/'var/backups/payment-plan-audit' if args.repo
               else Path.home()/'.mtrl-workbench/backups/payment-plan-audit')
    plan['annotation'] = annotate(plan, backups)
    plan['source_sha256'] = preview['source_sha256'] = plan['annotation']['sha256']
    for name, data in [('audit.json', plan), ('snapshot.json', snapshot), ('method-preview.json', preview)]:
        (output/name).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    if sha256(source.read_bytes()).hexdigest() != plan['source_sha256']:
        raise ValueError('标注后原文件再次变化，请重新核对')
    print(json.dumps({'summary': plan['summary'], 'output': str(output), 'annotation': plan['annotation'],
                      'preview_sha256': sha256((output/'method-preview.json').read_bytes()).hexdigest()}, ensure_ascii=False))


if __name__ == '__main__':
    try:
        main()
    except ValueError as exc:
        raise SystemExit(str(exc)) from None
    except Exception as exc:
        raise SystemExit(f'核对未完成：{type(exc).__name__}；请检查文件及本次输出；原表可能已标注，未写入系统数据库') from None
