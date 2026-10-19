"""Read explicitly selected machinery-report sheets; never save source workbooks."""
from __future__ import annotations

import argparse
from datetime import date, datetime
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import re


def norm(value):
    return re.sub(r"\s+", "", str(value or "")).replace("（", "(").replace("）", ")")


def field_key(label):
    text = norm(label)
    exact = {"序号": "sequence", "设备名称": "equipment", "规格型号": "specification",
             "单位": "unit", "数量": "quantity", "计划月数": "planned_months",
             "已包月数": "elapsed_months", "出租单位": "lessor", "审核人": "reviewer",
             "使用班组": "team", "车牌号": "plate", "备注": "remarks"}
    if text in exact:
        return exact[text]
    if text.startswith("合价"):
        return "gross_amount"
    if text.startswith("零租单价"):
        return "original_unit_price"
    if text.startswith("包月单价"):
        return "monthly_unit_price"
    if text.startswith("专票税率"):
        return "tax_description"
    if text.startswith("不含税"):
        return "reported_net_amount"
    if text == "租赁日期/天":
        return "rental_date_or_days"
    if text.startswith("租赁日期"):
        return "rental_date"
    if text.startswith("入场日期"):
        return "entry_date"
    if text.startswith("用途"):
        return "purpose"
    if text.startswith("班组长签字"):
        return "team_signature"
    return None


def json_value(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def letter(index):
    result = ""
    while index:
        index, digit = divmod(index - 1, 26)
        result = chr(65 + digit) + result
    return result


def rows_of(path, sheet):
    if path.suffix.lower() == ".xls":
        import xlrd
        book = xlrd.open_workbook(str(path), on_demand=True)
        try:
            tab = book.sheet_by_name(sheet)
            for r in range(tab.nrows):
                cells = []
                for c in range(tab.ncols):
                    cell = tab.cell(r, c)
                    value = cell.value if cell.ctype != xlrd.XL_CELL_EMPTY else None
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, book.datemode)
                    cells.append({"coordinate": f"{letter(c + 1)}{r + 1}",
                                  "value": json_value(value), "formula": None, "format": None})
                yield r + 1, cells
        finally:
            book.release_resources()
    elif path.suffix.lower() in {".xlsx", ".xlsm"}:
        import openpyxl
        formula_book = openpyxl.load_workbook(path, read_only=True, data_only=False)
        value_book = None
        try:
            value_book = openpyxl.load_workbook(path, read_only=True, data_only=True)
            tab, values = formula_book[sheet], value_book[sheet]
            if tab.max_row > 30000 or tab.max_column > 200:
                raise ValueError("Sheet extent exceeds inspection limit; inspect dimensions explicitly")
            for r, (formulas, cached) in enumerate(zip(tab.iter_rows(), values.iter_rows()), 1):
                cells = []
                for c, (original, cache) in enumerate(zip(formulas, cached), 1):
                    formula = getattr(original.value, "text", original.value) if original.data_type == "f" else None
                    cells.append({"coordinate": f"{letter(c)}{r}", "value": json_value(cache.value),
                                  "formula": formula, "format": getattr(original, "number_format", None)})
                yield r, cells
        finally:
            formula_book.close()
            if value_book is not None:
                value_book.close()
    else:
        raise ValueError(f"Unsupported source format: {path.suffix}")


def tax_check(gross, label, header):
    text = norm(label)
    if not isinstance(gross, (int, float)) or isinstance(gross, bool):
        return {"status": "待核", "reason": "合价缺少数值缓存"}
    if "普票" in text:
        return {"status": "可计算", "invoice": "普票", "net_check": str(Decimal(str(gross)))}
    if label is None or text == "":
        return {"status": "待核", "reason": "票种或税率缺失"}
    rate = None
    if isinstance(label, (int, float)) and not isinstance(label, bool) and "专票税率" in norm(header):
        rate = Decimal(str(label))
    elif "专票" in text or re.fullmatch(r"\d+(?:\.\d+)?%", text):
        percentages = re.findall(r"(\d+(?:\.\d+)?)%", text)
        if len(percentages) == 1:
            rate = Decimal(percentages[0]) / 100
    if rate is None or not 0 <= rate <= 1:
        return {"status": "待核", "reason": "票种或税率含义不明确"}
    return {"status": "可计算", "invoice": "专票", "rate": str(rate),
            "net_check": str(Decimal(str(gross)) / (1 + rate))}


def extract(plan):
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", plan.get("period", "")):
        raise ValueError("period must be YYYY-MM")
    if not isinstance(plan.get("sheets"), list) or not plan["sheets"]:
        raise ValueError("Select at least one reviewed source sheet")
    records, issues, sources = [], [], []
    selected = set()
    for selection in plan.get("sheets", []):
        path = Path(selection["path"]).resolve(strict=True)
        sheet = selection["sheet"]
        if selection.get("kind") not in {"零租", "包月"} or not selection.get("project") or not selection.get("reason"):
            raise ValueError("Each selection needs project, kind and a reviewed reason")
        key = (str(path).casefold(), sheet)
        if key in selected:
            raise ValueError(f"Selected sheet twice: {path}!{sheet}")
        selected.add(key)
        excluded = selection.get("exclude_rows", {})
        if not isinstance(excluded, dict) or any(not str(row).isdigit() or not reason for row, reason in excluded.items()):
            raise ValueError("exclude_rows needs row numbers and reviewed reasons")
        before = hashlib.sha256(path.read_bytes()).hexdigest()
        source = {"path": str(path), "sheet": sheet, "reason": selection["reason"], "sha256_before": before,
                  "formula_access": "cached values only" if path.suffix.lower() == ".xls" else "formula and cache",
                  "excluded_rows": excluded}
        sources.append(source)
        headers = None
        seen_header = False
        try:
            for row_number, cells in rows_of(path, sheet):
                labels = [norm(c["value"]) for c in cells]
                if "设备名称" in labels and any(t.startswith("合价") for t in labels):
                    headers = {}
                    for index, cell in enumerate(cells):
                        field = field_key(cell["value"])
                        if field:
                            if field in headers:
                                raise ValueError(f"Ambiguous duplicate header: {field} at row {row_number}")
                            headers[field] = (index, cell["value"])
                    seen_header = True
                    continue
                if headers is None:
                    continue
                if str(row_number) in excluded:
                    continue
                fields = {name: {**cells[index], "header": label} for name, (index, label) in headers.items()}
                equipment = fields["equipment"]["value"]
                amount = fields["gross_amount"]
                unit = fields.get("unit", {}).get("value")
                if not unit and "签字" in norm(equipment) and "签字" in norm(amount["value"]):
                    continue
                if norm(equipment) in {"合计", "合计:", "小计", "小计:", "零星小计", "总计", "总计:"}:
                    continue
                if not equipment:
                    if amount["formula"] or amount["value"] not in (None, ""):
                        issues.append({"path": str(path), "sheet": sheet, "row": row_number,
                                       "reason": "有合价但无设备名称，核实是否为汇总行或缺失名称的明细"})
                    continue
                # Identifying columns plus an amount distinguish detail from titles/signatures.
                if not unit:
                    if amount["formula"] or amount["value"] not in (None, ""):
                        issues.append({"path": str(path), "sheet": sheet, "row": row_number,
                                       "reason": "有名称和合价但单位缺失，核实是否为分组行或有效明细"})
                    continue
                if amount["value"] is None and not amount["formula"]:
                    issues.append({"path": str(path), "sheet": sheet, "row": row_number,
                                   "reason": "疑似明细但合价为空"})
                    continue
                tax_field = fields.get("tax_description", {})
                tax = tax_check(amount["value"], tax_field.get("value"), tax_field.get("header", ""))
                record = {"period": plan["period"], "project": selection["project"], "kind": selection["kind"],
                          "source": {"path": str(path), "sheet": sheet, "row": row_number}, "fields": fields,
                          "raw_cells": [cell for cell in cells if cell["value"] not in (None, "") or cell["formula"]],
                          "tax": tax}
                records.append(record)
                if tax["status"] == "待核":
                    issues.append({**record["source"], "reason": tax["reason"]})
                if "租赁分公司" in norm(fields.get("lessor", {}).get("value")):
                    issues.append({**record["source"], "reason": "租赁分公司内租，需从外租输出分离"})
            if not seen_header:
                issues.append({"path": str(path), "sheet": sheet, "reason": "未识别设备名称和合价明细表头"})
        except (ValueError, KeyError, ImportError, OSError) as exc:
            issues.append({"path": str(path), "sheet": sheet, "reason": str(exc)})
        finally:
            source["sha256_after"] = hashlib.sha256(path.read_bytes()).hexdigest()
            if source["sha256_after"] != before:
                raise RuntimeError(f"Source changed during extraction: {path}")
    return {"period": plan["period"], "records": records, "issues": issues, "sources": sources}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--plan", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    plan = json.loads(args.plan.read_text(encoding="utf-8-sig"))
    output = args.output.resolve()
    forbidden = [Path(s["path"]).resolve().parent for s in plan.get("sheets", [])]
    if plan.get("source_root"):
        forbidden.append(Path(plan["source_root"]).resolve())
    if any(output.is_relative_to(parent) for parent in forbidden):
        raise ValueError("Extraction output must be outside source directories")
    result = extract(plan)
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("x", encoding="utf-8") as handle:
        json.dump(result, handle, ensure_ascii=False, indent=2, allow_nan=False)
    print(json.dumps({"records": len(result["records"]), "issues": len(result["issues"]),
                      "source_sheets": len(result["sources"]), "output": str(output)}, ensure_ascii=False))
    return 0 if not result["issues"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
