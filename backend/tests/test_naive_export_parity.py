"""`TAIL-62`（2026-09-21）· **朴素导出三件的一致性判据**。

⭐ 为什么要有本件：`fl_export` / `jc_export` / `pj_export` 的实现已**下沉到 `naive_export`**
（三件只剩薄委托）⇒ ⛔ "换实现"最容易犯的错是**顺手改了形态**（列宽 / 序号 / 合计）
⇒ 本件把这些**钉住**：改坏一处就会红。

判据来源 = 下沉**之前**抓的黄金母本（`artifacts/golden/_golden_naive_before.json`），
这里把母本里的关键不变量写成断言（⛔ 不读那份 JSON —— 产物目录不该成为测试的依赖）。
"""
from __future__ import annotations

import io

import openpyxl

from app.report_export import fl_export as FL
from app.report_export import jc_export as JC
from app.report_export import naive_export as N
from app.report_export import pj_export as PJ


def _rows(columns, *, money_keys=frozenset(), seq_key=None, n: int = 3) -> list[dict]:
    """按列定义造样例行（金额给**字符串**，好验证"写进去是数字"）。"""
    out = []
    for i in range(1, n + 1):
        r: dict = {}
        for _, key in columns:
            if seq_key and key == seq_key:
                r[key] = None                     # 序号列的值由内核写，传什么都不算
            elif key in money_keys:
                r[key] = f"{i}23.45"
            else:
                r[key] = f"值{i}"
        out.append(r)
    return out


def _sheet(data: bytes):
    return openpyxl.load_workbook(io.BytesIO(data)).active


def _cells(ws) -> list[list]:
    return [[c.value for c in row] for row in ws.iter_rows()]


def test_fl_表头加粗冻结_金额写数字_末行合计():
    rows = _rows(FL.COLUMNS, money_keys=FL.MONEY_KEYS)
    ws = _sheet(FL.build_xlsx(rows))
    cells = _cells(ws)
    assert ws.title == FL.SHEET_NAME
    assert ws.freeze_panes == "A2"
    assert all(c.font and c.font.bold for c in ws[1])          # 表头加粗
    # 金额列必须是**数字**（⛔ 不是 "1,234.56" 那种文本）
    for _, key in FL.COLUMNS:
        if key in FL.MONEY_KEYS:
            i = [k for _, k in FL.COLUMNS].index(key)
            assert isinstance(cells[1][i], float), f"{key} 不是数字：{cells[1][i]!r}"
    assert cells[-1][0] == "合计"                                # 合计行在最后、首格写明
    # ⭐ 2026-10-11（`TAIL-283`）：**「写数字」的列 ≠ 「要合计」的列** —— 本表新增的「审批金额」
    #   是"同一张通知书在每笔上重复出现的对照数" ⇒ 写数字、但**不参与合计**（加起来是假数）。
    #   原断言从 `MONEY_KEYS` 里随便取一个当合计列 ⇒ 两个集合拆开后会**假红**（它其实测的是"哪个集合"）。
    assert FL.TOTAL_KEYS == frozenset({"amount"}), "合计口径变了（本表只该合付款金额）"
    for key in FL.TOTAL_KEYS:
        assert isinstance(cells[-1][[k for _, k in FL.COLUMNS].index(key)], float)
    assert cells[-1][[k for _, k in FL.COLUMNS].index("amount_approved")] in ("", None), \
        "审批金额被合计进去了（它是每笔重复的对照数）"


def test_jc_项目内序号_列宽上限28():
    rows = _rows(JC.COLUMNS, money_keys=JC.MONEY_KEYS, seq_key=JC.SEQ_KEY)
    for r, seq in zip(rows, [3, 1, None]):
        r[JC.SEQ_KEY] = seq
    ws = _sheet(JC.build_xlsx(rows))
    cells = _cells(ws)
    i = [k for _, k in JC.COLUMNS].index(JC.SEQ_KEY)
    assert [r[i] for r in cells[1:1 + len(rows)]] == [3, 1, "—"]   # 全量派生号，不随导出行序变
    # ⛔ 列宽上限是 28（不是内核默认 32）—— 统一成 32 会改掉本页导出形态
    assert max(w.width for w in ws.column_dimensions.values()) <= 28


def test_pj_没有合计行():
    """⛔ 页面没有合计行 ⇒ 导出**不能**凭空补一个（补了 = 假值）。"""
    rows = _rows(PJ.COLUMNS)
    ws = _sheet(PJ.build_xlsx(rows))
    cells = _cells(ws)
    assert not any(r and r[0] == "合计" for r in cells)
    assert len(cells) == 1 + len(rows)                           # 表头 + 数据，没有别的


def test_pj_ord列写int():
    rows = [{**r, "ord": str(i)} for i, r in enumerate(_rows(PJ.COLUMNS), start=1)]
    ws = _sheet(PJ.build_xlsx(rows))
    i = [k for _, k in PJ.COLUMNS].index("ord")
    assert [r[i] for r in _cells(ws)[1:]] == [1, 2, 3]           # 字符串 "1" ⇒ int 1


def test_三件与内核同输出_证明只有一个主人():
    """⭐ 这条是 `TAIL-62` 的**目的**：三件的输出必须**等于**内核直接产出
    （⇒ 谁改内核都改三处；谁在三件里偷偷另写一份，这里就会红）。"""
    for mod, extra in ((FL, {"money_keys": FL.MONEY_KEYS, "total_keys": FL.TOTAL_KEYS}),
                       (JC, {"money_keys": JC.MONEY_KEYS, "total_keys": JC.MONEY_KEYS,
                             "width_max": 28}),
                       (PJ, {"int_keys": PJ.NUMBER_KEYS})):
        rows = _rows(mod.COLUMNS, money_keys=extra.get("money_keys", frozenset()),
                     seq_key=extra.get("seq_key"))
        if mod is JC:
            for r in rows:
                r[JC.SEQ_KEY] = r.get(JC.SEQ_KEY) or "—"
                r["contract_no"] = r.get("contract_no") or "—"
        a = _cells(_sheet(mod.build_xlsx(rows)))
        b = _cells(_sheet(N.build_xlsx(rows, columns=mod.COLUMNS,
                                       sheet_name=mod.SHEET_NAME, **extra)))
        assert a == b, f"{mod.__name__} 与内核输出不一致"
