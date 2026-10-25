"""现行通用Excel读取的值类型、合并范围和错误契约。"""

from datetime import datetime, time
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from app.core.errors import MustReportError
from app.import_kit import readers
from app.import_kit.parse import amount, date_text


def book_bytes(rows, merges=(), *, title="台账", empty_first=False):
    """只构造隔离xlsx，合并范围由用例声明。"""
    wb = Workbook()
    ws = wb.active
    if empty_first:
        ws.title = "空表"
        ws = wb.create_sheet()
    ws.title = title
    for row in rows:
        ws.append(row)
    for merge in merges:
        ws.merge_cells(merge)
    buf = BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


def test_xlsx_values_and_shape_keep_each_callers_policy():
    data = book_bytes([[0, False, datetime(2026, 9, 15, 8, 30), time(8, 30), 45444, None]])
    generic, _ = readers.read_grid(data)
    assert generic == [[0, False, "2026-09-15", "08:30:00", 45444, None]]
    assert amount("1.005") == Decimal("1.01")
    assert date_text("20260915") == "20260915"


def test_merge_range_is_real_and_each_caller_keeps_its_extent():
    rows = [["供应商", "", "", "合法空列"], ["合同总价", "发票金额", "", ""]]
    rows += [[None] * 4 for _ in range(12)]
    rows += [["甲", 100, "", ""], ["乙", None, "", ""]]
    data = book_bytes(rows, ("A1:B1", "B15:B16"))
    generic, _ = readers.read_grid(data, fill_merged_rows=12)
    assert generic[0] == ["供应商", "供应商", None, "合法空列"]
    assert generic[15][1] is None, "通用数据区不得被整表展开"


def test_xls_grid_keeps_coordinates_dates_bool_and_release(monkeypatch):
    """引擎隔离替身（非真实xls），覆盖真实样表没有的日期/布尔/合并块。"""
    import xlrd

    matrix = [[xlrd.sheet.Cell(xlrd.XL_CELL_TEXT, "供应商"), xlrd.sheet.Cell(xlrd.XL_CELL_EMPTY, "")],
              [xlrd.sheet.Cell(xlrd.XL_CELL_DATE, 45444.5), xlrd.sheet.Cell(xlrd.XL_CELL_BOOLEAN, 0)],
              [xlrd.sheet.Cell(xlrd.XL_CELL_NUMBER, 0.0), xlrd.sheet.Cell(xlrd.XL_CELL_NUMBER, 12.5)]]
    sheet = SimpleNamespace(name="xls台账", nrows=3, ncols=2, merged_cells=[(0, 1, 0, 2)],
                            cell=lambda r, c: matrix[r][c])
    closed = []
    fake = SimpleNamespace(nsheets=1, datemode=0, sheet_by_index=lambda i: sheet,
                           sheet_by_name=lambda n: sheet, release_resources=lambda: closed.append(True))
    monkeypatch.setattr(xlrd, "open_workbook", lambda **kwargs: fake)
    data = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
    generic, _ = readers.read_grid(data, fill_merged_rows=1)
    assert generic == [["供应商", "供应商"], ["2024-06-01 12:00:00", "0"], ["0", "12.5"]]
    assert len(closed) == 1


def test_sheet_format_and_failure_contracts_stay_distinct():
    data = book_bytes([], title="空表")
    assert readers.list_sheets(data) == ["空表"]
    with pytest.raises(MustReportError, match="空的"):
        readers.read_grid(data)
    normal = book_bytes([["项目", "金额"], ["甲", 0]])
    assert readers.read_grid(normal, sheet="")[0][1] == ["甲", 0]
    with pytest.raises(MustReportError, match="没有叫"):
        readers.read_grid(data, sheet="不存在")
    assert readers.list_sheets(data) == ["空表"], "通用面按内容识别"
    with pytest.raises(MustReportError, match="既不是"):
        readers.read_grid(b"invalid")
