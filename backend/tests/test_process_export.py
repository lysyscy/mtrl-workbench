# -*- coding: utf-8 -*-
"""过程资料成品（`TAIL-04` 路径① **片 4a**）—— 一本工作簿 / `汇总表` 的离线断言。

形态（唯一主人 = 交付面 `demo/模板/模板清单.md` §5.3 ＋ §二-4 ＋ §四-5 ＋ **裁定⑩**）：
`A 序号`（**B 类**：系统行号，只给数据行编号）｜`B 时间`（`YYYY-MM` **文本**）｜
`C–G 线上入库（不含税）`｜`H–L 线下出库（不含税）`｜`M–N 分公司设备 / 内部设备`｜
`O–P 外租机械（含税）零租 / 包月`｜`Q–R 外租机械（不含税）零租 / 包月`；
尾块 = **4 行一块**：`小计`（16 列垂直求和）｜`合计：`（**四段横向**，四个合并格）｜`脚注 ×2`（`A:R` 合并）。

⭐ 断言分五类（各自独立成档 —— 红了就知道坏在哪一维）：
① **形态**：18 列 / **列序以模板为准**（临时设施 在 周转材料 **前**，裁定②）/ 标题 = `{展示名}材料机械消耗汇总表`；
② **逐格独立复算**（手写期望，⛔ 不问实现）；
③ **两段合计**（小计 = 数据区垂直求和 · 合计 = **四段**横向小计）＋ **尾块随行数派生**（脚注下移、
   下面不留行 / 不留壳）＋ 零残留；
④ **三张表同形态**（裁定⑩）：表①原始 / 表②调整后 / **差异表 = 逐格差额（调整后 − 原始）**
   —— ⚠️ 差额那一维**必须用合成矩阵验**（真机库里当前 0 条调整单 ⇒ 全 0 差额证明不了算法）；
⑤ **拒做**（一个项目都没选 / 项目没有数据 / 取不到展示名）。

⚠️ 取数与形态都走**唯一出口**（`domain.build_matrix` / `report_export.service.export_proc`）——
   ⛔ 不在测试里自己拼 spec 再调引擎（那会把出口的口径漏在断言之外：台账那片实测踩到过）。
"""

from __future__ import annotations

from decimal import Decimal

import openpyxl
import pytest

from app.core.errors import MustReportError
from app.mod06_proc import domain as D
from app.report_export import process_export as PX
from app.report_export import service as EX
from tests.acc_plugin import acc

SHEET = PX.SHEET
WH = "甲仓库"
DISP = "甲公司项目"
YM_A, YM_B, YM_C = "2026-07", "2026-08", "2026-09"        # C = 只有机械手录（材料**空态**）

#: ⭐ **手写期望**（16 列；列序 = 模板序）：`in_主辅材|工具|劳保|临时设施|周转材料` ＋
#: `out_…5` ＋ `mech_0…5` —— ⛔ 不从实现里取。
COLS16 = ("主辅材", "工具", "劳保", "临时设施", "周转材料")
ALL16 = ([f"in_{c}" for c in COLS16] + [f"out_{c}" for c in COLS16]
         + [f"mech_{i}" for i in range(6)])
#: 模板列字母（`C..R` 依次对应上面 16 列）—— 与 `process_export.COL_OF_KEY` **独立**写一遍，
#: 这样"列映射被改坏"能被断言抓住（⛔ 不引用实现里的那份）。
LETTERS16 = "CDEFGHIJKLMNOPQR"


def z(*vals) -> list:
    """16 列期望：给几个就填几个、其余按 `0.00`（**手写期望**用）。

    ⚠️ 口径提醒（域内核现行）：某月**有桶**（该月该来源有行）⇒ 该月缺的那些列是 `0.00`（不是留空）；
       **留空**只属于**空态**（`material_empty=True` ⇒ 材料 10 列全 `None`）。
    """
    out = [Decimal("0.00")] * 16
    for i, v in enumerate(vals):
        out[i] = None if v is None else Decimal(str(v))
    return out


RAW_EXPECT = [
    z("1000.00", "0.00", "0.00", "0.00", "0.00", "0.00", "0.00", "0.00", "0.00", "300.00"),  # 2026-07
    z("0.00", "50.00", "0.00", "200.00"),                                                    # 2026-08
    z(*([None] * 10), "111.00"),                                                             # 2026-09（材料**空态** ＋ 机械手录）
]
#: 调整后：`in_工具` 50 → 80；`in_临时设施` 200 → 150（**下调** ⇒ 差额为负）
ADJ_EXPECT = [
    z("1000.00", "0.00", "0.00", "0.00", "0.00", "0.00", "0.00", "0.00", "0.00", "300.00"),
    z("0.00", "80.00", "0.00", "150.00"),
    z(*([None] * 10), "111.00"),
]
#: ⭐ 差额（调整后 − 原始）—— **逐格**：只 `in_工具` +30、`in_临时设施` −50 有数，
#: 其余格**两边相等** ⇒ 差额 `0`；⚠️ **两边都空**（2026-09 的材料 10 列）⇒ **留空**（⛔ 不写 0）。
DIFF_EXPECT = [
    z(),                                          # 2026-07：两边一模一样 ⇒ 全 0
    z("0.00", "30.00", "0.00", "-50.00"),         # 2026-08：上调 ＋30 · 下调 −50
    z(*([None] * 10)),                            # 2026-09：材料留空；机械两边相等 ⇒ 0
]


def inv_rows() -> list[dict]:
    return [
        {"ym": YM_A, "warehouse": WH, "direction": "入库", "material_property": "主材",
         "amount": Decimal("1000.00")},
        {"ym": YM_A, "warehouse": WH, "direction": "出库", "material_property": "周转材料",
         "amount": Decimal("300.00")},
        {"ym": YM_B, "warehouse": WH, "direction": "入库", "material_property": "工具",
         "amount": Decimal("50.00"), "adj_amount": Decimal("80.00")},
        {"ym": YM_B, "warehouse": WH, "direction": "入库", "material_property": "临时设施",
         "amount": Decimal("200.00"), "adj_amount": Decimal("150.00")},
        # ⛔ 别的仓库 ⇒ 必须被排除（项目键 = 仓库名）
        {"ym": YM_B, "warehouse": "乙仓库", "direction": "入库", "material_property": "主材",
         "amount": Decimal("9999.00")},
        # ⛔ 认不出的物料性质 / 方向不是出入库 ⇒ 跳过（不臆造归属）
        {"ym": YM_B, "warehouse": WH, "direction": "入库", "material_property": "别的性质",
         "amount": Decimal("777.00")},
        {"ym": YM_B, "warehouse": WH, "direction": "调拨", "material_property": "主材",
         "amount": Decimal("888.00")},
    ]


def manual_rows() -> list[dict]:
    row = {"ym": YM_C, "warehouse": WH, "mech_0": Decimal("111.00")}
    for i in range(1, 6):
        row[f"mech_{i}"] = Decimal("0.00")
    return [row]


def matrix(adjusted: bool) -> list[dict]:
    return D.build_matrix(inv_rows=inv_rows(), manual_rows=manual_rows(),
                          history_rows=[], warehouse=WH, adjusted=adjusted, cutoff_ym="")


@pytest.fixture(autouse=True)
def _export_root(tmp_path, monkeypatch):
    """把落盘根改到tmp_path，并替换m06来源与批量cutoff读取；成品断言不变。"""
    from app.mod06_proc import repository as R
    from app.mod06_proc import service as M6

    monkeypatch.setattr(EX, "EXPORT_ROOT", tmp_path / "exports")
    monkeypatch.setattr(R, "load_inv_rows", lambda s: inv_rows())
    monkeypatch.setattr(R, "load_mech", lambda s: manual_rows())
    monkeypatch.setattr(R, "load_history", lambda s=None: [])
    monkeypatch.setattr(R, "load_cutoff", lambda s, warehouse="": "")
    monkeypatch.setattr(R, "load_cutoffs", lambda s: {})
    monkeypatch.setattr(M6, "displays_by_key", lambda s: {WH: DISP})


def export(tab: int, keys=(WH,), **kw):
    """走**唯一出口** `export_proc` ⇒ 产物路径（`None` session：取数已被桩掉）。"""
    return EX.export_proc(None, list(keys), tab, **kw)


def read(path) -> list[list]:
    wb = openpyxl.load_workbook(str(path), data_only=False)
    try:
        ws = wb[SHEET]
        return [[ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
                for r in range(1, ws.max_row + 1)]
    finally:
        wb.close()


def dec(v) -> Decimal:
    return Decimal(str(v)) if v not in (None, "") else Decimal("0.00")


def row16(rows: list[list], r: int) -> list:
    """第 `r` 行（1 基）的 **16 个数据列**（`C..R`）⇒ 原值。"""
    return [rows[r - 1][openpyxl.utils.column_index_from_string(c) - 1] for c in LETTERS16]


def same(got: list, want: list) -> str:
    for i, (g, w) in enumerate(zip(got, want)):
        if w is None:
            if g not in (None, ""):
                return f"第 {i + 1} 个数据列（{LETTERS16[i]}）该**留空**，实得 {g!r}"
        elif g is None or dec(g) != dec(w):
            return f"第 {i + 1} 个数据列（{LETTERS16[i]}）：{g!r} ≠ {w!r}"
    return ""


# ============================================================ ① 形态


@acc("ACC-RPT-EXP-001", "001-05")
def test_十八列且列序照模板_临时设施在周转材料前():
    """⭐ 裁定②「列序**以模板为准**」：模板里 **临时设施（F）在 周转材料（G）前** ——
    与域内 `CATEGORIES` 的顺序**不同**，⛔ 别按域内顺序排。"""
    rows = read(export(1))
    assert len(rows[1]) == 18, f"列数 = {len(rows[1])}"
    leaf = [rows[2][i] for i in range(0, 18)]
    assert leaf[2:7] == ["主辅材", "工具", "劳保", "临时设施", "周转材料"], leaf
    assert leaf[7:12] == ["主辅材", "工具", "劳保", "临时设施", "周转材料"], leaf
    assert rows[1][2] == "线上入库（不含税）" and rows[1][7] == "线下出库（不含税）", rows[1]
    assert rows[1][14] == "外租机械（含税）" and rows[1][16] == "外租机械（不含税）", rows[1]


@acc("ACC-RPT-EXP-001", "001-05")
def test_标题用展示名且文件名含展示名与表名():
    """标题与文件名都用**展示名**（用户 2026-09-18 裁）；文件名还带**表名**与年月。"""
    path = export(1, ym="2026-09")
    rows = read(path)
    assert str(rows[0][0]) == f"{DISP}{PX.TITLE_SUFFIX}", f"标题 = {rows[0][0]!r}"
    name = str(path)
    assert DISP in name and D.tab_name(1) in name and "2026-09" in name, f"文件名 = {name}"


# ============================================================ ② 逐格独立复算


@acc("ACC-RPT-EXP-001", "001-07")
def test_逐格等于手写期望_含材料空态留空():
    """表② 调整后逐格 = **手写期望**；⚠️ 材料空态那一月**留空**（⛔ 不写 0），机械列照写 0。"""
    rows = read(export(1))
    assert len(rows) >= 4 + len(ADJ_EXPECT)
    for i, want in enumerate(ADJ_EXPECT):
        diff = same(row16(rows, 4 + i), want)
        assert not diff, f"第 {i + 1} 行 · {diff}"
    assert [rows[3 + i][1] for i in range(len(ADJ_EXPECT))] == [YM_A, YM_B, YM_C]


@acc("ACC-RPT-EXP-001", "001-07")
def test_序号是系统行号且小计合计脚注行不编号():
    """「序号」= **B 类**（清单 §5.3）：只给**数据行**编号；小计 / 合计 / 脚注行**不编号**。"""
    rows = read(export(1))
    n = len(ADJ_EXPECT)
    assert [rows[3 + i][0] for i in range(n)] == list(range(1, n + 1)), "数据行序号 ≠ 1..N"
    # ⚠️ 只查 **小计 / 合计** 两行 —— 脚注行的 A 列**本来就有**脚注原文（它跨 `A:R` 两行合并，
    #    锚点就在 A），拿它一起查会把"模板固定文案"误判成编号。
    # ⚠️ 另外：`rows` 是**下标**（`rows[r-1]` = 第 r 行），⛔ 别把行号直接当下标用（本文件踩过一次）。
    for rr in (4 + n, 5 + n):                                # 小计 / 合计（1 基行号）
        assert rows[rr - 1][0] in (None, ""), f"第 {rr} 行（小计 / 合计）不该有编号：{rows[rr - 1][0]!r}"
    assert str(rows[6 + n - 1][0] or "").startswith("1、材料费用"), "脚注行的 A 列应是脚注原文"


@acc("ACC-RPT-EXP-001", "001-07")
def test_时间列是页面文本且数字格式为文本():
    """`模板清单` §四-4：时间 = `YYYY-MM` **文本**，且写前把该列数字格式设成文本。"""
    path = export(1)
    rows = read(path)
    wb = openpyxl.load_workbook(str(path))
    try:
        ws = wb[SHEET]
        fmts = {ws.cell(row=4 + i, column=2).number_format for i in range(len(ADJ_EXPECT))}
    finally:
        wb.close()
    assert [rows[3 + i][1] for i in range(len(ADJ_EXPECT))] == [YM_A, YM_B, YM_C]
    assert fmts == {"@"}, f"时间列数字格式 = {fmts}（应为 '@'）"


# ============================================================ ③ 两段合计 ＋ 尾块


@acc("ACC-RPT-EXP-001", "001-07")
def test_小计是数据区垂直求和且合计是四段横向之和():
    """⭐ **两段合计**（清单 §四-5 点名"最容易漏"）：
    `小计` = 数据区**垂直**求和（16 列全给）；`合计：` = **四段横向**小计之和（四个合并格）。"""
    rows = read(export(1))
    n = len(ADJ_EXPECT)
    sub_r, tot_r = 4 + n, 5 + n
    assert str(rows[sub_r - 1][1]) == "小计", rows[sub_r - 1][1]
    assert str(rows[tot_r - 1][1]) == "合计：", rows[tot_r - 1][1]
    for j, letter in enumerate(LETTERS16):
        want = sum((dec(ADJ_EXPECT[i][j]) for i in range(n)), Decimal("0.00"))
        got = dec(rows[sub_r - 1][openpyxl.utils.column_index_from_string(letter) - 1])
        assert got == want, f"小计 {letter} = {got} ≠ {want}"
    for c1, c2, anchor in ((3, 7, 3), (8, 12, 8), (13, 16, 13), (17, 18, 17)):
        seg = sum((dec(rows[sub_r - 1][c - 1]) for c in range(c1, c2 + 1)), Decimal("0.00"))
        got = dec(rows[tot_r - 1][anchor - 1])
        assert got == seg, f"合计 {openpyxl.utils.get_column_letter(anchor)} = {got} ≠ {seg}"
    # ⚠️ 合计只有**四个合并格**有值（其余格是合并影子 / 空的，⛔ 不该各写一份）
    for c in range(3, 19):
        covered = any(c1 <= c <= c2 for c1, c2, _a in ((3, 7, 3), (8, 12, 8), (13, 16, 13), (17, 18, 17)))
        if not covered:
            assert rows[tot_r - 1][c - 1] in (None, ""), f"合计行第 {c} 列不该有值"


@acc("ACC-RPT-EXP-001", "001-07")
def test_尾块随行数派生_脚注紧跟合计且下面不留残留():
    """⭐ 尾块 = **4 行一块**：脚注（`A:R` 跨两行合并）**紧跟**合计行 ⇒ 数据行数一变整块下移；
    ⛔ 尾块以下不许留内容 / 行壳（`TAIL-43` 同族）。"""
    path = export(1)
    rows = read(path)
    wb = openpyxl.load_workbook(str(path))
    try:
        ws = wb[SHEET]
        merges = [str(m) for m in ws.merged_cells.ranges]
        max_row = ws.max_row
    finally:
        wb.close()
    n = len(ADJ_EXPECT)
    tot_r = 5 + n
    assert len(rows) == n + 7, f"总行数 = {len(rows)}（应 = 标题1 ＋ 表头2 ＋ 数据{n} ＋ 尾块4）"
    # 脚注行 = 合计行 ＋ 1（`rows` 是下标 ⇒ 第 6+n 行 = `rows[5+n]`）
    assert str(rows[5 + n][0] or "").startswith("1、材料费用"), "脚注没有紧跟合计行"
    assert any(m.startswith(f"A{tot_r + 1}:") and m.endswith(f"R{tot_r + 2}") for m in merges), merges
    assert all(any(m.startswith(f"{openpyxl.utils.get_column_letter(c)}{tot_r}:") for m in merges)
               for c in (3, 8, 13, 17)), merges
    assert max_row == n + 7, f"max_row = {max_row}（尾块以下还有行壳）"


# ============================================================ ④ 三张表同形态（裁定⑩）


@acc("ACC-RPT-EXP-001", "001-06")
def test_表一表二同形态只有取数不同():
    """裁定⑩：三张表**形态完全相同**（同一模板 / 同一列映射 / 同一尾块），只有**数据**不同。"""
    r0, r1 = read(export(0)), read(export(1))
    assert len(r0) == len(r1), "两张表的总行数不同 ⇒ 形态分叉了"
    assert r0[0][0] == r1[0][0] and r0[1] == r1[1] and r0[2] == r1[2], "标题 / 表头不一致"
    for i, want in enumerate(RAW_EXPECT):
        diff = same(row16(r0, 4 + i), want)
        assert not diff, f"[表①] 第 {i + 1} 行 · {diff}"


@acc("ACC-RPT-EXP-001", "001-06")
def test_差异表逐格等于调整后减原始_含负数与空态():
    """⭐ **合成矩阵**验差额（真机库里当前 0 条调整单 ⇒ 全 0 差额证明不了算法）：
    逐格 = 调整后 − 原始；下调 ⇒ **负数**；某一格**两边都空** ⇒ **留空**（⛔ 不写 0）。"""
    a, r = matrix(True), matrix(False)
    assert D.difference_rows(r, a)[1]["in_工具"] == Decimal("30.00"), "上调差额"
    assert D.difference_rows(r, a)[1]["in_临时设施"] == Decimal("-50.00"), "下调差额应为负"
    assert D.difference_rows(r, a)[2]["in_主辅材"] is None, "两边都空 ⇒ 差额应留空（⛔ 不写 0）"
    for i, want in enumerate(DIFF_EXPECT):
        diff = same(row16(read(export(2)), 4 + i), want)
        assert not diff, f"[差异表] 第 {i + 1} 行 · {diff}"


@acc("ACC-RPT-EXP-001", "001-06")
def test_差异表与另两张同形态且文件名带表名():
    """差异表**沿用同一张模板**（18 列 / 同一尾块），文件名带表名（裁定⑩「最多文件名不一样」）。"""
    path = export(2)
    rows = read(path)
    assert len(rows[1]) == 18
    assert str(rows[0][0]) == f"{DISP}{PX.TITLE_SUFFIX}"
    assert D.tab_name(2) in str(path), f"文件名 = {path}"
    n = len(DIFF_EXPECT)
    assert str(rows[5 + n - 1][1]) == "合计：", "差异表也必须有尾块"


# ============================================================ ⑤ 拒做


@acc("ACC-RPT-EXP-001", "001-08")
def test_拒做一个项目都没选():
    with pytest.raises(MustReportError) as e:
        export(1, keys=())
    assert "一个项目都没选" in str(e.value)


@acc("ACC-RPT-EXP-001", "001-08")
def test_拒做项目没有数据():
    with pytest.raises(MustReportError) as e:
        export(1, keys=("查无此仓库",))
    assert "没有数据" in str(e.value)


@acc("ACC-RPT-EXP-001", "001-08")
def test_拒做取不到展示名(monkeypatch):
    """⛔ 不拿仓库名冒充展示名（用户 2026-09-18 裁）⇒ 取不到就**拒做并说清原因**。"""
    from app.mod06_proc import service as M6

    monkeypatch.setattr(M6, "displays_by_key", lambda s: {})   # 该项目没建档
    with pytest.raises(MustReportError) as e:
        export(1)
    assert "不在档案里" in str(e.value) and "展示名" in str(e.value)


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(pytest.main([__file__, "-v"]))
