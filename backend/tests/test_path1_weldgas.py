# -*- coding: utf-8 -*-
"""`TAIL-04` **路径①（模板填充）** 片 1：**焊材 / 气体提取成品导出** —— 离线断言（不连库、不连 i8）。

绑的判据：
- `ACC-X-EXP-001/001-08`（**导出内容须与当前筛选一致**）—— 成品的行 = 所筛选项目（项目状态 ＋ 档案顺序），
  与页面 `build_extract_data` **同一份口径**；
- `ACC-X-EXP-001/001-05`（**导出成品 / 提示语必须与模板实际形态一致**）—— 1 个工作簿 ＋ 2 个工作表
  （`焊材` / `气`），每表 15 列 = 项目 | 结余 | 1月…12月 | 合计；⛔ 提示语不得声称不符的结构；
- `ACC-X-EXP-001/001-09`（**数据区定位与尾块派生**）：起点认**标记文字** `DATA_START`（⛔ 不写死行号）；
  尾块随数据行数派生（n=0 时只剩「表头 ＋ 尾块」、不留空壳行）；
- `ACC-X-EXP-001/001-10`（**零残留**）：模板自带的示例数据（项目名 / 数字）一格都不得进成品；
- `ACC-X-EXP-001/001-11`（**以数值覆盖**）：显示区一律算好的**数值**（⛔ 不留公式）；算出来是 0 的格子**留空**；
- `ACC-X-EXP-001/001-12`（**拒做**）：拿不到报告期 / 数据为空 / 对账不过 ⇒ 显式报中文原因、**不产出文件**。

⚠️ 其余三条用 `@acc()` **显式声明不绑定**：
· 「结余」列当前**如实留空**（手填列、录入入口待建 —— 属**临时口径**，进判据反而会绊住以后）；
· `export_report_ym` 的四种时间模式（纯函数口径）；· 收发结存 zip 弹窗的接线（**胶水**，无独立口径）。

⚠️ 三处易踩的坑（本片实测）：
1. **模板里带着示例数据** —— 只写前 n 行、不清尾部 ⇒ 用户读到**示例 / 上月**的数（`TAIL-43` 同族）；
2. `openpyxl` 的 `insert_rows` / `delete_rows` **搬值与样式、不搬合并区** ⇒ 合并区要自己按 offset 平移；
3. `verify_filled` 必须**读两遍**工作簿（`data_only=True` 看值、`False` 看公式）——
   只用前者时「还有公式」那条**永远不报**（死代码）。
"""

from __future__ import annotations

import re
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from app.core.errors import MustReportError
from app.mod07_setl import extract as X
from app.mod07_setl.service import build_extract_data, build_extract_export, parse_extract_state
from app.mod07_setl.view import render_extract_page
from app.report_export import resources as R
from app.report_export import tpl_fill
from app.report_export import weldgas_export as W
from tests.acc_plugin import acc
from tests.m07_helpers import (
    ARCHIVED_WAREHOUSE,
    EXTRACT_EXPECT,
    FakeSession,
    extract_dataset,
)

CASE = "ACC-X-EXP-001"
ROOT = Path(__file__).resolve().parents[2]
SHELL = (ROOT / "backend" / "app" / "web" / "index.html").read_text(encoding="utf-8")
VIEW = (ROOT / "backend" / "app" / "mod07_setl" / "view.py").read_text(encoding="utf-8")


def _act(name: str) -> str:
    """取壳里某个动作体（`\\n <名>:` 起，**花括号配平** —— ⛔ 不用"下一个动作"当终点）。"""
    i = SHELL.find(f"\n {name}:")
    assert i >= 0, f"壳里找不到动作 `{name}:`"
    depth, k = 0, SHELL.find("{", i)
    assert k > 0, f"`{name}` 里找不到起始花括号"
    while k < len(SHELL):
        if SHELL[k] == "{":
            depth += 1
        elif SHELL[k] == "}":
            depth -= 1
            if depth == 0:
                return SHELL[i:k + 1]
        k += 1
    raise AssertionError(f"`{name}` 的花括号没配平")


def _plan(*, status: str = "在建", period: str = "2026-08") -> dict:
    return build_extract_export(FakeSession(extract_dataset()),
                               parse_extract_state({"status": status, "ext_period": period}))


def _build(tmp_path, plan: dict) -> str:
    return W.build_xlsx(plan["grids"], tmp_path, plan["report_ym"])


def _sheet(path, name: str) -> list[list]:
    wb = openpyxl.load_workbook(str(path), data_only=True)
    try:
        ws = wb[name]
        return [[ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
                for r in range(1, ws.max_row + 1)]
    finally:
        wb.close()


# ============================================================ 形态（001-05）


@acc(CASE, "001-05")
def test_成品形态与模板登记一致(tmp_path):
    """绑 `001-05`：**真读回成品**，它的工作表 / 列结构必须与 `模板-焊材气体.xlsx` 登记的一致。"""
    plan = _plan()
    out = Path(_build(tmp_path, plan))
    assert out.name == W.file_name(plan["report_ym"]) == "焊材气体提取表_2026年8月.xlsx", \
        f"命名跑偏：{out.name}"

    wb = openpyxl.load_workbook(out)
    try:
        assert list(wb.sheetnames) == ["焊材", "气"], f"工作表不对：{wb.sheetnames}"
        for sheet in ("焊材", "气"):
            ws = wb[sheet]
            # 表头**照模板保留**（⛔ 不改用户的外观）：结余 ＋ 1月…12月 ＋ 合计
            head = [ws.cell(row=1, column=c).value for c in range(1, 16)]
            assert head[1] == "结余" and head[2:14] == [f"{m}月" for m in range(1, 13)] \
                and head[14] == "合计", f"[{sheet}] 表头与模板不符：{head}"
            assert ws.max_column == 15, f"[{sheet}] 列数 ≠ 15：{ws.max_column}"
            # 数据区首行 = 标记文字所在行（模板 A2）—— 标记本身**不许留在产物里**
            assert ws.cell(row=2, column=1).value not in (None, "", tpl_fill.MARKER), \
                f"[{sheet}] A2 还是占位文字：{ws.cell(row=2, column=1).value!r}"
    finally:
        wb.close()


# ============================================================ 与筛选一致（001-08）


@acc(CASE, "001-08")
def test_成品行即所筛选项目且与页面同一份口径(tmp_path):
    """绑 `001-08`：行 = **当前筛选**下的项目（状态过滤 ＋ 档案顺序），与页面同源（⛔ 不各筛一份）。"""
    page = build_extract_data(FakeSession(extract_dataset()),
                              parse_extract_state({"status": "在建", "ext_period": "2026-08"}))
    plan = _plan()
    for blk in page["blocks"]:
        kind = blk["kind"]
        assert [r["warehouse"] for r in plan["grids"][kind]["rows"]] == \
            [r["warehouse"] for r in blk["rows"]], f"[{kind}] 导出与页面的行集合/行序不一致"

    out = Path(_build(tmp_path, plan))
    rows = _sheet(out, "焊材")
    n = len(plan["grids"]["焊材"]["rows"])
    assert rows[1 + n][0] is None and rows[1 + n][1] is None, "合计行的名称/结余格应为空（照模板）"
    # ⭐ 切到「归档」⇒ 该仓**没有提取数据** ⇒ 两类都 0 行；成品里数据区必须真的**空出来**
    arch = _plan(status="归档")
    assert all(not g["rows"] for g in arch["grids"].values()), "归档态不该有提取数据行"
    out2 = Path(_build(tmp_path, arch))
    for sheet in ("焊材", "气"):
        got = _sheet(out2, sheet)
        assert len(got) == 2, f"[{sheet}] 0 行时应当只剩「表头 ＋ 合计行」，实得 {len(got)} 行"
        # 合计行（表头正下方那一行）的月份列 = 各月列之和 = 0（模板里那格是 `=SUM(...)` ⇒ 空数据也是 0）
        assert all(v in (None, 0, Decimal("0.00")) for v in got[1][2:14]), \
            f"[{sheet}] 0 行时合计行的月份列只能是 0：{got[1][2:14]}"


# ============================================================ 逐格对账（独立手写期望）


@acc(CASE, "001-11")
def test_成品显示区逐格等于独立复算(tmp_path):
    """**两条独立路径**：① `m07_helpers.EXTRACT_EXPECT`（手写期望，不问实现）；
    ② 从成品显示区读回 —— 各月 / 合计 / 合计行必须逐格相同。"""
    plan = _plan()
    out = Path(_build(tmp_path, plan))
    months = [f"2026-{m:02d}" for m in range(1, 13)]
    for kind, sheet in W.SHEET_OF_KIND.items():
        rows = _sheet(out, sheet)
        whs = [r["warehouse"] for r in plan["grids"][kind]["rows"]]
        assert [rows[1 + i][0] for i in range(len(whs))] == whs, f"[{sheet}] 行序不对"
        for i, wh in enumerate(whs):
            row = rows[1 + i]
            for j, ym in enumerate(months):
                want = Decimal(EXTRACT_EXPECT[kind].get(wh, {}).get(ym, "0"))
                got = row[2 + j]
                assert (got is None) == (want == 0), \
                    f"[{sheet}] {wh} {ym}：{got!r}（0 应留空，⛔ 不写 0）"
                if want:
                    assert abs(Decimal(str(got)) - want) < Decimal("0.005"), \
                        f"[{sheet}] {wh} {ym}：{got} ≠ {want}"
            total = sum((Decimal(v) for v in EXTRACT_EXPECT[kind].get(wh, {}).values()),
                        Decimal("0.00"))
            assert abs(Decimal(str(row[14])) - total) < Decimal("0.005"), \
                f"[{sheet}] {wh} 合计 {row[14]} ≠ {total}"
        tail = rows[1 + len(whs)]
        for j, ym in enumerate(months):
            want = sum((Decimal(EXTRACT_EXPECT[kind].get(w, {}).get(ym, "0"))
                        for w in EXTRACT_EXPECT[kind]), Decimal("0"))
            got = tail[2 + j] or 0
            assert abs(Decimal(str(got)) - want) < Decimal("0.005"), \
                f"[{sheet}] 合计行 {ym}：{got} ≠ {want}"


@acc()
def test_结余列无来源时如实为空(tmp_path):
    """⚠️ 2026-09-18 用户裁定：**结余 = 历史结余、由用户手动录入** ⇒ 系统无来源。

    本轮**如实为空**（⛔ 不写 0、⛔ 不把模板里的示例值带走）；录入入口落地后由 `balances` 接缝喂进来。
    """
    plan = _plan()
    out = Path(_build(tmp_path, plan))
    for sheet in ("焊材", "气"):
        rows = _sheet(out, sheet)
        for row in rows[1:]:
            assert row[1] in (None, ""), f"[{sheet}] 结余列应留空，实得 {row[1]!r}"


# ============================================================ 引擎：尾块随行数派生 ＋ 零残留


@acc(CASE, "001-10")
def test_模板自带的示例值一格不留(tmp_path):
    """⭐ 与 `TAIL-43` 同族的守卫：模板里那份**示例数据**（项目名 / 数字）**一格都不许**进产物。

    ⚠️ 断言**不写死**模板里的具体名字 —— 现读模板、再逐个字号查产物（模板换了也不会假红）。
    """
    tpl = openpyxl.load_workbook(R.WELD_TEMPLATE, data_only=True)
    try:
        samples: dict[str, set] = {}
        for sheet in ("焊材", "气"):
            ws = tpl[sheet]
            area = tpl_fill.locate_data_area(ws, marker=tpl_fill.MARKER)
            names = {str(ws.cell(row=r, column=1).value or "").strip()
                     for r in range(area["data_start"], area["data_end"] + 1)}
            nums = set()
            for r in range(area["data_start"], area["data_end"] + 1):
                for c in range(2, area and ws.max_column + 1):
                    v = ws.cell(row=r, column=c).value
                    if isinstance(v, (int, float)):
                        nums.add(round(float(v), 2))
            samples[sheet] = (names - {""}, nums)
    finally:
        tpl.close()
    assert any(s[0] for s in samples.values()), "模板数据区里没读到项目名（取证前提不成立）"

    out = Path(_build(tmp_path, _plan()))
    wb = openpyxl.load_workbook(out, data_only=True)
    try:
        for sheet, (names, nums) in samples.items():
            ws = wb[sheet]
            got_names, got_nums = set(), set()
            for row in ws.iter_rows(min_row=2, values_only=True):
                got_names.add(str(row[0] or "").strip())
                got_nums |= {round(float(v), 2) for v in row[1:]
                             if isinstance(v, (int, float))}
            assert not (names & got_names), \
                f"[{sheet}] 模板示例项目名残留：{sorted(names & got_names)}"
            # ⚠️ 0 不算残留：合计列就是 `=SUM(B:N)`，空数据的总和本来就是 0（模板同表现）
            residue = (nums & got_nums) - {0.0}
            assert not residue, f"[{sheet}] 模板示例数字残留：{sorted(residue)}"
    finally:
        wb.close()


@acc(CASE, "001-09")
def test_数据区起点认标记文字而不是抄坐标():
    """⭐ 方案 A（《说明书》§7.1）：起点认 `DATA_START` 这串**文字** —— 把标记挪一行，定位就得跟着挪。

    ⛔ 抄坐标（把 2 写死）在真机上会**静默错位**：模板挪一次列/插一行，产物就整体串行。
    """
    wb = openpyxl.load_workbook(R.WELD_TEMPLATE)
    try:
        ws = wb["焊材"]
        assert tpl_fill.locate_data_area(ws)["data_start"] == 2, "本模板的标记该在 A2"
        ws["A3"] = tpl_fill.MARKER                       # 把标记整体下移一行
        ws["A2"] = "某项目"
        assert tpl_fill.locate_data_area(ws)["data_start"] == 3, "定位没跟着标记走（抄了坐标）"
        ws["A4"] = tpl_fill.MARKER                       # 两处标记 ⇒ 起点不唯一 ⇒ 必须报错
        with pytest.raises(ValueError) as e:
            tpl_fill.locate_data_area(ws)
        assert "不唯一" in str(e.value), f"没给中文原因：{e.value}"
    finally:
        wb.close()


@acc(CASE, "001-09")
def test_引擎_尾块随行数派生且零行也不留残骸(tmp_path):
    """引擎级：n = 0 / 1 / 7（= 样板行数）/ 12 四种行数都要**尾块紧贴、下面一格不留**。

    ⚠️ spec 必须**覆盖尾块的每一个汇总格** —— 漏一格就会被引擎的守卫拦下（实测：n=0 时
       尾块上移到第 2 行、原来第 9 行的 `=SUM(D2:D8)` 变成 `D2`，只写 C/O 两列 ⇒ 当场报错）。
       这正是那条守卫存在的理由：**汇总值必须由调用方算好后给**，⛔ 不许留个算不出来的公式。
    """
    template = R.WELD_TEMPLATE
    months = [chr(ord("C") + j) for j in range(12)]
    data_cells = {**{c: Decimal("1.00") for c in months}, "O": Decimal("1.00")}
    for n in (0, 1, 7, 12):
        spec = {"sheet": "焊材",
                "rows": [(f"项目{i:02d}", dict(data_cells)) for i in range(n)],
                "tail": {0: (None, {**{c: Decimal(str(n)) for c in months},
                                    "O": Decimal(str(n))})}}
        out = Path(tmp_path) / f"n{n}.xlsx"
        plan = tpl_fill.fill_workbook(template, out, [spec])
        assert tpl_fill.verify_filled(out, plan) == [], f"n={n} 自己的产物没过对账"
        ws = openpyxl.load_workbook(out, data_only=True)["焊材"]
        assert ws.cell(row=1, column=1).value == "焊材", "表头行不该被碰"      # 模板 R1 是表头
        want_last = 2 + n                                                  # 数据首行 2，尾块紧随其后
        assert ws.cell(row=want_last, column=3).value == float(n), \
            f"n={n} 尾块不在第 {want_last} 行：{ws.cell(row=want_last, column=3).value!r}"
        for r in range(want_last + 1, ws.max_row + 1):
            assert all(ws.cell(row=r, column=c).value is None for c in range(1, 16)), \
                f"n={n} 第 {r} 行（尾块以下）有残留"
        # ⚠️ 光没有"值"还不够：`delete_rows` 会留下**行属性空壳**（空 `<row>` 元素）——
        #    那会让 Excel 的 UsedRange 多算几行（用户看到数据下面还有几行空的）
        shells = sorted(r for r in ws.row_dimensions if r > want_last)
        assert not shells, f"n={n} 还有行属性空壳：{shells}"


@acc(CASE, "001-10")
def test_显示区对账抓得出残留(tmp_path):
    """绑 `001-10`（零残留）：尾块以下塞一个旧值 ⇒ `verify_filled` 必须报红。"""
    problems = _corrupt(tmp_path, _put_residue)
    assert any("尾块以下" in x for x in problems), f"没抓住残留：{problems[:3]}"


@acc(CASE, "001-11")
def test_显示区对账抓得出公式残留(tmp_path):
    """绑 `001-11`（以数值覆盖）：数据格里留公式 ⇒ `verify_filled` 必须报红。"""
    problems = _corrupt(tmp_path, _put_formula)
    assert any("还是公式" in x for x in problems), f"没抓住公式残留：{problems[:3]}"


def _put_residue(ws) -> None:
    ws.cell(row=ws.max_row + 2, column=1).value = "上一版的仓库"


def _put_formula(ws) -> None:
    ws.cell(row=2, column=13).value = "=SUM(B2:N2)"


def _corrupt(tmp_path, mutate) -> list[str]:
    """造一份产物 → 按 `mutate` 故意弄坏 → 返回守卫报出的问题（⛔ 不能是"看着对"）。"""
    plan = _plan()
    out = Path(_build(tmp_path, plan))
    good = tpl_fill.fill_workbook(R.WELD_TEMPLATE, out, W.build_specs(plan["grids"]))
    assert tpl_fill.verify_filled(out, good) == [], "干净的产物应当对账通过"
    wb = openpyxl.load_workbook(out)
    mutate(wb["焊材"])
    wb.save(out)
    wb.close()
    return tpl_fill.verify_filled(out, good)


@acc(CASE, "001-12")
def test_拿不到报告月时显式拒做(tmp_path):
    """⛔ 绝不静默交付：报告月拿不到（该筛选下没有数据）⇒ 中文原因、**不产出文件**。"""
    plan = _plan()
    with pytest.raises(MustReportError) as e:
        W.build_xlsx(plan["grids"], tmp_path, "")
    assert "报告月" in str(e.value), f"没给中文原因：{e.value}"
    assert not list(Path(tmp_path).rglob("*.xlsx")), "拒做时不该留下成品"


@acc()
def test_报告月口径_单月跨月年与全部():
    """报告月 = **止月**；「全部」⇒ 取数据里的最新一期（⛔ 不回落成"本月"）。"""
    from app.mod07_setl.service import export_report_ym

    assert export_report_ym("2026-08", ()) == "2026-08"
    assert export_report_ym("2026-06~2026-08", ()) == "2026-08"
    assert export_report_ym("2026", ()) == "2026-12"
    assert export_report_ym("", ("2026-06", "2026-09", "2026-08")) == "2026-09"
    assert export_report_ym("", ()) == "", "什么都没有 ⇒ 空（由实现层拒做）"


# ============================================================ 壳侧（两个页面都要有真保存）


@acc(CASE, "001-05")
def test_壳侧提取导出真下载且提示与模板形态一致(tmp_path):
    """绑 `001-05`：壳里的 `expExt` / `saveExt` 必须**真下载**，且提示语 = **模板的实际形态**。

    ⚠️ 断言只钉这两个动作**自己的作用域**（花括号配平区间）—— 别处的同名串不得喂饱它。
    """
    dialog = _act("expExt")
    save = _act("saveExt")
    assert "createObjectURL" in save and ".download" in save, "没真的触发下载（拿不到文件）"
    assert "/api/m07/extract/export?status=" in save, "`saveExt` 没打导出端点"
    assert "x-export-filename" in save, "壳侧没读后端带出的文件名（那就是自己拼了一份命名）"
    assert "'application/json'" in save, "没按 `content-type` 分流失败原因"
    assert "fExt.status" in save and "fExt.ym" in save, "筛选没带给后端（会导出成别的集合）"
    # ⭐ 提示语必须与模板实际形态一致：2 个工作表（焊材 / 气）＋ 列 = 项目|结余|1月…12月|合计
    assert "气" in dialog and "工作表" in dialog, f"提示没写工作表结构：{dialog[:120]}"
    for token in ("结余", "1月…12月", "合计"):
        assert token in dialog, f"提示缺成品的列：{token}"
    assert "不是</b>成品的列结构" in dialog or "不是" in dialog, "必须说清下表 ≠ 成品的列结构"


@acc(CASE, "001-05")
def test_服务端页的导出弹窗也是真下载(tmp_path):
    """绑 `001-05`：`/m07/extract`（导航真正指向的那一页）的导出弹窗**也得有真保存按钮**。

    ⚠️ 这是**真渲染**出来的 HTML（不是读源码字符串）：⛔ 只读源码会漏掉"渲染时没拼进去"这类问题。
    """
    plan = _plan()
    html = render_extract_page(build_extract_data(
        FakeSession(extract_dataset()),
        parse_extract_state({"status": "在建", "ext_period": "2026-08"})))
    assert "/api/m07/extract/export?status=" in html, "服务端页的导出弹窗没接端点"
    # ⚠️ 查 URL 时先把 HTML 实体还原 —— `esc()` 会把 `&` 写成 `&amp;`
    #    （HTML 属性里的合法写法，浏览器解析时还原；⛔ 但**不能因此不转义**：period 来自查询串）
    plain = html.replace("&amp;", "&")
    assert "/api/m07/extract/export?status=在建&period=2026-08" in plain, \
        "端点没带上当前筛选（会导出成别的集合）"
    assert "&amp;period=2026-08" in html, "URL 没过 `esc()` 出口（period 来自查询串 ⇒ 有注入口子）"
    assert re.search(r"dl\('[^']*extract/export[^']*'", html), "保存按钮没绑到下载函数"
    assert "x-export-filename" in html, "页脚脚本没读后端带出的文件名"
    assert "_t(" in html, "页脚脚本没带转义出口（后端原因里的尖括号会破坏页面）"
    for token in ("结余", "1月…12月", "工作表"):
        assert token in html, f"提示缺成品的形态：{token}"
    assert plan["report_ym"] == "2026-08"


@acc()
def test_服务端页的收发结存zip弹窗也补了真保存():
    """同一批补的：`/m07` 的「导出分仓库 zip」弹窗原来**只有关闭按钮**（点不出文件）。"""
    from app.mod07_setl.service import build_page_data, parse_state
    from app.mod07_setl.view import render_setl_page
    from tests.m07_helpers import dataset

    html = render_setl_page(build_page_data(FakeSession(dataset()),
                                            parse_state({"period": "2026-08"})))
    assert '<dialog class="dlg" id="zipDlg">' in html, "找不到 zip 弹窗"
    assert "/api/m07/settlement/export?period=2026-08" in html.replace("&amp;", "&"), \
        "zip 弹窗还是死按钮（没接导出端点 / 没带期间）"
    assert re.search(r"dl\('[^']*settlement/export[^']*'", html), "保存按钮没绑到下载函数"


__all__: list[str] = []
