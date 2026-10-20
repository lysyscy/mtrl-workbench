"""Guard source files, formulas, styles and worksheet structure during Excel edits."""
import hashlib
import importlib.util
from pathlib import Path
import tempfile
import unittest

from zipfile import ZipFile
import xml.etree.ElementTree as ET

import openpyxl
from openpyxl.styles import Font

SCRIPT = Path(__file__).resolve().parents[2] / "skills/machinery-fee-consolidation/scripts/apply_template_edits.py"
spec=importlib.util.spec_from_file_location("writer",SCRIPT)
writer=importlib.util.module_from_spec(spec);spec.loader.exec_module(writer)


class TemplateGuards(unittest.TestCase):
    def prepare(self, temp):
        root=Path(temp)/"sources";root.mkdir()
        source=root/"synthetic.xlsx"
        book=openpyxl.Workbook();book.active.title="表1"
        for name in ["表2","表3","表4"]:book.create_sheet(name)
        for sheet in book:
            sheet["A5"]=1;sheet["B5"]="合成原值";sheet["B5"].font=Font(name="宋体",bold=True)
            sheet["C5"]="=A5*2";sheet.row_dimensions[5].height=27
            sheet.column_dimensions["B"].width=22;sheet.merge_cells("A1:C1")
        book.save(source)
        plan={"template":str(source),"template_sha256":hashlib.sha256(source.read_bytes()).hexdigest(),
              "source_roots":[str(root)],"edits":[{"sheet":"表2","cell":"B5","type":"text","value":"新值<&>"},
                                                   {"sheet":"表2","cell":"C5","type":"formula_cache","value":2}]}
        return source,plan

    def test_preserves_other_parts_styles_and_formula(self):
        with tempfile.TemporaryDirectory() as temp:
            source,plan=self.prepare(temp);output=Path(temp)/"new.xlsx"
            writer.apply(plan,output)
            with ZipFile(source) as a,ZipFile(output) as b:
                parts=writer.sheet_parts(a)
                self.assertEqual([p for p in a.namelist() if a.read(p)!=b.read(p)],[parts["表2"]])
                # Generalize the old batch verifier: protect all sheet structure outside edited values.
                ns={"m":writer.MAIN}
                before,after=[ET.fromstring(z.read(parts["表2"])) for z in (a,b)]
                cells=[{c.attrib["r"]:c for c in tree.findall(".//m:c",ns)} for tree in (before,after)]
                self.assertEqual(cells[0].keys(),cells[1].keys())
                for coordinate in cells[0]:
                    left,right=cells[0][coordinate],cells[1][coordinate]
                    self.assertEqual(left.attrib.get("s"),right.attrib.get("s"))
                    formulas=[cell.find("m:f",ns) for cell in (left,right)]
                    self.assertEqual(*[ET.tostring(f) if f is not None else None for f in formulas])
                    if coordinate in {"B5","C5"}:
                        for cell in (left,right):
                            cell.attrib.pop("t",None)
                            for child in list(cell):cell.remove(child)
                self.assertEqual(ET.tostring(before),ET.tostring(after))
            book=openpyxl.load_workbook(output,data_only=False)
            self.assertEqual(book["表2"]["B5"].value,"新值<&>")
            self.assertEqual(book["表2"]["C5"].value,"=A5*2")
            self.assertEqual(book["表2"].row_dimensions[5].height,27)
            book.close()
            cache=openpyxl.load_workbook(output,data_only=True)
            self.assertEqual(cache["表2"]["C5"].value,2);cache.close()

    def test_protects_source_existing_output_headers_and_sheet1(self):
        with tempfile.TemporaryDirectory() as temp:
            source,plan=self.prepare(temp)
            with self.assertRaises(ValueError):writer.apply(plan,source.parent/"new.xlsx")
            output=Path(temp)/"new.xlsx";output.write_text("preserve",encoding="utf-8")
            with self.assertRaises(FileExistsError):writer.apply(plan,output)
            self.assertEqual(output.read_text(encoding="utf-8"),"preserve")
            output.unlink()
            plan["edits"][0]["sheet"]="表1"
            with self.assertRaises(ValueError):writer.apply(plan,output)
            self.assertFalse(output.exists())
            plan["edits"][0]["sheet"]="表2";plan["edits"][0]["cell"]="B4"
            with self.assertRaises(ValueError):writer.apply(plan,output)
            self.assertFalse(output.exists())

    def test_stale_source_formula_replacement_and_nan_refused(self):
        with tempfile.TemporaryDirectory() as temp:
            source,plan=self.prepare(temp);output=Path(temp)/"new.xlsx"
            plan["template_sha256"]="stale"
            with self.assertRaises(ValueError):writer.apply(plan,output)
            plan["template_sha256"]=hashlib.sha256(source.read_bytes()).hexdigest()
            plan["edits"]=[{"sheet":"表2","cell":"C5","type":"number","value":2}]
            with self.assertRaises(ValueError):writer.apply(plan,output)
            plan["edits"]=[{"sheet":"表2","cell":"A5","type":"number","value":float("nan")}]
            with self.assertRaises(ValueError):writer.apply(plan,output)
            self.assertFalse(output.exists())


if __name__=="__main__":unittest.main()
