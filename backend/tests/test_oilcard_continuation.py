"""接续真实服务：来源保真、隔离明确确认、重跑保护、补档不重记及成品导出。"""
from copy import deepcopy
from io import BytesIO
import json
from uuid import uuid4

import openpyxl
import pytest
from sqlalchemy import select

from app.core.db import session_scope
from app.core.errors import DataConflictError, MustReportError
from app.mod13_oilcard import continuation, invoices, ledger, reporting
from app.mod13_oilcard.models import OilContinuationSource, OilEntry, OilInvoice
from tests.test_oilcard_archives import scene, create  # noqa: F401
from tools.oneoff.oilcard_continuation import synthetic_manifest, build_manifest


def stage(manifest):
    with session_scope() as session:
        return continuation.stage_source(session, manifest)


def confirm(source_id, decisions, **kwargs):
    with session_scope() as session:
        return continuation.confirm_source(session, source_id, decisions, **kwargs)


def target(result, key):
    return next(row for row in result['rows'] if row['key'] == key)['confirmation']['targets'][0]['id']


@pytest.fixture
def enrolled(scene):
    client, _ = scene
    source_id = 'synthetic-463-test'
    with session_scope() as session:
        manifest, decisions, content, ids = synthetic_manifest(session, source_id)
        before = ledger.reconciliation(session)
        staged = continuation.stage_source(session, manifest)
        assert ledger.reconciliation(session) == before
    return client, source_id, manifest, decisions, content, ids, staged


def test_all_real_candidates_stay_pending_and_never_invent_facts(enrolled):
    client, source_id, manifest, _, _, _, staged = enrolled
    assert staged['confirmed_count'] == 0 and staged['pending_count'] == 13
    assert staged['pending_source_amounts'] == {'recharge': '15670.00', 'invoice': '10123.00', 'transfer': '30.00'}
    assert staged['pending_is_booked'] is False
    assert stage(deepcopy(manifest)) == staged
    assert client.get('/api/m13/reconciliation').json()['total_amount'] == '0.00'
    assert client.get('/api/m13/invoices').json()['total_amount'] == '0.00'
    pending = next(r for r in staged['rows'] if r['key'] == 'unknown-invoice')
    assert pending['raw'] == {'amount': '123', 'explicit_date': None, 'original_date_text': '原始数据'}
    with pytest.raises(MustReportError, match='发票号码'):
        confirm(source_id, [{'key': 'unknown-invoice', 'note': '缺票号与日期仍不能入账'}])
    assert stage(manifest) == staged


def test_confirmed_drill_repeated_source_pdf_allocation_reports_and_real_exports(enrolled, tmp_path, monkeypatch):
    client, source_id, manifest, decisions, content, ids, _ = enrolled
    first = confirm(source_id, decisions)
    invoice_id = target(first, 'invoice')
    invoice = client.get('/api/m13/invoices').json()['invoices'][0]
    assert invoice['id'] == invoice_id and invoice['amount'] == '10000.00'
    assert [a['amount'] for a in invoice['allocations']] == ['6000.00', '4000.00']
    pdf = {'key': 'pdf', 'invoice_id': invoice_id, 'version': invoice['version'], 'note': '明确补档到已记票，不加金额'}
    linked = confirm(source_id, [pdf], pdf_contents={'pdf': content})
    assert linked['confirmed_count'] == 12 and linked['pending_count'] == 1
    assert confirm(source_id, [*decisions, pdf], pdf_contents={'pdf': content}) == linked
    assert stage(manifest) == linked
    read_invoice = client.get('/api/m13/invoices').json()
    assert read_invoice['total_amount'] == '10000.00' and len(read_invoice['invoices']) == 1
    original = read_invoice['invoices'][0]['originals'][0]
    assert client.get(f"/api/m13/invoices/originals/{original['id']}").content == content
    reconciliation = client.get('/api/m13/reconciliation').json()
    assert (reconciliation['total_amount'], reconciliation['invoiced_amount'], reconciliation['difference_amount']) == ('15670.00', '10000.00', '5670.00')
    assert [(r['total_amount'], r['invoiced_amount']) for r in reconciliation['rows']] == [('15640.00', '6000.00'), ('30.00', '4000.00')]
    assert client.get('/api/m13/ledger').json()['total_amount'] == '15170.00'
    with session_scope() as session:
        october = reporting.linkage_month(session, project_id=ids['project_ids'][0], ym='2026-10')
        november = reporting.linkage_month(session, project_id=ids['project_ids'][0], ym='2026-11')
        assert october['recharge_amount'] == '10170.00' and october['custom_amount'] == '1830.00'
        assert october['fee_amount'] == '12000.00' and october['big']['amount'] == '13000.00'
        assert october['big']['recharge_amount'] == '10740.00' and october['big']['extra_amount'] == '430.00'
        assert november['fee_amount'] == '4900.00' and november['big']['amount'] == '17900.00'
        from app.report_export import service as exports
        monkeypatch.setattr(exports, 'EXPORT_ROOT', tmp_path)
        ledger_path = exports.export_ledger(session, ['__oilcard_463_a__'], ym='2026-10')
        big_path = exports.export_big(session, '2026-10')
    # Existing production exporters, independent expected cells/row values.
    workbook = openpyxl.load_workbook(ledger_path, data_only=True)
    try:
        sheet = workbook['特殊费用月报表']
        row = next(r for r in sheet.iter_rows() if r[0].value == '工程用油（油卡联动）')
        col = next(c.column for c in sheet[2] if c.value == '2026-10')
        assert row[col - 1].value == 12000
    finally:
        workbook.close()
    workbook = openpyxl.load_workbook(big_path, data_only=True)
    try:
        row = next(r for r in workbook.active.iter_rows() if any(c.value == '463演练A' for c in r))
        assert row[18].value == 13000
    finally:
        workbook.close()
    exported = client.get('/api/m13/invoices/export')
    assert exported.status_code == 200
    workbook = openpyxl.load_workbook(BytesIO(exported.content), data_only=True)
    try:
        assert workbook.active.max_row == 3  # one invoice with two allocations, then total
        assert workbook.active['G3'].value == 10000
    finally:
        workbook.close()


def test_changed_source_different_identity_and_confirmation_never_overwrite(enrolled):
    _, source_id, manifest, decisions, _, _, staged = enrolled
    changed = deepcopy(manifest)
    changed['rows'][0]['raw']['amount'] = '501'
    with pytest.raises(DataConflictError, match='来源的内容'):
        stage(changed)
    with pytest.raises(DataConflictError, match='另一稳定身份'):
        stage({**manifest, 'source_id': 'another-origin-name'})
    assert stage(manifest) == staged
    confirm(source_id, decisions[:1])
    with pytest.raises(DataConflictError, match='不同确认结论'):
        confirm(source_id, [{**decisions[0], 'note': '另一个结论'}])
    assert stage(manifest)['confirmed_count'] == 1


def test_manual_correction_blocks_reimport_without_overwriting_it(enrolled):
    client, source_id, manifest, decisions, _, _, _ = enrolled
    first = confirm(source_id, decisions[1:2])
    entry_id = target(first, 'prior')
    entry = client.get('/api/m13/ledger').json()['entries'][0]
    fields = {k: entry[k] for k in ('project_id', 'account_id', 'date_text', 'kind', 'amount', 'remark', 'version')}
    result = client.put(f'/api/m13/ledger/entries/{entry_id}', json={**fields, 'amount': '71', 'remark': '人工更正保留'})
    assert result.status_code == 200, result.text
    for action in (lambda: stage(manifest), lambda: confirm(source_id, decisions[1:2])):
        with pytest.raises(DataConflictError, match='已更正'):
            action()
    assert client.get('/api/m13/ledger').json()['entries'][0]['amount'] == '71.00'
    assert client.get('/api/m13/ledger').json()['entries'][0]['remark'] == '人工更正保留'


def test_pre_delete_column_invoice_confirmation_keeps_identity_and_detects_deletion(enrolled):
    _, source_id, manifest, decisions, _, _, _ = enrolled
    invoice_decision = [d for d in decisions if d['key'] == 'invoice']
    result = confirm(source_id, invoice_decision)
    invoice_id = target(result, 'invoice')
    with session_scope() as session:
        source = session.get(OilContinuationSource, source_id)
        saved = json.loads(source.confirmations_json)
        saved['invoice']['targets'][0]['snapshot'].pop('deleted')
        source.confirmations_json = json.dumps(saved)
    assert stage(manifest)['confirmed_count'] == 1
    assert confirm(source_id, invoice_decision)['confirmed_count'] == 1
    with session_scope() as session:
        assert json.loads(session.get(OilContinuationSource, source_id).confirmations_json) == saved
        session.get(OilInvoice, invoice_id).deleted = True
    with pytest.raises(DataConflictError, match='已更正'):
        stage(manifest)


def test_invalid_multi_allocation_and_undeclared_card_rollback_whole_confirmation(enrolled):
    client, source_id, manifest, decisions, _, _, staged = enrolled
    invoice = next(d for d in decisions if d['key'] == 'invoice')
    bad = deepcopy(invoice)
    bad['allocations'][1]['amount'] = '3999'
    with pytest.raises(MustReportError, match='分配合计'):
        confirm(source_id, [decisions[1], bad])
    assert client.get('/api/m13/ledger').json()['total_amount'] == '0.00'
    assert client.get('/api/m13/invoices').json()['invoices'] == []
    assert stage(manifest) == staged
    undeclared = {k: v for k, v in decisions[1].items() if k != 'card_id'}
    with pytest.raises(MustReportError, match='卡片快照'):
        confirm(source_id, [undeclared])


def test_unknown_transfer_only_enters_confirmed_opening_with_both_sides(enrolled):
    client, _, _, decisions, _, _, _ = enrolled
    transfer = next(d for d in decisions if d['key'] == 'transfer')
    source_id = 'unknown-transfer-source'
    manifest = {'source_id': source_id, 'rows': [{'key': '转账!2', 'kind': 'transfer',
        'raw': {'original_date_text': '原始数据', 'explicit_date': None, 'amount': '30', 'remark': '原平账文字'}}]}
    stage(manifest)
    decision = {**transfer, 'key': '转账!2'}
    with pytest.raises(MustReportError, match='截至月份'):
        confirm(source_id, [decision])
    with pytest.raises(MustReportError, match='业务性质'):
        confirm(source_id, [{**decision, 'cutoff_ym': '2025-11', 'treatment': ''}])
    result = confirm(source_id, [{**decision, 'cutoff_ym': '2025-11'}])
    assert len(result['rows'][0]['confirmation']['targets']) == 2
    assert client.get('/api/m13/ledger').json()['entries'] == []
    assert client.get('/api/m13/reconciliation').json()['total_amount'] == '0.00'
    with session_scope() as session:
        outgoing = transfer['out_project_id']
        assert reporting.big_amount(session, project_id=outgoing, ym='2025-10') == 0
        assert str(reporting.big_amount(session, project_id=outgoing, ym='2025-11')) == '-30.00'
        assert reporting.fee_month(session, project_id=outgoing, ym='2025-11')['recharge_amount'] == '0.00'


def test_pdf_requires_exact_hash_invoice_version_facts_and_unique_number(enrolled):
    client, source_id, _, decisions, content, _, _ = enrolled
    result = confirm(source_id, [d for d in decisions if d['key'] == 'invoice'])
    invoice_id = target(result, 'invoice')
    invoice = client.get('/api/m13/invoices').json()['invoices'][0]
    pdf = {'key': 'pdf', 'invoice_id': invoice_id, 'version': invoice['version'], 'note': '明确补档'}
    with pytest.raises(DataConflictError, match='hash'):
        confirm(source_id, [pdf], pdf_contents={'pdf': content + b'changed'})
    with pytest.raises(DataConflictError):
        confirm(source_id, [{**pdf, 'version': str(uuid4())}], pdf_contents={'pdf': content})
    with session_scope() as session:
        row = session.get(OilInvoice, invoice_id)
        row.amount = '9999.00'
    with pytest.raises(DataConflictError, match='已更正'):
        confirm(source_id, [pdf], pdf_contents={'pdf': content})
    with session_scope() as session:
        assert session.scalar(select(OilInvoice.id)) == invoice_id


def test_manifest_enforces_approved_ranges_and_keeps_full_original_markers(tmp_path):
    import json
    evidence = {'source-inventory.json': {'sha256': 'original-hash', 'sheets': {'油卡': {'range': 'A2:I49'}}, 'pdfs': []},
        'source-rows-private.json': {'source_sha256': 'original-hash', 'cards': [], 'invoice_amounts': [], 'transfers': [],
            'recharges': [{'source_row': 2, 'source_sheet': '油卡充值', 'amount': '1.00', 'original_date_text': '原始数据', 'explicit_date': None}]},
        'pdf-facts-private.json': {'pdfs': []}, 'legacy-totals.json': {'revision': 'old', 'projects': []},
        'project-candidates-private.json': {'candidates': []}}
    for name, value in evidence.items():
        (tmp_path / name).write_text(json.dumps(value, ensure_ascii=False), encoding='utf-8')
    manifest = build_manifest(tmp_path)
    assert manifest['source_id'] == 'fourth-company-oilcard-ledger'
    assert manifest['rows'][0]['key'] == '油卡充值!2' and manifest['rows'][0]['raw']['original_date_text'] == '原始数据'
    evidence['source-rows-private.json']['recharges'][0]['source_row'] = 67
    (tmp_path / 'source-rows-private.json').write_text(json.dumps(evidence['source-rows-private.json']), encoding='utf-8')
    with pytest.raises(ValueError, match='approved active ranges'):
        build_manifest(tmp_path)
