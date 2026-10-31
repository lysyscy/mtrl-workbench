"""Real receipt API, cross-upload identities, money and durable original recovery."""
from hashlib import sha256
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.core.db import session_scope
from app.mod13_oilcard.invoice_pdf import recognize_pdf
from app.mod13_oilcard.models import OilInvoiceAllocation
from tests.oilcard_invoice_samples import CARD, NUMBER, invoice_pdf
from tests.test_oilcard_archives import scene, create, edit  # noqa: F401
from tests.test_oilcard_ledger import entry, summary


def receive(client, content):
    response = client.post('/api/m13/invoices/receive', files={'file': ('original.pdf', content, 'application/pdf')})
    assert response.status_code == 200, response.text
    return response.json()


def invoices(client):
    response = client.get('/api/m13/invoices')
    assert response.status_code == 200, response.text
    return response.json()


def setup_account(client, project):
    card = create(client, 'cards', card_no=CARD)
    return create(client, 'accounts', name='合成收票账户', project_id=project, card_id=card['id'])


def test_receipt_actual_pdf_reconciliation_repeat_and_original(scene):
    client, (project, _) = scene
    account = setup_account(client, project)
    entry(client, account, '10000')
    entry(client, account, '200', kind='adjustment')
    content = invoice_pdf()
    first = receive(client, content)
    row, original = first['item'], first['original']
    assert row['status'] == 'registered'
    assert (row['invoice_no'], row['invoice_date'], row['amount'], row['card_no']) == (NUMBER, '2026-10-31', '3000.00', CARD)
    assert row['allocations'][0]['account_id'] == account['id']
    assert summary(client)['rows'][0]['total_amount'] == '10200.00'
    assert summary(client)['rows'][0]['invoiced_amount'] == '3000.00'
    assert summary(client)['rows'][0]['difference_amount'] == '7200.00'
    assert receive(client, content) == first
    duplicate = receive(client, invoice_pdf(variant='重复输入合成测试'))
    assert duplicate['original']['outcome'] == 'duplicate'
    assert duplicate['original']['id'] != original['id']
    assert duplicate['item']['originals'][0]['id'] == original['id']
    assert invoices(client)['total_amount'] == '3000.00'
    assert len(invoices(client)['invoices']) == 1
    digest = sha256(content).hexdigest()
    assert client.get(f'/api/m13/invoices/receipts/{digest}').json()['received'] is True
    assert client.get('/api/m13/invoices/receipts/' + '0' * 64).json() == {'received': False, 'original': None}
    downloaded = client.get(f"/api/m13/invoices/originals/{original['id']}?download=true")
    assert downloaded.content == content
    from urllib.parse import unquote
    filename = unquote(downloaded.headers['X-Export-Filename'])
    assert filename == f'{CARD[-5:]} 甲 3000.00 {NUMBER}.pdf'
    assert downloaded.headers['Content-Type'] == 'application/pdf'
    assert 'attachment' in downloaded.headers['Content-Disposition']
    assert 'inline' in client.get(f"/api/m13/invoices/originals/{original['id']}").headers['Content-Disposition']
    client.cookies.clear()
    assert client.post('/api/m13/invoices/receive', files={'file': ('x.pdf', content)}).status_code == 401
    assert invoices(client)['total_amount'] == '3000.00'


@pytest.mark.parametrize('scenario', ['unknown', 'short', 'ambiguous', 'archived', 'unreliable'])
def test_uncertain_card_or_layout_never_guesses_account(scene, scenario):
    client, (project, _) = scene
    account = setup_account(client, project)
    options = {}
    if scenario == 'unknown': options['card'] = '1000111400099999999'
    if scenario == 'short': options['card'] = CARD[-5:]
    if scenario == 'ambiguous': create(client, 'accounts', name='同卡第二账户', project_id=project, card_id=account['card_id'])
    if scenario == 'archived': edit(client, 'accounts', account, status='archived')
    if scenario == 'unreliable': options['subtotal'] = '2600.00'
    row = receive(client, invoice_pdf(**options))['item']
    assert row['status'] == 'pending' and row['reason']
    assert not row['allocations']
    assert invoices(client)['total_amount'] == '0.00'
    assert summary(client)['invoiced_amount'] == '0.00'
    if scenario == 'short': assert row['card_no'] is None


def test_invoice_number_conflict_keeps_both_originals_and_confirmed_money(scene):
    client, (project, _) = scene
    setup_account(client, project)
    first_content = invoice_pdf()
    first = receive(client, first_content)
    conflict_content = invoice_pdf(amount='3100.00', subtotal='2743.36', tax='356.64')
    conflict = receive(client, conflict_content)
    assert conflict['original']['outcome'] == 'conflict'
    assert conflict['item']['id'] == first['item']['id']
    assert conflict['item']['amount'] == '3000.00'
    assert conflict['item']['allocations'] == first['item']['allocations']
    assert invoices(client)['total_amount'] == '3000.00'
    assert summary(client)['invoiced_amount'] == '3000.00'
    for result, content in ((first, first_content), (conflict, conflict_content)):
        assert client.get(f"/api/m13/invoices/originals/{result['original']['id']}").content == content
    from urllib.parse import unquote
    conflict_name = unquote(client.get(f"/api/m13/invoices/originals/{conflict['original']['id']}?download=true").headers['X-Export-Filename'])
    assert conflict_name == f"{CARD[-5:]} 待确认 3100.00 {NUMBER} 待核原件-{conflict['original']['id']}.pdf"
    main_name = unquote(client.get(f"/api/m13/invoices/originals/{first['original']['id']}?download=true").headers['X-Export-Filename'])
    assert main_name == f'{CARD[-5:]} 甲 3000.00 {NUMBER}.pdf'
    assert receive(client, conflict_content) == conflict


@pytest.mark.parametrize('changed', [{'buyer': '另一家合成测试购买方'}, {'subtotal': '2600.00', 'tax': '400.00'}])
def test_same_four_facts_but_other_visible_content_changes_is_conflict(scene, changed):
    client, (project, _) = scene
    setup_account(client, project)
    first = receive(client, invoice_pdf())
    metadata_only = receive(client, invoice_pdf(variant='PDF metadata only'))
    assert metadata_only['original']['outcome'] == 'duplicate'
    conflict = receive(client, invoice_pdf(**changed))
    assert conflict['original']['outcome'] == 'conflict'
    assert conflict['item']['allocations'] == first['item']['allocations']
    assert conflict['item']['amount'] == '3000.00'
    assert invoices(client)['total_amount'] == '3000.00'
    assert len(conflict['item']['originals']) == 3


def test_note_safe_undo_version_and_all_saved_export(scene):
    client, (project, _) = scene
    setup_account(client, project)
    first = receive(client, invoice_pdf())['item']
    pending = receive(client, invoice_pdf(number='20261031000000000002', card='12345'))['item']
    result = client.put(f"/api/m13/invoices/{first['id']}/remark", json={'version': first['version'], 'remark': '发票独立备注'})
    assert result.status_code == 200, result.text
    assert result.headers['X-Operation-History'] == 'recorded'
    assert client.put(f"/api/m13/invoices/{first['id']}/remark", json={'version': first['version'], 'remark': '覆盖'}).status_code == 409
    operation = client.get('/api/history/status').json()['undo']
    assert client.post('/api/history/undo', json={'operation_id': operation['id']}).json()['ok']
    assert invoices(client)['invoices'][0]['remark'] == ''
    assert client.get(f"/api/m13/invoices/originals/{first['originals'][0]['id']}").content == invoice_pdf()
    assert client.post('/api/history/redo', json={'operation_id': operation['id']}).json()['ok']
    assert invoices(client)['invoices'][0]['remark'] == '发票独立备注'
    exported = client.get('/api/m13/invoices/export')
    assert exported.status_code == 200
    sheet = load_workbook(BytesIO(exported.content)).active
    assert sheet.max_row == 4
    assert sheet['C2'].value == NUMBER and sheet['C2'].data_type == 's'
    assert sheet['H2'].value == '发票独立备注'
    assert sheet['C3'].value == pending['invoice_no']
    assert sheet['G3'].value == 0 and sheet['G4'].value == 3000


def test_invalid_and_unreadable_pdf_keep_boundaries(scene):
    client, (project, _) = scene
    setup_account(client, project)
    assert client.post('/api/m13/invoices/receive', files={'file': ('x.pdf', b'not pdf')}).status_code == 400
    assert not invoices(client)['invoices']
    broken = b'%PDF-1.7\nmalformed synthetic PDF'
    row = receive(client, broken)['item']
    assert row['status'] == 'pending' and row['amount'] is None
    assert client.get(f"/api/m13/invoices/originals/{row['originals'][0]['id']}").content == broken


def test_registered_allocation_mismatch_refuses_partial_money(scene):
    client, (project, _) = scene
    setup_account(client, project)
    row = receive(client, invoice_pdf())['item']
    with session_scope() as session:
        allocation = session.scalar(select(OilInvoiceAllocation).where(OilInvoiceAllocation.invoice_id == row['id']))
        allocation.amount = '2999.00'
    assert client.get('/api/m13/invoices').status_code == 400
    assert client.get('/api/m13/reconciliation').status_code == 400


def test_pdf_bank_number_and_two_text_modes():
    result = recognize_pdf(invoice_pdf())
    assert {k: v for k, v in result.items() if k != 'text_sha256'} == {'invoice_no': NUMBER, 'invoice_date': '2026-10-31', 'amount': '3000.00', 'card_no': CARD, 'reliable': True, 'reason': ''}
    assert len(result['text_sha256']) == 64
    assert not recognize_pdf(invoice_pdf(card='12345'))['reliable']
    assert not recognize_pdf(invoice_pdf(card='', bank=CARD))['reliable']
