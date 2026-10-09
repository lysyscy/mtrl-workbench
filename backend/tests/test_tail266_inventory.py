"""TAIL-266：库存原单入库金额、核算月、年度边界及独立结余的共同出口回归。"""
from decimal import Decimal as D
from app.mod07_setl import service
from app.mod07_setl.models import SetlDetail, WeldGasBalance
from app.mod08_inv.models import InvDetail
from tests.m07_helpers import FakeSession
from tests.acc_plugin import acc


def inventory_data():
    rows = []
    for i,(ym,name,amount,direction) in enumerate([
        ('2025-12','焊条J422','100','入库'),('2026-01','焊丝','10','入库'),
        ('2026-01','焊条','15','入库'),('2026-09','焊条','20','入库'),
        ('2027-01','焊条','900','入库'),('2026-01','焊丝','800','出库'),
        ('2026-02','氧气','30','入库'),('2026-02','氧气瓶','700','入库'),
    ],1):
        rows.append(InvDetail(id=i,warehouse_name='甲仓',ym=ym,material_name=name,
                              amount=D(amount),tax_total=D('9999'),direction=direction,
                              doc_no=f'隔离{i}',seq='1',doc_date='2030-01-01'))
    return {InvDetail: rows, SetlDetail: [SetlDetail(id=1,ym='2026-01',warehouse_name='甲仓',
              resource_name='焊丝',in_amt=D('88888'))],
            WeldGasBalance: [WeldGasBalance(id=1,kind='焊材',warehouse_name='甲仓',amount=D('456'))]}


@acc('ACC-SETL-EXT-001', '001-02')
def test_inventory_year_months_future_and_balance_are_independent():
    data=service.build_extract_export(FakeSession(inventory_data()), {'status':'全部','period':'2026'})
    weld=data['annual_grids']['焊材']['rows'][0]
    gas=data['annual_grids']['气体']['rows'][0]
    assert weld['prior'] == D('100')
    assert weld['cells'] == [D('25'),D('0'),D('0'),D('0'),D('0'),D('0'),D('0'),D('0'),D('20'),D('0'),D('0'),D('0')]
    assert weld['total'] == D('145')
    assert gas['cells'][1] == D('30') and gas['total'] == D('30')
    assert data['grids']['焊材']['rows'][0]['balance'] == D('456')
    assert data['grids']['焊材']['rows'][0]['total'] == D('45')


@acc()
def test_all_source_consumers_share_inventory_amounts():
    from app.mod07_setl import big_service
    s=FakeSession(inventory_data())
    page=service.build_extract_data(s, {'status':'全部','period':'2026'})
    assert page['annual']['焊材']['rows'][0]['total'] == D('145')
    assert big_service._extract_amounts(s)['焊材']['甲仓']['2026-01'] == D('25')


@acc()
def test_real_export_keeps_original_sheets_and_adds_annual_cents(tmp_path):
    from app.report_export.weldgas_export import build_xlsx
    from openpyxl import load_workbook
    dataset = inventory_data()
    dataset[InvDetail][1].amount = D('10.13')
    data = service.build_extract_export(FakeSession(dataset), {'status':'全部','period':'2026'})
    out = build_xlsx(data['grids'], tmp_path, data['report_ym'], annual_grids=data['annual_grids'])
    wb = load_workbook(out, data_only=True)
    assert wb.sheetnames == ['焊材','气','焊材年度汇总','气体年度汇总']
    annual = list(wb['焊材年度汇总'].values)
    assert annual[1][:3] == ('甲仓',100,25.13)
    assert annual[1][-1] == 145.13
    original = [row for row in wb['焊材'].values if row[0] == '甲仓'][0]
    assert original[1] == 456 and original[2] == 25.13 and original[14] == 45.13
    wb.close()


@acc('ACC-SETL-EXT-001', '001-02')
@acc('ACC-SETL-EXT-001', '001-11')
def test_stored_inventory_months_and_balance_share_page_export_and_big():
    from app.core.db import session_scope
    from app.mod01_projects import repository as projects
    from app.mod01_projects.models import TYPE_NORMAL
    from app.mod07_setl import big_service
    # conftest redirects session_scope to TEST_DATABASE_URL; no production writes.
    with session_scope() as s:
        projects.create_project(s, project_type=TYPE_NORMAL, name_full='年度库存隔离测试', warehouse_name='__TAIL266__', phid='')
        dataset = inventory_data()
        for row in [*dataset[InvDetail], *dataset[WeldGasBalance]]:
            row.id = None
            row.warehouse_name = '__TAIL266__'
            s.add(row)
        s.flush();s.expire_all()
        state = {'status':'全部','period':'2026'}
        page = service.build_extract_data(s, state)
        exported = service.build_extract_export(s, state)
        own = lambda grid: next(row for row in grid['焊材']['rows'] if row['warehouse'] == '__TAIL266__')
        assert own(page['annual']) == own(exported['annual_grids'])
        assert own(page['annual'])['total'] == D('145')
        assert own(exported['grids'])['balance'] == D('456')
        amounts = big_service._extract_amounts(s)['焊材']['__TAIL266__']
        assert amounts['2026-01'] == D('25') and sum(amounts.values()) == D('1045')
        s.rollback()
