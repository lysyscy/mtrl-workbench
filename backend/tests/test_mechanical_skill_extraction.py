"""Behavior checks using clearly synthetic workbooks, separate from real evidence."""
import importlib.util
from decimal import Decimal
from datetime import datetime
from pathlib import Path
import tempfile
import unittest


import openpyxl

SCRIPT = Path(__file__).resolve().parents[2] / "skills/machinery-fee-consolidation/scripts/extract_reports.py"
spec = importlib.util.spec_from_file_location("reader", SCRIPT)
reader = importlib.util.module_from_spec(spec)
spec.loader.exec_module(reader)


class ExtractionBehavior(unittest.TestCase):
    def test_moved_columns_missing_sequence_repeated_header_and_repeated_business(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "fixture.xlsx"
            book = openpyxl.Workbook()
            tab = book.active
            tab.title = "设备"
            headers = ["序号", "设备名称", "规格型号", "单位", "数量", "零租单价（元）", "合价（元）",
                       "专票税率", None, "租赁日期", "出租单位"]
            tab.append(["合成样本，不是真实项目"])
            tab.append(headers)
            tab.append([None, "起重机", "25T", "台班", "2", 500, 1000, "1%普票，不含油", None, "8.21", "A单位"])
            tab.append(headers)
            tab.append([None, "起重机", "25T", "台班", "2", 500, 1000, "1%普票，不含油", None, "8.21", "A单位"])
            tab.append([1, "起重机", "50T", "台班", 1, 1130, 1130, .13, None, "9.10", "A单位"])
            tab.append([None, "合计", None, None, None, None, 3130])
            book.save(path)
            result = reader.extract({"period": "2026-09", "sheets": [{"path": str(path), "sheet": "设备",
                                       "project": "合成", "kind": "包月", "reason": "测试分类与字段独立"}]})
            self.assertEqual(result["issues"], [])
            self.assertEqual(len(result["records"]), 3)
            first = result["records"][0]
            self.assertEqual(first["fields"]["quantity"]["value"], "2")
            self.assertEqual(first["fields"]["rental_date"]["coordinate"], "J3")
            self.assertEqual(first["tax"]["net_check"], "1000")
            self.assertNotIn("planned_months", first["fields"])
            self.assertEqual(Decimal(result["records"][2]["tax"]["net_check"]), Decimal("1000"))

    def test_missing_amount_cache_missing_tax_and_internal_lease_remain_visible(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "fixture.xlsx"
            book = openpyxl.Workbook()
            tab = book.active
            tab.append(["设备名称", "单位", "合价", "专票税率", "出租单位"])
            tab.append(["吊车", "台班", "=2*500", .13, "A单位"])
            tab.append(["吊车", "台班", 1000, None, "A单位"])
            tab.append(["绞磨机", "台", 30, "普票", "租赁分公司"])
            book.save(path)
            selection = {"path": str(path), "sheet": "Sheet", "project": "合成", "kind": "包月", "reason": "测试异常"}
            result = reader.extract({"period": "2026-09", "sheets": [selection]})
            self.assertEqual(len(result["records"]), 3)
            self.assertEqual(len(result["issues"]), 3)
            self.assertIsNone(result["records"][0]["fields"]["gross_amount"]["value"])
            self.assertEqual(result["records"][1]["tax"]["status"], "待核")
            self.assertIn("内租", result["issues"][2]["reason"])
            with self.assertRaises(ValueError):
                reader.extract({"period": "2026-09", "sheets": [selection, selection]})

    def test_empty_plan_and_ambiguous_tax_are_not_successful_zero(self):
        with self.assertRaises(ValueError):
            reader.extract({"period": "2026-09", "sheets": []})
        self.assertEqual(reader.tax_check(1000, "3%或6%专票", "专票税率")["status"], "待核")
        self.assertEqual(reader.tax_check(1000, "普票", "专票税率")["net_check"], "1000")

    def test_short_numeric_dates_keep_general_and_real_dates_keep_their_type(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "dates.xlsx"
            book = openpyxl.Workbook()
            tab = book.active
            tab.append(["设备名称", "单位", "合价", "专票税率", "租赁日期"])
            tab.append(["吊车", "台班", 113, .13, 8.21])
            tab.append(["吊车", "台班", 113, .13, datetime(2026, 8, 21)])
            tab["E3"].number_format = "yyyy-mm-dd"
            book.save(path)
            result = reader.extract({"period": "2026-09", "sheets": [
                {"path": str(path), "sheet": "Sheet", "project": "合成", "kind": "零租", "reason": "日期类型验证"}]})
            self.assertEqual(result["issues"], [])
            short, real = [r["fields"]["rental_date"] for r in result["records"]]
            self.assertEqual(short["value"], 8.21)
            self.assertEqual(short["format"], "General")
            self.assertEqual(real["value"], "2026-08-21T00:00:00")
            self.assertEqual(real["format"], "yyyy-mm-dd")
            self.assertEqual(result["sources"][0]["sha256_before"], result["sources"][0]["sha256_after"])


if __name__ == "__main__":
    unittest.main()
