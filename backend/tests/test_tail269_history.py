"""TAIL-269：历史金额、2025空月、模板与大表共用年度合计。"""
from decimal import Decimal as D
from app.mod07_setl import extract as X, inventory_extract as I, big_service as B
from tests.m07_helpers import FakeSession
from tests.test_tail266_inventory import inventory_data


def test_history_year_boundary_and_all_consumers():
    amounts={'焊材':{'甲仓':{'2025-01':D('999'), '2025-11':D('888'),
        '2025-12':D('100'), '2026-01':D('25'), '2027-01':D('900')}}, '气体':{}}
    balances={'焊材':{'甲仓':D('456')}}
    grid=X.extract_template_grid(amounts,'焊材',warehouses=['甲仓'],year='2025',balances=balances)
    assert grid['rows'][0]['cells'] == [None]*11+[D('100')]
    assert grid['rows'][0]['total'] == D('556')
    row=I.annual_grid(amounts,'焊材',warehouses=['甲仓'],year='2025',balances=balances)['rows'][0]
    assert row['cells']==grid['rows'][0]['cells'] and row['total']==D('556')
    totals=B._amounts(FakeSession(inventory_data()), ['甲仓'], '2026-01')['甲仓']
    assert totals['hw_cum']==D('501') and totals['qt_cum']==D('30')
    assert totals['hw_mon']==D('25')
    for history,want in [(None,D('25')), (D('0'),D('25')), (D('456'),D('481'))]:
        r=X.extract_template_grid(amounts,'焊材',warehouses=['甲仓'],year='2026',
            balances={'焊材':{'甲仓':history}})['rows'][0]
        assert r['balance']==history and r['total']==want


def test_2025_export_is_fixed_fifteen_columns(tmp_path):
    from app.report_export.weldgas_export import build_xlsx
    from openpyxl import load_workbook
    amounts={'焊材':{'甲仓':{'2025-01':D('999'), '2025-12':D('100')}}, '气体':{}}
    grids={kind:X.extract_template_grid(amounts,kind,warehouses=['甲仓'],year='2025',
        balances={'焊材':{'甲仓':D('456')}}) for kind in X.KINDS}
    path=build_xlsx(grids,tmp_path,'2025-12')
    wb=load_workbook(path,data_only=True)
    for ws in wb:
        assert ws.max_column==15 and ws['B1'].value=='2025-11及之前'
        assert [ws.cell(2,col).value for col in range(3,14)]==[None]*11
        assert [ws.cell(ws.max_row,col).value for col in range(3,14)]==[None]*11
    assert wb['焊材']['N2'].value==100 and wb['焊材']['O2'].value==556
    wb.close()
