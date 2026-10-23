"""普通导出的实际文件范围：有序 ID、完整贸易列，以及空范围不扩成全量。"""
import base64
from contextlib import contextmanager
from io import BytesIO

import openpyxl


def test_ordinary_export_keeps_order_trade_columns_and_empty_scope(monkeypatch):
    from app.core.db import session_scope
    from app.mod01_projects.repository import create_project
    from app.mod11_jc import api
    from app.mod11_jc.models import JcContract, JcInvoice, JcPayment, TAG_IN

    with session_scope() as session:
        try:
            project = create_project(session, project_type='normal', name_full='__418导出__', warehouse_name='__418仓__', phid='')
            session.flush()
            rows = [JcContract(project_id=project.id, material_name=name, tag=TAG_IN,
                               contract_date='2026-10-01', contract_amt=100, trade_co='贸易公司', trade_contract_amt=120)
                    for name in ['甲材料', '乙材料', '筛掉的材料']]
            session.add_all(rows)
            session.flush()
            for row in rows:
                session.add(JcInvoice(contract_id=row.id, kind='supplier', invoice_date='2026-10-01', amount=100))
                session.add(JcInvoice(contract_id=row.id, kind='trade', invoice_date='2026-10-01', amount=120))
                session.add(JcPayment(contract_id=row.id, pay_date='2026-10-02', amount=20, source='manual'))
            session.flush()

            @contextmanager
            def scope():
                yield session
            monkeypatch.setattr(api, 'session_scope', scope)

            def read(ids):
                result = api.api_export(api.ExportIn(ids=ids))
                return openpyxl.load_workbook(BytesIO(base64.b64decode(result['content_b64'])), data_only=True)

            book = read([rows[1].id, rows[0].id])
            sheet = book.active
            header = [cell.value for cell in sheet[1]]
            assert sheet.max_row == 4  # 表头、两条合同、合计。
            assert [sheet.cell(n, header.index('材料名') + 1).value for n in [2, 3]] == ['乙材料', '甲材料']
            for label, expected in [('贸易公司名', '贸易公司'), ('贸易公司合同金额', 120), ('贸易公司发票金额', 120),
                                    ('结算金额', 120), ('付款金额', 20), ('欠款金额', 100)]:
                assert sheet.cell(2, header.index(label) + 1).value == expected
            assert sheet.cell(4, header.index('欠款金额') + 1).value == 200
            book.close()
            empty = read([])
            assert not any('材料' in str(cell.value) for row in list(empty.active.rows)[1:] for cell in row)
            empty.close()
        finally:
            session.rollback()
