"""User-approved monthly exports: formulas, text dates and complete month columns."""
from decimal import Decimal

import openpyxl
import pytest

from app.report_export import mechanical_export as M, process_export as P


def data():
    return {'ym': '2026-09', 'projects': [{'warehouse': '甲', 'name': '甲项目'}],
            'effective': [{'warehouse': '甲', 'ym': '2026-09', 'values': {'mech_2': '10', 'mech_3': '0'}}],
            'annual_effective': [
                {'warehouse': '甲', 'ym': '2024-07', 'values': {'mech_2': '20'}},
                {'warehouse': '甲', 'ym': '2026-09', 'values': {'mech_2': '10', 'mech_3': '0'}},
                {'warehouse': '乙', 'ym': '2023-01', 'values': {'mech_2': '99'}},
                {'warehouse': '甲', 'ym': '2026-10', 'values': {'mech_2': '999'}}],
            'summary': [{'warehouse': '甲', 'ym': '2026-09', 'cumulative': {'zero': '30', 'monthly': '0', 'internal': '0'}}],
            'histories': [{'warehouse': '甲', 'i6': '10', 'annual_rows': []}],
            'details': [{'warehouse': '甲', 'category': 'zero', 'totals': {'gross': '10', 'net': '10'},
                         'rows': [{'device': '=HYPERLINK("bad")', 'amount': '10', 'net': '10',
                                   'date': '9.1-9.15', 'memo': '车牌原文'}]}],
            'i8': [{'warehouse': '甲', 'totals': {'usage': '20', 'project': '5', 'team': '15'}}]}


def test_monthly_formulas_have_caches_and_preview_has_results(tmp_path):
    path = M.export(data(), tmp_path)
    formulas = openpyxl.load_workbook(path)
    values = openpyxl.load_workbook(path, data_only=True)
    summary, zero, monthly, annual = formulas.worksheets
    assert summary['J7'].data_type == 'f'
    assert values.worksheets[0]['J7'].value == 10
    assert values.worksheets[0]['E7'].value == 30
    assert {'E4:I4', 'J4:N4'} <= set(map(str, summary.merged_cells.ranges))
    assert zero['J6'].value == '9.1-9.15' and zero['J6'].number_format == '@'
    assert zero['P6'].value == '车牌原文' and zero['B6'].data_type == 's'
    assert zero['G7'].data_type == 'f' and values.worksheets[1]['G7'].value == 10
    assert monthly['H6'].data_type == 'f'
    assert monthly['H5'].value == '=""'
    assert values.worksheets[2]['H5'].value is None
    assert [annual.cell(1, c).value for c in range(3, 5)] == ['2024年7月', '2026年9月']
    assert annual.max_column == 10 and 'SUM(C2:D2)' in annual['E2'].value
    assert 'E6-F6' in annual['I6'].value
    assert values.worksheets[3]['I6'].value == 0
    assert annual['C2'].border.left.style and annual['C2'].border.right.style
    assert all('往年合计' != c.value for c in annual[1])
    preview_values = [c['value'] for s in M.preview(data())['sheets'] for row in s['rows'] for c in row]
    assert '10' in preview_values and '=SUM(J6:J6)' not in preview_values


@pytest.mark.parametrize('n', [0, 1, 3, 9, 12])
def test_process_merges_formulas_and_blank_zero_for_any_row_count(tmp_path, n):
    rows = [{'ym': f'2026-{i+1:02d}', 'mech_2': Decimal('10'), 'mech_1': 0} for i in range(n)]
    path = P.build_xlsx(rows, tmp_path, display='甲', tab=1, ym='2026-09')
    raw = openpyxl.load_workbook(path).active
    values = openpyxl.load_workbook(path, data_only=True).active
    subtotal, total = 4 + n, 5 + n
    assert {f'M{total}:P{total}', f'Q{total}:R{total}'} <= set(map(str, raw.merged_cells.ranges))
    assert raw[f'O{subtotal}'].data_type == 'f'
    assert raw[f'M{total}'].data_type == 'f'
    assert values[f'M{total}'].value == (10 * n if n else None)
    assert values[f'N{subtotal}'].value is None
    if n:
        assert values['N4'].value is None and values['O4'].value == 10
    assert str(raw[f'A{total+1}'].value).startswith('1、材料费用')


def test_process_formula_includes_turnover_material_after_template_column_reordering(tmp_path):
    path = P.build_xlsx([{'ym': '2026-09', 'in_周转材料': 7}], tmp_path,
                        display='甲', tab=1, ym='2026-09')
    raw = openpyxl.load_workbook(path).active
    values = openpyxl.load_workbook(path, data_only=True).active
    assert 'SUM(C5:G5)' in raw['C6'].value
    assert values['C6'].value == 7


def test_generated_report_can_be_read_back_without_losing_merged_project_rows(tmp_path):
    from app.mod06_proc.mech_report import read_months
    path = M.export(data(), tmp_path)
    rows = read_months(path.read_bytes(), ['2024-07', '2026-09'])
    month = next(r for r in rows if r['核算月'] == '2026-09')
    assert month['报表项目名'] == '甲项目'
    assert month['外租零租含税'] == '10.00'
    assert month['外租包月含税'] == '0.00'
    assert month['外租零租不含税'] == '10.00'
