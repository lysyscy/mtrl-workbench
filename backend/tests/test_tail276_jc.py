"""历史汇总沿用集采出口，关联付款动态取统计金额、日期取实际流水。"""
from decimal import Decimal
from io import BytesIO
from contextlib import contextmanager
import base64

from openpyxl import load_workbook
from sqlalchemy import event

from app.core.db import session_scope
from app.mod01_projects import repository as projects
from app.mod01_projects import router as project_api
from app.mod01_projects.schemas import CreateProjectRequest, ProjectExportIn
from app.mod03_notice.models import PaymentNotice
from app.mod03_notice.notice_actions import apply_stat_change
from app.mod05_flow import repository as flows
from app.mod05_flow import service as flow_service
from app.mod11_jc import repository as R, service as S


def test_history_summary_and_dynamic_stat_payment(monkeypatch):
    with session_scope() as s:
        @contextmanager
        def borrowed_session():
            yield s
        monkeypatch.setattr(project_api, 'session_scope', borrowed_session)
        created = project_api.create_project(CreateProjectRequest(
            project_type='normal', name_full='历史项目测试', disp='历史项目测试'))
        p = projects.get_project(s, created.id)
        assert p.warehouse_name.startswith('未建仓-')
        assert any(x.key == p.warehouse_name and x.display == p.disp
                   for x in project_api.candidates(include_archived=True))
        exported = project_api.export_projects(ProjectExportIn(ids=[p.id]))
        project_book = load_workbook(BytesIO(base64.b64decode(exported['content_b64'])), data_only=True)
        assert exported['rows'] == 1 and p.name_full in str(list(project_book.active.values))
        c = R.insert_contract(s, project_id=p.id, material_name=p.name_full,
                              remark='历史未结金额；核对截至日：2026-09-30；')
        R.insert_invoice(s, contract_id=c.id, kind='supplier', invoice_date='2026-09-30', amount=100000)
        baseline = S.debt_of_ledger(s)
        n = PaymentNotice(ph_id='__tail276_dynamic__', ownership_id=p.id, category='jc',
                          amount_approved=80000, stat_paid=0, doc_date='2026-09-10')
        s.add(n)
        s.flush()
        assert S.act_link_notice(s, notice_id=n.id, contract_id=c.id)['ok']
        link = R.payment_of_notice(s, n.id)
        assert link.amount == 0 and link.pay_date == ''
        # 存量旧值也不能重新成为计算来源。
        link.amount = Decimal('80000')
        link.pay_date = '2026-09-10'
        s.flush()
        assert S.build_detail(s, c.id)['debt'] == Decimal('100000')
        assert R.payments_of(s, c.id)[0]['date'] == ''
        f1 = flows.create_flow(s, notice_id=n.id, project_id=p.id, amount=30000, pay_date='2026-10-03')
        assert S.build_detail(s, c.id)['debt'] == Decimal('70000')
        f2 = flows.create_flow(s, notice_id=n.id, project_id=p.id, amount=20000, pay_date='2026-10-05')
        assert S.build_detail(s, c.id)['pay_sum'] == Decimal('50000')
        assert R.payments_of(s, c.id)[0]['date'] == '2026-10-05'
        flows.revoke(s, f2, reason='测试撤销')
        s.flush()
        assert S.build_detail(s, c.id)['debt'] == Decimal('70000')
        assert R.payments_of(s, c.id)[0]['date'] == '2026-10-03'
        apply_stat_change(s, notice_id=n.id, field='stat_paid', value='45000', changed_by='测试')
        assert flows.effective_total_of(s, n.id) == Decimal('30000')
        assert S.build_detail(s, c.id)['debt'] == Decimal('55000')
        assert S.debt_of_ledger(s) == baseline - Decimal('45000')
        # MOD-05仍通过MOD-11注入，不产生第二份计算。
        assert flow_service.build_debt_live(s) == flow_service.build_debt(s, jc_debt=S.debt_of_ledger(s))
        c2 = R.insert_contract(s, project_id=p.id, material_name='普通合同', contract_date='2026-10-01')
        S.act_link_notice(s, notice_id=n.id, contract_id=c2.id)
        assert R.payments_of(s, c.id) == []
        assert R.payments_of(s, c2.id)[0]['amount'] == Decimal('45000')
        S.act_unlink_notice(s, notice_id=n.id)
        assert s.get(type(f1), f1.id).status == 'effective'
        manual = R.insert_payment(s, contract_id=c.id, pay_date='2026-10-08', amount=123)
        assert R.payments_of(s, c.id)[0]['amount'] == Decimal('123')
        assert R.payments_of(s, c.id)[0]['date'] == manual.pay_date
        queries = []
        def count(*args):
            queries.append(args[2])
        event.listen(s.bind, 'before_cursor_execute', count)
        try:
            R.payments_of_many(s, [c.id, c2.id])
            assert len(queries) == 1
        finally:
            event.remove(s.bind, 'before_cursor_execute', count)
        book = load_workbook(BytesIO(S.export_xlsx(s, ids=[c.id])), data_only=True)
        assert '历史项目测试' in str(list(book.active.values))
        s.rollback()
