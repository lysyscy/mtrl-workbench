"""逐条历史行、未填日期付款与导出同源；退役累计调整不再写入。"""
from decimal import Decimal
import json

from app.mod01_projects.models import Project
from app.mod03_notice.models import PaymentNotice
from app.mod05_flow import service, repository
from app.mod05_flow.models import OpeningBalance
from app.report_export import ledger_export
from tests.test_proc_split_cutoffs import session


def test_history_manual_auto_and_export_are_independent(session, tmp_path, monkeypatch):
    p = Project(project_type='normal', name_full='__tail274__',
                warehouse_name='未建仓-274-test', disp='历史项目')
    session.add(p)
    session.flush()
    for cat in ('material', 'mechanical'):
        opening = OpeningBalance(category=cat, project_id=p.id, cutoff_ym='2025-11',
                                 settle_amount=10000, paid_amount=4000)
        notice = PaymentNotice(ph_id='__tail274__'+cat, bill_no='__tail274__'+cat,
                               doc_date='', category=cat, ownership_id=p.id,
                               stat_settle=99999, stat_paid=2000, excluded=False, fields={})
        session.add_all([opening, notice])
        session.flush()
        repository.create_flow(session, notice_id=notice.id, project_id=p.id,
                               amount=2000, pay_date='2026-10-10')

        def side():
            return next(s for s in service.build_ledger(session, project_id=p.id,
                        end_ym='2025-12')['sides'] if s['category'] == cat)

        data = side()
        assert [r['row_kind'] for r in data['rows'][:3]] == ['undated', 'manual', 'month']
        assert [r['payment'] for r in data['rows'][:3]] == ['2000.00', '4000.00', '0.00']
        assert [r['other_expense'] for r in data['rows'][:3]] == ['0.00', '10000.00', '0.00']
        spec = ledger_export.ledger_spec(ledger_export.SHEET_OF_CATEGORY[cat], data, display='历史项目')
        assert spec['tail'][0][1]['D'] == Decimal('6000')
        assert spec['tail'][0][1]['E'] == Decimal('4000')
        assert '未填' in data['rows'][0]['month_label']
        assert [r[1]['B'] for r in spec['rows']] == [data['rows'][0]['month_label'], '2025-11', '2025-12']
        assert sum(Decimal(r['payment']) for r in data['rows'] if r['counts_in_total']) == Decimal('6000')
        assert data['rows'][1]['record_id'] == opening.id
        # 核对完整导出调用者写出的实际单元格，非仅拼装spec。
        from app.report_export import service as export_service
        import openpyxl
        monkeypatch.setattr(export_service, 'EXPORT_ROOT', tmp_path)
        path = export_service.export_ledger(session, [p.warehouse_name], '2025-12')
        wb = openpyxl.load_workbook(path, data_only=True)
        try:
            ws = wb[ledger_export.SHEET_OF_CATEGORY[cat]]
            exported = {r[1].value: r[3].value for r in ws.iter_rows()}
            assert exported[data['rows'][0]['month_label']] == 2000
            assert exported['2025-11'] == 4000
            assert not any(' · ' in str(label) for label in exported)
            tail = next(r for r in ws.iter_rows() if r[1].value == '合计：')
            assert tail[3].value == 6000
        finally:
            wb.close()


def test_empty_archived_project_keeps_identity_and_both_ledgers(session, tmp_path, monkeypatch):
    p = Project(project_type='normal', name_full='__tail274_empty__',
                warehouse_name='未建仓-274-empty', disp='无仓库项目')
    session.add(p)
    session.flush()
    data = service.build_ledger(session, project_id=p.id, end_ym='2025-12')
    assert data['project_id'] == p.id
    option = next(x for x in data['project_options'] if x['id'] == p.id)
    assert option['warehouse_name'] == p.warehouse_name
    assert [s['category'] for s in data['sides']] == ['material', 'mechanical']
    for side in data['sides']:
        assert side['rows'] == [] and side['display_rows'] == []
    from app.report_export import service as export_service
    import openpyxl
    monkeypatch.setattr(export_service, 'EXPORT_ROOT', tmp_path)
    path = export_service.export_ledger(session, [p.warehouse_name], '2025-12')
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        assert wb.sheetnames == ['辅材', '机械', '特殊费用月报表']
        for name in ('辅材', '机械'):
            labels = [row[1].value for row in wb[name].iter_rows()]
            assert '2025-11及之前' not in labels
            assert not any(' · ' in str(label) for label in labels)
    finally:
        wb.close()


def test_zero_automatic_payment_still_has_history_segment(session):
    p = Project(project_type='normal', name_full='__tail274_zero__',
                warehouse_name='未建仓-274-zero', disp='零付款项目')
    session.add(p)
    session.flush()
    session.add_all([
        OpeningBalance(category='material', project_id=p.id, cutoff_ym='2025-11',
                       settle_amount=10, paid_amount=0),
        PaymentNotice(ph_id='__tail274_zero__', bill_no='__tail274_zero__',
                      category='material', ownership_id=p.id, doc_date='',
                      stat_paid=0, stat_settle=0, excluded=False, fields={}),
        PaymentNotice(ph_id='__tail274_month__', bill_no='__tail274_month__',
                      category='material', ownership_id=p.id, doc_date='2026-01-10',
                      stat_paid=700, stat_settle=1000, excluded=False, fields={}),
    ])
    session.flush()
    side = service.build_ledger(session, project_id=p.id, end_ym='2026-01')['sides'][0]
    assert [r['row_kind'] for r in side['rows'][:3]] == ['undated', 'manual', 'month']
    assert side['rows'][0]['payment'] == '0.00'
    assert side['rows'][0]['row_kind'] == 'undated'
    assert side['rows'][-1]['payment'] == '700.00'
    spec = ledger_export.ledger_spec('辅材', side, display='零付款项目')
    assert spec['tail'][0][1]['D'] == Decimal('700')


def test_retired_total_payment_route_preserves_individual_history_and_source(session, monkeypatch):
    from contextlib import contextmanager
    from app.mod05_flow import router
    p = Project(project_type='normal', name_full='__tail274_edit__',
                warehouse_name='未建仓-274-edit', disp='合计编辑')
    session.add(p)
    session.flush()
    n = PaymentNotice(ph_id='__tail274_edit__', bill_no='__tail274_edit__',
                      category='material', ownership_id=p.id, doc_date='',
                      stat_paid=2000, stat_settle=0, excluded=False, fields={})
    session.add(n)
    repository.upsert_opening(session, category='material', project_id=p.id,
                              cutoff_ym='2025-11', settle_amount=9000, paid_amount=1000, remark='保留')
    session.flush()
    @contextmanager
    def scope():
        yield session
    monkeypatch.setattr(router, 'session_scope', scope)
    result = router.api_opening_payment(router.OpeningPaymentIn(
        category='material', project_id=p.id, payment='6000'))
    assert result.status_code == 410 and json.loads(result.body)['ok'] is False
    opening = repository.list_openings(session, project_id=p.id, category='material')[0]
    assert opening.settle_amount == Decimal('9000') and opening.paid_amount == Decimal('1000')
    assert opening.remark == '保留' and n.stat_paid == Decimal('2000')
