"""集采打印版的成品边界；不依赖正式数据库或本机打印机。"""
from io import BytesIO

import openpyxl
import pytest

from app.report_export import jc_print_export as export


def test_每项目独立版式_打印区随长表增长_金额与文本保真():
    rows = [{"project_seq": n, "date": "2026-10-14", "material": "=1+1",
             "supplier": "供应商", "contract_amt": None, "inv_sum": "100.25",
             "settle": "100.25", "pay_sum": "20.00", "debt": "80.25",
             "remark": "60天承兑", "arch_no": "档001"} for n in range(1, 81)]
    rows[0]["remark"] = "付款条件及补充备注" * 20
    blob = export.build_xlsx([
        {"display": "甲/项目", "trade": False, "rows": rows},
        {"display": "甲:项目", "trade": True,
         "rows": [{**rows[0], "trade_co": "贸易公司", "trade_contract_amt": "110.25",
                   "trade_inv_sum": "110.25", "settle": "110.25", "debt": "90.25"}]},
    ])
    wb = openpyxl.load_workbook(BytesIO(blob))
    assert len(wb.worksheets) == 2
    assert len(set(s.title.lower() for s in wb)) == 2
    normal, trade = wb.worksheets
    assert normal["A1"].value == "甲/项目集中采购付款清单"
    assert normal["A4"].value == 1 and normal["A83"].value == 80
    assert normal["C4"].value == "=1+1" and normal["C4"].data_type == "s"
    assert normal["E4"].value is None
    assert normal["F84"].value == 8020 and normal["I84"].value == 6420
    assert normal["N4"].value == "档001"
    assert normal.row_dimensions[4].height > normal.row_dimensions[5].height
    assert str(normal.print_area).endswith("!$A$1:$J$84")
    assert str(trade.print_area).endswith("!$A$1:$M$5")
    assert trade["D4"].value == "贸易公司" and trade["H4"].value == 110.25
    assert trade["J5"].value == 110.25
    for ws in wb:
        assert ws.page_setup.orientation == "landscape"
        assert ws.page_setup.paperSize == 9  # OOXML 的 A4 纸张编号。
        assert ws.page_setup.fitToWidth == 1 and ws.page_setup.fitToHeight == 0
        assert ws.sheet_properties.pageSetUpPr.fitToPage is True
        assert ws.print_title_rows == "$1:$3"
        assert "&P" in ws.oddFooter.center.text and "&N" in ws.oddFooter.center.text
        assert ws.max_row == (84 if ws is normal else 5)
        assert not any(c.data_type == "f" for row in ws for c in row)
        assert "付款方式" not in [c.value for row in ws.iter_rows(max_row=3) for c in row]


def test_成品核对识别真正公式而允许等号开头的文本():
    from app.report_export import tpl_fill
    from app.report_export.save import workbook_bytes

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "序号"
    ws["A2"] = 1
    ws["B2"] = "=1+1"
    ws["B2"].data_type = "s"
    plan = {"sheets": [{"sheet": ws.title, "max_col": 2, "data_start": 2,
                       "key_col": "A", "data_rows": [(1, {"B": "=1+1"})], "tail": [], "last_row": 2}]}
    assert tpl_fill.verify_filled(workbook_bytes(wb), plan) == []
    ws["B2"] = "=1+1"
    assert any("还是公式" in p for p in tpl_fill.verify_filled(workbook_bytes(wb), plan))


def test_独立接口按全部纳入合同导出_判版不受不纳入行影响(monkeypatch):
    import base64
    from contextlib import contextmanager
    from app.core.db import session_scope
    from app.main import create_app
    from app.mod01_projects.repository import create_project
    from app.mod03_notice import models as _notice_models  # noqa: F401
    from app.mod11_jc import api, service
    from app.mod11_jc.models import JcContract, JcInvoice, JcPayment, TAG_IN, TAG_OUT

    assert "/api/m11/print-export" in create_app().openapi()["paths"]
    with session_scope() as s:
        try:
            for wh in ('__336新__', '__336老__', '__336无合同__'):
                p = create_project(s, project_type='normal', name_full=wh, warehouse_name=wh, phid='')
                s.flush()
                if wh.endswith('无合同__'):
                    continue
                row = JcContract(project_id=p.id, material_name=wh, supplier='测试供应商', contract_date='2026-10-14', tag=TAG_IN, contract_amt=100)
                s.add(row)
                s.flush()
                s.add(JcInvoice(contract_id=row.id, kind='supplier', invoice_date='2026-10-14', amount=100))
                s.add(JcPayment(contract_id=row.id, pay_date='2026-10-14', amount=20, source='manual'))
                if wh == '__336老__':
                    s.add(JcContract(project_id=p.id, material_name='不导出的贸易合同', supplier='贸易供应商', contract_date='2026-01-01', tag=TAG_OUT, trade_co='贸易公司'))
            s.flush()

            @contextmanager
            def scope():
                yield s
            monkeypatch.setattr(api, 'session_scope', scope)
            out = api.api_print_export(api.PrintExportIn(projects=['__336新__', '__336老__']))
            wb = openpyxl.load_workbook(BytesIO(base64.b64decode(out['content_b64'])))
            assert len(wb.worksheets) == 2
            new, old = wb.worksheets
            assert new.max_row == old.max_row == 5
            assert new['A4'].value == old['A4'].value == 1
            assert new['I4'].value == 80 and old['L4'].value == 80
            assert new['J4'].value is None and old['D4'].value is None
            assert '$J$5' in str(new.print_area) and '$M$5' in str(old.print_area)
            from app.core.errors import MustReportError
            for projects in ([], ['不存在的项目'], ['__336无合同__']):
                with pytest.raises(MustReportError):
                    service.export_print_xlsx(s, projects=projects)
        finally:
            s.rollback()
