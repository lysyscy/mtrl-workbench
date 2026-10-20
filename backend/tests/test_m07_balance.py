# -*- coding: utf-8 -*-
"""`MOD-07` 提取成品「结余」列 —— **手工录入 · 点格即改**（2026-09-18 用户裁定「甲」）离线断言。

口径（主人）：
- 用户 **2026-09-18 裁定**：结余 = **历史结余 · 由用户手工录入** ⇒ 系统**没有任何来源**
  （⛔ 不是收发结存的 `end_amt` 派生 —— 那条口径已被裁定作废）；
- `docs/requirements/report-templates.md` **§五**：该列 = **B 类**（模板列 ↔ 页面列的登记处）；
- 形态 = **甲**：提取页主表加一列「结余」＋ **点格即改**；
  ⭐ 粒度 = **类别 × 仓库 · 不按月**（用户维护的是「当前的结存」，改一次覆盖一次）。

⭐ **判据点（2026-09-18 用户批「立」）**：挂在 `ACC-SETL-EXT-001` 下 **5 点** ——
`001-07` 未录入 ⇒ 成品留空（⛔ 不写 0）· 该列不参与合计行 ｜ `001-08` **`0` 是合法录入值** ｜
`001-09` 录入入口 = **点格即改** · 粒度 = **类别 × 仓库 · 不按月** · 清空 ⇒ 回到未录入 ｜
`001-10` 非法输入（非数字 / 负数 / 认不出的类别 / 空仓库）**一律拒、不落库、原值不变** ｜
`001-11` 已录入值**随导出带上**（导出与页面**同一份取数出口**）。
（⚠️ **加点 = 只动 `验收基线.json` 一处**；「四处同步」是**新增用例**的规矩。）

⭐ 三条最容易写错的：
1. **未录入 ≠ 0**：未录入 ⇒ 库里**没有那一行** ⇒ 导出**如实留空**；`0` 是**合法录入值**（写 0）。
2. **不按月**：唯一键 = `(类别, 仓库)` ⇒ 重复保存是**覆盖**，不会长出第二行。
3. **取数出口只有一份**（`repository`）—— 用**身份断言**钉住，⛔ 不是「两边都能调」那种假绿。
"""

from __future__ import annotations

import inspect
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from app.mod07_setl import extract as X
from app.mod07_setl import repository as R
from app.mod07_setl import service as S
from app.mod07_setl.models import WeldGasBalance
from app.mod07_setl.repository import load_balances
from app.mod07_setl.service import build_extract_export, parse_extract_state, save_balance_cell
from app.report_export import weldgas_export as W
from tests.m07_helpers import (
    BALANCE_EXPECT,
    FakeSession,
    balance_rows,
    extract_dataset,
    indep_balance,
)

MODDIR = Path(__file__).resolve().parents[1] / "app" / "mod07_setl"
CASE = "ACC-SETL-EXT-001"
PARAMS = {"status": "在建", "ext_period": "2026-08"}


def _session(*, with_balances: bool = True) -> FakeSession:
    return FakeSession(extract_dataset(with_balances=with_balances))


def _plan(*, with_balances: bool = True) -> dict:
    return build_extract_export(_session(with_balances=with_balances),
                                parse_extract_state(PARAMS))


def _build(tmp_path, plan: dict) -> str:
    return W.build_xlsx(plan["grids"], tmp_path, plan["report_ym"])


def _sheet(path, name: str) -> list[list]:
    """读回一个 sheet 的全部格（值模式）。"""
    wb = openpyxl.load_workbook(str(path), data_only=True)
    try:
        ws = wb[name]
        return [[ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
                for r in range(1, ws.max_row + 1)]
    finally:
        wb.close()


def _balance_cells(tmp_path, plan: dict) -> dict[str, list]:
    """成品的「结余」列（B 列）数据行 —— `{类别: [逐行值]}`。⚠️ 行数从**计划**取，不写死。"""
    out = Path(_build(tmp_path, plan))
    got: dict[str, list] = {}
    for kind, sheet in W.SHEET_OF_KIND.items():
        rows = _sheet(out, sheet)
        n = len(plan["grids"][kind]["rows"])
        got[kind] = [rows[1 + i][1] for i in range(n)]
    return got


def _warehouses(plan: dict, kind: str) -> list[str]:
    return [r["warehouse"] for r in plan["grids"][kind]["rows"]]


# ============================================================ 取数：未录入 ≠ 0

def test_未录入的键不出现而不是0():
    """⭐ `load_balances` 的**形状**：未录入的 `(类别, 仓库)` **键不出现**（域值 `None`）。

    ⛔ 若给它补个 0，「没填」与「填了 0」就再也分不开 —— 用户会读到一句没写过的「结余 = 0」。
    """
    bals = load_balances(_session())
    assert "乙仓" not in (bals.get("气体") or {}), "气体 · 乙仓 未录入，键不该出现（补 0 = 冒充确实为 0）"
    assert (bals.get("气体") or {}).get("甲仓") == Decimal("88.50")
    assert (bals.get("焊材") or {}).get("甲仓") == Decimal("1200.00")


def test_0是合法录入值会照常出现():
    """⭐ `0` 与「未录入」是**两件事**：`0` 是合法录入值 ⇒ 它**必须出现在字典里**（值 = 0）。"""
    bals = load_balances(_session())
    assert "乙仓" in (bals.get("焊材") or {}), "焊材 · 乙仓 录的是 0 ⇒ 键必须在"
    assert bals["焊材"]["乙仓"] == Decimal("0.00")


def test_取数出口只有一份实现():
    """⭐ 要证明「只有一份」，就得断言**同一性**（`is`），⛔ 不是「两边都能调」（各包一层照样过）。"""
    assert S.load_balances is R.load_balances, "服务层又包了一层 ⇒ 改动仍会分叉"
    assert S.save_balance is R.save_balance, "服务层又包了一层 ⇒ 改动仍会分叉"
    # 反向守卫：本模块只许 `repository` 定义它们（防「下次又抄一份」）
    for f in MODDIR.glob("*.py"):
        text = f.read_text(encoding="utf-8")
        for name in ("load_balances", "save_balance"):
            assert text.count(f"def {name}(") == (1 if f.name == "repository.py" else 0), \
                f"`{f.name}` 里又定义了一份 `{name}`（一份内容只许一个主人）"


# ============================================================ 导出：带值 / 留空 / 0 照写

def test_录入后导出按类别仓库带上该值(tmp_path):
    """「录入生效」这一维：成品的「结余」列逐格 = **独立手写期望表**（不问实现）。"""
    plan = _plan()
    cells = _balance_cells(tmp_path, plan)
    for kind in W.SHEET_OF_KIND:
        for wh, value in zip(_warehouses(plan, kind), cells[kind]):
            want = indep_balance(kind, wh)
            if want is None:
                continue
            assert value is not None, f"[{kind}] {wh} 录了 {want} 却没写进成品"
            assert abs(Decimal(str(value)) - want) < Decimal("0.005"), \
                f"[{kind}] {wh}：{value} ≠ 独立复算 {want}"


def test_未录入时导出该格如实留空(tmp_path):
    """⭐ 与 `TAIL-43` 同族的那条：**没有来源就留空** —— ⛔ 不写 0、⛔ 不带模板示例值。"""
    plan = _plan(with_balances=False)                     # 库里一条结余都没有
    assert all(r["balance"] is None
               for g in plan["grids"].values() for r in g["rows"]), "网格里不该凭空长出结余"
    for kind, values in _balance_cells(tmp_path, plan).items():
        assert all(v in (None, "") for v in values), f"[{kind}] 未录入却写出了值：{values}"


def test_0照写0而不是留空(tmp_path):
    """⭐ `0` 是**合法录入值** ⇒ 成品里那一格必须是 `0`（与「未录入的空白」可区分）。"""
    plan = _plan()
    cells = _balance_cells(tmp_path, plan)
    got = cells["焊材"][_warehouses(plan, "焊材").index("乙仓")]
    assert got is not None, "录了 0 却被写成空（把 0 与「未录入」混为一谈了）"
    assert abs(Decimal(str(got)) - Decimal("0.00")) < Decimal("0.005"), f"0 被写成了 {got!r}"


def test_结余列不参与合计行(tmp_path):
    """照模板：**合计行**的「结余」格留空（各仓结余相加没有业务含义 —— 模板即如此）。"""
    plan = _plan()
    out = Path(_build(tmp_path, plan))
    for kind, sheet in W.SHEET_OF_KIND.items():
        rows = _sheet(out, sheet)
        n = len(_warehouses(plan, kind))
        assert rows[1 + n][1] in (None, ""), \
            f"[{sheet}] 合计行的结余格应当留空（照模板），实得 {rows[1 + n][1]!r}"


# ============================================================ 写动作（点格即改）

def test_保存一格只动这一格且粒度不按月():
    """粒度 = **类别 × 仓库**：改一格不影响别的格。"""
    s = _session()
    r = save_balance_cell(s, kind="焊材", warehouse="甲仓", raw="999.90")
    assert r["ok"] is True and r["value"] == Decimal("999.90")
    bals = load_balances(s)
    assert bals["焊材"]["甲仓"] == Decimal("999.90")
    assert bals["焊材"]["乙仓"] == Decimal("0.00"), "同类别别的仓库被带改了"
    assert bals["气体"]["甲仓"] == Decimal("88.50"), "别的类别被带改了"


def test_重复保存是覆盖不长出第二行():
    """⭐ **不按月**的结构性判据：同一格存两次 ⇒ 覆盖（行数不变）。"""
    s = _session()
    before = len(s.by_model[WeldGasBalance])
    save_balance_cell(s, kind="焊材", warehouse="甲仓", raw="1.00")
    save_balance_cell(s, kind="焊材", warehouse="甲仓", raw="2.00")
    assert len(s.by_model[WeldGasBalance]) == before, \
        "同一格存两次长出了第二行 ⇒ 键不是「类别 × 仓库」"
    assert load_balances(s)["焊材"]["甲仓"] == Decimal("2.00")


def test_写动作签名里没有月份参数():
    """⭐ **不按月**的接口层面判据：签名里不该出现月份（出现 = 粒度偷偷变成按月）。"""
    params = set(inspect.signature(save_balance_cell).parameters)
    assert not ({"month", "ym", "period"} & params), \
        f"结余的粒度是「类别 × 仓库 · 不按月」，签名里却出现了月份参数：{sorted(params)}"


def test_清空后回到未录入():
    """清空 ⇒ **回到未录入**（导出里如实留空），⛔ 不是写成 0。"""
    s = _session()
    r = save_balance_cell(s, kind="焊材", warehouse="甲仓", raw="")
    assert r["ok"] is True and r["value"] is None
    assert r["message"] == X.MSG_BAL_CLEARED
    assert "甲仓" not in (load_balances(s).get("焊材") or {}), "清空后键还在 ⇒ 会被读成 0"

    # 清空之后导出：该格**如实留空**
    s2 = _session()
    save_balance_cell(s2, kind="焊材", warehouse="甲仓", raw="")
    plan = build_extract_export(s2, parse_extract_state(PARAMS))
    whs = _warehouses(plan, "焊材")
    assert plan["grids"]["焊材"]["rows"][whs.index("甲仓")]["balance"] is None


def test_非数字被拒不写库原值不变():
    s = _session()
    before = load_balances(s)
    r = save_balance_cell(s, kind="焊材", warehouse="甲仓", raw="abc")
    assert r["ok"] is False and r["message"] == X.MSG_BAL_BAD, f"没给中文原因：{r}"
    assert load_balances(s) == before, "非法输入却改了库（⛔ 也不许静默变 0）"


def test_负数被拒不写库原值不变():
    s = _session()
    before = load_balances(s)
    r = save_balance_cell(s, kind="焊材", warehouse="甲仓", raw="-1.00")
    assert r["ok"] is False and r["message"] == X.MSG_BAL_BAD, f"没给中文原因：{r}"
    assert load_balances(s) == before


def test_认不出的类别被拒():
    s = _session()
    r = save_balance_cell(s, kind="电缆", warehouse="甲仓", raw="1.00")
    assert r["ok"] is False and r["message"] == X.MSG_BAL_NO_KIND
    assert load_balances(s) == load_balances(_session()), "被拒的类别却落了库"


def test_空仓库被拒():
    s = _session()
    r = save_balance_cell(s, kind="焊材", warehouse="   ", raw="1.00")
    assert r["ok"] is False and r["message"] == X.MSG_BAL_NO_WH
    assert not (load_balances(s).get("焊材") or {}).get(""), "空仓库名被落了库"


def test_千分位与空格的输入照常解析():
    """金额解析**唯一出口 = `core.money`**（剥千分位 / 去空白）⇒ 界面上粘贴带逗号的数也能存。"""
    s = _session()
    r = save_balance_cell(s, kind="焊材", warehouse="甲仓", raw=" 1,234.56 ")
    assert r["ok"] is True and r["value"] == Decimal("1234.56"), r


# ============================================================ 装配（零件齐 ≠ 接线上了）

def test_导出真的去读了库里的结余():
    """⭐ 与「零件齐、装配缺」同款的反向守卫：`build_extract_export` 必须**真的调用**取数出口，
    且在**组装网格之前** —— ⛔ 漏了这一环会让整条链「看着对、永远留空」。"""
    src = inspect.getsource(S.build_extract_export)
    i_read = src.find("load_balances(s)")
    i_grid = src.find("extract_template_grid(")
    assert i_read >= 0, "`build_extract_export` 没读结余（零件在、装配缺）"
    assert i_grid > i_read, "读结余排在组装网格之后 ⇒ 网格里拿到的还是上一版"
    assert "balances=balances" in src, "读到了却没传给 `extract_template_grid`"
    assert "balances=None" not in src, "还留着写死留空那一行 —— 录入值永远进不了成品"


def test_桩自己成立():
    """前提自检：样本里「有录入 / 只有 0 / 完全未录入」三种情形都在（否则上面的断言会失真）。"""
    assert indep_balance("焊材", "甲仓") == Decimal("1200.00")
    assert indep_balance("焊材", "乙仓") == Decimal("0.00")
    assert indep_balance("焊材", "丙仓") is None
    assert indep_balance("气体", "乙仓") is None
    assert len(balance_rows()) == sum(len(v) for v in BALANCE_EXPECT.values())


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(pytest.main([__file__, "-v"]))
