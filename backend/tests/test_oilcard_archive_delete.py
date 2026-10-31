"""Card retirement and local invoice removal retain identity, money and safe undo."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook
import pytest

from app.core import history
from app.core import session as SESS
from app.core.db import session_scope
from app.mod13_oilcard.models import OilCard, OilInvoice
from tests.oilcard_invoice_samples import CARD, invoice_pdf
from tests.test_oilcard_archives import scene, create, edit, read  # noqa: F401
from tests.test_oilcard_invoice_adjustments import allocation, allocate, correct, credit, replay
from tests.test_oilcard_invoices import invoices, receive, setup_account
from tests.test_oilcard_ledger import entry, ledger


def remove(client, row):
    return client.post(f"/api/m13/invoices/{row['id']}/delete", json={'version': row['version']})


def test_card_archive_preserves_existing_pairs_money_and_undo_but_blocks_new_candidates(scene):
    client, (project, _) = scene
    primary = create(client, 'cards', kind='primary')
    card = create(client, 'cards', card_no=CARD, parent_card_id=primary['id'])
    account = create(client, 'accounts', name='仍保留旧卡的账户', project_id=project, card_id=card['id'])
    other = create(client, 'accounts', name='另一个账户', project_id=project)
    saved_entry = entry(client, account, '100')
    received = receive(client, invoice_pdf())['item']
    archived = edit(client, 'cards', card, status='archived')
    operation = client.get('/api/history/status').json()['undo']
    replay(client, 'undo', operation)
    assert next(c for c in read(client)['cards'] if c['id'] == card['id']) == card
    replay(client, 'redo', operation)
    assert ledger(client)['entries'] == [saved_entry]
    assert invoices(client)['total_amount'] == '3000.00'
    assert edit(client, 'accounts', account, remark='旧关系仍可维护')['card_id'] == card['id']
    payload = {k: v for k, v in archived.items() if k not in ('id', 'request_id', 'status')}
    legacy = client.put(f"/api/m13/cards/{card['id']}", json={**payload, 'remark': '旧客户端省略状态'})
    assert legacy.status_code == 200 and legacy.json()['item']['status'] == 'archived'
    assert client.post('/api/m13/accounts', json={'request_id': str(uuid4()), 'name': '新关系',
        'project_id': project, 'card_id': card['id']}).status_code == 400
    assert client.post(f"/api/m13/accounts/{other['id']}/replace-card", json={
        'request_id': str(uuid4()), 'version': other['version'], 'card_id': card['id']}).status_code == 400
    pending = receive(client, invoice_pdf(number='20261031000000000002'))['item']
    assert pending['status'] == 'pending' and pending['allocations'] == []
    assert invoices(client)['total_amount'] == '3000.00'
    primary = edit(client, 'cards', primary, status='archived')
    current = next(c for c in read(client)['cards'] if c['id'] == card['id'])
    assert edit(client, 'cards', current, remark='保留已归档主卡关系')['parent_card_id'] == primary['id']
    assert client.post('/api/m13/cards', json={'request_id': str(uuid4()), 'parent_card_id': primary['id']}).status_code == 400
    workbook = load_workbook(BytesIO(client.get('/api/m13/export').content))
    rows = list(workbook['实体油卡'].values)
    assert len(rows) == 3 and rows[0][-2:] == ('状态', '备注')
    assert [row[-2] for row in rows[1:]] == ['归档', '归档']
    assert next(c for c in read(client)['accounts'] if c['id'] == account['id'])['status'] == 'active'
    current = next(c for c in read(client)['cards'] if c['id'] == card['id'])
    assert edit(client, 'cards', current, status='active')['status'] == 'active'
    assert received['allocations'][0]['card_id'] == card['id']


def test_delete_whole_invoice_excludes_totals_exports_keeps_original_aliases_and_undo(scene):
    client, (project, other) = scene
    first = setup_account(client, project)
    second = create(client, 'accounts', name='分配乙', project_id=other)
    content = invoice_pdf()
    original = receive(client, content)['item']
    corrected = correct(client, original, [allocation(first, '2000'), allocation(second, '1000')],
        invoice_no='20261031000000000009')
    assert corrected.status_code == 200, corrected.text
    row = corrected.json()['item']
    assert remove(client, original).status_code == 409
    response = remove(client, row)
    assert response.status_code == 200 and response.headers['X-Operation-History'] == 'recorded'
    deleted = response.json()['item']
    assert deleted['deleted'] is True and deleted['status'] == 'registered'
    assert deleted['allocations'] == row['allocations'] and deleted['originals'] == row['originals']
    operation = client.get('/api/history/status').json()['undo']
    assert operation['count'] == 1
    assert invoices(client) == {'invoices': [], 'total_amount': '0.00'}
    assert client.get('/api/m13/reconciliation').json()['invoiced_amount'] == '0.00'
    workbook = load_workbook(BytesIO(client.get('/api/m13/export').content))
    assert list(workbook['发票台账'].values)[1][0] == '合计'
    assert list(workbook['发票分配'].values)[1][0] == '合计'
    source_id = row['originals'][0]['id']
    assert client.get(f'/api/m13/invoices/originals/{source_id}').content == content
    receipt = client.get('/api/m13/invoices/receipts/' + sha256(content).hexdigest()).json()
    assert receipt['received'] is True and receipt['receipt_status'] == 'deleted'
    assert '不会再次登记金额' in receipt['receipt_message']
    assert client.put(f"/api/m13/invoices/{row['id']}/remark", json={
        'version': deleted['version'], 'remark': '不允许改已删除票'}).status_code == 400
    assert allocate(client, deleted, [allocation(first, '3000')]).status_code == 400
    assert correct(client, deleted, [allocation(first, '3000')]).status_code == 400
    client.cookies.set(SESS.COOKIE_NAME, SESS.make('__oilcard_other_actor__'))
    for pdf in (content, invoice_pdf(variant='same old number new file'),
            invoice_pdf(number='20261031000000000009', buyer='同票号冲突证据')):
        repeated = receive(client, pdf)
        assert repeated['item']['id'] == row['id'] and repeated['item']['deleted'] is True
        assert repeated['receipt_status'] == 'deleted'
    assert invoices(client)['total_amount'] == '0.00'
    client.cookies.set(SESS.COOKIE_NAME, SESS.make('__oilcard_actor__'))
    replay(client, 'undo', operation)
    restored = invoices(client)['invoices'][0]
    assert not restored['deleted'] and restored['allocations'] == row['allocations']
    assert invoices(client)['total_amount'] == '3000.00'
    replay(client, 'redo', operation)
    assert invoices(client)['total_amount'] == '0.00'


def test_delete_credit_requires_order_restores_net_and_keeps_parent_conflict_guard(scene):
    client, (project, _) = scene
    account = setup_account(client, project)
    original = receive(client, invoice_pdf())['item']
    red = receive(client, invoice_pdf(number='20261031000000000002', amount='-1000.00',
        subtotal='-884.96', tax='-115.04'))['item']
    result = credit(client, red, original, account)
    assert result.status_code == 200, result.text
    red = result.json()['item']
    original = invoices(client)['invoices'][0]
    refused = remove(client, original)
    assert refused.status_code == 400 and '先逐张删除关联红票' in refused.json()['message']
    assert invoices(client)['total_amount'] == '2000.00'
    assert remove(client, red).status_code == 200
    operation = client.get('/api/history/status').json()['undo']
    assert operation['count'] == 2 and invoices(client)['total_amount'] == '3000.00'
    replay(client, 'undo', operation)
    assert invoices(client)['total_amount'] == '2000.00'
    replay(client, 'redo', operation)
    original = invoices(client)['invoices'][0]
    client.cookies.set(SESS.COOKIE_NAME, SESS.make('__oilcard_other_actor__'))
    changed = client.put(f"/api/m13/invoices/{original['id']}/remark", json={
        'version': original['version'], 'remark': '其他账号后续修改'})
    assert changed.status_code == 200
    client.cookies.set(SESS.COOKIE_NAME, SESS.make('__oilcard_actor__'))
    assert client.post('/api/history/undo', json={'operation_id': operation['id']}).status_code == 409
    assert remove(client, changed.json()['item']).status_code == 200
    assert invoices(client)['total_amount'] == '0.00'


def legacy_history(operation_id):
    """Represent immutable snapshots recorded before the two additive columns existed."""
    with session_scope() as session:
        operation = session.get(history.OperationHistory, operation_id)
        changes = deepcopy(operation.changes)
        for change in changes:
            for table, values in [(change['table'], change['before']), (change['table'], change['after']),
                    *((parent['table'], parent['data']) for parent in change.get('parents', []))]:
                if values is not None and table in ('oil_card', 'oil_invoice'):
                    values.pop('status' if table == 'oil_card' else 'deleted')
        operation.changes = changes


def test_pre_column_history_and_parent_guards_replay_without_rewriting_snapshots(scene):
    client, (project, _) = scene
    account = setup_account(client, project)
    row = receive(client, invoice_pdf())['item']
    # The invoice's old snapshot and the allocation's old card guard both replay.
    second = create(client, 'accounts', name='旧历史另一账户', project_id=project, card_id=account['card_id'])
    noted = client.put(f"/api/m13/invoices/{row['id']}/remark", json={
        'version': row['version'], 'remark': '连续的旧历史'})
    assert noted.status_code == 200
    earlier = client.get('/api/history/status').json()['undo']
    result = allocate(client, noted.json()['item'], [allocation(second, '3000')])
    assert result.status_code == 200
    operation = client.get('/api/history/status').json()['undo']
    legacy_history(earlier['id'])
    legacy_history(operation['id'])
    replay(client, 'undo', operation)
    assert invoices(client)['invoices'][0]['allocations'] == row['allocations']
    replay(client, 'undo', earlier)
    assert invoices(client)['invoices'][0] == row
    replay(client, 'redo', earlier)
    replay(client, 'redo', operation)
    assert invoices(client)['invoices'][0]['allocations'][0]['account_id'] == second['id']
    with session_scope() as session:
        changes = session.get(history.OperationHistory, operation['id']).changes
        for change in changes:
            if change['table'] == 'oil_invoice':
                assert 'deleted' not in change['after']
            for parent in change['parents']:
                if parent['table'] == 'oil_card':
                    assert 'status' not in parent['data']


@pytest.mark.parametrize('table', ['oil_card', 'oil_invoice'])
def test_pre_column_history_rejects_actual_nondefault_archive_or_deletion(scene, table):
    client, (project, _) = scene
    account = setup_account(client, project)
    row = receive(client, invoice_pdf())['item']
    second = create(client, 'accounts', name='旧历史另一账户', project_id=project, card_id=account['card_id'])
    response = allocate(client, row, [allocation(second, '3000')])
    assert response.status_code == 200
    operation = client.get('/api/history/status').json()['undo']
    legacy_history(operation['id'])
    with session_scope() as session:
        if table == 'oil_card':
            session.get(OilCard, account['card_id']).status = 'archived'
        else:
            session.get(OilInvoice, row['id']).deleted = True
    assert client.post('/api/history/undo', json={'operation_id': operation['id']}).status_code == 409
