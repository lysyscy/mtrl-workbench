"""过程报表来源、资格、单次读取与成品：独立金额预期，回滚只在测试库事务内验证。"""
from __future__ import annotations

from collections import Counter
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4
from zipfile import ZipFile

import openpyxl
import pytest
from sqlalchemy.orm import Session

from app.core.db import get_engine
from app.core.errors import MustReportError
from app.mod06_proc import repository as R, service as M6
from app.report_export import service as EX

INV, HIST, MECH, MIX, UNKNOWN, EMPTY = (
    "库存仓", "历史仓", "机械仓", "混合仓", "未登记仓", "无数据档案仓")


@pytest.fixture
def sources(monkeypatch, tmp_path):
    """只提供假来源和展示名；导出真实工作簿到临时目录，不访问i8。"""
    inv = [
        {"id": 101, "warehouse": INV, "ym": "2026-01", "direction": "入库",
         "material_property": "主材", "amount": Decimal("11"), "adj_amount": Decimal("15")},
        {"id": 201, "warehouse": MIX, "ym": "2026-01", "direction": "入库",
         "material_property": "主材", "amount": Decimal("10"), "adj_amount": Decimal("13")},
        {"id": 202, "warehouse": MIX, "ym": "2026-01", "direction": "入库",
         "material_property": "主材", "amount": Decimal("20"), "adj_amount": Decimal("17")},
        {"warehouse": UNKNOWN, "ym": "2026-01", "direction": "入库",
         "material_property": "主材", "amount": Decimal("19")},
    ]
    history = [{"warehouse": HIST, "ym": "2025-10", "in_主辅材": Decimal("17")},
               {"warehouse": MIX, "ym": "2025-12", "in_主辅材": Decimal("6")}]
    manual = [{"warehouse": MECH, "ym": "2026-02", "mech_0": Decimal("7")},
              {"warehouse": MIX, "ym": "2026-03", "mech_1": Decimal("0")}]
    cutoffs = {INV: "2026-01", HIST: "2025-11", MECH: "2025-11", MIX: "2025-11"}
    displays = {INV: "库存展示", HIST: "历史展示", MECH: "机械展示",
                MIX: "混合展示", EMPTY: "空档案展示"}
    calls = Counter()

    def loader(name, value):
        def load(*args, **kwargs):
            calls[name] += 1
            return value
        return load

    monkeypatch.setattr(R, "load_inv_rows", loader("inventory", inv))
    monkeypatch.setattr(R, "load_mech", loader("mechanical", manual))
    monkeypatch.setattr(R, "load_history", loader("history", history))
    monkeypatch.setattr(M6, "load_cutoffs", loader("material_cutoffs", cutoffs))
    monkeypatch.setattr(M6, "load_batches", loader("batches", []))
    monkeypatch.setattr(M6, "displays_by_key", loader("displays", displays))

    def cutoff(session, *, warehouse):
        calls["single_cutoff"] += 1
        return cutoffs.get(warehouse, "")

    monkeypatch.setattr(R, "load_cutoff", cutoff)
    monkeypatch.setattr(EX, "EXPORT_ROOT", tmp_path / "exports")
    return SimpleNamespace(calls=calls, displays=displays, cutoffs=cutoffs)


def export_ok(warehouses, tab=0):
    """有有效数据的项目必须能导出，拒绝时给出明确行为断言红。"""
    try:
        return EX.export_proc(None, warehouses, tab, ym="2026-03")
    except MustReportError as exc:
        raise AssertionError(f"合法报表来源被导出拒绝：{exc}") from None


def workbook(path_or_bytes):
    """读真实xlsx的标题与月度数据，不从矩阵实现计算期望。"""
    source = BytesIO(path_or_bytes) if isinstance(path_or_bytes, bytes) else path_or_bytes
    book = openpyxl.load_workbook(source, data_only=True)
    try:
        sheet = book["汇总表"]
        rows = {}
        for line in range(4, sheet.max_row + 1):
            month = sheet.cell(line, 2).value
            if isinstance(month, str) and len(month) == 7 and month[4] == "-":
                rows[month] = [sheet.cell(line, col).value for col in range(3, 19)]
        return sheet.cell(1, 1).value, rows
    finally:
        book.close()


@pytest.mark.parametrize("tab", [0, 1, 2])
@pytest.mark.parametrize("warehouse,month,raw,adjusted", [
    (INV, "2026-01", 11, 15), (HIST, "2025-10", 17, 17), (MECH, "2026-02", 7, 7)])
def test_each_effective_source_exports_independent_amounts(sources, warehouse, month, raw, adjusted, tab):
    page = M6.build_page_data(None, M6.parse_state({"pj": warehouse, "tab": str(tab)}))
    assert warehouse in page["projects"] and page["warehouse"] == warehouse
    path = export_ok([warehouse], tab)
    title, rows = workbook(path)
    assert title == sources.displays[warehouse] + "材料机械消耗汇总表"
    assert set(rows) == {month}
    value = raw if tab == 0 else adjusted if tab == 1 else adjusted - raw
    if warehouse == MECH:
        assert rows[month][:10] == [None] * 10
        assert rows[month][10] == (value or None)
        book = openpyxl.load_workbook(path, data_only=True)
        try:
            subtotal = next(row for row in book["汇总表"].iter_rows() if row[1].value == "小计")
            assert [cell.value for cell in subtotal[2:12]] == [None] * 10
            assert subtotal[12].value == (value or None)
        finally:
            book.close()
    else:
        assert rows[month][0] == (value or None)
    assert page["cutoff_ym"] == sources.cutoffs[warehouse]


def test_zip_contents_names_and_sources_loaded_once(sources):
    path = export_ok([INV, HIST, MECH], 1)
    assert path.name == "过程资料_2026-03.zip"
    names = {"库存展示_过程资料_表② 调整后表_2026-03.xlsx": ("2026-01", 0, 15),
             "历史展示_过程资料_表② 调整后表_2026-03.xlsx": ("2025-10", 0, 17),
             "机械展示_过程资料_表② 调整后表_2026-03.xlsx": ("2026-02", 10, 7)}
    with ZipFile(path) as archive:
        assert set(archive.namelist()) == set(names)
        for name, (month, col, amount) in names.items():
            _, rows = workbook(archive.read(name))
            assert rows[month][col] == amount
    assert {key: sources.calls[key] for key in
            ("inventory", "mechanical", "history", "material_cutoffs", "displays")} == {
                "inventory": 1, "mechanical": 1, "history": 1, "material_cutoffs": 1, "displays": 1}
    assert sources.calls["single_cutoff"] == 0
    # 机械截止只管导入；报表不再读它来遮挡已存金额。
    assert sources.calls["mechanical_cutoffs"] == 0


def test_page_shape_document_differences_and_monthly_export_are_distinct(sources):
    page = M6.build_page_data(None, M6.parse_state({"pj": MIX, "tab": "2"}))
    assert set(page) == {"state", "projects", "warehouse", "rows_raw", "rows_adjusted", "rows",
                         "totals", "segments", "grand_total", "diffs", "diff_totals", "cutoff_ym",
                         "history_batches", "mech_months", "can_edit_history", "export"}
    assert page["rows"] is None and len(page["diffs"]) == 2
    assert {row["id"]: row["delta"] for row in page["diffs"]} == {201: Decimal("3"), 202: Decimal("-3")}
    assert page["rows_raw"][0]["in_主辅材"] == Decimal("6")  # 截止之后的历史金额仍可见。
    assert page["rows_raw"][2]["material_empty"] is True
    assert page["rows_raw"][2]["mech_1"] == Decimal("0")
    _, rows = workbook(export_ok([MIX], 2))
    assert set(rows) == {"2025-12", "2026-01", "2026-03"}
    assert rows["2026-01"][0] is None  # 差额相抵为0，导出按确认口径留空。
    assert rows["2026-03"][:10] == [None] * 10
    assert rows["2026-03"][11] is None
    matrices = M6.build_matrices(None)
    assert set(matrices) == {"projects", "matrices"}
    assert set(matrices["matrices"][MIX]) == {"raw", "adjusted", "cutoff_ym", "totals", "segments", "diffs", "diff_totals"}
    assert matrices["matrices"][MIX]["diffs"] == page["diffs"]
    assert matrices["matrices"][MIX]["diff_totals"] == page["diff_totals"]
    assert matrices["matrices"][MECH]["diffs"] == []
    assert matrices["matrices"][MECH]["diff_totals"]["before"] is None
    assert EMPTY not in matrices["projects"]


def test_empty_and_unregistered_errors_remain_export_rules(sources):
    for warehouses, message in (([], "一个项目都没选"), ([EMPTY], "没有数据"), ([UNKNOWN], "不在档案里")):
        with pytest.raises(MustReportError, match=message):
            EX.export_proc(None, warehouses, 1, ym="2026-03")
    sources.displays.pop(MECH)
    with pytest.raises(MustReportError, match="展示名"):
        EX.export_proc(None, [MECH], 1, ym="2026-03")


def test_mechanical_only_subtotals_explicitly_clear_template_formulas(sources):
    from app.report_export import process_export as P

    matrix = M6.build_matrices(None, warehouses=[MECH])["matrices"][MECH]
    subtotal = P.body_spec(matrix["raw"])["tail"][P.OFF_SUB][1]
    assert set(subtotal) == {P.TIME_COL, *(P.COL_OF_KEY[key] for key in M6.D.DATA_KEYS)}
    assert [subtotal[P.COL_OF_KEY[key]] for key in M6.D.DATA_KEYS[:10]] == [""] * 10
    assert subtotal["M"] == 7


def test_bucketed_matrices_keep_precedence_null_zero_order_and_document_trace(monkeypatch):
    inv = [
        {"id": 9, "warehouse": "A", "ym": "2025-11", "direction": "入库",
         "material_property": "主材", "amount": "10.01", "adj_amount": "0",
         "doc_no": "单据", "seq": "2", "material": "材料", "adj_remark": "原因",
         "adj_operator": "操作人", "adj_at": "2026-01-02 03:04"},
        {"id": 8, "warehouse": "B", "ym": "2026-01", "direction": "出库",
         "material_property": "主材", "amount": None},
        {"id": 7, "warehouse": "A", "ym": "2025-11", "direction": "出库",
         "material_property": "未知", "amount": None, "adj_amount": "2.01"},
    ]
    history = [{"warehouse": "A", "ym": "2025-11", "in_主辅材": "999"},
               {"warehouse": "A", "ym": "2025-10", "in_主辅材": "0"}]
    manual = [{"warehouse": "C", "ym": "2026-01", "mech_0": "0"}]
    for name, value in (("load_inv_rows", inv), ("load_history", history), ("load_mech", manual)):
        monkeypatch.setattr(R, name, lambda s, value=value: value)
    monkeypatch.setattr(M6, "load_cutoffs", lambda s: {"A": "2025-11"})
    original = M6.D.build_matrix
    seen = []

    def matrix(**kw):
        assert all(str(row.get("warehouse") or "") == kw["warehouse"]
                   for name in ("inv_rows", "history_rows", "manual_rows") for row in kw[name])
        seen.append((kw["warehouse"], len(kw["inv_rows"])))
        result = original(**kw)
        legacy = original(**{**kw, "inv_rows": inv, "history_rows": history, "manual_rows": manual})
        assert result == legacy
        return result

    monkeypatch.setattr(M6.D, "build_matrix", matrix)
    data = M6.build_matrices(None, warehouses=["C", "A", "C", "不存在"])
    assert data["projects"] == ["A", "B", "C"]
    assert list(data["matrices"]) == ["C", "A"]
    assert seen == [("C", 0), ("C", 0), ("A", 2), ("A", 2)]
    a = data["matrices"]["A"]
    assert [row["id"] for row in a["diffs"]] == [9, 7]
    assert a["diffs"][0] == {"id": 9, "doc_no": "单据", "seq": "2", "material": "材料",
                              "warehouse": "A", "before": Decimal("10.01"), "after": Decimal("0"),
                              "delta": Decimal("-10.01"), "remark": "原因", "operator": "操作人",
                              "at": "2026-01-02 03:04"}
    assert a["diff_totals"]["delta"] == Decimal("-8.00")
    assert a["raw"][0]["in_主辅材"] == 0 and a["raw"][1]["in_主辅材"] == Decimal("10.01")
    assert a["adjusted"][1]["in_主辅材"] == 0
    assert data["matrices"]["C"]["raw"][0]["in_主辅材"] is None


@pytest.fixture
def pg_session():
    """复用仓库外层事务/SAVEPOINT配方，仅清隔离库的过程资料表且整体回滚。"""
    from app import models_registry  # noqa: F401
    from app.mod06_proc.models import (ProcMechImportChange, ProcMechManual, ProcHistorySnapshot,
                                       ProcHistoryBatch, ProcCutoff, ProcMaterialManual)

    engine = get_engine()
    assert engine.url.database and engine.url.database.endswith("_test")
    with engine.connect() as connection:
        outer = connection.begin()
        try:
            ProcMaterialManual.__table__.create(connection, checkfirst=True)
            for model in (ProcMechImportChange, ProcMechManual, ProcHistorySnapshot, ProcHistoryBatch, ProcCutoff):
                connection.execute(model.__table__.delete())
            with Session(bind=connection, join_transaction_mode="create_savepoint") as session:
                yield session
        finally:
            outer.rollback()


def test_inactive_historical_batch_keeps_trace_and_manual_only_project_still_exports(pg_session, monkeypatch, tmp_path):
    from app.mod06_proc.models import ProcHistoryBatch, ProcHistorySnapshot
    warehouse = '__report_' + uuid4().hex
    monkeypatch.setattr(R, 'load_inv_rows', lambda s: [])
    monkeypatch.setattr(M6, 'displays_by_key', lambda s: {warehouse: '历史验证展示'})
    monkeypatch.setattr(EX, 'EXPORT_ROOT', tmp_path / 'exports')
    R.save_mech(pg_session, warehouse=warehouse, ym='2026-03', values={'mech_1': 7})
    batch = ProcHistoryBatch(batch_no='__report_batch_' + uuid4().hex, warehouse_name=warehouse, status='active')
    snapshot = ProcHistorySnapshot(batch_no=batch.batch_no, warehouse_name=warehouse, ym='2026-01', in_0=10)
    pg_session.add_all([batch, snapshot])
    pg_session.flush()
    _, rows = workbook(EX.export_proc(pg_session, [warehouse], 0, ym='2026-03'))
    assert rows['2026-01'][0] == 10
    batch.status = 'rolled_back'
    pg_session.flush()
    assert R.load_history(pg_session) == []
    assert len(R.load_history(pg_session, active_only=False)) == 1
    _, rows = workbook(EX.export_proc(pg_session, [warehouse], 0, ym='2026-03'))
    assert set(rows) == {'2026-03'} and rows['2026-03'][:10] == [None] * 10
    assert rows['2026-03'][11] == 7
