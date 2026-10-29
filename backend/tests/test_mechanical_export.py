"""Protect original four-sheet layout, frozen values and truthful annual mapping."""

from copy import deepcopy
from decimal import Decimal
from openpyxl import load_workbook
import pytest

from app.report_export import mechanical_export as export
from app.core.errors import MustReportError


def sample():
    return {'ym': '2026-09', 'projects': [{'warehouse': '甲', 'name': '甲项目'}],
            'effective': [{'warehouse': '甲', 'ym': '2026-09', 'values': {'mech_2': '100.00', 'mech_3': '0.00'}}],
            'annual_effective': [{'warehouse': '甲', 'ym': '2026-09', 'values': {'mech_2': '100.00', 'mech_3': '0.00'}}],
            'summary': [{'warehouse': '甲', 'ym': '2026-09', 'start_year': 2024, 'start_month': 3, 'cumulative': {'zero': '150.00', 'monthly': '0.00', 'internal': '0.00', 'company': None}}],
            'histories': [{'warehouse': '甲', 'i6': '10.00', 'annual_rows': [], 'meta_at': {'2026-09': {'year': 2025, 'month': 1}}}],
            'details': [{'warehouse': '甲', 'ym': '2026-09', 'category': 'zero', 'issues': [],
                         'totals': {'gross': '99.00', 'net': '96.1165048543689'},
                         'rows': [{'device': '=HYPERLINK("bad")', 'amount': '99.00', 'rate': '99.00', 'net': '96.1165048543689',
                                   'qty': '2小时', 'tax': '3%专票', 'date': '8.21', 'plate': '原文', 'memo': '备注'}]}],
            'i8': [], 'reconciliations': []}


def test_four_original_sheets_values_layout_and_missing_state(tmp_path):
    data = sample()
    wb = export.build(data)
    assert wb.sheetnames == ['工程公司汇总表', '零租设备月报表', '包月外租设备月报', '年度汇总']
    summary, zero, monthly, annual = wb.worksheets
    assert summary['G5'].value == '累计周材费' and summary['L5'].value == '当月周材费'
    assert summary['G6'].value is None and summary['L6'].value is None
    assert summary['J6'].value == Decimal('100.00')
    assert summary['E6'].value == Decimal('150.00')
    assert summary['C6'].value == '2024'
    assert summary['D6'].value == '3'
    assert zero['E6'].value == '2小时'
    assert zero['J6'].value == '8.21'
    assert zero['B6'].data_type == 's'
    assert zero['I6'].value == Decimal('96.1165048543689')
    assert monthly['O4'].value == '不含税价'
    assert monthly['A6'].value == '合计'
    assert all(str(s.print_area) and s.print_title_rows for s in wb)
    assert all(not s.row_dimensions[row].hidden for s in wb for row in range(5 if s is zero or s is monthly else 2, s.max_row + 1))
    assert zero.freeze_panes is None and monthly.freeze_panes is None
    wb.close()


def test_preview_matches_frozen_file_and_next_export_does_not_modify_first(tmp_path):
    data = sample()
    preview = export.preview(data)
    assert [sheet['name'] for sheet in preview['sheets']] == ['工程公司汇总表', '零租设备月报表', '包月外租设备月报', '年度汇总']
    assert [sheet['project_rows'] for sheet in preview['sheets']] == [[{'warehouse':'甲','row':row}] for row in [6,5,5,2]]
    path = export.export(data, tmp_path)
    before = path.read_bytes()
    changed = deepcopy(data)
    changed['effective'][0]['values']['mech_2'] = '200.00'
    second = export.export(changed, tmp_path)
    assert second != path and path.read_bytes() == before
    wb = load_workbook(path, data_only=True)
    assert Decimal(str(wb.worksheets[0]['J6'].value)) == Decimal('100.00')
    values = [cell['value'] for row in preview['sheets'][0]['rows'] for cell in row]
    assert '100.00' in values
    assert not any('忽略' in str(cell.value) for sheet in wb for row in sheet for cell in row)
    wb.close()


def test_monthly_hidden_net_column_survives_empty_template_cells():
    data = sample()
    data['details'].append({'warehouse': '甲', 'ym': '2026-09', 'category': 'monthly',
                            'totals': {'gross': '100.00', 'net': '100.00'}, 'issues': [],
                            'rows': [{'device': '升降平台', 'amount': '100', 'net': '100',
                                      'tax': '普票', 'memo': '原备注', 'date': '8月1日'}]})
    wb = export.build(data)
    monthly = wb.worksheets[2]
    assert monthly['O6'].value == Decimal('100')
    assert monthly['M6'].value == '原备注'
    assert monthly['I6'].value == '8月1日'
    assert monthly['O4'].value == '不含税价'
    wb.close()


def test_unknown_net_does_not_export_a_partial_project_or_sheet_total():
    data=sample()
    data['details'][0]['rows'].append({'device':'待补税率设备','amount':'50.00','net':None})
    data['details'][0]['totals']={'gross':'149.00','net':None}
    wb=export.build(data)
    zero=wb.worksheets[1]
    assert zero['G5'].value.result==Decimal('149.00')
    assert zero['I5'].value.result is None
    assert zero.cell(zero.max_row,9).value.result is None
    wb.close()


def annual_sample(year):
    data = sample()
    data['ym'] = f'{year}-01'
    data['histories'][0].update(base_month='2026-09', annual_rows=[{
        'category': 'zero', 'cells': {'C': '1000', 'D': '10', 'O': '20', 'P': '30', 'AA': '40', 'AH': '99'},
        'source': {'file': '原2026年度报表.xlsx', 'year': 2026}}])
    data['annual_effective'] = [{'warehouse': '甲', 'ym': '2027-01', 'values': {'mech_2': '50'}}]
    return data


@pytest.mark.parametrize('year', [2025, 2026, 2027])
def test_annual_columns_follow_real_year_without_relabeling_source_values(year):
    wb = export.build(annual_sample(year))
    annual = wb['年度汇总']
    months = {c.value: annual.cell(2, c.column).value for c in annual[1]
              if isinstance(c.value, str) and '年' in c.value}
    assert months['2023年2月'] == Decimal('99')
    assert months['2025年1月'] == Decimal('10')
    assert ('2025年12月' in months) == (year >= 2026)
    assert ('2026年1月' in months) == (year >= 2026)
    assert ('2026年12月' in months) == (year >= 2027)
    if year == 2027:
        assert months['2026年1月'] == Decimal('30')
        assert months['2026年12月'] == Decimal('40')
        assert months['2027年1月'] == Decimal('50')
    assert '往年合计' not in [c.value for c in annual[1]]
    wb.close()


def test_rollover_older_total_uses_complete_previous_year_and_confirmed_revision_once():
    data = annual_sample(2027)
    cells = data['histories'][0]['annual_rows'][0]['cells']
    for month in range(1, 13):
        cells[export.get_column_letter(3 + month)] = str(month)
    data['annual_effective'].append({'warehouse': '甲', 'ym': '2025-01', 'values': {'mech_2': '2'}})
    wb = export.build(data)
    annual = wb['年度汇总']
    col = next(c.column for c in annual[1] if c.value == '2025年1月')
    assert annual.cell(2, col).value == Decimal('2')
    # The opening lump sum is never silently allocated to a made-up month.
    assert annual['C2'].value == Decimal('99')
    wb.close()


def test_annual_source_without_reliable_year_is_rejected_before_mislabeling():
    data = annual_sample(2027)
    data['histories'][0].pop('base_month')
    data['histories'][0]['annual_rows'][0]['source'].pop('year')
    with pytest.raises(MustReportError, match='年份依据'):
        export.build(data)


def test_missing_big_table_start_date_does_not_fall_back_to_history():
    data = sample()
    data['summary'][0].update(start_year=None, start_month=None)
    wb = export.build(data)
    assert wb.worksheets[0]['C6'].value == '' and wb.worksheets[0]['D6'].value == ''
    wb.close()


def test_explicit_clear_overrides_history_but_absent_category_keeps_reference():
    data=sample()
    hist=data['histories'][0]
    hist['reference_months']={'2026-09':{'mech_2':'103','mech_3':'206'}}
    effective=data['annual_effective'][0]
    effective['values']={'mech_2':None,'mech_3':None}
    effective['sources']={'mech_2':{'mode':'manual','basis':{'amount':None}}}
    assert export.annual_values(data,'甲',hist,{},'zero')=={}
    assert export.annual_values(data,'甲',hist,{},'monthly')=={'2026-09':Decimal('206')}


@pytest.mark.parametrize('month_count', [28, 29])
def test_annual_totals_and_usage_settlement_difference_follow_month_columns(tmp_path, month_count):
    data = sample()
    months = [f'{2024 + (5 + index) // 12}-{(5 + index) % 12 + 1:02d}' for index in range(month_count)]
    data['ym'] = months[-1]
    data['projects'].append({'warehouse': '乙', 'name': '乙项目'})
    data['annual_effective'] = [
        {'warehouse': warehouse, 'ym': month, 'values': {'mech_2': value}}
        for warehouse, value in [('甲', '100'), ('乙', '200')] for month in months]
    data['histories'] = [{'warehouse': '甲', 'i6': '100'}, {'warehouse': '乙', 'i6': '200'}]
    data['i8'] = [
        {'warehouse': '甲', 'totals': {'usage': '300', 'project': '120', 'team': '30'}},
        {'warehouse': '乙', 'totals': {'usage': '500', 'project': '400', 'team': '50'}}]
    path = export.export(data, tmp_path)
    formulas = load_workbook(path)
    cached = load_workbook(path, data_only=True)
    annual, values = formulas['年度汇总'], cached['年度汇总']
    columns = {cell.value: cell.column_letter for cell in annual[1]}
    difference = columns['班组使用-消耗结算']
    assert difference == ('AK' if month_count == 28 else 'AL')
    usage, settlement = columns['i8'], columns['信息化消耗结算']
    assert f'{usage}6-{settlement}6' in annual[f'{difference}6'].value
    assert values[f'{difference}6'].value == 150
    assert values[f'{difference}11'].value == 50
    assert values[f'{difference}12'].value == 200
    assert annual['C12'].value == '=SUBTOTAL(9,C2:C11)/2'
    assert values['C12'].value == 300
    cumulative = columns['合计']
    assert annual[f'{cumulative}12'].value == f'=SUBTOTAL(9,{cumulative}2:{cumulative}11)/2'
    assert values[f'{cumulative}12'].value == 300 * month_count
    assert annual[f'{usage}12'].value == f'=SUBTOTAL(9,{usage}2:{usage}11)'
    assert values[f'{usage}12'].value == 800
    formulas.close()
    cached.close()


def test_usage_settlement_difference_keeps_unknown_inputs_blank(tmp_path):
    data = sample()
    data['i8'] = [{'warehouse': '甲', 'totals': {'usage': '300', 'project': '120', 'team': None}}]
    path = export.export(data, tmp_path)
    wb = load_workbook(path, data_only=True)
    annual = wb['年度汇总']
    difference = next(cell.column for cell in annual[1] if cell.value == '班组使用-消耗结算')
    assert annual.cell(6, difference).value is None
    assert annual.cell(7, difference).value is None
    wb.close()
