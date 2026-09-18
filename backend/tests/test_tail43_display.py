# -*- coding: utf-8 -*-
"""`TAIL-43`：**模板法成品的两张透视「显示区」必须重建** —— 离线断言。

**病**（真机实测根因）：模板法只重建「数据 sheet ＋ 缓存 ＋ 透视定义」，**不重建显示区**
（`xl/worksheets/sheetN.xml` 里 Excel 存下来的静态渲染值）。本月数据**比模板少**时，
透视的 `location` 会**收缩**，模板旧渲染行落到新 `location` **之外** ⇒
**任何应用都不会去清它们**（Excel 实测：新区域照渲染、区域外旧值原样留着）
⇒ 用户打开成品读到的是**模板那个月**的数。
⚠️ 顺带纠正一条：`refreshOnLoad` 只保证刷新**缓存**（ECMA-376 `CT_PivotCacheDefinition`），
**不保证**清掉 `location` 之外的残留 ⇒ 「靠打开即刷新来顶」**靠不住**。

**断言分三段**
1. **口径不是照印象写的**：用**模板自身**复算它的每一个显示格（模板缓存与显示同月）⇒ 逐格对上；
2. **修好了**：新成品无僵尸行、该空则空、三路一致；
3. **守卫有区分力**：把显示区换回模板旧值 / 往收发结存塞一行旧值 ⇒ **必须报红**。

⚠️ **本片尚无专属判据点**（`TAIL-43` 的落点写的是「需新立判据」）：全部断言用 `@acc()` 显式声明
**不绑定** ⇒ ⚠️ 别把它们当成"已判据化"；立新用例要走 §四 四处同步（需一句话授权）。
"""

from __future__ import annotations

import io
import re
import zipfile
from html import unescape
from pathlib import Path

import openpyxl
import pandas as pd

from app.report_export import kcdj_export as KE
from app.report_export import setl_export as SE
from app.report_export import tpl_engine as TE
from tests.acc_plugin import acc

ROOT = Path(__file__).resolve().parents[2]
TPL = ROOT / "backend" / "app" / "resources" / "templates" / "模板_库存单据明细表.xlsx"
NS = TE.NS


# ----------------------------------------------------------------- 造源 / 读产物
def _synth(with_safe: bool, whs=("甲仓", "乙仓"), safe_amt=100.0, safe_both=False) -> bytes:
    """最小 i8 库存单据源（28 列）。`with_safe=False` ⇒ 一行安全文明施工费都没有。

    `safe_both=True` ⇒ 安全费**同时**落在「入库×数据组1」与「出库×数据组1」两列上 ——
    用来钉「两张透视的列结构不同、⛔ 不许拿过程资料那套 13 列去算安全费那张」（实测踩到过：
    只落一列时**看着全对**）。
    """
    from app.report_export.schema_cols import KCDJ_COLS

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "报表1"
    ws.append(["库存单据明细表"] + [""] * 27)
    ws.append([KCDJ_COLS[i] for i in range(28)])

    def _row(seq, prop, direction, wh, cbs, amt):
        r = [""] * 28
        r[0], r[1], r[2] = seq, "2026", "9"
        r[4], r[9], r[10], r[11] = "某项目", "个", "规格", prop
        r[12], r[13], r[14], r[15], r[16] = "10", "1.0", "1.0", amt, amt
        r[17], r[18], r[27] = direction, wh, cbs
        return r

    safe = KE.SAFE_CBS if with_safe else ""
    n = 1
    for wh in whs:
        ws.append(_row(str(n), "主材", "入库", wh, safe, str(safe_amt)))       # 数据组1 / 入库
        n += 1
        ws.append(_row(str(n), "工具", "入库", wh, "", str(safe_amt * 2)))     # 工具 / 入库
        n += 1
        ws.append(_row(str(n), "辅材", "出库", wh, safe if safe_both else "",
                       str(safe_amt * 4)))                                     # 数据组1 / 出库
        n += 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _build(td: Path, src: bytes) -> Path:
    return Path(KE.build_xlsx(src, td / KE.FILE_PREFIX / "2026年9月"))


def _rows(path, part):
    xml = zipfile.ZipFile(path).read(part).decode("utf-8")
    return {int(m.group(1)): m.group(0)
            for m in re.finditer(r'<row r="(\d+)"[^>]*>.*?</row>', xml, re.S)}


def _num(row_xml, ref):
    m = re.search(r'<c r="%s"(?:\s[^>]*)?>(.*?)</c>' % ref, row_xml or "", re.S)
    if not m:
        return None
    v = re.search(r"<v>([^<]*)</v>", m.group(1))
    return float(v.group(1)) if v else None


def _label(path, rows, ref):
    raw = zipfile.ZipFile(path).read("xl/sharedStrings.xml").decode("utf-8")
    strs = [unescape("".join(re.findall(r"<t[^>]*>([^<]*)</t>", si)))
            for si in re.findall(r"<si>(.*?)</si>", raw, re.S)]
    m = re.search(r'<c r="%s"[^>]*t="s"[^>]*>(.*?)</c>' % ref,
                  rows[int(re.sub(r"\D", "", ref))], re.S)
    v = re.search(r"<v>(\d+)</v>", m.group(1)) if m else None
    return strs[int(v.group(1))] if v else None


def _cache_data(path):
    """成品/模板的**缓存** → 与 `报表1` 同构的元组（"路径 B"复算用）。"""
    import xml.etree.ElementTree as ET

    z = zipfile.ZipFile(path)
    try:
        cd = ET.fromstring(z.read("xl/pivotCache/pivotCacheDefinition1.xml"))
        cr = ET.fromstring(z.read("xl/pivotCache/pivotCacheRecords1.xml"))
    finally:
        z.close()
    fvals, fgrp = [], []
    for f in cd.findall("{%s}cacheFields/{%s}cacheField" % (NS, NS)):
        si = f.find("{%s}sharedItems" % NS)
        fvals.append([(float(c.get("v")) if c.tag.endswith("n") else str(c.get("v")))
                      for c in si] if si is not None else [])
        fg = f.find("{%s}fieldGroup" % NS)
        gi = fg.find("{%s}groupItems" % NS) if fg is not None else None
        dd = fg.find("{%s}discretePr" % NS) if fg is not None else None
        fgrp.append(None if (fg is None or gi is None or dd is None)
                    else {"items": [c.get("v") for c in gi],
                          "map": [int(c.get("v")) for c in dd], "base": int(fg.get("base"))})
    data = []
    for r in cr.findall("{%s}r" % NS):
        xs = [(int(x.get("v")) if x.get("v") is not None else None)
              for x in r.findall("{%s}x" % NS)]
        row = []
        for fi in range(len(fvals)):
            g = fgrp[fi]
            if g is not None:
                bk = xs[g["base"]] if g["base"] < len(xs) else None
                gidx = g["map"][bk] if bk is not None and bk < len(g["map"]) else None
                row.append(g["items"][gidx] if gidx is not None else "")
                continue
            k = xs[fi] if fi < len(xs) else None
            row.append(fvals[fi][k] if k is not None and k < len(fvals[fi]) else "")
        data.append(tuple(row))
    return data


# ================================================================= 1. 口径标定
@acc()
def test_模板自身可标定_缓存复算逐格等于模板显示区():
    """⭐ 口径标定：模板的**缓存**与**显示**是同一个月的 ⇒ 用缓存复算它的每一个显示格，必须逐格对上。

    ⛔ 这一条是「列映射不是照印象写的」的**唯一证据** —— 没有它，后面所有断言都建立在猜测上。
    """
    specs, layout = KE.display_structure()
    cols = layout["columns"]
    assert [c["letter"] for c in cols] == list("BCDEFGHIJKLMN"), "列序跑偏"
    assert len(cols) == 13, "13 列 = 2 方向 ×(5 分组 ＋ 汇总) ＋ 总计"
    data = _cache_data(TPL)
    by, total = KE.aggregate_grid(data, cols, layout["nat2group"])
    s1 = _rows(TPL, "xl/worksheets/sheet1.xml")

    # ① 逐格：模板 7 个仓库 × 13 列
    mis = []
    for r in range(7, 14):
        lb = _label(TPL, s1, "A%d" % r)
        assert lb, "模板第 %d 行没有仓库标签" % r
        for c in cols:
            exp = by.get(lb, {}).get(c["letter"])
            got = _num(s1[r], "%s%d" % (c["letter"], r))
            if (exp is None) != (got is None) or (exp is not None and abs(exp - got) > 0.01):
                mis.append((lb, c["letter"], got, exp))
    assert not mis, "模板显示区逐格复算对不上 ⇒ 列映射搞错了：%s" % mis[:5]

    # ② 总计行
    for c in cols:
        exp = total.get(c["letter"])
        if exp is None:
            continue
        got = _num(s1[14], "%s14" % c["letter"])
        assert got is not None and abs(exp - got) < 0.01, \
            "模板总计行 %s 对不上：%s vs %s" % (c["letter"], got, exp)

    # ③ 第二张透视（安全文明施工费）同一口径标定
    s2 = _rows(TPL, "xl/worksheets/sheet2.xml")
    sby, stot = KE.aggregate_grid(data, specs[1]["columns"], layout["nat2group"],
                                  filter_col=KE.COL_CBS, filter_val=KE.SAFE_CBS)
    assert sby, "模板里居然没有安全文明施工费的分仓数据 ⇒ 标定样本不成立"
    for r in range(4, 9):
        lb = _label(TPL, s2, "A%d" % r)
        got, exp = _num(s2[r], "B%d" % r), sby.get(lb, {}).get("B")
        assert (exp is None) == (got is None) and (exp is None or abs(exp - got) < 0.01), \
            "模板「安全文明施工费」%s 对不上：%s vs %s" % (lb, got, exp)
    assert abs((_num(s2[9], "B9") or 0) - stot["B"]) < 0.01, "模板该透视总计对不上"


# ================================================================= 2. 修好了
@acc()
def test_显示区无僵尸行且行区间随仓库数收缩(tmp_path):
    """⭐ 病灶回归：本月仓库**比模板少** ⇒ 旧行必须**消失**（⛔ 不是留着空白）。"""
    out = _build(tmp_path, _synth(with_safe=False, whs=("甲仓", "乙仓")))
    r1 = _rows(out, "xl/worksheets/sheet1.xml")
    assert sorted(r1) == list(range(3, 7)) + [7, 8, 9], \
        "过程资料行号不是 3..6 ＋ 7..9 ⇒ 有僵尸旧行残留：%s" % sorted(r1)
    assert _label(out, r1, "A7") == "甲仓" and _label(out, r1, "A8") == "乙仓"
    assert _label(out, r1, "A9") == "总计"
    # 甲仓：入库汇总 = 主材100 ＋ 工具200 = 300；出库汇总 = 辅材400；总计 700
    assert _num(r1[7], "B7") == 100.0 and _num(r1[7], "C7") == 200.0
    assert _num(r1[7], "H7") == 400.0
    assert _num(r1[7], "G7") == 300.0, "入库汇总错：%s" % _num(r1[7], "G7")
    assert _num(r1[7], "M7") == 400.0, "出库汇总错：%s" % _num(r1[7], "M7")
    assert _num(r1[7], "N7") == 700.0, "总计错：%s" % _num(r1[7], "N7")


@acc()
def test_该月无该费用时透视如实为空且不留模板旧值(tmp_path):
    """⭐ 用户看到的那一幕：本月**没有**安全文明施工费 ⇒ 那张透视必须**如实为空**。

    ⛔ 三不许：不许显示 0（0 会被读成「确实为 0」）· 不许留模板旧值 · 不许留旧行。
    """
    out = _build(tmp_path, _synth(with_safe=False))
    r2 = _rows(out, "xl/worksheets/sheet2.xml")
    assert sorted(r2) == [1, 3, 4], \
        "安全文明施工费有僵尸旧行（模板旧值残留）：%s" % sorted(r2)
    assert _label(out, r2, "A3") == "仓库" and _label(out, r2, "A4") == "总计"
    assert _num(r2[4], "B4") is None, \
        "该月没有该费用 ⇒ 总计格必须**空**，⛔ 不是 %s" % _num(r2[4], "B4")
    # 页筛选行照旧保留（模板口径：多选 h="1" 选出 安全文明施工费）
    assert _label(out, r2, "A1") == "cbs名称" and _label(out, r2, "B1") == KE.SAFE_CBS


@acc()
def test_有该费用时逐仓与总计都写出来(tmp_path):
    """反向对照（证明上一条不是恒绿）：本月**确有**该费用 ⇒ 行数与值都要长出来。"""
    out = _build(tmp_path, _synth(with_safe=True, safe_amt=111.5))
    r2 = _rows(out, "xl/worksheets/sheet2.xml")
    assert sorted(r2) == [1, 3, 4, 5, 6], "行数不对：%s" % sorted(r2)
    assert _label(out, r2, "A4") == "甲仓" and _num(r2[4], "B4") == 111.5
    assert _label(out, r2, "A5") == "乙仓" and _num(r2[5], "B5") == 111.5
    assert _label(out, r2, "A6") == "总计" and _num(r2[6], "B6") == 223.0


@acc()
def test_页筛选保持模板形态不得改写成pageField_item(tmp_path):
    """⭐ 回答「`<pageField fld="27"/>` 缺 `item` 要不要补成 `item=N`」：**不许补**。

    依据（实测）：模板用 `axisPage` ＋ `multipleItemSelectionAllowed="1"` ＋
    `<item h="1">`（隐藏 = 未选中）表达**多选**式的页筛选，`<pageField>` 上**不带 `item`**。
    `item="N"` 是**单选**语义 ⇒ 改了就是换了一种筛选行为。
    """
    z = zipfile.ZipFile(TPL)
    try:
        cd = z.read("xl/pivotCache/pivotCacheDefinition1.xml").decode("utf-8")
        pt2 = z.read("xl/pivotTables/pivotTable2.xml").decode("utf-8")
    finally:
        z.close()
    assert TE.pivot_page_filter(cd, pt2, KE.COL_CBS) == [KE.SAFE_CBS], \
        "从模板读出的页筛选选中项不是安全文明施工费 ⇒ 模板口径变了"
    assert 'multipleItemSelectionAllowed="1"' in pt2, "模板页筛选不是多选形态？"
    out = _build(tmp_path, _synth(with_safe=True))
    xml = zipfile.ZipFile(out).read("xl/pivotTables/pivotTable2.xml").decode("utf-8")
    page = re.search(r"<pageFields.*?</pageFields>", xml, re.S)
    assert page and page.group(0) == '<pageFields count="1"><pageField fld="27"/></pageFields>', \
        "页筛选被改写了（应逐字保持模板形态）：%s" % (page.group(0) if page else None)


@acc()
def test_三路一致_报表1与缓存与显示区逐格相同(tmp_path):
    """⭐ **三条独立路径**：`报表1`(写进去的明细) → 缓存 → 显示区静态值，逐格相同。"""
    out = _build(tmp_path, _synth(with_safe=True))
    _specs_, layout = KE.display_structure()
    cols = layout["columns"]
    wb = openpyxl.load_workbook(out, read_only=True, data_only=True)
    try:
        data = [tuple(r) for r in wb["报表1"].iter_rows(values_only=True)][1:]
    finally:
        wb.close()
    src_by, src_total = KE.aggregate_grid(data, cols, layout["nat2group"])
    cache_by, _ = KE.aggregate_grid(_cache_data(out), cols, layout["nat2group"])
    assert src_by == cache_by, "报表1 与缓存两条路径不一致"
    r1 = _rows(out, "xl/worksheets/sheet1.xml")
    # ⚠️ 逐行**按表里读到的标签**取，⛔ 不假定字典顺序（行序由引擎定，两者不必同序）
    for r in range(7, 7 + len(src_by)):
        lb = _label(out, r1, "A%d" % r)
        assert lb in src_by, "显示区出现了源里没有的仓库：%s" % lb
        for c in cols:
            got = _num(r1[r], "%s%d" % (c["letter"], r))
            exp = src_by[lb].get(c["letter"])
            assert (exp is None) == (got is None) and (exp is None or abs(exp - got) < 0.01), \
                "显示区 %s%d 与源不一致：%s vs %s" % (c["letter"], r, got, exp)
    assert _num(r1[7 + len(src_by)], "N%d" % (7 + len(src_by))) == src_total.get("N")


# ================================================================= 3. 守卫有区分力
@acc()
def test_反向_把显示区换回模板旧值守卫必须报红(tmp_path):
    """⭐⭐ **反向灵敏度**：把新成品的显示区**换回模板的旧渲染值** ⇒ 守卫必须报红。

    ⛔ 不报红就说明这把尺子恒绿（`TAIL-43` 当初正是"守不住"才漏出去的）。
    """
    out = _build(tmp_path, _synth(with_safe=True))
    bad = tmp_path / "坏_2026年9月.xlsx"
    tpl, zin = zipfile.ZipFile(TPL), zipfile.ZipFile(out)
    try:
        with zipfile.ZipFile(bad, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename.startswith("xl/worksheets/") and \
                        item.filename.endswith(("sheet1.xml", "sheet2.xml")):
                    zout.writestr(item, tpl.read(item.filename))       # 换回模板旧渲染值
                else:
                    zout.writestr(item, zin.read(item.filename))
    finally:
        zin.close()
        tpl.close()
    specs, _layout = KE.display_structure()
    r1 = _rows(bad, "xl/worksheets/sheet1.xml")
    labels = [lb for lb in (_label(bad, r1, "A%d" % r) for r in range(7, 14)) if lb]
    problems = TE.verify_display(str(bad), [
        dict(specs[0], rows=[(w, {}) for w in labels]),
        dict(specs[1], rows=[(w, {}) for w in ("甲仓", "乙仓")]),
    ])
    assert problems, "把显示区换回模板旧值、守卫却全绿 ⇒ 这把尺子是假绿"
    assert any("僵尸旧值" in p for p in problems), \
        "没报出「显示区之外还有行」——正是本片病灶；实报：%s" % problems[:4]
    assert any("vs 缓存重算" in p for p in problems), "没报出「格子的数与缓存对不上」"


@acc()
def test_安全费那张按规定自己的列聚合不许挪用过程资料的列(tmp_path):
    """⭐⭐ **静默错列回归**：两张透视的列结构**不同**（过程资料 13 列 / 安全费 1 列）。

    实测踩到：拿过程资料那套 13 列去算安全费的值 ⇒ `B` 列变成「入库 × 数据组1」；
    若样本恰好只落在该列上，**看着全对、其实是错列**。这里把安全费**同时**放进
    「入库×数据组1」与「出库×数据组1」⇒ 只取单列的实现必然算成一半。
    """
    out = _build(tmp_path, _synth(with_safe=True, safe_amt=100.0, safe_both=True))
    r2 = _rows(out, "xl/worksheets/sheet2.xml")
    # 甲仓：入库 100 ＋ 出库 400 = 500；乙仓同；总计 1000
    assert _num(r2[4], "B4") == 500.0, "安全费只算了一个方向 ⇒ 挪用了过程资料的列：%s" % _num(r2[4], "B4")
    assert _num(r2[5], "B5") == 500.0
    assert _num(r2[6], "B6") == 1000.0, "总计错：%s" % _num(r2[6], "B6")
    # 同一次构建里，「过程资料」那张必须**不受影响**（它才是 13 列口径）
    r1 = _rows(out, "xl/worksheets/sheet1.xml")
    assert _num(r1[7], "B7") == 100.0 and _num(r1[7], "H7") == 400.0
    assert _num(r1[7], "N7") == 100.0 + 200.0 + 400.0


@acc()
def test_收发结存的显示区也不许留旧行(tmp_path):
    """⭐ **一并复核的那份**（收发结存同一族做法）：`Sheet1` 就是它的显示区。

    ⚠️ 它**原实现已经会写显示区**（openpyxl 逐行写），所以 2026-09 没暴露 ——
    但**只写新行、不清旧行** ⇒ 仓库变少时旧行残留（模板原本写到第 11 行）。
    """
    df = pd.DataFrame([{"仓库": "甲仓", "资源分类": "x", "资源编码": "1", "资源名称": "y",
                        "规格型号": "", "计量单位": "个", "物料性质": "主材",
                        "期初数量": 0, "期初金额": 0.0, "收数量": 1, "收单价": 1.0,
                        "收金额": 1.0, "发数量": 0, "发单价": 0.0, "发金额": 0.0,
                        "结存数量": 1, "结存金额": 1.0}], columns=list(SE.SFJC_HEADER))
    z = Path(SE.build_zip(df, tmp_path / "out", "2026年9月"))
    zf = zipfile.ZipFile(z)
    try:
        name = [n for n in zf.namelist() if n.endswith("收发结存表.xlsx")][0]
        main = tmp_path / "main.xlsx"
        main.write_bytes(zf.read(name))
    finally:
        zf.close()
    ws = openpyxl.load_workbook(main)["Sheet1"]
    # 行 3 = 表头、4 = 甲仓、5 = 总计 ⇒ 不得再有第 6 行（模板原本有到第 11 行）
    assert ws.max_row == 5, \
        "Sheet1 第 6 行及以后还在（模板旧渲染值）⇒ 用户会读到上个月的数：max_row=%d" % ws.max_row
    assert ws.cell(row=5, column=1).value == "总计"

    # 守卫的**区分力**：人为塞一行"上一版的仓库" ⇒ `verify()` 必须报出来
    pv = df.groupby("仓库")[SE.MONEY_FIELDS].sum()
    assert SE.verify(str(main), df, pv, pv.sum(), ["甲仓"]) == [], "干净成品不该报问题"
    wb = openpyxl.load_workbook(main)
    wb["Sheet1"].cell(row=6, column=1).value = "上一版的仓库"
    wb["Sheet1"].cell(row=6, column=2).value = 999.0
    wb.save(main)
    problems = SE.verify(str(main), df, pv, pv.sum(), ["甲仓"])
    assert any("僵尸旧值" in p for p in problems), \
        "人为塞进去的旧行没被报出来 ⇒ 这把尺子守不住：%s" % problems


@acc()
def test_新仓库行顺序确定且不随哈希种子变(tmp_path):
    """⚠️ 顺手修的一处隐患：模板里没有的新仓库，行顺序原先来自 `set` 迭代顺序
    （字符串哈希随 `PYTHONHASHSEED` 变）⇒ 现在它**直接决定显示区每行的位置**，必须确定。
    """
    orders = []
    for i in range(2):
        out = _build(tmp_path / ("a%d" % i),
                     _synth(with_safe=False, whs=("新生仓", "甲仓", "乙仓")))
        r1 = _rows(out, "xl/worksheets/sheet1.xml")
        orders.append([_label(out, r1, "A%d" % r) for r in (7, 8, 9)])
    assert orders[0] == orders[1] == ["新生仓", "甲仓", "乙仓"], \
        "新仓库行顺序不确定（模板里没有的按**首次出现**追加）：%s" % orders
