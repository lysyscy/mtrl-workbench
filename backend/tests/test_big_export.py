# -*- coding: utf-8 -*-
"""活动大表成品（`TAIL-04` 路径① **片 4b**）—— 一本工作簿 / `Sheet1` 的离线断言。

形态（唯一主人 = 交付面 `demo/模板/模板清单.md` **§5.5** ＋ §二-7 ＋ §四-5 ＋ §六-2/3）：
`A 序号` ｜ `B 项目名称` ｜ `C 开工时间`（**页面文本** `YYYY-MM`）｜ `D~G` 辅材使用（不含税）4 列 ｜
`H~K` 外租使用（含税）4 列 ｜ `L 合计` ｜ `M 机械费测算` ｜ `N 使用率` ｜ `O 产值完成率` ｜
`P 合同（预估）价` ｜ `Q 产值累计（含税）` ｜ `R 主材累计（含税）` ｜ `S 工程用油` ｜ `T 备注`
—— ⭐ **20 列全 A 类、同名同序**（§5.5 唯一完全一致的表）⇒ **列映射 = 恒等**；
三级表头 `R2~R4` · `DATA_START` = **A5** · 数据一行一个项目 · 末行 = **合计行**（`A:C` 合并放 `合计：`）。

⭐ 断言分四类：
① **形态**：sheet / 20 列 / 恒等映射（**逐格**与域内核出口比）/ 标题随核算月重写；
② **格与格式**：序号 1..N（合计行不编号）· 开工时间 = 文本 ＋ 该列**数字格式 = 文本**；
③ **合计行**：可合计的 14 列给数值、**比率列 / 备注留空**、**不可得的总额显式留空**（⛔ 不写 0）；
④ ⭐⭐ **四处公式「以数值覆盖」**：数据行 `L/N/O` ＋ 合计行 14 个 `=SUM(...)` ⇒ 全表**残留公式 = 0**；
   尾块以下不留残留 / 行壳；一行都没有 / 模板标题形态不认识 ⇒ **拒做**。

⚠️ 取数口打桩（`big_service.build_big_data`）—— 行装配那一段在 `big_domain` 里**真跑**
   （`assemble_rows` / `totals_row` 都是纯函数）⇒ 断言的期望值**手写**、⛔ 不从实现里取。
"""

from __future__ import annotations

from decimal import Decimal

import openpyxl
import pytest

from app.core.errors import MustReportError
from app.mod07_setl import big_domain as B
from app.report_export import big_export as BX
from app.report_export import service as EX
from tests.acc_plugin import acc

SHEET = "Sheet1"
YM = "2026-08"
DS = 5                                     # `DATA_START` = A5 ⇒ 数据首行 = 第 5 行（1 基）

#: 两个项目（一行一个）—— 手写期望就按这两行算
PROJ = [{"warehouse": "甲仓库", "display": "甲项目"}, {"warehouse": "乙仓库", "display": "乙项目"}]
#: 自动列取数（`_amounts` 的形状）：甲全给、乙只给焊材累计
AMOUNTS = {
    "甲仓库": {"hw_cum": 1, "hw_mon": 2, "qt_cum": 3, "qt_mon": 4,
               "rent_cum": 5, "rent_mon": 6, "lease_cum": 7, "lease_mon": 8, "mat_cum": None},
    "乙仓库": {"hw_cum": 100, "hw_mon": None, "qt_cum": None, "qt_mon": None,
               "rent_cum": None, "rent_mon": None, "lease_cum": None, "lease_mon": None,
               "mat_cum": None},
}
#: 手填：甲给了开工时间（`YYYY-MM`）、机械费测算、备注 ⇒ 期望里应看到它们。
#: ⚠️ 形状 = `{年月: {仓库名: {手填 key: 值}}}`（`big_domain.assemble_rows` 的 `own_by_ym` 入口 ——
#:    第一版我写成了 `{仓库: {年月: …}}`，测出来开工时间是 None ⇒ 形状不认识就**静默全空**）。
MANUAL = {"2026-08": {"甲仓库": {"open_ym": "2025-09", "mech_est": Decimal("1000"),
                                "note": "塔吊 1 台"}}}

#: ⭐ **手写期望**（20 列，按 §5.5 列序）—— 甲项目：
#: 1 序号 ｜ 甲项目 ｜ 2025-09 ｜ 1 ｜ 2 ｜ 3 ｜ 4 ｜ 5 ｜ 6 ｜ 7 ｜ 8 ｜ **12**（=rent+lease 派生）
#: ｜ 1000 ｜ **0.012**（使用率 = 12 ÷ 1000 —— 手写的就得**自己算出来**，⛔ 别抄实现）
#: ｜ **None**（产值完成率：分子分母都 0 ⇒ **不可得**，页面显「—」）
#: ｜ **None** 合同价 ｜ **None** 产值累计 ｜ **None** 主材累计（取数口未接入）｜ **None** 工程用油（未录）
#: ｜ 塔吊 1 台
ROW_A = [1, "甲项目", "2025-09", 1, 2, 3, 4, 5, 6, 7, 8, 12,
         1000, "0.012", None, None, None, None, None, "塔吊 1 台"]
#: 乙项目：自动列只给了焊材累计 100 ⇒ 其余自动列 = **None ⇒ 留空**（不可得 ≠ 0）；
#: 「合计」= 外租 ＋ 设备租赁（两列都不可得 ⇒ 域内核给 `0.00`，与页面同源）；手填全未录 ⇒ 留空。
ROW_B = [2, "乙项目", None, 100, None, None, None, None, None, None, None, 0,
         None, None, None, None, None, None, None, None]


def payload(*, rows: int = 2) -> dict:
    """真跑域内核的行装配 ⇒ `build_big_data` 的形状（打桩的返回值）。"""
    projects = PROJ[:rows]
    amounts = {p["warehouse"]: AMOUNTS[p["warehouse"]] for p in projects}
    built = B.assemble_rows(projects, amounts, MANUAL, month=YM)
    return {"month": YM, "rows": built, "totals": B.totals_row(built)}


@pytest.fixture(autouse=True)
def _seams(tmp_path, monkeypatch):
    """落盘根指到 `tmp_path` ＋ 打桩唯一取数口 `big_service.build_big_data`。"""
    from app.mod07_setl import big_service as BS

    monkeypatch.setattr(EX, "EXPORT_ROOT", tmp_path / "exports")
    monkeypatch.setattr(BS, "build_big_data", lambda s, state=None: payload())
    monkeypatch.setattr(BS, "parse_state", lambda p=None: {"month": YM})


def export() -> object:
    return EX.export_big(None, YM)


def read(path) -> list[list]:
    wb = openpyxl.load_workbook(str(path), data_only=False)
    try:
        ws = wb[SHEET]
        return [[ws.cell(row=r, column=c).value for c in range(1, 21)]
                for r in range(1, ws.max_row + 1)]
    finally:
        wb.close()


def dec(v) -> Decimal:
    return Decimal(str(v)) if v not in (None, "") else Decimal("0.00")


def same(got: list, want: list) -> str:
    for i, (g, w) in enumerate(zip(got, want)):
        if w is None:
            if g not in (None, ""):
                return f"第 {i + 1} 列（{BX.COL_OF_KEY[B.BIG_COLUMNS[i]['key']]}）该**留空**，实得 {g!r}"
        elif g is None or (isinstance(w, (int, float, Decimal)) and not isinstance(w, str)
                           and dec(g) != dec(w)):
            return f"第 {i + 1} 列（{BX.COL_OF_KEY[B.BIG_COLUMNS[i]['key']]}）：{g!r} ≠ {w!r}"
        elif isinstance(w, str) and str(g) != w:
            return f"第 {i + 1} 列（{BX.COL_OF_KEY[B.BIG_COLUMNS[i]['key']]}）：{g!r} ≠ {w!r}"
    return ""


# ============================================================ ① 形态


@acc("ACC-RPT-BIG-001", "001-16")
def test_一张表二十列且三级表头照模板():
    """一本工作簿 / `Sheet1` / 20 列；三级表头 `R2~R4` **原样沿用模板**（⛔ 不重画）。"""
    rows = read(export())
    assert len(rows[0]) == 20 and str(rows[0][0]).startswith("第四工程公司在建项目辅材及机械费统计表")
    assert rows[0][19] == "单位：元", f"`T1` 应是「单位：元」，实得 {rows[0][19]!r}"
    assert [rows[1][i] for i in (0, 1)] == ["序号", "项目名称"]
    assert rows[1][3] == "辅材使用（不含税）" and rows[1][7] == "外租使用（含税）"
    assert rows[2][3] == "焊材" and rows[2][5] == "三气", rows[2]
    assert rows[3][3] == "累计" and rows[3][4] == "当月", rows[3]


@acc("ACC-RPT-BIG-001", "001-16")
def test_列映射是恒等逐格等于域内核出口():
    """⭐ §5.5「20 列**同名同序**」（唯一完全一致的表）⇒ **第 i 列 ↔ `BIG_COLUMNS[i]`**：
    逐格与域内核的 `cell_value` 出口比（⛔ 不抄 20 个列字母，也就不会漂）。"""
    rows = read(export())
    for i, row in enumerate(payload()["rows"]):
        want = [B.cell_value(row, col["key"])[0] for col in B.BIG_COLUMNS]
        diff = same(rows[DS - 1 + i], want)                  # ⚠️ 金额按**数值**比（`12` vs `Decimal("12.00")`）
        assert not diff, f"第 {i + 1} 行 · {diff}"


@acc("ACC-RPT-BIG-001", "001-16")
def test_标题随核算月重写且当月末日按日历算(tmp_path):
    """清单 §六-3：标题里**写死的日期**由系统替换 —— 认**括注形态**、月末日按**日历**算。"""
    from app.report_export import big_export

    tpl = "第四工程公司在建项目辅材及机械费统计表（数据截止8月31日）（8月份数据）"
    assert big_export.title_of(tpl, "2026-02") == \
        "第四工程公司在建项目辅材及机械费统计表（数据截止2月28日）（2月份数据）"
    assert big_export.title_of(tpl, "2026-04") == \
        "第四工程公司在建项目辅材及机械费统计表（数据截止4月30日）（4月份数据）"
    assert "（数据截止8月31日）（8月份数据）" in str(read(export())[0][0]), "成品标题没跟着核算月走"


@acc("ACC-RPT-BIG-001", "001-16")
def test_成品文件名含报表名与期间标签():
    assert BX.file_name(YM) == "经济活动分析大表_2026年8月.xlsx", BX.file_name(YM)


# ============================================================ ② 格与格式


@acc("ACC-RPT-BIG-001", "001-17")
def test_两行逐格等于手写期望():
    """20 列逐格 = **手写期望**（含派生的「合计」= 外租 ＋ 设备租赁、比率列**不可得 ⇒ 留空**）。"""
    rows = read(export())
    for i, want in enumerate((ROW_A, ROW_B)):
        diff = same(rows[DS - 1 + i], want)
        assert not diff, f"第 {i + 1} 行 · {diff}"


@acc("ACC-RPT-BIG-001", "001-17")
def test_序号是系统行号且合计行不编号():
    rows = read(export())
    assert [rows[DS - 1 + i][0] for i in range(2)] == [1, 2], "序号 ≠ 1..N"
    assert str(rows[DS + 1][0]) == BX.TOTAL_LABEL, f"合计行首格 = {rows[DS + 1][0]!r}"


@acc("ACC-RPT-BIG-001", "001-17")
def test_开工时间是页面文本且数字格式为文本():
    """清单 §四-4 同款：开工时间 = `YYYY-MM` 文本；写前把该列数字格式设成文本（模板那列是数值）。"""
    path = export()
    rows = read(path)
    assert rows[DS - 1][2] == "2025-09", f"开工时间 = {rows[DS - 1][2]!r}"
    wb = openpyxl.load_workbook(str(path))
    try:
        ws = wb[SHEET]
        assert {ws.cell(row=DS + i, column=3).number_format for i in (0, 1)} == {"@"}
    finally:
        wb.close()


# ============================================================ ③ 合计行


@acc("ACC-RPT-BIG-001", "001-18")
def test_合计行只给可合计的十四列_比率与备注留空():
    """合计行 = `TOTAL_KEYS` 的 14 列给数值；**比率列（使用率 / 产值完成率）与备注留空**（照模板：
    模板的 `R13` 只在 `D..S` 的 14 个格上写了 `=SUM`，`N/O/T` 本来就是空的）。"""
    rows = read(export())
    tot = rows[DS + 1]
    keys = {BX.COL_OF_KEY[k] for k in B.TOTAL_KEYS}
    for j, col in enumerate(B.BIG_COLUMNS):
        letter = BX.COL_OF_KEY[col["key"]]
        if letter in keys or j == 0:
            continue
        assert tot[j] in (None, ""), f"合计行 {letter}（{col['key']}）该留空，实得 {tot[j]!r}"
    assert dec(tot[3]) == Decimal("101.00"), f"焊材累计合计 = {tot[3]!r}"     # 1 + 100
    assert dec(tot[11]) == Decimal("12.00"), f"合计列合计 = {tot[11]!r}"       # 甲 12 ＋ 乙 0
    assert tot[13] in (None, "") and tot[14] in (None, ""), "比率列不合计"


@acc("ACC-RPT-BIG-001", "001-18")
def test_不可得的总额显式留空不写0():
    """⭐ **不可得 ≠ 0**（清单 §六-2）：主材累计的取数口未接入 ⇒ 合计行那格**留空**，
    ⛔ 不写 0 —— 而模板那格是 `=SUM(R5:R12)` ⇒ 必须**显式清掉**（引擎的「空串 = 显式留空」）。"""
    rows = read(export())
    assert rows[DS + 1][17] in (None, ""), f"主材累计合计该留空，实得 {rows[DS + 1][17]!r}"
    assert ROW_A[17] is None and ROW_B[17] is None, "行上的主材累计也该留空"


# ============================================================ ④ 公式覆盖 / 零残留 / 拒做


@acc("ACC-RPT-BIG-001", "001-19")
def test_全表残留公式为零且尾块以下不留残留():
    """⭐⭐ 清单 §六-2：模板公式**不参与求值** ⇒ 数据行的 `L=H+J` / `N=L/M` / `O=Q/P`
    与合计行 14 个 `=SUM(...)` 一律**以数值覆盖**；⛔ 尾块以下不留残留 / 行壳。"""
    path = export()
    wb = openpyxl.load_workbook(str(path), data_only=False)
    try:
        ws = wb[SHEET]
        bad = [(r, c) for r in range(DS, ws.max_row + 1) for c in range(1, 21)
               if isinstance(ws.cell(row=r, column=c).value, str)
               and str(ws.cell(row=r, column=c).value).startswith("=")]
        assert not bad, f"还有公式没覆盖：{bad[:6]}"
        # 数据 2 行（第 5、6 行）＋ 合计 1 行（第 7 行）⇒ max_row = DS + n - 1 + 1 = DS + 2
        assert ws.max_row == DS + 2, f"max_row = {ws.max_row}（应 = 数据 2 行 ＋ 合计 1 行）"
    finally:
        wb.close()
    assert len(read(path)) == DS + 2


@acc("ACC-RPT-BIG-001", "001-19")
def test_拒做一行都没有():
    from app.mod07_setl import big_service as BS

    BS.build_big_data = lambda s, state=None: {"month": YM, "rows": [], "totals": {}}
    try:
        with pytest.raises(MustReportError) as e:
            export()
        assert "一行都没有" in str(e.value)
    finally:
        pass


@acc("ACC-RPT-BIG-001", "001-19")
def test_拒做模板标题形态不认识():
    """标题里找不到那两个**日期括注** ⇒ **拒做**（清单 §六-3 登记了这两处由系统替换 ⇒ 认形态、不猜）。"""
    with pytest.raises(MustReportError) as e:
        BX.title_of("某张没有日期括注的表", YM)
    assert "日期括注" in str(e.value)


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(pytest.main([__file__, "-v"]))
