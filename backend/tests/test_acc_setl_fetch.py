"""`MOD-07` 收发结存 · **抓取通道**（`API-02`）离线自检（**不连库、不碰网络**）。

对应判据 `ACC-SETL-FETCH-001`（10 点，`PRD/验收/验收基线.json`）：
- `001-01` 入参 = 「起止年月」**四个独立值** ＋ 项目参数 key = **`pc`**（⛔ 不是 `xmid`、不是单据日期区间）
- `001-02` 三项标识 ＋ **`reportName` = 「收发结存表」**（**必须传表名**，与库存单据的 `null` 相反）
- `001-03` 抓取**只允许单月**
- `001-04` 解析**按导出文件表头文字认列**（不写死下标、列序被打乱也认得出）；**关键列缺失 ⇒ 报错不猜**
- `001-05` **拆合并单元格填充** ＋ **剔除「合计」行** ＋ 整行空跳过；**有金额却没仓库名 ⇒ 报错**（不静默丢行）
- `001-06` 数字列：千分位 / 空值容错；**非数字 ⇒ 报错并带行号**
- `001-07` 落库**幂等**：同月重抓不产生重复行 · 当月变少**删旧自动行** · **人工行保留**
- `001-08` **写—读闭环**：抓完该月，页面就读得到（同一份会话）
- `001-09` 抓前**二次确认**（D9）· 失败**显式报错待人处理 ＋ 不重试**（导出调用**只发生一次**；
  ⚠️ 2026-09-22 纠正：本域无手动导入通道，原「降级为手动导入」是悬空承诺）
- `001-10` **陌生仓库**（档案里没有）⇒ **自动建档 / 扶正**（§2.4.2 链路 B），结果**逐条列出**；
  取不到仓库名单 ⇒ **降级不阻断**（数据照常落库）—— 原「⛔ 不自动建档」与 §2.4.2 相反，**已作废**（2026-09-17）
"""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO

import pytest

from app.core.config import I8_REPORTS
from app.core.errors import MustReportError
from app.mod02_i8 import report_export
from app.mod02_i8.models import RunRecord
from app.mod07_setl import fetch_service as FS
from app.mod07_setl import importer
from app.mod07_setl.models import SOURCE_AUTO, SetlDetail
from app.mod07_setl.repository import upsert_month
from app.mod07_setl.service import build_page_data, parse_state
from tests.acc_plugin import acc
from tests.m07_helpers import FakeSession

CASE = "ACC-SETL-FETCH-001"

# i8 导出表头 17 列（顺序 = 导出的真实顺序）
COLUMNS: tuple[str, ...] = (
    "仓库", "资源分类", "资源编码", "资源名称", "规格型号", "计量单位", "物料性质",
    "期初数量", "期初金额", "收数量", "收单价", "收金额",
    "发数量", "发单价", "发金额", "结存数量", "结存金额",
)
# 两行表头：**第 1 行 = 分组名（横向合并）· 第 2 行 = 真实列名**（与真机导出一致，2026-09-17 实测）
GROUPS: dict[str, tuple[str, ...]] = {
    "期初统计": ("期初数量", "期初金额"),
    "收统计": ("收数量", "收单价", "收金额"),
    "发统计": ("发数量", "发单价", "发金额"),
    "结存统计": ("结存数量", "结存金额"),
}


# ============================================================ 造 xlsx（离线）

def build_xlsx(rows: list[list], *, order: tuple[str, ...] | None = None,
               merged: bool = True) -> bytes:
    """造一份 i8 形态的导出：`merged=True` = **两行合并表头**，`False` = 单行表头。"""
    from openpyxl import Workbook

    order = tuple(order or COLUMNS)
    wb = Workbook()
    ws = wb.active
    ws.title = "报表1"

    if merged:
        # 第 1 行 = 分组名（该组首列有值，其余留空后横向合并）；第 2 行 = **真实列名**（照真机导出）
        head1 = [next((g for g, m in GROUPS.items() if c in m), c) for c in order]
        ws.append(head1)
        ws.append(list(order))
        # ⚠️ 必须**先写后并** —— openpyxl 对已合并区域的写会报 cell 只读
        for i, c in enumerate(order, start=1):
            g = next((g for g, m in GROUPS.items() if c in m), "")
            if not g:
                ws.merge_cells(start_row=1, start_column=i, end_row=2, end_column=i)
            else:
                span = len(GROUPS[g])
                first = order.index(GROUPS[g][0]) + 1
                ws.merge_cells(start_row=1, start_column=first, end_row=1,
                               end_column=first + span - 1)
    else:
        ws.append(list(order))

    for row in rows:
        ws.append(row)
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _values(warehouse: str = "甲仓", code: str = "R01", *, begin: str = "1,000.00",
            in_amt: str = "500.50", out_amt: str = "200.25", end_amt: str = "1,300.25") -> dict:
    return {
        "仓库": warehouse, "资源分类": "材料", "资源编码": code, "资源名称": "焊条",
        "规格型号": "J422", "计量单位": "kg", "物料性质": "主材",
        "期初数量": "10", "期初金额": begin,
        "收数量": "5", "收单价": "100.10", "收金额": in_amt,
        "发数量": "2", "发单价": "100.13", "发金额": out_amt,
        "结存数量": "13", "结存金额": end_amt,
    }


def _line(values: dict, order: tuple[str, ...] = COLUMNS) -> list:
    return [values.get(c, "") for c in order]


def _total_line(*, begin="1,000.00", in_amt="500.50", out_amt="200.25", end_amt="1,300.25") -> list:
    v = _values("合计", "—", begin=begin, in_amt=in_amt, out_amt=out_amt, end_amt=end_amt)
    return _line(v)


def _xlsx_one(*, order: tuple[str, ...] = COLUMNS, merged: bool = True) -> bytes:
    """一份最小的正常导出：1 条明细 ＋ 1 条「合计」＋ 1 个整行空行。"""
    return build_xlsx([_line(_values(), order), _total_line(), [""] * 17],
                      order=order, merged=merged)


def _fake_client(wh_rows=None):
    """假 i8 client：`close()` ＋ **仓库名单**（`API-05`，链路 B 建档要用；默认空名单）。

    ⚠️ 默认空名单 ⇒ 既有用例的档案**一行都不会被新增**（陌生仓库全走「判不了」分支），
       故不影响它们；要测自动建档就传 `wh_rows=[(仓库名, 所属项目全名), …]`。
    """
    class _Client:
        def close(self):
            pass

        def post_json(self, _path, data=None):       # noqa: ANN001, ARG002
            return {"Record": [{"Whname": w, "PhidPc_EXName": o, "Nullify": "0"}
                               for w, o in (wh_rows or [])]}

    return _Client()


def _broken_client():
    """假 client：取仓库名单必炸（测「降级不阻断」）。"""
    class _Client:
        def close(self):
            pass

        def post_json(self, _path, data=None):       # noqa: ANN001, ARG002
            raise RuntimeError("i8 502")

    return _Client()


def _project(pid: int, warehouse: str, phid: str):
    from app.mod01_projects.models import Project

    return Project(id=pid, warehouse_name=warehouse, name_full=f"{warehouse}项目",
                   disp=warehouse, phid=phid)


# ============================================================ 判据点

@acc(CASE, "001-01")
def test_入参为起止年月四个独立值且项目参数为pc():
    body = report_export.build_body(spec_key="setl", year_month="2026-08", proj_ids="1,2")
    params = {p["paramKey"]: p for p in body["queryParams"][0]["sheetQueryParams"]}

    assert set(params) == {"a", "b", "aa", "bb", "pc"}, f"入参键不对：{sorted(params)}"
    assert (params["a"]["paramValue"], params["b"]["paramValue"]) == ("2026", "8")
    assert (params["aa"]["paramValue"], params["bb"]["paramValue"]) == ("2026", "8")   # 单月 ⇒ 起止同月
    assert [params[k]["label"] for k in ("a", "b", "aa", "bb")] == [
        "开始核算年", "开始核算月", "结束核算年", "结束核算月"]
    # 项目：paramType=0 / fieldType="0" / 值为 id 列表
    assert params["pc"]["paramType"] == "0" and params["pc"]["fieldType"] == "0"
    assert params["pc"]["paramValue"] == "1,2"
    # ⛔ 不得混入 API-01 的「单据日期区间」，也不得用库存单据的项目 key
    text = str(body)
    assert "BillDt" not in text and "queryfilter" not in text and "start_date" not in text
    assert "xmid" not in text


@acc(CASE, "001-02")
def test_三项标识齐备且收发结存必须传表名():
    spec = I8_REPORTS["setl"]
    assert spec["rep_id"] == "3700000000000085"
    assert spec["dataset_id"] == "3700000000000178"
    assert spec["proj_param"] == "pc"
    # ⭐ 与库存单据**相反**：本表必须传表名
    assert spec["report_name"] == "收发结存表"
    assert I8_REPORTS["inv"]["report_name"] is None

    body = report_export.build_body(spec_key="setl", year_month="2026-08", proj_ids="1")
    assert body["reportName"] == "收发结存表" and body["reportId"] == spec["rep_id"]
    assert body["dataSetIds"] == [] and body["exportType"] == "excel"

    # 预检：登记不全 / 标记未定稿 ⇒ **当场拒**（不猜取值、不发请求）
    assert report_export.require_ready("setl") is spec
    original = dict(I8_REPORTS["setl"])
    try:
        I8_REPORTS["setl"] = {k: v for k, v in original.items() if k != "report_name"}
        with pytest.raises(MustReportError):
            report_export.require_ready("setl")
        I8_REPORTS["setl"] = {**original, "report_name_undetermined": True}
        with pytest.raises(MustReportError):
            report_export.require_ready("setl")
    finally:
        I8_REPORTS["setl"] = original


@acc(CASE, "001-03")
def test_抓取只允许单月():
    assert report_export.single_month("2026-08") == ("2026", "8", "2026", "8")
    assert FS.build_plan("2026-08")["ym"] == "2026-08"
    for bad in ("2026-01~2026-06", "2026", "", "2026-13"):
        with pytest.raises(MustReportError):
            FS.build_plan(bad)


@acc(CASE, "001-04")
def test_按表头文字认列且关键列缺失报错不猜():
    # ① 列序**被打乱**、且是单行表头 ⇒ 照样认得出（证明不依赖下标）
    shuffled = tuple(reversed(COLUMNS))
    rows = importer.parse_rows(build_xlsx([_line(_values(), shuffled)], order=shuffled,
                                          merged=False), ym="2026-08")
    assert len(rows) == 1
    got = rows[0]
    assert got["warehouse_name"] == "甲仓" and got["resource_code"] == "R01"
    assert got["resource_name"] == "焊条" and got["spec"] == "J422"
    assert got["begin_amt"] == Decimal("1000.00") and got["end_amt"] == Decimal("1300.25")

    # ② 缺关键列（去掉「收金额」）⇒ 报错不猜
    thin = tuple(c for c in COLUMNS if c != "收金额")
    with pytest.raises(MustReportError) as exc:
        importer.parse_rows(build_xlsx([_line(_values(), thin)], order=thin, merged=False),
                            ym="2026-08")
    assert "收金额" in str(exc.value)


@acc(CASE, "001-05")
def test_拆合并表头并剔除合计行():
    rows = importer.parse_rows(_xlsx_one(), ym="2026-08")
    # 合并表头认得出：第 1 行分组名（`期初统计`…）＋ 第 2 行真实列名（`期初金额`…）⇒ **取第 2 行**
    # （2026-09-17 真机导出实测的形状；原先按"上下拼接"实现，真机一跑就认不出关键列）
    assert len(rows) == 1, f"应只剩 1 条明细，实得 {len(rows)}"
    assert rows[0]["warehouse_name"] == "甲仓"
    assert "合计" not in {r["warehouse_name"] for r in rows}
    # ⭐ 认列必须**取表头块最下面一行**（叶子行）—— 第 1 行的 `期初统计` 只是分组名
    index, first = importer.find_header_block(importer.read_grid(_xlsx_one()))
    assert {"期初金额", "收金额", "发金额", "结存金额"} <= set(index), "合并表头认列取错行"
    assert first == 2, "数据首行应是表头块之后（第 2 行表头 ⇒ 数据从第 3 行起）"

    # 有金额却没仓库名 ⇒ **报错**（不静默丢行）
    orphan = _line(_values(warehouse=""))
    with pytest.raises(MustReportError) as exc:
        importer.parse_rows(build_xlsx([orphan]), ym="2026-08")
    assert "仓库" in str(exc.value)


@acc(CASE, "001-06")
def test_数字列千分位与空值容错且非数字报错带行号():
    rows = importer.parse_rows(_xlsx_one(), ym="2026-08")
    # ⚠️ 千分位由 `core.money.to_money` 统一剥（金额口径的全站唯一出口）⇒ 这里验的是
    #    **端到端**：导出里带逗号的金额，走到域内就是规范金额。（这条断言没法靠"局部变异"证区分力，
    #    因为剥逗号只有一个真出口；能变异的是下面那条"非数字必须报错"。）
    assert rows[0]["begin_amt"] == Decimal("1000.00")      # 千分位已去
    assert rows[0]["in_price"] == Decimal("100.10")

    blank = _values()
    blank["期初金额"] = ""                                  # 空 ⇒ 0（不报错）
    rows = importer.parse_rows(build_xlsx([_line(blank)]), ym="2026-08")
    assert rows[0]["begin_amt"] == Decimal("0.00")

    bad = _values()
    bad["收金额"] = "五百"
    with pytest.raises(MustReportError) as exc:
        importer.parse_rows(build_xlsx([_line(bad)]), ym="2026-08")
    assert "收金额" in str(exc.value) and "第 3 行" in str(exc.value)


@acc(CASE, "001-07")
def test_落库幂等且人工行不被覆盖():
    def _detail(row_id, code, amount, source=SOURCE_AUTO):
        return SetlDetail(id=row_id, ym="2026-08", warehouse_name="甲仓", resource_code=code,
                          resource_name="物料", spec="—", unit="kg", material_property="主材",
                          begin_qty=Decimal("0"), begin_amt=Decimal(amount),
                          in_amt=Decimal("0"), out_amt=Decimal("0"), end_amt=Decimal("0"),
                          source=source)

    s = FakeSession({SetlDetail: [_detail(1, "R01", "1.00"),
                                  _detail(2, "R98", "9.99"),          # 本次不再出现 ⇒ 该删
                                  _detail(3, "R99", "7.77", "manual")]})  # 人工行 ⇒ 必须留
    stats = upsert_month(s, ym="2026-08", rows=[
        {"warehouse_name": "甲仓", "resource_code": "R01", "begin_amt": Decimal("2.00")},
        {"warehouse_name": "甲仓", "resource_code": "R02", "begin_amt": Decimal("3.00")},
    ])
    assert stats == {"inserted": 1, "updated": 1, "deleted": 1, "kept_manual": 1}, stats
    live = {d.resource_code: d for d in s.by_model[SetlDetail]}
    assert set(live) == {"R01", "R02", "R99"}                  # R98 已删、R99 保留、无重复行
    assert live["R01"].begin_amt == Decimal("2.00")            # 匹配行只更新源字段
    assert live["R99"].begin_amt == Decimal("7.77")            # 人工行原样

    # 同一批再导一次 ⇒ **不产生重复行**
    again = upsert_month(s, ym="2026-08", rows=[
        {"warehouse_name": "甲仓", "resource_code": "R01", "begin_amt": Decimal("2.00")},
        {"warehouse_name": "甲仓", "resource_code": "R02", "begin_amt": Decimal("3.00")},
        {"warehouse_name": "甲仓", "resource_code": "R99", "begin_amt": Decimal("3.00")},
    ])
    assert again["inserted"] == 0 and len(s.by_model[SetlDetail]) == 3


@acc(CASE, "001-08")
def test_抓取后页面即读到该月数据():
    from app.mod01_projects.models import Project

    s = FakeSession({Project: [_project(1, "甲仓", "P1")]})
    out = FS.run_fetch(s, year_month="2026-08", confirmed=True, operator="测试",
                       client_factory=_fake_client,
                       export=lambda _c, _ym, _ids: build_xlsx([
                           _line(_values(code="R01")), _line(_values(code="R02")),
                           _total_line()]))
    assert out["ok"] is True and out["rows"] == 2, out

    data = build_page_data(s, parse_state({"period": "2026-08"}))
    assert [r["warehouse"] for r in data["rows"]] == ["甲仓"]
    assert data["rows"][0]["rows"] == 2                       # 行数 = 该月明细行数
    assert data["counts"] == {"warehouses": 1, "rows": 2, "files": 2, "ym": "2026-08"}


@acc(CASE, "001-09")
def test_二次确认且失败降级不重试():
    s = FakeSession({})
    with pytest.raises(MustReportError) as exc:
        FS.run_fetch(s, year_month="2026-08", confirmed=False, operator="测试")
    assert "强踢" in str(exc.value)                            # 明写会强踢（D9）
    assert not s.by_model.get(RunRecord), "用法错误**不该**写执行记录"

    from app.mod01_projects.models import Project

    s = FakeSession({Project: [_project(1, "甲仓", "P1")]})
    calls: list[str] = []

    def _boom(client, ym, ids):
        calls.append(ym)
        raise RuntimeError("i8 未登录")

    out = FS.run_fetch(s, year_month="2026-08", confirmed=True, operator="测试",
                       client_factory=_fake_client, export=_boom)
    assert out["ok"] is False and out["degraded"] is True
    # ⚠️ 2026-09-22（`TAIL-69`）：本域**从未建成**手动导入通道 —— 判词从「降级为手动导入」
    #    纠正为「**显式报错待人处理**」（⛔ 判据不删，改指真实现 —— 原判词是悬空承诺）。
    assert "i8 接口不可用" in out["message"] and "excel_import_available" not in out
    assert calls == ["2026-08"], "失败**不得重试**（导出调用只应发生一次）"
    records = s.by_model.get(RunRecord) or []
    assert len(records) == 1 and records[0].status == "failed"  # 失败要留执行记录


@acc(CASE, "001-10")
def test_陌生仓库自动建档扶正并逐条列出():
    """⭐ 口径（2026-09-17 按《说明书》§2.4.2 订正）：陌生仓库 ⇒ **自动建档 / 扶正** —— 无感、不要求用户操作。

    ⚠️ 本条原来钉的是「**不自动建档**」，那与 §2.4.2「**B 材料侧扶正**：陌生仓库 → 同步仓库名单 →
       全名回查占位行 → **扶正改名** → 重跑导入」**相反** ⇒ 该写法已作废（规则主人 =《说明书》）。
    """
    from app.mod01_projects.models import Project
    from app.mod01_projects.repository import all_projects

    s = FakeSession({Project: [_project(1, "甲仓", "P1"),
                               _project(2, "未建仓-1", "P2")]})      # 2 号 = 占位行，等扶正
    out = FS.run_fetch(s, year_month="2026-08", confirmed=True, operator="测试",
                       client_factory=lambda: _fake_client(
                           [("乙仓", "未建仓-1项目"), ("丙仓", "丁项目")]),
                       export=lambda _c, _ym, _ids: build_xlsx([
                           _line(_values(warehouse="甲仓", code="R01")),
                           _line(_values(warehouse="乙仓", code="R01")),
                           _line(_values(warehouse="丙仓", code="R01"))]))
    assert out["ok"] is True
    assert out["unknown_warehouses"] == sorted(["乙仓", "丙仓"])

    rec = out["recovery"]
    # ✅ 能扶正就扶正：占位行「未建仓-1」→ 真实仓库名「乙仓」（展示名一并跟着换）
    assert (rec["promote"], rec["create"]) == (1, 0), rec
    row2 = {p.id: p for p in all_projects(s)}[2]
    assert (row2.warehouse_name, row2.disp) == ("乙仓", "乙仓")
    assert any("未建仓-1" in x and "乙仓" in x for x in rec["promoted"]), "结果要**逐条**说出扶正了什么"
    # ⛔ 判不了的不臆造：「丙仓」的所属项目「丁项目」还没建档 ⇒ 列出来要人看，**不建**
    assert rec["unresolved"] == 1
    assert "丙仓" == rec["unresolved_items"][0]["warehouse_name"]
    # 数据照常落库（建档的事再怎么样也不能丢数据）
    data = build_page_data(s, parse_state({"period": "2026-08"}))
    assert [r["warehouse"] for r in data["rows"]] == ["丙仓", "乙仓", "甲仓"]

    # ⚠️ 取不到仓库名单 ⇒ **降级不阻断**：抓取照样成功、数据照常落库，只在回执里说明
    s2 = FakeSession({Project: [_project(1, "甲仓", "P1")]})
    out2 = FS.run_fetch(s2, year_month="2026-08", confirmed=True, operator="测试",
                        client_factory=_broken_client,
                        export=lambda _c, _ym, _ids: build_xlsx(
                            [_line(_values(warehouse="乙仓", code="R01"))]))
    assert out2["ok"] is True and out2["rows"] == 1, "建档取数失败**不得**把抓取判死"
    assert out2["recovery"]["degraded"] is True
    assert "可稍后重抓" in out2["recovery"]["message"]
    assert [p.warehouse_name for p in all_projects(s2)] == ["甲仓"], "降级时一行也不许乱建"
    assert [r["warehouse"] for r in build_page_data(
        s2, parse_state({"period": "2026-08"}))["rows"]] == ["乙仓"]

