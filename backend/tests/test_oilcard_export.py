"""核真实下载工作簿，取数范围及原文不依赖页面筛选。"""
from io import BytesIO
from urllib.parse import unquote
from uuid import uuid4

from openpyxl import load_workbook

from app.core.db import session_scope
from app.core.money import to_money
from app.mod13_oilcard.ledger import stage_opening
from tests.oilcard_invoice_samples import CARD, NUMBER, invoice_pdf
from tests.test_oilcard_archives import scene, create, edit  # noqa: F401
from tests.test_oilcard_invoice_adjustments import allocation, allocate, credit
from tests.test_oilcard_invoices import receive
from tests.test_oilcard_ledger import entry


def test_full_workbook_keeps_all_saved_rows_original_dates_and_text_card_numbers(scene):
    client, (project, other) = scene
    card = create(client, 'cards', card_no='0001234567890123456789')
    first = create(client, 'accounts', project_id=project, name='甲用卡', card_id=card['id'])
    second = create(client, 'accounts', project_id=other, name='乙用卡')
    recharge = entry(client, first, '10000', date_text='原始数据', remark='=保留原文')
    entry(client, first, '200', kind='adjustment', date_text='2026-10-31')
    entry(client, second, '-20', kind='adjustment', date_text='')
    refund = client.post('/api/m13/ledger/refunds', json={'request_id': str(uuid4()),
        'original_id': recharge['id'], 'amount': '125.50', 'date_text': '退款待核', 'remark': '实际退款'}).json()['item']
    result = client.get('/api/m13/ledger/export?kind=adjustment&project_id=0')
    assert result.status_code == 200 and result.headers['X-Export-Filename'].endswith('.xlsx')
    assert 'attachment' in result.headers['Content-Disposition']
    sheet = load_workbook(BytesIO(result.content), data_only=False).active
    rows = list(sheet.values)
    assert len(rows) == 6
    assert rows[0] == ('记录编号', '日期', '类型', '项目', '用卡账户', '实体卡号', '金额', '备注', '关联原充值', '转账关联编号')
    assert [row[1] for row in rows[1:-1]] == ['原始数据', '2026-10-31', None, '退款待核']
    assert [to_money(row[6]) for row in rows[1:-1]] == [to_money('10000'), to_money('200'), to_money('-20'), to_money('-125.50')]
    assert rows[1][5] == '0001234567890123456789' and sheet['F2'].data_type == 's'
    assert rows[1][7] == '=保留原文' and sheet['H2'].data_type == 's'
    assert rows[4][0] == refund['id'] and rows[4][8] == recharge['id']
    assert rows[-1][0] == '合计' and to_money(rows[-1][6]) == to_money('10054.50')


def test_undo_removed_entry_is_not_exported_and_decimals_match_authority(scene):
    client, (project, _) = scene
    account = create(client, 'accounts', project_id=project, name='金额核对')
    entry(client, account, '0.10')
    entry(client, account, '0.20')
    removed = entry(client, account, '50')
    operation = client.get('/api/history/status').json()['undo']
    assert client.post('/api/history/undo', json={'operation_id': operation['id']}).json()['ok']
    sheet = load_workbook(BytesIO(client.get('/api/m13/ledger/export').content), data_only=True).active
    rows = list(sheet.values)
    assert len(rows) == 4 and removed['id'] not in [row[0] for row in rows[1:-1]]
    assert to_money(rows[-1][6]) == to_money('0.30')
    total = load_workbook(BytesIO(client.get('/api/m13/export').content), data_only=True)
    rows = list(total['充值台账'].values)
    assert removed['id'] not in [row[0] for row in rows[1:-1]]
    assert to_money(rows[-1][6]) == to_money('0.30')
    assert to_money(list(total['核对汇总'].values)[-1][11]) == to_money('0.30')


def _rows(sheet, *, total=True):
    values = list(sheet.values)
    return [dict(zip(values[0], row)) for row in (values[1:-1] if total else values[1:])]


def test_total_workbook_full_history_openings_invoice_states_and_split_amounts(scene):
    client, (project, other) = scene
    primary = create(client, 'cards', card_no='00001234567890123456', kind='primary', remark='=主卡独立备注')
    card = create(client, 'cards', card_no=CARD, parent_card_id=primary['id'], remark='副卡原备注')
    first = create(client, 'accounts', project_id=project, name='历史核算甲', card_id=card['id'], remark='账户原备注')
    second = create(client, 'accounts', project_id=other, name='归档核算乙', status='archived')
    recharge = entry(client, first, '10000', date_text='原始数据', remark='=充值原备注')
    entry(client, first, '200', kind='adjustment', date_text='2026-10-31')
    entry(client, second, '-20', kind='adjustment', date_text='')
    transfer = client.post('/api/m13/ledger/transfers', json={'request_id': str(uuid4()),
        'out_project_id': project, 'out_account_id': first['id'], 'in_project_id': other,
        'in_account_id': second['id'], 'amount': '150', 'date_text': '转账日期待核',
        'out_remark': '转出原备注', 'in_remark': '转入原备注'})
    assert transfer.status_code == 200, transfer.text
    with session_scope() as session:
        stage_opening(session, source_key='isolated:confirmed-opening', source_note='已确认原始期初依据',
            project_id=project, account_id=first['id'], amount='500', original_date_text='原始数据',
            cutoff_ym='2025-12', confirmed=True, remark='期初原备注')
        stage_opening(session, source_key='isolated:pending-opening', source_note='尚待确认来源',
            project_id=other, account_id=second['id'], amount='999', original_date_text='早期日期不明', remark='不得计入')
    original = receive(client, invoice_pdf(amount='10000.00', subtotal='8849.56', tax='1150.44'))['item']
    result = allocate(client, original, [allocation(first, '6000'), allocation(second, '4000')])
    assert result.status_code == 200, result.text
    original = result.json()['item']
    repeated = receive(client, invoice_pdf(amount='10000.00', subtotal='8849.56', tax='1150.44', variant='same invoice'))
    assert repeated['item']['id'] == original['id']
    red = receive(client, invoice_pdf(number='00000000000000000002', amount='-1000.00', subtotal='-884.96', tax='-115.04'))['item']
    assert credit(client, red, original, first).status_code == 200
    pending = receive(client, invoice_pdf(number='00000000000000000003', card='12345'))['item']
    voided = receive(client, invoice_pdf(number='00000000000000000004', amount='700.00', subtotal='619.47', tax='80.53'))['item']
    result = client.post(f"/api/m13/invoices/{voided['id']}/void", json={'version': voided['version'],
        'date': '2026-10-31', 'note': '真实作废业务的合成依据'})
    assert result.status_code == 200, result.text
    replacement = create(client, 'cards', card_no='00000000000000000009')
    edit(client, 'accounts', first, project_id=other, card_id=replacement['id'], status='archived')
    assert client.post(f'/api/m13/linkage/{project}/enable').status_code == 200
    assert client.post(f'/api/m13/linkage/{project}/big/enable').status_code == 200
    assert client.put(f'/api/m13/linkage/{project}/2026-10/fee', json={'amount': '-20.25', 'remark': '当月自定义原备注'}).status_code == 200
    assert client.put(f'/api/m13/linkage/{project}/2026-10/big', json={'amount': '1000', 'remark': '截至月快照原备注'}).status_code == 200

    response = client.get('/api/m13/export?status=active&project_id=0&account_id=0&ym=2020-01')
    assert response.status_code == 200, response.text
    assert unquote(response.headers['X-Export-Filename']) == '油卡总台账.xlsx'
    assert 'attachment' in response.headers['Content-Disposition']
    assert response.headers['Cache-Control'] == 'no-store'
    workbook = load_workbook(BytesIO(response.content), data_only=False)
    assert workbook.sheetnames == ['核对汇总', '充值台账', '发票台账', '发票分配', '用卡账户',
        '实体油卡', '期初核对', '账户关联历史', '用油自定义']
    summary_rows = _rows(workbook['核对汇总'])
    historical = next(row for row in summary_rows if row['账户编号'] == first['id'] and row['归属'] == '历史归属')
    assert (historical['核算项目'], historical['实体卡号'], historical['主卡号']) == ('甲', CARD, primary['card_no'])
    assert historical['账户状态'] == '归档' and historical['主卡备注'] == '=主卡独立备注'
    assert historical['账户备注'] == '账户原备注'
    assert to_money(historical['已确认期初']) == to_money('500')
    assert to_money(historical['充值发生额']) == to_money('10050')
    authority = client.get('/api/m13/reconciliation').json()
    totals = dict(zip(next(workbook['核对汇总'].values), list(workbook['核对汇总'].values)[-1]))
    for label, key, expected in [('累计充值', 'total_amount', '10680'), ('已登记开票', 'invoiced_amount', '9000'),
            ('差额', 'difference_amount', '1680')]:
        assert to_money(totals[label]) == to_money(authority[key]) == to_money(expected)
    recharge_rows = _rows(workbook['充值台账'])
    assert len(recharge_rows) == 5
    assert next(row for row in recharge_rows if row['记录编号'] == recharge['id'])['日期'] == '原始数据'
    assert workbook['充值台账']['H2'].value == '=充值原备注' and workbook['充值台账']['H2'].data_type == 's'
    assert [to_money(row['金额']) for row in recharge_rows[-2:]] == [to_money('-150'), to_money('150')]
    assert to_money(list(workbook['充值台账'].values)[-1][6]) == to_money('10180')
    invoice_rows = _rows(workbook['发票台账'])
    assert len(invoice_rows) == 4
    original_row = next(row for row in invoice_rows if row['记录编号'] == original['id'])
    assert original_row['发票号'] == NUMBER and to_money(original_row['价税合计']) == to_money('10000')
    assert to_money(original_row['已登记金额']) == to_money('10000') and original_row['原件份数'] == 2
    assert next(row for row in invoice_rows if row['记录编号'] == pending['id'])['状态'] == '待确认'
    assert next(row for row in invoice_rows if row['记录编号'] == voided['id'])['状态'] == '作废'
    assert sorted((row['票据关系'], str(to_money(row['已登记金额']))) for row in invoice_rows) == [
        ('作废', '0.00'), ('普通发票', '0.00'), ('普通发票', '10000.00'), ('红冲', '-1000.00')]
    assert workbook['发票台账']['C3'].value == '00000000000000000002' and workbook['发票台账']['C3'].data_type == 's'
    parts = [row for row in _rows(workbook['发票分配']) if row['发票记录编号'] == original['id']]
    assert [(row['核算项目'], to_money(row['分配金额'])) for row in parts] == [('甲', to_money('6000')), ('乙', to_money('4000'))]
    assert to_money(list(workbook['发票分配'].values)[-1][-1]) == to_money('9000')
    account_rows = _rows(workbook['用卡账户'], total=False)
    assert len(account_rows) == 2 and all(row['状态'] == '归档' for row in account_rows)
    assert next(row for row in account_rows if row['账户编号'] == first['id'])['当前实体卡号'] == replacement['card_no']
    assert workbook['实体油卡']['B2'].value == primary['card_no'] and workbook['实体油卡']['B2'].data_type == 's'
    openings = _rows(workbook['期初核对'])
    assert [(row['原日期标记'], row['截至月份'], row['确认状态'], to_money(row['已确认计入金额'])) for row in openings] == [
        ('原始数据', '2025-12', '已确认', to_money('500')), ('早期日期不明', None, '待确认', to_money('0'))]
    assert to_money(list(workbook['期初核对'].values)[-1][-2]) == to_money('500')
    history = [row for row in _rows(workbook['账户关联历史'], total=False) if row['账户编号'] == first['id']]
    assert [(row['关联项目'], row['实体卡号']) for row in history] == [('甲', CARD), ('乙', replacement['card_no'])]
    customs = _rows(workbook['用油自定义'], total=False)
    assert [(row['金额类型'], to_money(row['自定义金额']), row['备注']) for row in customs] == [
        ('特殊费用当月自定义额', to_money('-20.25'), '当月自定义原备注'), ('大表截至月额外快照', to_money('1000'), '截至月快照原备注')]
    assert all(cell.data_type != 'f' for sheet in workbook for row in sheet for cell in row)
    assert all(cell.number_format == '#,##0.00' for sheet in workbook for row in sheet for cell in row
        if isinstance(cell.value, (int, float)) and cell.column_letter in ('J', 'K', 'L', 'M', 'N') and sheet.title == '核对汇总')
