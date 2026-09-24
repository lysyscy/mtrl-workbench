"""`m06` 过程资料历史导入 · **真文件形态**的行为判据（`TAIL-75` B3 · 素材 = 用户那三份「材料机械消耗汇总表」）。

⭐ 真文件形态（2026-09-22 实测）：
- **两级表头**：r1 是标题（或空）、**r2 分组头（合并格）＋ r3 叶子**、数据 r4 起；
- 分组头是**合并格** ⇒ 只写在锚格（`外租机械（不含税）` 只写在 O2，P2 空）⇒ 不铺开就拼不出那一列；
- **父/子在项目间互换**：慧能 = `含税` › `外租设备（零租）`；晋江/南平 = `外租机械（含税）` › `零租`；
- 「**时间**」列 = 核算月，值是 **Excel 序列号**（`45444` ⇒ 2024-06）或**浮点年月**（`2025.09`；`2025.1` = 2025-10）；
- 尾部有**合计/小计行**；有 6 个机械列（**认出但⛔不导入**）。

⚠️ 与本目录 `test_page_m06.py` 分工：那份守**通用链**（认列/异常/落库/回滚），本份守**真文件逼出来的四件**。
"""

from __future__ import annotations

import io
from decimal import Decimal

import pytest
from openpyxl import Workbook

from app.core.errors import MustReportError
from app.mod06_proc import importer
from tests.acc_plugin import acc

#: 慧能那份：分组头 = 「含税/不含税」，叶子 = 「外租设备（零租）/（包月）」
HUI = [
    ("序号", "时间", "线上入库", "", "", "", "线下出库", "", "", "",
     "分公司设备", "内部设备", "含税", "", "不含税", ""),
    ("", "", "主辅材", "工具", "劳保", "周转材料", "主辅材", "工具", "劳保", "周转材料",
     "", "", "外租设备（零租）", "外租设备（包月）", "外租设备（零租）", "外租设备（包月）"),
]
#: 晋江/南平那份：分组头 = 「外租机械（含税）/（不含税）」，叶子 = 「零租/包月」
JJ = [
    ("序号", "时间", "线上入库（不含税）", "", "", "", "线下出库（不含税）", "", "", "",
     "分公司设备", "内部设备", "外租机械（含税）", "", "外租机械（不含税）", ""),
    ("", "", "主辅材", "工具", "劳保", "周转材料", "主辅材", "工具", "劳保", "周转材料",
     "", "", "零租", "包月", "零租", "包月"),
]
#: 分组头**是合并格**（真文件就是这样）⇒ 只有锚格有值
MERGES = [("C2", "F2"), ("G2", "J2"), ("M2", "N2"), ("O2", "P2"), ("A2", "A3"), ("B2", "B3")]


def _book(header: list[tuple], rows: list[tuple], *, extra_sheet: bool = False,
          merges: list[tuple[str, str]] | None = None) -> bytes:
    """⚠️ **真文件是 r1 标题 / r2 分组头 / r3 叶子** —— 夹具必须照这个行号，否则合并区会落在错的层上
    （自测踩过：少了标题行 ⇒ 分组头的合并区落到叶子上 ⇒ 分组头读不到 ⇒ 整表认不出）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "汇总表"
    if extra_sheet:                                  # 第 1 张放个"别的表"，考挑工作表
        other = wb.create_sheet("填表说明", 0)
        other.append(["本项目过程资料填报说明"])
        other.append(["时间", "内容"])
    ws.append(["项目材料设备工机具采购、消耗汇总表"])      # r1 标题
    for row in header:
        ws.append(row)
    for row in rows:
        ws.append(row)
    for a, b in (merges if merges is not None else MERGES):
        ws.merge_cells(f"{a}:{b}")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _data(ym, *, main=1000, tool=20):
    """一行数据（16 列，对齐上面的表头）。

    ⚠️ 「没填」用**空**表示（⛔ 不用 `0`）—— 真文件里空与 0 是两件事（空 ⇒ 不写这一列；0 ⇒ 明确为 0）。
    """
    return (1, ym, main, tool, "", "", "", "", "", "", 300, 400, 500, 600, 700, 800)


@acc("ACC-PROC-IMPORT-001", "001-01")
def test_两级表头拼名认列_父子在项目间互换也认得出():
    """⭐ 慧能（`含税`›`外租设备（零租）`）与晋江（`外租机械（含税）`›`零租`）**两种顺序都要认**。"""
    jj = importer.parse_sheet(_book(JJ, [_data(45871)]))
    assert jj.ignored_mech == ["租分公司设备", "内租设备", "外租零租含税", "外租包月含税",
                               "外租零租不含税", "外租包月不含税"], "机械 6 列要**全部认出来**（只为提示）"
    assert jj.unknown_headers == ["序号"]          # ⛔ 不把两级表头的两半各报一条（那是噪音）
    hui = importer.parse_sheet(_book(HUI, [_data(2025.11)]))
    assert hui.ignored_mech == jj.ignored_mech, "父/子互换的那份，机械 6 列也要全认出来"
    assert hui.unknown_headers == ["序号"]
    assert hui.header_levels == (2, 3) and hui.header_row == 2


@acc("ACC-PROC-IMPORT-001", "001-02")
def test_时间列_Excel序列号与浮点年月都认得出():
    """⭐ 「时间」= 核算月：`45444`（序列号 ⇒ 2024-06）· `2025.09` · `2025.1`（浮点丢了末位 0 ⇒ 10 月）。"""
    plan = importer.parse_sheet(_book(JJ, [_data(45444), _data(45475), _data(45507)]))
    assert plan.months == ["2024-06", "2024-07", "2024-08"], "Excel 日期序列号要按 1900 日期系统换算"
    float_plan = importer.parse_sheet(_book(HUI, [_data(2025.09), _data(2025.1), _data(2025.11)]))
    assert float_plan.months == ["2025-09", "2025-10", "2025-11"], "浮点年月 2025.1 = 2025-10"


@acc("ACC-PROC-IMPORT-001", "001-03")
def test_尾部合计行不计入也不报异常():
    """⭐ 真文件尾部是「合计/小计」行 ⇒ **不进 rows、也不进 anomalies**（只留一条提示）。"""
    plan = importer.parse_sheet(_book(JJ, [_data(45871), ("", "小计", 100, 20, 0, 0, 0, 0, 0, 0,
                                                          0, 0, 0, 0, 0, 0),
                                       ("", "合计：", 200, 40, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0)]))
    assert plan.writable == 1 and not plan.anomalies
    # 提示**攒成一条**（两行合计 ⇒ "行 5、6"；南平那份有 10 行空行，一行一条会刷屏）
    assert len(plan.notes) == 1 and "合计/小计行" in plan.notes[0] and "行 5、6" in plan.notes[0]


@acc("ACC-PROC-IMPORT-001", "001-04")
def test_机械列只提示不导入_数据行里不许出现机械金额():
    """⚠️ 口径：机械 6 列**全段手录**（旁路表）⇒ 导入时只提示；⛔ 混进快照就是口径破口。"""
    plan = importer.parse_sheet(_book(JJ, [_data(45871)]))
    row = plan.rows[0]
    assert not [k for k in row if k.startswith("mech_")], f"机械列不许进快照：{row}"
    assert set(row) == {"ym", "in_主辅材", "in_工具"}, "只有真填了值的材料列进快照"


@acc("ACC-PROC-IMPORT-001", "001-05")
def test_非数据行一律跳过并提示_不当异常报():
    """⭐ 用户 2026-09-22 裁定 ③④：空行就是没数据 · 数据区下面的**手写说明行**不是数据 ·
    合计块的下半行也不是 ⇒ 三样都**跳过 ＋ 逐条提示**（⛔ 不静默、⛔ 不当异常）；**有数据但月份认不出**才报异常。"""
    plan = importer.parse_sheet(_book(JJ, [
        _data(45871),
        ("", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""),          # 空行
        (2, "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""),           # 空行（只有序号）
        ("", "6月总出库也多了53.54", "", "", "", "", "", "", "", "", "", "", "", "", "", ""),  # 说明行
        ("", "合计：", 200, 40, "", "", "", "", "", "", "", "", "", "", "", ""),      # 合计行
        ("", "", 5750574.56, "", "", "", "", "", "", "", "", "", "", "", "", ""),   # 合计块下半行
    ]))
    assert plan.writable == 1 and not plan.anomalies, plan.anomalies
    joined = " ".join(plan.notes)
    assert "空行" in joined and "说明行" in joined and "合计" in joined, plan.notes
    # 有数据但月份认不出 ⇒ **必须报异常**（这条不能被上面的"跳过"冲淡）
    with pytest.raises(MustReportError) as bad:
        importer.parse_sheet(_book(JJ, [_data("八月")]))
    assert "行 4" in str(bad.value) and "核算月认不出" in str(bad.value)


@acc("ACC-PROC-IMPORT-001", "001-06")
def test_负数照收_导入层不改用户的定稿数字():
    """⭐ 用户 2026-09-22 裁定 ①：真文件里有**红字/冲销**（慧能「出库工具 = -7579.56」）⇒ **收**，
    值原样入库（⛔ 导入层不替它改数、⛔ 不因此丢整行）。"""
    plan = importer.parse_sheet(_book(HUI, [("", 2025.11, 704704.42, 84308.46, 24787.27, 102700,
                                             794325.17, -7579.56, 22353.64, 102700, 10214.33,
                                             "", 29950, 276942, 29077.67, 269050.48)]))
    row = plan.rows[0]
    assert row["ym"] == "2025-11" and row["out_工具"] == Decimal("-7579.56")
    assert not plan.anomalies, plan.anomalies


@acc("ACC-PROC-IMPORT-001", "001-07")
def test_标段行_中文年月带后缀照认_值不换算():
    """⭐ 用户 2026-09-22 裁定 ②：「2025年9月（一标）」是**真数据**（大设备，数没错）⇒
    认月（2025-09）· **值原样不动** · 后缀进提示；同月多行（一标/三标）都进 rows（页面侧相加）。"""
    plan = importer.parse_sheet(_book(JJ, [
        (16, "2025年9月（一标）", 33724630.66, 22078.7, 3962, "", 91933.76, 2884.2, 475.8, "",
         1275, "", 56700, 2400, 53260.17, 2123.89),
        (17, "2025年9月（三标）", "", "", "", "", 16861321.44, 16071, 2066.4, "", "", "", "", "", "", ""),
    ]))
    assert [r["ym"] for r in plan.rows] == ["2025-09", "2025-09"], "两标段都进（同月 ⇒ 页面侧相加）"
    assert plan.rows[0]["in_主辅材"] == Decimal("33724630.66"), "值不许被换算"
    assert plan.rows[1]["out_主辅材"] == Decimal("16861321.44")
    assert any("一标" in n for n in plan.notes), plan.notes
    assert not plan.anomalies


@acc()
def test_多工作表_按哪张认得出表头挑():
    """真文件里第 1 张可能是「填表说明」这类别的表 ⇒ 按"认得出表头"挑，⛔ 不猜位置。"""
    plan = importer.parse_sheet(_book(JJ, [_data(45871)], extra_sheet=True))
    assert plan.sheet == "汇总表"
