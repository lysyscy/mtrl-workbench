"""机械报表读回验真：历史月份、明确零与明细聚合精度。"""

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook

from app.mod06_proc import mech_report as reader


def book():
    wb = Workbook()
    s = wb.active
    s.title = "年度汇总"
    s.append(["项目", "费用", "往年合计", "1月", datetime(2026, 9, 1), datetime(2026, 8, 1),
              2023.9, 2023.1, 2024.1, 2024.9, 2024.1, "合计"])
    s.append(["报表A", "当月零租费", 999999, 17, 100, 200, 9, 10, 1, 9, 10, 888888])
    s.append(["报表A", "当月包月费", None, None, 300])
    s.append(["报表A", "当月内租费", None, None, 30])
    s.append(["报表B", "当月零租费", None, None, 0])
    s.append(["灵宝一期项目", "当月零租费", None, None, 999])
    s.append(["合计", "当月零租费", None, None, 999999])
    for name, col in [("零租设备月报表", 9), ("包月外租设备月报", 15)]:
        s = wb.create_sheet(name)
        s.cell(3, 1, "2026年9月")
        s.cell(5, 1, "一"); s.cell(5, 2, "报表A"); s.cell(5, col, 9999)
        for row, value in [(6, 12.345), (7, 12.345)]:
            s.cell(row, 1, row-5); s.cell(row, 2, "汽车吊"); s.cell(row, col, value)
        s.cell(8, 1, "合计"); s.cell(8, 2, "合计"); s.cell(8, col, 9999)
        s.cell(9, 2, "无序号设备"); s.cell(9, col, 1.005)
        s.cell(10, 1, "本页小计"); s.cell(10, col, 9999)
    # 不创建工程公司汇总表；解析必须仅依赖已确认的三张来源。
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_month_sources_history_and_aggregate_rounding():
    data = book()
    rows = {r["报表项目名"]: r for r in reader.read_months(data, ["2026-09"])}
    assert rows["报表A"] == {"核算月": "2026-09", "报表项目名": "报表A", "内租设备": "30.00",
                            "外租零租含税": "100.00", "外租包月含税": "300.00",
                            "外租零租不含税": "25.70", "外租包月不含税": "25.70"}
    assert rows["报表B"]["外租零租含税"] == "0.00"
    assert "内租设备" not in rows["报表B"]
    old = reader.read_months(data, ["2023-10", "2024-01", "2024-10", "2025-01", "2026-08"])
    amounts = {r["核算月"]: r["外租零租含税"] for r in old if r["报表项目名"] == "报表A"}
    assert amounts == {"2023-10": "10.00", "2024-01": "1.00", "2024-10": "10.00",
                       "2025-01": "17.00", "2026-08": "200.00"}
    assert all("外租零租不含税" not in r for r in old)
