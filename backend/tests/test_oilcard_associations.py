"""换卡、同实体卡号纠错与所选历史快照更正的真实事务保护。"""
from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook
import pytest
from sqlalchemy import event

from app.core import session as SESS
from app.mod13_oilcard.models import OilAccountAssignment
from tests.test_oilcard_archives import scene, create, edit, read  # noqa: F401
from tests.test_oilcard_ledger import entry, ledger, summary, update
from tests.test_oilcard_transfers import payload, transfer, pair


def replace(client, account, **values):
    fields = {'request_id': str(uuid4()), 'version': account['version'], 'card_id': None,
        'new_card': {'card_no': '0000000000000000457', 'parent_card_id': None, 'remark': '新卡独立备注'}, **values}
    response = client.post(f"/api/m13/accounts/{account['id']}/replace-card", json=fields)
    assert response.status_code == 200, response.text
    assert response.headers.get('X-Operation-History') == 'recorded'
    return fields, response.json()


def correct(client, row, **values):
    fields = {field: row[field] for field in ('version', 'project_id', 'account_id', 'card_id')}
    response = client.put(f"/api/m13/ledger/entries/{row['id']}/attribution", json={**fields, **values})
    assert response.status_code == 200, response.text
    assert response.headers.get('X-Operation-History') == 'recorded'
    return response.json()['item']


def replay(client, direction, operation):
    response = client.post(f'/api/history/{direction}', json={'operation_id': operation['id']})
    assert response.status_code == 200, response.text
    assert response.json()['ok']


def test_replacement_preserves_shared_card_old_money_refund_and_notes(scene):
    client, (project, other) = scene
    primary = create(client, 'cards', kind='primary', remark='主卡独立备注')
    old = create(client, 'cards', card_no='0000000000000000456', parent_card_id=primary['id'], remark='旧卡备注')
    account = create(client, 'accounts', name='换卡账户', project_id=project, card_id=old['id'], remark='账户独立备注')
    shared = create(client, 'accounts', name='同卡另一账户', project_id=project, card_id=old['id'])
    before = entry(client, account, '1000', date_text='原始数据', remark='旧账备注')
    _, result = replace(client, account)
    replaced, assignment = result['item'], result['assignment']
    new_id = replaced['card_id']
    assert replaced == {**account, 'card_id': new_id, 'version': replaced['version']}
    assert assignment['change_kind'] == 'replacement' and assignment['previous_card_id'] == old['id']
    assert assignment['card_id'] == new_id and assignment['changed_by']
    assert next(row for row in read(client)['accounts'] if row['id'] == shared['id']) == shared
    assert next(row for row in read(client)['cards'] if row['id'] == old['id']) == old
    assert next(row for row in read(client)['cards'] if row['id'] == primary['id']) == primary
    after = entry(client, replaced, '200', remark='新卡充值')
    assert before['card_id'] == old['id'] and after['card_id'] == new_id
    refund = client.post('/api/m13/ledger/refunds', json={'request_id': str(uuid4()), 'original_id': before['id'], 'amount': '10', 'remark': '原卡退款'}).json()['item']
    assert refund['card_id'] == old['id'] and refund['project_id'] == project
    moved = edit(client, 'accounts', replaced, project_id=other, status='archived')
    assert update(client, before, remark='仅更正旧账备注')['card_id'] == old['id']
    assert update(client, refund, date_text='2026-10-31')['card_id'] == old['id']
    assert summary(client)['total_amount'] == '1190.00'
    assert any(row['card_id'] == old['id'] and not row['current'] for row in summary(client)['rows'] if row['account_id'] == account['id'])
    trail = client.get(f"/api/m13/accounts/{account['id']}/assignments").json()
    assert [row['change_kind'] for row in trail] == ['assignment', 'replacement', 'assignment']
    assert trail[1]['previous_card_id'] == old['id']
    edit(client, 'accounts', moved, status='active')
    assert len(client.get(f"/api/m13/accounts/{account['id']}/assignments").json()) == 3
    assert next(row for row in read(client)['accounts'] if row['id'] == account['id'])['remark'] == '账户独立备注'


def test_replacement_idempotency_versions_undo_redo_and_failure_atomicity(scene):
    client, (project, _) = scene
    old = create(client, 'cards')
    account = create(client, 'accounts', name='安全换卡', project_id=project, card_id=old['id'])
    fields, result = replace(client, account)
    operation = client.get('/api/history/status').json()['undo']
    assert operation['count'] == 3
    assert client.post(f"/api/m13/accounts/{account['id']}/replace-card", json=fields).json() == result
    assert client.post(f"/api/m13/accounts/{account['id']}/replace-card", json={**fields, 'new_card': {'card_no': 'changed'}}).status_code == 409
    assert len(read(client)['cards']) == 2
    replay(client, 'undo', operation)
    assert read(client)['cards'] == [old] and read(client)['accounts'] == [account]
    assert len(client.get(f"/api/m13/accounts/{account['id']}/assignments").json()) == 1
    replay(client, 'redo', operation)
    assert read(client)['accounts'] == [result['item']]
    assert client.post(f"/api/m13/accounts/{account['id']}/replace-card", json={**fields, 'request_id': str(uuid4())}).status_code == 409
    now = result['item']
    current = read(client)
    def fail(*args):
        raise RuntimeError('synthetic replacement failure')
    event.listen(OilAccountAssignment, 'after_insert', fail)
    try:
        with pytest.raises(RuntimeError, match='synthetic replacement failure'):
            replace(client, now, new_card={'card_no': '0000458'})
    finally:
        event.remove(OilAccountAssignment, 'after_insert', fail)
    assert read(client) == current
    _, restored = replace(client, now, card_id=old['id'], new_card=None)
    assert restored['item']['card_id'] == old['id']
    operation = client.get('/api/history/status').json()['undo']
    assert operation['count'] == 2
    replay(client, 'undo', operation)
    assert read(client)['accounts'] == [now]


def test_number_correction_keeps_entity_and_explicit_history_correction_undo_export(scene):
    client, (project, other) = scene
    old = create(client, 'cards', card_no='0000000000000000001', remark='卡片备注')
    account = create(client, 'accounts', name='旧账账户', project_id=project, card_id=old['id'], remark='账户备注')
    original = entry(client, account, '100', remark='账务备注')
    corrected_card = edit(client, 'cards', old, card_no='0000000000000000002')
    assert corrected_card['id'] == old['id'] and len(read(client)['cards']) == 1
    assert ledger(client)['entries'] == [original]
    _, replacement = replace(client, account)
    moved = edit(client, 'accounts', replacement['item'], project_id=other, status='archived')
    # 旧项目/账户仍可保留，只明确更正卡片；不会套用账户当前的新项目。
    changed = correct(client, original, card_id=replacement['item']['card_id'])
    assert changed == {**original, 'card_id': replacement['item']['card_id'], 'version': changed['version']}
    assert read(client)['accounts'] == [moved]
    operation = client.get('/api/history/status').json()['undo']
    assert operation['count'] == 1
    replay(client, 'undo', operation)
    assert ledger(client)['entries'] == [original]
    replay(client, 'redo', operation)
    assert ledger(client)['entries'] == [changed]
    sheet = load_workbook(BytesIO(client.get('/api/m13/ledger/export').content), data_only=True).active
    assert sheet.cell(2, 6).value == '0000000000000000457'
    assert sheet.cell(2, 8).value == '账务备注'
    stale = {field: original[field] for field in ('version', 'project_id', 'account_id', 'card_id')}
    assert client.put(f"/api/m13/ledger/entries/{original['id']}/attribution", json=stale).status_code == 409
    assert ledger(client)['entries'] == [changed]
    primary = create(client, 'cards', kind='primary')
    bad = {**stale, 'version': changed['version'], 'card_id': primary['id']}
    assert client.put(f"/api/m13/ledger/entries/{original['id']}/attribution", json=bad).status_code == 400
    client.cookies.clear()
    assert client.put(f"/api/m13/ledger/entries/{original['id']}/attribution", json=stale).status_code == 401


@pytest.mark.parametrize('conflict', ['external_reference', 'later_account_edit'])
def test_replacement_undo_still_rejects_external_reference_or_later_edit_without_partial_changes(scene, conflict):
    client, (project, _) = scene
    old = create(client, 'cards')
    account = create(client, 'accounts', name='换卡保护', project_id=project, card_id=old['id'])
    _, replacement = replace(client, account)
    operation = client.get('/api/history/status').json()['undo']
    client.cookies.set(SESS.COOKIE_NAME, SESS.make('__replacement_other_actor__'))
    if conflict == 'external_reference':
        create(client, 'accounts', name='新卡另有使用者', project_id=project, card_id=replacement['item']['card_id'])
    else:
        edit(client, 'accounts', replacement['item'], remark='后续改动必须保留')
    before = read(client)
    trail = client.get(f"/api/m13/accounts/{account['id']}/assignments").json()
    client.cookies.set(SESS.COOKIE_NAME, SESS.make('__oilcard_actor__'))
    result = client.post('/api/history/undo', json={'operation_id': operation['id']})
    assert result.status_code == 409
    assert read(client) == before
    assert client.get(f"/api/m13/accounts/{account['id']}/assignments").json() == trail
    assert client.get('/api/history/status').json()['undo'] == operation


def test_transfer_history_correction_only_one_side_and_pair_versions_replay_together(scene):
    client, (project, other) = scene
    card = create(client, 'cards')
    source = create(client, 'accounts', name='转出', project_id=project, card_id=card['id'])
    target = create(client, 'accounts', name='转入', project_id=other, card_id=card['id'])
    extra = create(client, 'accounts', name='归档的正确账户', project_id=other, status='archived')
    out, incoming = transfer(client, payload(source, target))
    corrected = correct(client, out, project_id=other, account_id=extra['id'], card_id=None)
    changed_out, changed_in = pair(client, out['transfer_id'])
    assert corrected == changed_out
    assert corrected['kind'] == 'transfer_out' and corrected['amount'] == '-1000.00'
    assert changed_out['version'] == changed_in['version'] != out['version']
    assert {key: changed_in[key] for key in changed_in if key != 'version'} == {key: incoming[key] for key in incoming if key != 'version'}
    assert ledger(client)['total_amount'] == '0.00'
    operation = client.get('/api/history/status').json()['undo']
    assert operation['count'] == 2
    replay(client, 'undo', operation)
    assert pair(client, out['transfer_id']) == [out, incoming]
    replay(client, 'redo', operation)
    assert pair(client, out['transfer_id']) == [changed_out, changed_in]
    same = {field: changed_out[field] for field in ('version', 'project_id', 'account_id', 'card_id')}
    assert client.put(f"/api/m13/ledger/entries/{out['id']}/attribution", json={**same, 'account_id': target['id']}).status_code == 400
    assert pair(client, out['transfer_id']) == [changed_out, changed_in]
