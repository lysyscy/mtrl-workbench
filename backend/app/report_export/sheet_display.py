# -*- coding: utf-8 -*-
"""透视**显示区**的重建与对账 —— `TAIL-43`。

**它解决什么**（2026-09-18 真机实测根因）：模板法原先只重建「数据 sheet ＋ 缓存 ＋ 透视定义」，
**不重建显示区**（`xl/worksheets/sheetN.xml` 里 Excel 存下来的**静态单元格值**）。
本月数据**比模板少**时，透视的 `location` 会**收缩**，模板旧渲染行就落到新 `location`
**之外** ⇒ **任何应用都不会去清它们**（Excel 实测：新区域照渲染、区域外旧值原样留着）
⇒ 用户打开成品**读到的是模板那个月**的数（本片就是这么暴露的）。

⇒ 唯一可靠的做法 = **自己把显示区重建出来**，⛔ 不依赖"打开即刷新"是否被应用认：
   · `render_pivot_sheet`：整块换掉 `<sheetData>` —— **新显示区之外一格都不留**（清僵尸行的关键）；
   · `verify_display`：**显示区 vs 缓存重算**（独立于「数据 ↔ 缓存」的**第二条**对账）。

⚠️ 实测纠正两条（别走回头路）：
   1. **`refreshOnLoad` 只保证刷新缓存**（ECMA-376 `CT_PivotCacheDefinition`），
      **不保证**清掉 `location` 之外的残留 ⇒ "靠打开即刷新来顶"**靠不住**；
   2. **口径一律从模板读**（列序 / 分组映射 / 页筛选选中项）—— ⛔ 不在代码里另抄一份。
      页筛选是「`multipleItemSelectionAllowed="1"` ＋ `<item h="1">` 标未选」的**多选**形态，
      ⛔ **不许**改写成 `<pageField item="N"/>`（那是**单选**语义，会改变筛选行为）。
"""
import re
import zipfile
from html import unescape
from openpyxl.utils import column_index_from_string, get_column_letter

from app.core.errors import MustReportError
from app.report_export.xl_common import NS, fmt_num, norm_num


# ================= 显示区重建（`TAIL-43` · 2026-09-18） =================
# ⚠️⚠️ **为什么必须有这一节**（真机实测根因，⛔ 别删别绕过）：
#   模板法原先只重建「数据 sheet ＋ 缓存 ＋ 透视定义」，**不重建透视的「显示区」**
#   —— 即 `xl/worksheets/sheetN.xml` 里由 Excel 保存下来的**静态单元格值**。
#   当本月数据**比模板少**时，透视的 `location` 会**收缩**，模板的旧渲染行就落到了
#   新 `location` **之外** ⇒ **任何应用都不会去清它们**（Excel 实测：新区域照渲染，
#   区域外的旧单元格原样留着）⇒ 用户打开成品**读到的是模板那个月的数**。
#   ⇒ 唯一可靠的做法 = **自己把显示区重建出来**（不依赖"打开即刷新"是否被应用认）。
#   ⚠️ 顺带一条实测纠正：`refreshOnLoad` 只保证「刷新**缓存**」（ECMA-376
#      `CT_PivotCacheDefinition`），**不保证**清掉 `location` 之外的残留单元格。

def _sheet_rows(sheet_xml):
    """`{行号: 该行原始 XML}`（行本身原样复用，⛔ 不重写页筛选/表头行）。"""
    return {int(m.group(1)): m.group(0)
            for m in re.finditer(r'<row r="(\d+)"[^>]*>.*?</row>', sheet_xml, re.S)}


def _shared_items_of(fields, idx):
    """第 `idx` 个 `cacheField` 的 `<sharedItems>` 取值列表（保持顺序 = 下标即值）。"""
    si = re.search(r'<sharedItems([^>]*?)>(.*?)</sharedItems>', fields[idx][1], re.S)
    if not si:
        return []
    return [unescape(m.group(1)) for m in
            re.finditer(r'<(?:n|s|b|d)\s[^>]*v="([^"]*)"', si.group(2))]


def _pivot_field_xml(tpl_pt, idx):
    blk = re.search(r'<pivotFields count="\d+">(.*?)</pivotFields>', tpl_pt, re.S)
    if not blk:
        return ''
    items = re.findall(r'<pivotField\b[^>]*?(?:/>|>.*?</pivotField>)', blk.group(1), re.S)
    return items[idx] if idx < len(items) else ''


def _item_order(tpl_pt, idx):
    """`pivotField[idx]` 的 `<items>` 里各项下标顺序（跳过 `t="default"` 的空白项）。"""
    pf = _pivot_field_xml(tpl_pt, idx)
    m = re.search(r'<items[^>]*>(.*?)</items>', pf, re.S)
    if not m:
        return []
    out = []
    for im in re.finditer(r'<item\b([^>]*)/>', m.group(1)):
        if 't="default"' in im.group(1):
            continue
        x = re.search(r'\bx="(\d+)"', im.group(1))
        if x:
            out.append(int(x.group(1)))
    return out


def pivot_column_specs(tpl_pcd, tpl_pt, dir_field, group_field):
    """从**模板自身**读出透视的显示列结构（纯函数 · 无 IO）。

    口径来源（⛔ 不另立一份，双主人必漂）：
      · **方向项顺序** = `pivotTable` 里方向字段 `<items>` 的 `<item x>` 顺序
      · **分组项顺序** ＝ 同上（分组字段）；分组名 = `cacheField` 的
        `<fieldGroup base><groupItems>`，基础字段项 → 分组名 = `<discretePr>`
      · **列序** = 〔每个方向：各组… ＋ 该方向汇总〕… ＋ 总计
        （汇总 = 该字段的**小计**，Excel 自动渲染 ⇒ 模板 `colItems` 里就有它）

    ⚠️ **自我校验**：算出来的列数必须等于模板 `colItems` 的项数（含 grand）
    —— 不等说明**模板结构变了** ⇒ 必须重新核对显示区，⛔ 不许照着旧假设瞎写。
    """
    fields = re.findall(r'<cacheField\b([^>]*)>(.*?)</cacheField>', tpl_pcd, re.S)
    dir_vals = _shared_items_of(fields, dir_field)
    directions = [dir_vals[i] for i in _item_order(tpl_pt, dir_field) if i < len(dir_vals)]

    fg = re.search(r'<fieldGroup\b([^>]*)>(.*?)</fieldGroup>', fields[group_field][1], re.S)
    if not fg:
        raise MustReportError(
            '模板里「分组字段 %d」没有 `<fieldGroup>` ⇒ 无法还原显示区列序。' % group_field)
    base = re.search(r'\bbase="(\d+)"', fg.group(1))
    gi = re.search(r'<groupItems[^>]*>(.*?)</groupItems>', fg.group(2), re.S)
    dp = re.search(r'<discretePr[^>]*>(.*?)</discretePr>', fg.group(2), re.S)
    if not (base and gi and dp):
        raise MustReportError('模板里「分组字段 %d」的 `<fieldGroup>` 不完整（缺 base/groupItems/discretePr）。'
                              % group_field)
    base_idx = int(base.group(1))
    group_items = [unescape(x) for x in re.findall(r'<s v="([^"]*)"', gi.group(1))]
    dmap = [int(x) for x in re.findall(r'<x v="(\d+)"', dp.group(1))]
    base_vals = _shared_items_of(fields, base_idx)
    nat2group = {}
    for i, g in enumerate(dmap):
        if i < len(base_vals) and g < len(group_items):
            nat2group[base_vals[i]] = group_items[g]
    groups = [group_items[i] for i in _item_order(tpl_pt, group_field) if i < len(group_items)]

    columns = []
    letter_no = 2                                     # 标签列 = A
    for d in directions:
        for g in groups:
            columns.append({'letter': get_column_letter(letter_no), 'mode': 'cell',
                            'dir': d, 'group': g})
            letter_no += 1
        columns.append({'letter': get_column_letter(letter_no), 'mode': 'dir_total',
                        'dir': d, 'group': None})
        letter_no += 1
    columns.append({'letter': get_column_letter(letter_no), 'mode': 'grand',
                    'dir': None, 'group': None})

    m = re.search(r'<colItems count="(\d+)"', tpl_pt)
    tpl_cols = int(m.group(1)) if m else None
    if tpl_cols != len(columns):
        raise MustReportError(
            '模板透视的列结构与引擎的推导不一致：模板 `colItems`=%s，推导出 %d 列'
            '（方向 %s × 分组 %s）⇒ **模板已被替换**，显示区重建必须重新核对后再放行。'
            % (tpl_cols, len(columns), directions, groups))
    return {'directions': directions, 'groups': groups, 'nat2group': nat2group,
            'columns': columns}


def pivot_page_filter(tpl_pcd, tpl_pt, page_field):
    """读透视**页筛选的选中项**（口径来源 = 模板 · 纯函数）。

    ⚠️ 依据实测：模板用 `axisPage` ＋ `multipleItemSelectionAllowed="1"` ＋ 各
    `<item h="1">`（**隐藏 = 未选中**）表达"页筛选只留一个值"，`<pageField>` 上
    **不带 `item`**。⇒ ⛔ **不该**改写成 `<pageField fld="N" item="k"/>` ——
    那是**单选**语义，与模板存法不同（多选 vs 单选），会改变筛选行为。
    返回：被选中的值列表；该透视没有页筛选 ⇒ `None`。
    """
    pf = _pivot_field_xml(tpl_pt, page_field)
    if 'axis="axisPage"' not in pf:
        return None
    m = re.search(r'<items[^>]*>(.*?)</items>', pf, re.S)
    if not m:
        return None
    si = re.search(r'<sharedItems([^>]*?)>(.*?)</sharedItems>', _fields_of(tpl_pcd)[page_field][1],
                   re.S)
    vals = [unescape(x) for x in re.findall(r'<(?:n|s|b|d)\s[^>]*v="([^"]*)"', si.group(2))] \
        if si else []
    picked = []
    for im in re.finditer(r'<item\b([^>]*)/>', m.group(1)):
        attrs = im.group(1)
        if 't="default"' in attrs or re.search(r'\bh="1"', attrs):
            continue
        x = re.search(r'\bx="(\d+)"', attrs)
        if x and int(x.group(1)) < len(vals):
            picked.append(vals[int(x.group(1))])
    return picked


def _fields_of(tpl_pcd):
    return re.findall(r'<cacheField\b([^>]*)>(.*?)</cacheField>', tpl_pcd, re.S)


def render_pivot_sheet(tpl_sheet, spec, str_idx):
    """重建一张**透视 sheet** 的显示区（纯函数 · 无 IO）→ 新 sheet XML。

    ⭐ 关键是**「其余行全部丢弃」**：新显示区之外一格都不留 —— 这正是
    `TAIL-43` 的病灶（模板旧渲染值落在新 `location` 之外、没人清）。

    `spec`：
      · `keep_rows`     原样保留的行号（页筛选行等）
      · `header_rows`   `(首, 末)` —— 表头行**原样保留**（列标题/分组标题照模板）
      · `first_data_row` 第一个数据行号
      · `label_col`     行标签列（默认 `A`）
      · `columns`       数据列字母（`pivot_column_specs` 给的顺序）
      · `rows`          `[(标签, {列字母: 值}), …]`；**值为 `None` = 该格无匹配记录 ⇒ 留空**
                        （⛔ 不写 0 —— 0 会被读成「确实为 0」，与「没有」是两回事）
      · `total`         `(标签, {列字母: 值})` 总计行
    """
    label_col = spec.get('label_col', 'A')
    cols = [c['letter'] if isinstance(c, dict) else c for c in spec['columns']]
    keep = set(spec.get('keep_rows', ()))
    h0, h1 = spec['header_rows']
    hdr = set(range(h0, h1 + 1))
    old = _sheet_rows(tpl_sheet)

    def cell_xml(ref, value, is_label=False):
        if value is None or (isinstance(value, str) and value.strip() == ''):
            return ''
        if is_label or isinstance(value, str):
            s = str(value)
            if s not in str_idx:
                raise MustReportError(
                    '显示区重建缺字符串 `%s` ⇒ sharedStrings 未收录（⛔ 不静默留空）。' % s)
            return '<c r="%s" t="s"><v>%d</v></c>' % (ref, str_idx[s])
        n = norm_num(value)
        if n is None:
            return ''
        return '<c r="%s"><v>%s</v></c>' % (ref, fmt_num(round(n, 2)))

    max_col = 1
    for c in cols:
        max_col = max(max_col, column_index_from_string(c))
    max_col = max(max_col, column_index_from_string(label_col))

    parts = []
    for r in sorted(keep | hdr):
        if r in old:
            parts.append(old[r])
    first = spec['first_data_row']
    planned = list(spec['rows']) + [spec['total']]
    for i, (label, cells) in enumerate(planned):
        r = first + i
        body = cell_xml('%s%d' % (label_col, r), label, True)
        body += ''.join(cell_xml('%s%d' % (c, r), cells.get(c)) for c in cols)
        parts.append('<row r="%d" spans="1:%d">%s</row>' % (r, max_col, body))
    last = first + len(planned) - 1

    start = min(keep) if any(keep) else h0
    new = re.sub(r'<sheetData>.*?</sheetData>',
                 lambda mo: '<sheetData>%s</sheetData>' % ''.join(parts), tpl_sheet, flags=re.S)
    new = re.sub(r'<dimension ref="[^"]*"',
                 '<dimension ref="%s%d:%s%d"' % (label_col, start,
                                                 get_column_letter(max_col), last), new)
    return new


def verify_display(out_path, displays):
    """**显示区对账**：从成品的**缓存**重算 vs 成品**显示区的静态值**（独立第二条路径）。

    ⚠️ 它专抓 `TAIL-43` 这类病 —— 原有的「数据 ↔ 缓存」对账**看不见显示区**
    （显示区是 Excel 存的渲染结果，跟缓存是两回事）。
    两类问题各查一遍：① **僵尸行/多余行**（新显示区之外还有行）；② **格子里的数对不上缓存**。
    """
    import xml.etree.ElementTree as ET
    problems = []
    zf = zipfile.ZipFile(out_path)
    if 'xl/pivotCache/pivotCacheDefinition1.xml' not in zf.namelist():
        zf.close()
        return problems
    cd = ET.fromstring(zf.read('xl/pivotCache/pivotCacheDefinition1.xml'))
    fields = cd.findall('{%s}cacheFields/{%s}cacheField' % (NS, NS))
    fvals, fgrp = [], []
    for f in fields:
        si = f.find('{%s}sharedItems' % NS)
        vals = []
        if si is not None:
            for ch in si:
                vals.append(float(ch.get('v')) if ch.tag.endswith('n') else str(ch.get('v')))
        fvals.append(vals)
        fg = f.find('{%s}fieldGroup' % NS)
        # ⚠️ 模板里**基础字段自己**也带 `<fieldGroup par="N"/>`（空壳，只是"我属于 N 这个分组"）
        #    ⇒ 只有**同时**具备 `base`/`groupItems`/`discretePr` 的才算分组字段。
        gi = fg.find('{%s}groupItems' % NS) if fg is not None else None
        dd = fg.find('{%s}discretePr' % NS) if fg is not None else None
        if fg is None or gi is None or dd is None:
            fgrp.append(None)
        else:
            fgrp.append({'items': [c.get('v') for c in gi],
                         'map': [int(c.get('v')) for c in dd],
                         'base': int(fg.get('base'))})
    cr = ET.fromstring(zf.read('xl/pivotCache/pivotCacheRecords1.xml'))
    recs = []
    for r in cr.findall('{%s}r' % NS):
        xs = r.findall('{%s}x' % NS)
        idxs = []
        for x in xs:
            v = x.get('v')
            idxs.append(int(v) if v is not None else None)
        recs.append(idxs)

    def decode(xs, f):
        if f >= len(fgrp):
            return None
        g = fgrp[f]
        if g is not None:
            bk = xs[g['base']] if g['base'] < len(xs) else None
            gi = g['map'][bk] if bk is not None and bk < len(g['map']) else None
            return g['items'][gi] if gi is not None and gi < len(g['items']) else None
        k = xs[f] if f < len(xs) else None
        if k is None or k >= len(fvals[f]):
            return None
        return fvals[f][k]

    for spec in displays:
        name = spec['sheet_file']
        if name not in zf.namelist():
            problems.append('缺透视 sheet：%s' % name)
            continue
        sheet = zf.read(name).decode('utf-8')
        rows = _sheet_rows(sheet)
        first = spec['first_data_row']
        want_last = first + len(spec['rows'])            # 含总计行
        stray = [r for r in rows if r > want_last]
        if stray:
            problems.append('%s 显示区之外还有 %d 行（僵尸旧值）：%s'
                            % (name, len(stray), sorted(stray)[:8]))
        agg, matched = {}, {}
        for xs in recs:
            if spec.get('filter'):
                ff, fv = spec['filter']
                if decode(xs, ff) != fv:
                    continue
            label = decode(xs, spec['label_field'])
            d = decode(xs, spec['dir_field']) if spec.get('dir_field') is not None else None
            # ⚠️ 分组字段**本身**就带 `<fieldGroup>` ⇒ `decode` 已经把值解析成**分组名**了，
            #    ⛔ 别再经 `nat2group` 映射一次（那会把「数据组1」这类"分组名 ≠ 基础值名"的
            #    整组抹成 None —— 实测踩到：只有 主材/辅材 那一组对不上，其余靠"同名"侥幸通过）。
            g = decode(xs, spec['group_field']) if spec.get('group_field') is not None else None
            amt = norm_num(decode(xs, spec['value_field']))
            if amt is None:
                continue
            # ⚠️ 用 **set** 而不是 tuple：`g is None` 时 `(d, g)` 与 `(d, None)` 是同一个桶，
            #    用 tuple 连加两次 ⇒ 汇总会凭空翻一倍（实测踩到）。
            for key in {(label, d, g), (label, d, None), (label, None, None)}:
                agg[key] = round(agg.get(key, 0.0) + amt, 2)
                matched[key] = matched.get(key, 0) + 1
        for i, (label, _cells) in enumerate(spec['rows']):
            r = first + i
            row = rows.get(r, '')
            for c in spec['columns']:
                key = (label, c['dir'], c['group'])
                exp = agg.get(key) if matched.get(key) else None
                got = _cell_number(row, '%s%d' % (c['letter'], r))
                if exp is None and got is None:
                    continue
                if exp is None or got is None or abs(exp - got) > 0.01:
                    problems.append('%s 显示区 %s%d：%s vs 缓存重算 %s'
                                    % (name, c['letter'], r, got, exp))
    zf.close()
    return problems


def _cell_number(row_xml, ref):
    m = re.search(r'<c r="%s"(?:\s[^>]*)?>(.*?)</c>' % ref, row_xml or '', re.S)
    if not m:
        return None
    v = re.search(r'<v>([^<]*)</v>', m.group(1))
    if not v:
        return None
    return norm_num(unescape(v.group(1)))
