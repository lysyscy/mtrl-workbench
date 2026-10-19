# -*- coding: utf-8 -*-
"""路径①（**模板填充**）的通用引擎 —— `openpyxl` 读模板 → 写数据区 → **尾块随行数派生**。

## 依据（口径回主人处读，本处只落形态）

- 《材料科业务全流程说明书》**§7.1**：① 数据区起点按**标记文字** `DATA_START` 定位（⛔ 不抄坐标）；
  ② **尾部行块随行数派生**（⛔ 不删行让位）。
- 交付面 `docs/requirements/report-templates.md`（**唯一登记处**）：§四-5 实测做法（样板数据区 → 插入/删除 →
  尾部行块按 offset 整体下移 → 改写公式区间与合并）· §六-2（模板里的公式**不参与求值**
  ⇒ 口径数字由系统算好、**以数值覆盖**）。

## ⚠️ 与路径②（`tpl_engine` · `zip` 字节替换）**是两类做法**

那条保模板的字节、只换「数据 / 缓存 / 透视」；这条用 `openpyxl` **重写整个 sheet**
（表头/样式/列宽/合并沿用模板，**值由调用方给**）。

## 分工（同 `tpl_engine` 的规矩）

本引擎**只认结构与位移、不认业务** —— 哪一行放什么、哪一列是什么，全部由调用方（实现层）给。
⇒ 域内核算数、引擎摆位、实现层给映射。

## ⭐ 三条硬性（都从 §7.1 与清单 §四-5 来，⛔ 不许简化）

1. **认标记文字**定位数据区起点（实测：只改列位置就会错 ⇒ 坐标一律不抄）。
2. **尾块随行数派生**：先把样板数据区对齐到 n 行（**保留首行当样式母版**）⇒ 写值 ⇒
   尾块（合计/小计/脚注/合并块）**按 offset 整体平移**（⚠️ `openpyxl` 的 `insert_rows` /
   `delete_rows` 搬**值与样式**、**但不搬合并区**，实测 3.1.5 ⇒ 合并区必须自己平移）。
3. ⭐⭐ **「新数据区之外不留残留」守卫** —— 与 `TAIL-43` **同族**：模板里带着示例数据，
   只写前 n 行、不清尾部 ⇒ 用户会读到**示例 / 上月**的数（实测踩过）。
"""

from __future__ import annotations

import re
from io import BytesIO
from copy import copy
from decimal import Decimal
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter
from app.report_export.save import save_workbook

#: 数据区起点的**标记文字**（方案 A：认文字，不认坐标/定义名称）
MARKER = "DATA_START"

#: 写文本前先设的**数字格式**（`模板清单` §四-4：模板那几列原本是**日期 / 数值格式**，
#: 不先改就会被 Excel 反解 —— `2026-07` 被吃成日期）
TEXT_FMT = "@"


def set_text_format(ws, plan: dict, col: str) -> int:
    """把 `col` 列的**数据行**格设成**文本格式** ⇒ 返回改了几格。

    ⭐ 为什么在引擎里：**同一个坑已经踩过两次** —— 台账的「月份」列、过程资料的「时间」列
       （`模板清单` §四-4 那一条裁定管的就是这两列）⇒ 口径提到引擎，⛔ 实现层别再各写一份。
    ⚠️ 必须在 `fill_sheet` **之后**调用（它要 `plan` 里的 `data_start` / `data_rows`）。
    """
    ci = column_index_from_string(col)
    for i in range(len(plan["data_rows"])):
        ws.cell(row=plan["data_start"] + i, column=ci).number_format = TEXT_FMT
    return len(plan["data_rows"])

#: 「回指数据区」的汇总公式：形如 `=SUM(C3:C18)` / `=SUBTOTAL(9,O2:O8)`，
#: ⚠️ 必须**跨多行**且**起于数据区首行** —— 否则数据行里的 `=SUM(B2:N2)` 会被误判成尾块（实测）。
_SUM_REF = re.compile(r"(?P<c1>\$?[A-Z]{1,3})\$?(?P<r1>\d+)\s*:\s*(?P<c2>\$?[A-Z]{1,3})\$?(?P<r2>\d+)")


def _row_values(ws, r, max_col):
    return [ws.cell(row=r, column=c).value for c in range(1, max_col + 1)]


def keep_only_columns(ws, *, keep: str, drop) -> dict:
    """把模板里的若干列**压成一列**：删掉 `drop`（列字母），只留 `keep`（列字母）。

    ⭐ 为什么需要它：`模板清单` **§5.2** 对台账「特殊费用月报表」的裁定（①丙）= **月份列由系统按
       「当前核算月」生成**（模板的 `3月份`…`8月份` 只是示例、会被替换）⇒ 单月导出时要把那几列压成一列。
       ⛔ 这不是"多写一份填充逻辑" —— 行维度仍交 `fill_sheet`，本函数**只挪列**（结构级、不认业务）。

    ⚠️ 三条**实测**（与 `insert_rows` / `delete_rows` 是同一族坑）：
      1. `delete_cols` **不搬合并区** ⇒ 跨列合并（如标题 `A1:H1`）必须**自己按新几何重算**；
      2. `delete_cols` 会留下**列属性空壳**（末列以右的 `column_dimensions`）⇒ 删完清掉；
      3. ⭐ **必须在 `fill_sheet` 之前调用** —— 数据区 / 尾块的定位都基于列，先填后删会把对账计划全打乱。

    返回 `{"kept": 列字母, "dropped": [...], "max_col": 新列数, "merges": [(新合并区字符串), …]}`。
    """
    from openpyxl.utils import column_index_from_string, get_column_letter

    ki = column_index_from_string(keep)
    di = sorted({column_index_from_string(c) for c in drop})
    if not di:
        return {"kept": keep, "dropped": [], "max_col": ws.max_column,
                "merges": [str(r) for r in ws.merged_cells.ranges]}
    if min(di) <= ki:
        raise ValueError("只支持把**右侧**的列并在左侧那一列上（`keep` = %s 必须在 `drop` 左边）" % keep)

    before = [(r.min_col, r.min_row, r.max_col, r.max_row) for r in ws.merged_cells.ranges]
    for rng in [str(r) for r in ws.merged_cells.ranges]:      # 先全拆，删完按新几何重并
        ws.unmerge_cells(rng)

    for ci in reversed(di):                                   # ⚠️ **从右往左**删（高索引先做 ⇒ 低索引不受影响）
        ws.delete_cols(ci, 1)

    merges: list[str] = []
    for mn, mr, mx, xr in before:
        new_mn = mn - sum(1 for ci in di if ci < mn)
        new_mx = mx - sum(1 for ci in di if ci <= mx)
        if new_mx < new_mn:
            continue                                          # 整块都在被删的列里 ⇒ 无处可并
        ws.merge_cells(start_column=new_mn, start_row=mr, end_column=new_mx, end_row=xr)
        merges.append(
            "%s%d:%s%d" % (get_column_letter(new_mn), mr, get_column_letter(new_mx), xr))

    for c in [x for x in ws.column_dimensions if x and column_index_from_string(x) > ws.max_column]:
        ws.column_dimensions.pop(c, None)                    # 清掉末列以右的列属性空壳

    return {"kept": keep, "dropped": [get_column_letter(ci) for ci in di],
            "max_col": ws.max_column, "merges": merges}


def locate_data_area(ws, *, marker: str = MARKER, key_col: str = "A") -> dict:
    """定位**数据区**与**尾块**。

    - **起点** = `key_col` 列里**逐字等于** `marker` 的那一格（⛔ 不认坐标）；
    - **终点** = 下方**第一个**「回指数据区」的汇总行（其公式区间跨多行且起于数据区首行）
      ⇒ **尾块 = 它及其以下**；找不到（该 sheet 尾部没有汇总行）⇒ 数据区直到 `max_row`、无尾块。
    - ⛔ 合并区落在**数据区内部** ⇒ **显式报错**（形态不认识 —— 比产出错文件好）。
    """
    key_idx = column_index_from_string(key_col)
    hits = [r for r in range(1, ws.max_row + 1)
            if str(ws.cell(row=r, column=key_idx).value or "").strip() == marker]
    if not hits:
        raise ValueError(
            "模板的 `%s` 列里找不到标记文字 `%s` ⇒ 数据区起点无从定位。"
            "（方案 A：认文字、不抄坐标；清单 §二 登记了每个 sheet 该填哪一格）" % (key_col, marker))
    if len(hits) > 1:
        raise ValueError("标记文字 `%s` 在 `%s` 列出现 %d 次（%s）⇒ 起点不唯一。"
                         % (marker, key_col, len(hits), hits))
    ds = hits[0]

    de, tail_start = ws.max_row, None
    for r in range(ds + 1, ws.max_row + 1):
        if any(_is_area_total(v, ds) for v in _row_values(ws, r, ws.max_column)):
            de, tail_start = r - 1, r
            break
    if tail_start is None:
        tail_start = ws.max_row + 1

    for rng in ws.merged_cells.ranges:                       # 数据区内部不许有合并（形态不认识）
        if ds <= rng.min_row <= de or ds <= rng.max_row <= de:
            raise ValueError("模板的数据区（第 %d~%d 行）里有合并区 `%s` ⇒ 形态不认识，"
                             "本引擎拒绝据此产出文件。" % (ds, de, rng))
    return {"data_start": ds, "data_end": de, "tail_start": tail_start,
            "sample_rows": de - ds + 1, "key_col": key_col}


def _is_area_total(value, data_start: int) -> bool:
    """该格是不是「回指数据区」的汇总公式（⇒ 它所在行 = 尾块首行）。"""
    if not isinstance(value, str) or not value.startswith("="):
        return False
    for m in _SUM_REF.finditer(value):
        if int(m.group("r1")) == data_start and int(m.group("r2")) > data_start:
            return True
    return False


def _copy_row_style(ws, src_row: int, dst_row: int, max_col: int) -> None:
    """把样板行的**样式 / 数字格式 / 行高**抄到新行（⛔ `insert_rows` 不会带样式）。"""
    for c in range(1, max_col + 1):
        dst_row_obj = ws.cell(row=dst_row, column=c)
        src = ws.cell(row=src_row, column=c)
        dst_row_obj._style = copy(src._style)
    src_dim, dst_dim = ws.row_dimensions.get(src_row), ws.row_dimensions.get(dst_row)
    if src_dim is not None and src_dim.height is not None:
        if dst_dim is None:
            dst_dim = ws.row_dimensions[dst_row]
        dst_dim.height = src_dim.height


def _shift_merges(ws, from_row: int, offset: int) -> None:
    """把 `from_row` 及以下的合并区按 `offset` 平移（⚠️ `openpyxl` 自己**不搬**，实测 3.1.5）。

    ⚠️⚠️ **别用 `ws.unmerge_cells(str(rng))`**（2026-09-18 实测踩到 · 片 4a 第一次真用本函数时）：
      它会对区间里**每一个**格执行 `del self._cells[(row, col)]`，**不管那格有没有被创建过**
       ⇒ 只要区间里有一个"从没被访问过的影子格"就抛 `KeyError: (r, c)`
      （实测：过程资料 `Q13:R13` 的 `R13`）。⇒ 直接改 `merged_cells.ranges`（**纯几何登记**，
      不碰单元格本体）：旧范围移除、新范围登记。
      ⭐ 通法：**合并区的"几何"与"单元格本体"是两套账** —— 挪位置只该动前者。
    """
    moved = [rng for rng in list(ws.merged_cells.ranges) if rng.min_row >= from_row]
    for rng in moved:
        box = (rng.min_col, rng.min_row + offset, rng.max_col, rng.max_row + offset)
        ws.merged_cells.ranges.remove(rng)
        ws.merge_cells(start_column=box[0], start_row=box[1], end_column=box[2], end_row=box[3])


def tail_end_row(ws, tail_start: int, max_col: int) -> int:
    """尾块的**终点行** = 从尾块首行往下**最后一行还有内容**（有值 或 被合并区覆盖到）。

    ⚠️ 为什么必须按结构算：尾块里可能有**不归调用方覆盖**的行 —— 过程资料的**固定脚注**
       （`A14:R15` 跨两行合并块，清单 §四-5）就是。只按"调用方覆盖了哪几行"算终点的话，
       脚注会被 `verify_filled` 当**残留**报红（实测踩到：`汇总表!第 7 行（尾块以下）还有内容`）。
    ⭐ 认结构、不抄行数：脚注几行**不问调用方**，问 sheet 自己（值 ＋ 合并区几何）。

    没东西 ⇒ 返回 `tail_start - 1`（= 数据区末行，调用方据此得"没有尾块"）。
    """
    end = tail_start - 1
    for r in range(tail_start, ws.max_row + 1):
        if any(ws.cell(row=r, column=c).value is not None for c in range(1, max_col + 1)):
            end = r
    for rng in ws.merged_cells.ranges:
        if rng.min_row >= tail_start:
            end = max(end, rng.max_row)
    return end


def _clear_row(ws, r: int, max_col: int) -> None:
    """**整行清值**（只清值、不动样式）—— 防"只写前 n 行"式的残留。"""
    for c in range(1, max_col + 1):
        ws.cell(row=r, column=c).value = None


def _trim_row_dims(ws, last_row: int) -> None:
    """删掉 `last_row` 以下的**空行壳**（行属性 ＋ 空单元格登记）。

    ⚠️ 壳有**两种**（2026-09-18 实测踩到两种，过程资料 `UsedRange` 报成 `A1:R11` 而内容只到第 8 行）：
      ① `delete_rows` 搬走值与样式后**留下空的 `<row>` 元素** ⇒ `row_dimensions` 里还挂着它；
      ② 模板里本来就有 `<row>` 元素、**没有值**的空行（过程资料模板 R16~R18 就是）——
         它们会跟着 `delete_rows` **上移**，成为"**有行无值**"的壳。
    两种都会让 Excel 的 `UsedRange` 多算几行 ⇒ 用户看到"数据下面还空着几行"、还会进打印范围
    （与 `TAIL-43` 同族：**"新数据区之外不留残留"**）。
    """
    if ws.max_row > last_row:
        ws.delete_rows(last_row + 1, ws.max_row - last_row)      # 先删行（连空单元格登记一起）
    for r in [x for x in ws.row_dimensions if x > last_row]:
        ws.row_dimensions.pop(r, None)                          # 再清行属性（`delete_rows` 不保证清它）


def _check_tail_rows(ws, tail_start: int, written: set, max_col: int) -> None:
    """⚠️ 尾块里**没被我方覆盖的格** ⇒ **只许放文本**（脚注之类）。

    出现**公式**或**数字** ⇒ 显式报错：那要么是"忘了给汇总值"、要么是模板自带的示例数字残留
    —— 两种都会让用户读到算不出来 / 不是本月的数（`TAIL-43` 同族）。
    """
    for r in range(tail_start, ws.max_row + 1):
        for c, v in enumerate(_row_values(ws, r, max_col), start=1):
            if (r, c) in written:
                continue
            if v is None or (isinstance(v, str) and v.strip() == ""):
                continue
            if isinstance(v, str) and v.startswith("="):
                raise ValueError("尾块 `%s%d` 有本引擎未覆盖的公式 ⇒ 拒绝产出"
                                 "（汇总值必须由调用方算好后给）。" % (get_column_letter(c), r))
            if isinstance(v, (int, float)):
                raise ValueError("尾块 `%s%d` 有本引擎未覆盖的数字 %r ⇒ 拒绝产出"
                                 "（模板自带的示例值不许留在产物里）。"
                                 % (get_column_letter(c), r, v))


def fill_sheet(ws, spec: dict) -> dict:
    """按 `spec` 重写一个 sheet 的数据区 ＋ 尾块 ⇒ 返回**对账计划**（供 `verify_filled` 读回）。

    `spec`：
      - `sheet`    —— sheet 名（仅登记用）
      - `rows`     —— `[(label, {列字母: 值})]`，**顺序 = 行序**；`label=None` ⇒ 首格留空
      - `tail`     —— `{尾块行偏移: (label, {列字母: 值})}`，**稀疏**（未提到的尾行原样保留，见守卫）
      - 可选：`marker` / `key_col`（默认 `DATA_START` / `A`）

    ⚠️ 值一律**以数值覆盖**（清单 §六-2：模板公式不参与求值）；`None` ⇒ 该格**留空**
       （⛔ 不写 0 —— 0 会被读成"确实为 0"）。
    """
    key_col = spec.get("key_col", "A")
    key_idx = column_index_from_string(key_col)
    max_col = ws.max_column
    area = locate_data_area(ws, marker=spec.get("marker", MARKER), key_col=key_col)
    ds, sample_n, tail_start = area["data_start"], area["sample_rows"], area["tail_start"]
    rows = list(spec["rows"])
    n = len(rows)
    offset = n - sample_n

    # ① 样板数据区对齐到 n 行（**保留首行当样式母版**）
    if offset > 0:
        ws.insert_rows(ds + sample_n, offset)
        for r in range(ds + sample_n, ds + n):
            _copy_row_style(ws, ds, r, max_col)
    elif offset < 0:
        ws.delete_rows(ds + n, -offset)
    if offset:
        _shift_merges(ws, tail_start, offset)

    # ② 写数据行（整行先清值 ⇒ 新数据区之外一格不留）
    written: set = set()
    for i, (label, cells) in enumerate(rows):
        r = ds + i
        _clear_row(ws, r, max_col)
        if label is not None:
            ws.cell(row=r, column=key_idx).value = label
        for col, value in (cells or {}).items():
            if value is not None:
                ws.cell(row=r, column=column_index_from_string(col)).value = value
                written.add((r, column_index_from_string(col)))

    # ③ 尾块：按 offset 平移（openpyxl 已搬值与样式）＋ 覆盖汇总值
    tail = dict(spec.get("tail") or {})
    for off, (label, cells) in tail.items():
        r = ds + n + int(off)
        if label is not None:
            ws.cell(row=r, column=key_idx).value = label
        for col, value in (cells or {}).items():
            ci = column_index_from_string(col)
            if value == "":                                  # ⭐ **空串 = 显式留空**（清掉模板那格的公式 / 旧值）
                ws.cell(row=r, column=ci).value = None
                written.add((r, ci))
            elif value is not None:
                ws.cell(row=r, column=ci).value = value
                written.add((r, ci))
    _check_tail_rows(ws, ds + n, written, max_col)
    # 最后一行 = **尾块终点**（按结构算：覆盖到的行 ∪ 还有内容的行 ∪ 合并区覆盖到的行）；
    # 没有尾块时 = 数据末行；连数据行都没有 ⇒ 表头行（ds-1）
    tail_end = tail_end_row(ws, ds + n, max_col)
    last_row = max([ds + n - 1, ds - 1, tail_end] + [ds + n + int(off) for off in tail])
    _trim_row_dims(ws, last_row)

    return {"sheet": ws.title, "key_col": key_col, "max_col": max_col,
            "data_start": ds, "data_rows": [(lb, dict(c or {})) for lb, c in rows],
            "tail_start": ds + n, "last_row": last_row,
            "tail": [(ds + n + int(off), lb, dict(c or {})) for off, (lb, c) in tail.items()]}


def fill_workbook(template, out_path, specs) -> dict:
    """读模板 → 逐个 sheet 填充 → 存到 `out_path` ⇒ 返回对账计划。

    ⚠️ 一律用**可写模式**读（保留公式文本以便守卫发现漏覆盖）；不会改动模板本体。
    """
    wb = openpyxl.load_workbook(str(template))
    try:
        plan = [fill_sheet(wb[s["sheet"]], s) for s in specs]
        save_workbook(wb, out_path)
    finally:
        wb.close()
    return {"path": str(out_path), "sheets": plan}


def verify_filled(path, plan) -> list[str]:
    """**产物显示区对账**（第二道：读回产物 vs 调用方给的期望值）。

    抓的是「同一份代码自己解析自己」查不出的那类问题：
      ① 数据行逐格 ≠ 期望（含**该空的没空**）；
      ② 数据行数/位置不对；
      ③ ⛔ **尾块以下还有内容**（模板旧渲染值残留 —— `TAIL-43` 同族病灶）；
      ④ ⛔ 显示区里还有**公式**（值覆盖不彻底 ⇒ 打开时是算不出来 / 陈旧的值）。

    ⚠️ 必须**读两遍**：`data_only=True` 只给值（公式格读成 `None`）、`data_only=False` 才看得到公式
       —— 实测踩到：只用前者时 ③④ 两条**永远不会报**（死代码，比没有还坏）。
    """
    # 文件路径和内存成品共用同一套对账，两个解析器各自持有输入流。
    problems: list[str] = []
    wb = openpyxl.load_workbook(BytesIO(path) if isinstance(path, bytes) else str(path), data_only=True)
    wf = openpyxl.load_workbook(BytesIO(path) if isinstance(path, bytes) else str(path), data_only=False)
    try:
        for sp in plan["sheets"]:
            ws, raw = wb[sp["sheet"]], wf[sp["sheet"]]
            max_col = sp["max_col"]
            ds, n = sp["data_start"], len(sp["data_rows"])
            key_idx = column_index_from_string(sp["key_col"])
            for i, (label, cells) in enumerate(sp["data_rows"]):
                r = ds + i
                if label is not None:
                    got = ws.cell(row=r, column=key_idx).value
                    if str(got or "") != str(label):
                        problems.append("%s!%s%d 名称不符：%r ≠ %r"
                                        % (sp["sheet"], sp["key_col"], r, got, label))
                for c in range(1, max_col + 1):
                    if c == key_idx:                     # 名称列已单独比对（⛔ 别用"格值"再比一次）
                        continue
                    letter = get_column_letter(c)
                    exp = cells.get(letter)
                    got = ws.cell(row=r, column=c).value
                    if not _same(got, exp):
                        problems.append("%s!%s%d 与期望不符：%r ≠ %r"
                                        % (sp["sheet"], letter, r, got, exp))
            for r, _lb, cells in sp["tail"]:
                for letter, exp in cells.items():
                    got = ws.cell(row=r, column=column_index_from_string(letter)).value
                    if not _same(got, exp):
                        problems.append("%s!%s%d（尾块）与期望不符：%r ≠ %r"
                                        % (sp["sheet"], letter, r, got, exp))
            last = sp["last_row"]          # 单一出口：`fill_sheet` 已按**结构**算好（含脚注那种不覆盖的行）
            for r in range(last + 1, raw.max_row + 1):          # ③ 残留（含公式残留）
                if any(raw.cell(row=r, column=c).value is not None for c in range(1, max_col + 1)):
                    problems.append("%s!第 %d 行（尾块以下）还有内容 ⇒ **模板旧值残留**"
                                    % (sp["sheet"], r))
            # ③b 空壳：行属性留在 last 以下 ⇒ Excel 的 UsedRange 会多算几行（看着像"下面还有内容"）
            shells = sorted(r for r in raw.row_dimensions if r > last)
            if shells:
                problems.append("%s!第 %s 行还留着**行属性空壳** ⇒ 打开时数据下面多出几行空的"
                                % (sp["sheet"], shells[:6]))
            # ③c 行壳：`max_row` 超过末行（"有行无值"的空行）⇒ `UsedRange` 同样会多算
            if raw.max_row > last:
                problems.append("%s!末行以下还有**行壳**（max_row=%d > 末行 %d）⇒ "
                                "打开时 UsedRange 会多算几行（与 TAIL-43 同族）"
                                % (sp["sheet"], raw.max_row, last))
            for r in range(ds, last + 1):                       # ④ 值覆盖不彻底
                for c in range(1, max_col + 1):
                    if raw.cell(row=r, column=c).data_type == "f":
                        problems.append("%s!%s%d 还是公式 ⇒ 值覆盖不彻底"
                                        % (sp["sheet"], get_column_letter(c), r))
    finally:
        wb.close()
        wf.close()
    return problems


def _same(got, exp) -> bool:
    """产物的值与期望是否一致（数字按 2 位小数、容差 0；空 = `None` / 空串）。

    ⚠️ 期望值是**金额**（域内核给的是 `Decimal`）⇒ 数值比较必须认 `Decimal`
       —— 只判 `int/float` 会让 `Decimal("0.00")` 掉进字符串比较、报出「0 ≠ 0.00」的假红（实测踩到）。
    """
    if exp is None or exp == "":
        # ⚠️ 空串与 `None` 同义 = **该格留空**（`""` 是调用方表达"显式留空"的写法，见 `fill_sheet` 尾块）
        if got is None:
            return True
        return isinstance(got, str) and got.strip() == ""
    if got is None:
        return False
    if isinstance(exp, (int, float, Decimal)) and not isinstance(exp, bool):
        try:
            return abs(float(got) - float(exp)) < 0.005
        except (TypeError, ValueError):
            return False
    return str(got) == str(exp)


__all__ = ["MARKER", "TEXT_FMT", "set_text_format", "keep_only_columns", "locate_data_area",
           "fill_sheet", "fill_workbook", "verify_filled"]
