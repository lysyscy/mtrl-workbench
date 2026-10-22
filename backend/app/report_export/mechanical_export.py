"""Saved mechanical snapshots in the approved four-sheet layout, with cached formulas."""
from copy import copy
from decimal import Decimal, InvalidOperation
from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook
from openpyxl.styles import Border, Side
from openpyxl.utils import get_column_letter

from app.core.errors import MustReportError
from app.report_export.resources import MECHANICAL_TEMPLATE, require_template
from app.report_export.save import Formula, save_workbook, set_formula

MONEY_FIELDS = {'amount', 'rate', 'net', 'qty', 'plan', 'used'}
ZERO_FIELDS = ('device', 'spec', 'unit', 'qty', 'rate', 'amount', 'tax', 'net',
               'date', 'use', 'team', 'reviewer', 'vendor', 'plate', 'memo')
MONTHLY_FIELDS = ('device', 'spec', 'unit', 'plan', 'used', 'rate', 'amount',
                  'date', 'vendor', 'reviewer', 'tax', 'memo')
CATEGORY_COLUMNS = {'zero': 'mech_2', 'monthly': 'mech_3', 'internal': 'mech_1'}


def numeric(value):
    if value in (None, ''):
        return None
    try:
        result = Decimal(str(value))
        if not result.is_finite():
            raise InvalidOperation
        return result
    except (InvalidOperation, ValueError) as exc:
        raise MustReportError('报表含有无效金额，请回编辑页核对') from exc


def put(sheet, row, col, value, *, number=False):
    cell = sheet.cell(row, col)
    if number:
        cell.value = numeric(value)
    else:
        cell.value = '' if value is None else str(value)
        cell.data_type = 's'
    return cell


def append_row(sheet, row, values, styles, numeric_cols=()):
    for col, value in enumerate(values, 1):
        cell = put(sheet, row, col, value, number=col in numeric_cols)
        cell._style = copy(styles[col - 1])
    sheet.row_dimensions[row].height = 22.5
    sheet.row_dimensions[row].hidden = False
    sheet.row_dimensions[row].outlineLevel = 0


def amount(data, wh, ym, column):
    return next((r['values'].get(column) for r in data.get('annual_effective', data['effective'])
                 if r['warehouse'] == wh and r['ym'] == ym), None)


def history(data, wh):
    return next((h for h in data['histories'] if h['warehouse'] == wh), {})


def cumulative(data, wh, column, ym):
    """Consume the same confirmed cumulative used by the workbench."""
    category = {'mech_0': 'company', 'mech_1': 'internal', 'mech_2': 'zero', 'mech_3': 'monthly'}.get(column)
    summary = next((r for r in data.get('summary', []) if r['warehouse'] == wh and r['ym'] == ym), None)
    return numeric(summary['cumulative'].get(category)) if summary else None


def annual_values(data, wh, hist, source, category):
    """Map reliable historical columns to actual months, then overlay adopted month values."""
    cells = source.get('cells', {})
    if not cells:
        source_year = None
    else:
        source_year = source.get('source', {}).get('year') or str(hist.get('base_month') or '')[:4]
        if not str(source_year).isdigit() or len(str(source_year)) != 4:
            raise MustReportError('年度历史缺少年份依据，请补齐来源年份后再导出')
        source_year = int(source_year)
    values = {}
    if source_year:
        for month in range(1, 13):
            values[f'{source_year - 1}-{month:02d}'] = cells.get(get_column_letter(3 + month))
            values[f'{source_year}-{month:02d}'] = cells.get(get_column_letter(15 + month))
        # AH:BC are explicitly dated archive columns in the original company workbook.
        archive = [f'2023-{m:02d}' for m in range(2, 13)] + [f'2024-{m:02d}' for m in (1, 2, *range(4, 13))]
        for col, month in enumerate(archive, 34):
            if cells.get(get_column_letter(col)) is not None:
                values[month] = cells[get_column_letter(col)]
    column = CATEGORY_COLUMNS.get(category)
    if column:
        references = hist.get('reference_at', {}).get(data['ym'], hist.get('reference_months', {}))
        for months in (references, hist.get('original_months', {})):
            for month, fields in months.items():
                if values.get(month) is None and fields.get(column) is not None:
                    values[month] = fields[column]
        for row in data.get('annual_effective', data['effective']):
            if row['warehouse'] == wh and row['values'].get(column) is not None:
                values[row['ym']] = row['values'][column]
    return {month: numeric(value) for month, value in values.items()
            if month <= data['ym'] and value is not None}


def sum_formula(sheet, row, col, rows, *, empty_zero=False):
    rows = list(rows)
    cells = [sheet.cell(r, col) for r in rows]
    values = [c.value.result if isinstance(c.value, Formula) else numeric(c.value) for c in cells]
    result = sum((v for v in values if v is not None), Decimal(0))
    if not empty_zero and not any(v is not None for v in values):
        result = None
    if not rows:
        set_formula(sheet.cell(row, col), Formula('=0' if empty_zero else '=""', result))
        return result
    if rows and rows == list(range(rows[0], rows[-1] + 1)):
        refs = f'{get_column_letter(col)}{rows[0]}:{get_column_letter(col)}{rows[-1]}'
    else:
        refs = ','.join(c.coordinate for c in cells) or '0'
    expression = f'SUM({refs})'
    set_formula(sheet.cell(row, col), Formula(
        f'={expression}' if empty_zero else f'=IF(COUNT({refs})=0,"",{expression})', result))
    return result


def build(data):
    ym = data.get('ym', '')
    if len(ym) != 7 or not data.get('projects'):
        raise MustReportError('请选择报表月份和至少一个项目')
    wb = load_workbook(require_template(MECHANICAL_TEMPLATE))
    warnings = []
    styles = {i: {r: [copy(sheet.cell(r, c)._style) for c in range(1, width + 1)]
                  for r in rows} for i, (sheet, rows, width) in enumerate(zip(
                      wb, ((6, 7), (5, 6, 7), (5, 6, 7), range(1, 8)), (15, 17, 15, 35)))}
    for sheet in wb:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str) and '2026年9月' in cell.value:
                    cell.value = cell.value.replace('2026年9月', f'{ym[:4]}年{int(ym[5:])}月')
    for sheet, header in zip(wb.worksheets[:3], (5, 4, 4)):
        sheet.delete_rows(header + 1, sheet.max_row - header)
    summary = wb.worksheets[0]
    summary.row_dimensions[5].height = 33
    for i, project in enumerate(data['projects'], 6):
        wh = project['warehouse']
        values = [i - 5, project['name'], None, None,
                  cumulative(data, wh, 'mech_2', ym), cumulative(data, wh, 'mech_3', ym),
                  None, cumulative(data, wh, 'mech_1', ym), cumulative(data, wh, 'mech_0', ym),
                  amount(data, wh, ym, 'mech_2'), amount(data, wh, ym, 'mech_3'), None,
                  amount(data, wh, ym, 'mech_1'), amount(data, wh, ym, 'mech_0'), None]
        metadata = next((r for r in data.get('summary', []) if r['warehouse'] == wh and r['ym'] == ym), {})
        values[2:4] = [metadata.get('start_year'), metadata.get('start_month')]
        append_row(summary, i, values, styles[0][6], range(5, 15))
    total_row = 6 + len(data['projects'])
    append_row(summary, total_row, [None] * 15, styles[0][7])
    put(summary, total_row, 1, '合计')
    summary.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=2)
    for col in range(5, 15):
        sum_formula(summary, total_row, col, range(6, total_row))
    for sheet_index, category, fields, total_col, net_col in (
            (1, 'zero', ZERO_FIELDS, 7, 9), (2, 'monthly', MONTHLY_FIELDS, 8, 15)):
        sheet = wb.worksheets[sheet_index]
        row_no = 5
        groups = []
        for group_no, project in enumerate(data['projects'], 1):
            wh = project['warehouse']
            detail = next((r for r in data['details'] if r['warehouse'] == wh and r['category'] == category), None)
            rows = detail['rows'] if detail else []
            gross = detail['totals']['gross'] if detail else None
            net = detail['totals']['net'] if detail else None
            group_row = row_no
            groups.append(group_row)
            append_row(sheet, row_no, [None] * len(styles[sheet_index][5]), styles[sheet_index][5])
            put(sheet, row_no, 1, group_no, number=True)
            put(sheet, row_no, 2, project['name'])
            sheet.merge_cells(start_row=row_no, start_column=2, end_row=row_no, end_column=3)
            put(sheet, row_no, total_col, gross, number=True)
            put(sheet, row_no, net_col, net, number=True)
            row_no += 1
            for n, row in enumerate(rows, 1):
                values = [n] + [row.get(field) for field in fields]
                if category == 'monthly':
                    # M preserves the original report note; plate remains separate in the system.
                    values += [None, row.get('net')]
                numeric_cols = [1] + [i + 2 for i, field in enumerate(fields) if field in {'amount', 'rate', 'net'}]
                for i, field in enumerate(fields, 2):
                    if field in {'qty', 'plan', 'used'} and row.get(field) not in (None, ''):
                        try:
                            numeric(row[field])
                            numeric_cols.append(i)
                        except MustReportError:
                            pass
                if category == 'monthly':
                    numeric_cols += [15]
                append_row(sheet, row_no, values, styles[sheet_index][6], numeric_cols)
                if category == 'zero':
                    sheet.cell(row_no, 10).number_format = '@'
                row_no += 1
            for col in (total_col, net_col):
                sum_formula(sheet, group_row, col, range(group_row + 1, row_no),
                            empty_zero=bool(detail and not rows and numeric(detail['totals'].get('gross')) == 0))
            if not rows:
                warnings.append(f'{project["name"]}：{"零租" if category == "zero" else "包月"}明细未取得')
            if detail and detail.get('issues'):
                warnings.append(f'{project["name"]}：{len(detail["issues"])}行字段待补')
            effective = amount(data, wh, ym, CATEGORY_COLUMNS[category])
            if numeric(effective) != numeric(gross):
                warnings.append(f'{project["name"]}：{"零租" if category == "zero" else "包月"}明细合计与已采用金额不同')
            elif detail and effective is not None and sheet.cell(group_row, total_col).value.result == numeric(effective):
                reference = f"'{sheet.title}'!{get_column_letter(total_col)}{group_row}"
                set_formula(summary.cell(group_no + 5, 10 if category == 'zero' else 11),
                            Formula(f'={reference}', numeric(effective)))
        append_row(sheet, row_no, [None] * len(styles[sheet_index][7]), styles[sheet_index][7])
        put(sheet, row_no, 1, '合计')
        sheet.merge_cells(start_row=row_no, start_column=1, end_row=row_no, end_column=6)
        for col in (total_col, net_col):
            sum_formula(sheet, row_no, col, groups)
    annual = wb.worksheets[3]
    annual.page_setup.orientation = 'landscape'
    annual.page_setup.paperSize = '8'
    annual.row_dimensions[1].height = 33
    categories = [('zero', '当月零租费'), ('monthly', '当月包月费'),
                  ('weekly', '当月周材费'), ('internal', '当月内租费')]
    month_values = {}
    for project in data['projects']:
        wh, hist = project['warehouse'], history(data, project['warehouse'])
        for category, _label in categories:
            source = next((r for r in hist.get('annual_rows', []) if r['category'] == category), {})
            month_values[wh, category] = annual_values(data, wh, hist, source, category)
            opening = numeric(source.get('cells', {}).get('C'))
            source_year = source.get('source', {}).get('year') or str(hist.get('base_month') or '')[:4]
            if opening is not None:
                archive = sum((v for m, v in month_values[wh, category].items()
                               if m < f'{int(source_year) - 1}-01'), Decimal(0))
                if abs(opening - archive) >= Decimal('0.01'):
                    warnings.append(f'{project["name"]}：原往年合计尚未全部拆到可靠月份，请补齐历史来源')
    months = sorted({m for values in month_values.values() for m in values}) or [ym]
    sum_col = len(months) + 3
    width = sum_col + 5
    annual._cells.clear()
    annual.row_dimensions.clear()
    annual.column_dimensions.clear()
    headers = ['项目', '费用'] + [f'{m[:4]}年{int(m[5:])}月' for m in months] + ['合计', '总班组使用', 'i6', 'i8', '差额', '信息化消耗结算']
    def annual_styles(template_row):
        old = styles[3][template_row]
        return old[:2] + [old[2]] * len(months) + old[29:35]
    append_row(annual, 1, headers, annual_styles(1))
    annual.row_dimensions[1].height = 30
    for col in range(1, width + 1):
        annual.column_dimensions[get_column_letter(col)].width = 20 if col == 1 or col >= sum_col else 15 if col == 2 else 12
    annual.freeze_panes = 'C2'
    row_no = 2
    totals = []
    for project_index, project in enumerate(data['projects']):
        wh, hist = project['warehouse'], history(data, project['warehouse'])
        start = row_no
        for template_row, (category, label) in enumerate(categories, 2):
            append_row(annual, row_no, [None] * width, annual_styles(template_row))
            put(annual, row_no, 2, label)
            for col, month in enumerate(months, 3):
                put(annual, row_no, col, month_values[wh, category].get(month), number=True)
            known = cumulative(data, wh, CATEGORY_COLUMNS[category], ym) if category in CATEGORY_COLUMNS else None
            references = f'C{row_no}:{get_column_letter(sum_col - 1)}{row_no}'
            values = list(month_values[wh, category].values())
            result = sum(values, Decimal(0)) if values else Decimal(0) if known == 0 else None
            set_formula(annual.cell(row_no, sum_col), Formula(
                f'=SUM({references})' if known == 0 else f'=IF(COUNT({references})=0,"",SUM({references}))', result))
            if known is not None and (result is None or abs(known - result) >= Decimal('0.01')):
                warnings.append(f'{project["name"]}：{label}累计与可追溯月份合计不同，请补齐历史月份')
            elif known is not None:
                col = {'zero': 5, 'monthly': 6, 'internal': 8}.get(category)
                if col:
                    reference = f"'{annual.title}'!{get_column_letter(sum_col)}{row_no}"
                    set_formula(summary.cell(project_index + 6, col), Formula(f'={reference}', result))
            row_no += 1
        append_row(annual, row_no, [None] * width, annual_styles(6))
        put(annual, start, 1, project['name'])
        annual.merge_cells(start_row=start, start_column=1, end_row=row_no, end_column=1)
        put(annual, row_no, 2, '合计')
        for col in range(3, sum_col + 1):
            sum_formula(annual, row_no, col, range(start, row_no))
        snapshot = next((r for r in data['i8'] if r['warehouse'] == wh), {})
        usage, i6 = numeric(snapshot.get('totals', {}).get('usage')), numeric(hist.get('i6'))
        put(annual, row_no, sum_col + 2, i6, number=True)
        put(annual, row_no, sum_col + 3, usage, number=True)
        right = i6 + usage if i6 is not None and usage is not None else None
        right_refs = f'{get_column_letter(sum_col + 2)}{row_no}:{get_column_letter(sum_col + 3)}{row_no}'
        set_formula(annual.cell(row_no, sum_col + 1), Formula(f'=IF(COUNT({right_refs})<2,"",SUM({right_refs}))', right))
        left = annual.cell(row_no, sum_col).value.result
        left_ref, right_ref = f'{get_column_letter(sum_col)}{row_no}', f'{get_column_letter(sum_col + 1)}{row_no}'
        set_formula(annual.cell(row_no, sum_col + 4), Formula(
            f'=IF(COUNT({left_ref},{right_ref})<2,"",{left_ref}-{right_ref})',
            left - right if left is not None and right is not None else None))
        parts = [numeric(snapshot.get('totals', {}).get(k)) for k in ('project', 'team')]
        put(annual, row_no, sum_col + 5, sum(parts, Decimal(0)) if all(v is not None for v in parts) else None, number=True)
        totals.append(row_no)
        row_no += 1
    append_row(annual, row_no, [None] * width, annual_styles(7))
    put(annual, row_no, 1, '总计')
    annual.merge_cells(start_row=row_no, start_column=1, end_row=row_no, end_column=2)
    for col in range(3, width + 1):
        sum_formula(annual, row_no, col, totals)
    # Rebuild caches after summary rows have acquired their cross-sheet formulas.
    for col in range(5, 15):
        sum_formula(summary, total_row, col, range(6, total_row))
    border = Border(**{side: Side(style='thin', color='000000') for side in ('left', 'right', 'top', 'bottom')})
    for row in annual.iter_rows(min_col=3, max_col=width, max_row=row_no):
        for cell in row:
            cell.border = copy(border)
    for sheet, max_col in zip(wb, (15, 16, 15, width)):
        sheet.print_area = f'A1:{get_column_letter(max_col)}{sheet.max_row}'
    return wb, warnings


def preview(data):
    wb, warnings = build(data)
    sheets = []
    for sheet in wb.worksheets:
        merged_starts, hidden = {}, set()
        for item in sheet.merged_cells.ranges:
            merged_starts[(item.min_row, item.min_col)] = (item.max_row - item.min_row + 1, item.max_col - item.min_col + 1)
            hidden.update((r, c) for r in range(item.min_row, item.max_row + 1) for c in range(item.min_col, item.max_col + 1)
                          if (r, c) != (item.min_row, item.min_col))
        rows = []
        for row in sheet:
            cells = []
            for cell in row:
                if (cell.row, cell.column) in hidden:
                    continue
                rs, cs = merged_starts.get((cell.row, cell.column), (1, 1))
                value = cell.value
                if isinstance(value, Formula):
                    value = value.result
                cells.append({'value': '' if value is None else str(value), 'rowspan': rs, 'colspan': cs, 'column': cell.column})
            rows.append(cells)
        sheets.append({'name': sheet.title, 'rows': rows})
    wb.close()
    return {'ok': True, 'ym': data['ym'], 'projects': data['projects'], 'sheets': sheets, 'warnings': warnings}


def export(data, output_dir):
    wb, _ = build(data)
    return save_workbook(wb, Path(output_dir) / f'机械费报表-{data["ym"]}-{uuid4().hex[:12]}.xlsx')
