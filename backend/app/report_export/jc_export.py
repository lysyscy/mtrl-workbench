"""`MOD-11` 集采台账的 **Excel 导出**（`M5-9`）。

⚠️ **口径（用户 2026-09-20 定）**：

- **行** = **页面上当前看到的那些**（调用方传 id 列表，见 `mod11_jc.service.export_xlsx`）
  ⇒ 前端筛选口径与后端 **100% 一致**。
  ⭐ **为什么不传"筛选条件"让后端自己筛**：页面的筛选是**前端**算的（还含「合同金额为空/非空」
    这种 `build_list` 里没有的条件）⇒ 传条件就得在后端**再实现一遍** ⇒ **两套口径必打架**（本项目的旧账）。
- **列** = **全列**（**含**「贸易公司」那 3 列）—— 用户明确选 **(B)**：
  屏幕上的「隐藏列」**只是屏幕的事**，导出永远给人**完整的表**（Excel 里显隐/列宽是他自己的事）。
  ⭐ 所以这里是「**行跟筛选、列不跟隐藏**」—— 两句不矛盾。

⚠️ **不做"看起来像个表"的假合计**：只对**金额列**求和；材料名 / 供应商 / 备注那些**不合计**
  （对文本求和没有意义，硬凑一个数只会让人以为那是业务指标）。
⚠️ **「勾选 / 操作」两列不导出**（一个是屏幕控件、一个是按钮）。
"""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO

#: 导出的列（**顺序 = 页面表头顺序**去掉「勾选 / 操作」）⇒ 导出的表与屏幕**并排看能对上**。
#: 第二项 = 行字典里的键（`mod11_jc.service.build_list` 的输出直接可用）。
COLUMNS: tuple[tuple[str, str], ...] = (
    ("序号", "id"),
    ("合同日期", "date"),
    ("材料名", "material"),
    ("供应商", "supplier"),
    ("合同金额", "contract_amt"),
    ("发票金额", "inv_sum"),
    ("贸易公司名", "trade_co"),
    ("贸易公司合同金额", "trade_contract_amt"),
    ("贸易公司发票金额", "trade_inv_sum"),
    ("结算金额", "settle"),
    ("付款金额", "pay_sum"),
    ("欠款金额", "debt"),
    ("付款方式", "method"),
    ("备注", "remark"),
    ("三态", "tag"),
    ("合同存档情况", "arch"),
    ("存档号", "arch_no"),
    ("寄回日期", "ret_date"),
)

#: 金额列（**只对这些列合计**）—— 见模块头「不做假合计」。
MONEY_KEYS: frozenset[str] = frozenset({
    "contract_amt", "inv_sum", "trade_contract_amt", "trade_inv_sum",
    "settle", "pay_sum", "debt",
})

SHEET_NAME = "集采台账"


def _money(v) -> float | None:
    """金额 ⇒ `float`（⛔ 空的**留空**，不写 0 —— "没这个数"与"是 0"是两件事）。"""
    if v in (None, ""):
        return None
    try:
        return float(Decimal(str(v)))
    except Exception:                       # noqa: BLE001 —— 认不出就当空（⛔ 不抛，别为一行毁整表）
        return None


def build_xlsx(rows: list[dict]) -> bytes:
    """行字典列表 ⇒ 一本 xlsx 的**字节**（不落盘：调用方决定给谁）。

    ⚠️ 表头**加粗**＋**冻结首行**：几十行的表滚起来没有冻结行，人对不上哪列是哪列。
    ⚠️ 金额列写**数字**（⛔ 不写字符串）：这样用户在 Excel 里能直接求和/排序
       —— 写成 `"1234.56"` 那种"看着是数、其实是文本"的表，是本项目最讨厌的一类假象。
    """
    import openpyxl
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    ws.append([label for label, _ in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for r in rows or ():
        ws.append([
            (_money(r.get(key)) if key in MONEY_KEYS else r.get(key))
            for _, key in COLUMNS
        ])

    # 合计行（**只合金额列**）—— 放最后，并在首列写明是合计（⛔ 不让人误当成一行数据）
    if rows:
        total = ["" for _ in COLUMNS]
        total[0] = "合计"
        for i, (_, key) in enumerate(COLUMNS):
            if key in MONEY_KEYS:
                vals = [_money(r.get(key)) for r in rows]
                total[i] = float(sum(v for v in vals if v is not None))
        ws.append(total)
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)

    # 列宽：按内容给个**可读的**初值（⛔ 不写死窄宽 —— 那是"表挤"的另一半原因）
    for i, (label, _) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, min(28, len(label) * 2 + 4))

    ws.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


__all__ = ["COLUMNS", "MONEY_KEYS", "SHEET_NAME", "build_xlsx"]
