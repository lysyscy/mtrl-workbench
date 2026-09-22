"""**朴素导出**的共用件（`TAIL-04` · `REQ-R1` · §7.1）—— 「列表页导出 Excel」那一档。

## 这一档是什么

`TAIL-04` 的导出分**两档**（⛔ 别混做 —— 见 `report_export/__init__.py`）：

| 档 | 做法 | 谁在用 |
|---|---|---|
| **路径② 模板法** | 拿用户手作的成品 `xlsx` 当基底做 `zip` 字节替换（保留真透视表） | 收发结存 · 库存单据明细表 |
| **路径① 模板填充** | 读模板 → 认 `DATA_START` → 填数（`tpl_fill.py`） | 四个模板（焊材气体 / 逐月欠款台账 / 过程资料 / 活动大表） |
| ⭐ **朴素导出**（本件） | 算一个行列表 → `openpyxl` 从零写一本 xlsx | **列表页**的「导出 Excel（R1）」按钮 |

## 口径（§7.1 —— 每个调用方都必须遵守，本件不替它判）

- **行** = **页面上当前看到的那些**（调用方传 **id 列表**，**按传入顺序**出）
  ⇒ 前端筛选口径与后端 **100% 一致**。
  ⭐ **为什么不传"筛选条件"让后端自己筛**：页面的筛选是**前端**算的（还含"关键词模糊搜"
    这类后端没有的条件）⇒ 传条件就得在后端**再实现一遍** ⇒ **两套口径必打架**（本项目的旧账）。
- **列** = **页面表头去掉「勾选 / 操作」那一列**（那是控件，不是数据）⇒ 导出的表与屏幕**并排看能对上**。
- **取值** = **页面显示的那个值**（⛔ 不是原始外键 / 英文枚举）。
  ⚠️ **金额例外**（本项目既定口径，见 `jc_export` / `fl_export` / `pj_export`）：
     屏幕上的**千分位是排版**，Excel 里写**数字**才能求和/排序 ⇒ 金额列一律写 `float`。
     写成 `"1,234.56"` 那种"看着是数、其实是文本"的表，是本项目最讨厌的一类假象。

## ⚠️ 合计行（**可选的，但一旦要就必须与页面同口径**）

`total_keys` = **哪几列要合计**（⛔ **不是"哪些列是金额"**）—— 两者**故意分开**：
「已付款金额 / 剩余付款金额」（`MOD-04`）是金额、要写数字，但页面**不合计它们**
（合计只算「i8 流程金额 / 本月预付款金额」两列）⇒ 混成一个集合就会凭空多出两个合计格。

`total_rows` = **参与合计的行**：
- `None`（默认）⇒ 传入的**全部**行；
- ⭐ **给空序列 ⇒ 不写合计行** —— 这是 `MOD-03` 的真口径：页面合计**只算「非剔除」行**
  （`sumRows=list.filter(ntLive)`），切到「分类 = 剔除」时该视图全是剔除行 ⇒ **页面没有合计行**
  （不给一排 0，免得看着像坏了）⇒ 导出也**必须没有**（⛔ 凭空补一个 = 假值）。

## 与另外三件的关系（⚠️ 先说清，免得下一个人困惑）

`jc_export.py` / `fl_export.py` / `pj_export.py` **各自带一份同源实现**（它们是本件的**前身**，
逐页做的时候一页一份）。⛔ **本次没有回头去拆它们** —— 那三件已被各自的断言钉住，
拆分是**另一件事**（改了要重跑它们的判据，收益只是少几行重复）。
⇒ **规矩：新页一律走本件**；那三件**要拆另开一轮**，⛔ 别在本件里"顺手"改它们。
"""

from __future__ import annotations

from decimal import Decimal

from app.report_export.save import workbook_bytes      # ⭐ `TAIL-68`：序列化唯一出口

#: 合计行首格写的字（与页面 `sumRowOf` 的 `lead` 默认值**同一个词** —— 页面上就是「合计」）。
TOTAL_LEAD = "合计"


def money(v) -> float | None:
    """金额 ⇒ `float`；**空 ⇒ `None`**（⛔ 不写 0 —— "没这个数"与"是 0"是两件事）。

    ⚠️ 认不出的值**当空处理、不抛**：⛔ 不为一行脏数据毁掉整张表
       （真抛出的话用户拿到的是"导出失败"，而不是"这一格空着"——后者才是事实）。
    """
    if v in (None, ""):
        return None
    try:
        return float(Decimal(str(v)))
    except Exception:                       # noqa: BLE001 —— 见上
        return None


def build_xlsx(
    rows: list[dict],
    *,
    columns: tuple[tuple[str, str], ...],
    sheet_name: str,
    money_keys: frozenset[str] | set[str] = frozenset(),
    total_keys: frozenset[str] | set[str] = frozenset(),
    total_rows: list[dict] | None = None,
    seq_key: str | None = None,
    int_keys: frozenset[str] | set[str] = frozenset(),
    width_max: int = 32,
) -> bytes:
    """行字典列表 ⇒ 一本 xlsx 的**字节**（不落盘：调用方决定给谁）。

    ⚠️ 表头**加粗**＋**冻结首行**：几十行的表滚起来没有冻结行，人对不上哪列是哪列。
    ⚠️ `money_keys` 里的列写**数字**（⛔ 不写字符串）—— 见模块头「取值」那条。
    ⚠️ 合计行**放最后**并加粗，首格写明「合计」（⛔ 不让人把它误当成一行数据）。

    ⭐ `TAIL-62`（2026-09-21）：为收编 `jc_export` / `pj_export` **新增三个可选参数**
      （⛔ **默认值与旧行为完全一致** ⇒ 已经在用本件的调用方**零影响**）：
      · `seq_key` = 这一列写**行号**（1..n）—— 集采台账的「序号」列；
      · `int_keys` = 这几列写 `int`（认不出就原样）—— 项目档案的「排序」列；
      · `width_max` = 列宽上限（集采台账口径是 **28**，其余 32）⇒ ⛔ 别统一成 32，
        那会**改掉集采台账的导出形态**（本轮是"换实现"，不是"改口径"）。
    """
    import openpyxl
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    ws.append([label for label, _ in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    def _cell(r: dict, key: str, n: int):
        if seq_key and key == seq_key:
            return n                                  # ⭐ 序号列（= 屏幕上的行号）
        v = r.get(key)
        if key in money_keys:
            v = money(v)
        elif key in int_keys and v not in (None, ""):
            try:
                v = int(v)
            except (TypeError, ValueError):
                pass                                  # 认不出就原样（⛔ 不为一行脏数据毁整表）
        return "" if v is None else v

    for n, r in enumerate(rows or (), start=1):
        ws.append([_cell(r, key, n) for _, key in columns])

    # 合计行 —— `total_rows is None` ⇒ 全部行；给空序列 ⇒ **不写**（与页面同口径，见模块头）
    if total_keys and (rows or ()) :
        base = rows if total_rows is None else total_rows
        if base:
            total = ["" for _ in columns]
            total[0] = TOTAL_LEAD
            for i, (_, key) in enumerate(columns):
                if key in total_keys:
                    vals = [money(r.get(key)) for r in base]
                    total[i] = float(sum(v for v in vals if v is not None))
            ws.append(total)
            for cell in ws[ws.max_row]:
                cell.font = Font(bold=True)

    # 列宽：按表头给个**可读的**初值（⛔ 不写死窄宽 —— 那是"表挤"的另一半原因）
    for i, (label, _) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, min(width_max, len(label) * 2 + 4))

    ws.freeze_panes = "A2"

    return workbook_bytes(wb)          # ⭐ `TAIL-68`：序列化也走唯一出口（全站不再有第二处 `wb.save(`）


__all__ = ["TOTAL_LEAD", "build_xlsx", "money"]
