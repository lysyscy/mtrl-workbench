"""集采打印版：项目分表，复用模板填充；字段取值与金额仍由集采域提供。"""
from __future__ import annotations

from copy import copy
import math
import re
import unicodedata

import openpyxl
from openpyxl.utils import get_column_letter

from app.core.errors import MustReportError
from app.core.money import sum_money, to_money
from app.mod11_jc.domain import SUM_KEYS
from app.report_export import resources, tpl_fill
from app.report_export.save import workbook_bytes

NORMAL_KEYS = ("project_seq", "date", "material", "supplier", "contract_amt", "inv_sum",
               "settle", "pay_sum", "debt", "remark")
TRADE_KEYS = NORMAL_KEYS[:3] + ("trade_co",) + NORMAL_KEYS[3:6] + (
    "trade_contract_amt", "trade_inv_sum") + NORMAL_KEYS[6:]
EXTRA_KEYS = ("contract_no", "tag", "arch", "arch_no", "ret_date")


def build_xlsx(sections: list[dict]) -> bytes:
    """每个 section = display、trade、完整纳入 rows；返回可打印的工作簿。"""
    if not sections:
        raise MustReportError("请先选择要导出的项目")
    wb = openpyxl.load_workbook(resources.require_template(resources.JC_PRINT_TEMPLATE))
    templates = list(wb.worksheets)
    plans = []
    try:
        for section in sections:
            keys = TRADE_KEYS if section["trade"] else NORMAL_KEYS
            ws = wb.copy_worksheet(wb["有贸易公司" if section["trade"] else "无贸易公司"])
            name = re.sub(r"[\x00-\x1f\\/*?:\[\]]", "_", section["display"]).strip("'")
            name = (name or "项目")[:31]
            candidate, suffix = name, 1
            while candidate.casefold() in {s.title.casefold() for s in wb}:
                ending = f"({suffix})"
                candidate = name[:31-len(ending)] + ending
                suffix += 1
            ws.title = candidate
            ws["A1"] = section["display"] + "集中采购付款清单"
            ws["A1"].data_type = "s"
            columns = {key: get_column_letter(n) for n, key in enumerate(keys + EXTRA_KEYS, 1)}
            data = []
            for row in section["rows"]:
                cells = {columns[k]: (to_money(row[k]) if k in SUM_KEYS else row[k])
                         for k in columns if row.get(k) not in (None, "")}
                data.append((row["project_seq"], cells))
            totals = {columns[k]: sum_money(r.get(k) or 0 for r in section["rows"])
                      for k in keys if k in SUM_KEYS}
            plan = tpl_fill.fill_sheet(ws, {"rows": data, "tail": {0: ("合计", totals)}})
            plans.append(plan)
            for r in range(plan["data_start"], plan["tail_start"]):
                height = ws.row_dimensions[r].height or 15
                for key, col in columns.items():
                    cell = ws[f"{col}{r}"]
                    if key not in SUM_KEYS and cell.value is not None:
                        cell.data_type = "s" if isinstance(cell.value, str) else cell.data_type
                    if key in keys and isinstance(cell.value, str) and cell.alignment.wrap_text:
                        # shortcut: 行高按宋体字宽保守估算，更换字体或出现截字时改用字体实测。
                        width = max(1, ws.column_dimensions[col].width * 0.85)
                        lines = sum(max(1, math.ceil(sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1
                                                       for ch in line) / width))
                                    for line in cell.value.split("\n"))
                        height = max(height, lines * (cell.font.sz or 10) * 1.5)
                ws.row_dimensions[r].height = height
            source = wb["有贸易公司" if section["trade"] else "无贸易公司"]
            ws.sheet_properties.pageSetUpPr = copy(source.sheet_properties.pageSetUpPr)
            ws.HeaderFooter = copy(source.HeaderFooter)
            ws.print_title_rows = "1:3"
            ws.print_area = f"A1:{get_column_letter(len(keys))}{plan['last_row']}"
            ws.freeze_panes = "D4"
        for ws in templates:
            wb.remove(ws)
        blob = workbook_bytes(wb)
    finally:
        wb.close()
    problems = tpl_fill.verify_filled(blob, {"sheets": plans})
    if problems:
        raise MustReportError("打印台账核对未通过，请联系维护人：" + "；".join(problems[:3]))
    return blob
