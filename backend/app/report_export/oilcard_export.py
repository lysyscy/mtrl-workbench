"""全量油卡充值台账；取数与金额合计归业务服务，序列化沿共用出口。"""
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from app.report_export.save import workbook_bytes

COLUMNS = (('记录编号', 'id'), ('日期', 'date_text'), ('类型', 'kind_label'),
           ('项目', 'project'), ('用卡账户', 'account'), ('实体卡号', 'card'),
           ('金额', 'amount'), ('备注', 'remark'), ('关联原充值', 'related_entry_id'), ('转账关联编号', 'transfer_id'))


def build(snapshot):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '充值台账'
    sheet.append([label for label, _ in COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center')
    for row_index, row in enumerate(snapshot['rows'], 2):
        for column, (_, key) in enumerate(COLUMNS, 1):
            value = row[key]
            cell = sheet.cell(row_index, column)
            if key == 'amount':
                cell.value = Decimal(value)
                cell.number_format = '#,##0.00'
            elif key in ('id', 'related_entry_id'):
                cell.value = value
            else:
                cell.value = value
                cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
    total_row = sheet.max_row + 1
    sheet.cell(total_row, 1, '合计')
    sheet.cell(total_row, 7, Decimal(snapshot['total_amount'])).number_format = '#,##0.00'
    for cell in sheet[total_row]:
        cell.font = Font(bold=True)
    for column, width in enumerate((12, 22, 10, 28, 24, 30, 18, 40, 16, 40), 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:J{max(1, total_row - 1)}'
    return workbook_bytes(workbook)


def build_invoices(snapshot):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '发票台账'
    sheet.append(['记录编号', '开票日期', '发票号', '项目', '用卡账户', '价税合计', '已登记金额', '备注', '待核事项'])
    for row in snapshot['rows']:
        registered = Decimal(row['amount']) if row['status'] == 'registered' else Decimal('0.00')
        issues = [row['reason'], *(o['reason'] for o in row['originals'] if o['outcome'] == 'conflict')]
        sheet.append([row['id'], row['invoice_date'], row['invoice_no'], row['project'], row['account'],
            Decimal(row['amount']) if row['amount'] is not None else None, registered, row['remark'], '；'.join(x for x in issues if x)])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = 's'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        for column in (6, 7):
            sheet.cell(sheet.max_row, column).number_format = '#,##0.00'
    sheet.append(['合计', None, None, None, None, None, Decimal(snapshot['total_amount'])])
    sheet.cell(sheet.max_row, 7).number_format = '#,##0.00'
    for index in (1, sheet.max_row):
        for cell in sheet[index]:
            cell.font = Font(bold=True)
    for column, width in enumerate((12, 16, 28, 28, 24, 18, 18, 40, 60), 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:I{max(1, sheet.max_row - 1)}'
    return workbook_bytes(workbook)
