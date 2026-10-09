"""从年度汇总及两张明细表读取机械费；选定单月，缺失金额保留原值。"""

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
import re

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from app.core.errors import MustReportError
from app.core.money import amount_text, require_non_negative, to_money

LABELS = {"mech_1": "内租设备", "mech_2": "外租零租含税", "mech_3": "外租包月含税",
          "mech_4": "外租零租不含税", "mech_5": "外租包月不含税"}
HEADERS = ["核算月", "报表项目名", *LABELS.values()]
SUMMARY_FIELDS = {"当月内租费": "mech_1", "当月零租费": "mech_2", "当月包月费": "mech_3"}


def _month_columns(sheet):
    """日期头定年份；已确认的历史小数头按从左到右月份解释重复的.1。"""
    result = {}
    dated = []
    for cell in sheet[1]:
        value = cell.value
        dt = value if isinstance(value, (date, datetime)) else (
            from_excel(value) if isinstance(value, (int, float)) and 40000 < value < 60000 else None)
        if dt:
            result[cell.column] = dt.strftime("%Y-%m")
            dated.append(dt.year)
    if not dated:
        raise MustReportError("年度汇总缺少年月表头，请核对后重新解析")
    previous = {}
    for cell in sheet[1]:
        value = str(cell.value or "").strip()
        month = re.fullmatch(r"(\d{1,2})月", value)
        history = re.fullmatch(r"(20\d{2})\.(\d{1,2})", value)
        if month:
            result[cell.column] = f"{min(dated)-1:04d}-{int(month[1]):02d}"
        elif history:
            year, number = int(history[1]), int(history[2])
            if number == 1 and previous.get(year, 0) >= 9:
                number = 10
            if not 1 <= number <= 12 or number <= previous.get(year, 0):
                raise MustReportError("年度汇总历史月份顺序认不出，请核对表头")
            result[cell.column] = f"{year:04d}-{number:02d}"
            previous[year] = number
    return result


def _report_month(sheet):
    months = {f"{m[1]}-{int(m[2]):02d}" for row in sheet.iter_rows(min_row=1, max_row=4)
              for cell in row if isinstance(cell.value, str)
              for m in [re.search(r"(20\d{2})年\s*(\d{1,2})月", cell.value)] if m}
    if len(months) != 1:
        raise MustReportError(f"「{sheet.title}」报表月份认不出，请核对后重新解析")
    return next(iter(months))


def _value(sheet, formulas, row, column):
    value = sheet.cell(row, column).value
    if value is None and formulas.cell(row, column).data_type == "f":
        raise MustReportError(f"「{sheet.title}」第{row}行的公式没有计算结果，请用表格软件计算并保存后重试")
    if value is None or isinstance(value, str) and not value.strip():
        return None
    try:
        # 先经金额出口校验，再保留明细精度相加；只在项目合计时舍入，避免逐行舍入丢分。
        require_non_negative(str(value), field=f"{sheet.title}第{row}行")
        return Decimal(amount_text(value))
    except ValueError as exc:
        raise MustReportError(f"「{sheet.title}」第{row}行金额不正确，请修正后重新解析") from exc


def read_months(data, selected):
    """读指定年月；历史脚本与日常解析共用，历史不含税不倒算。"""
    selected = set(selected)
    if not selected or any(not re.fullmatch(r"20\d{2}-(0[1-9]|1[0-2])", ym) for ym in selected):
        raise MustReportError("请先选择导入月份，格式为YYYY-MM")
    wb = load_workbook(BytesIO(data), data_only=True)
    formulas = load_workbook(BytesIO(data), data_only=False)
    try:
        names = ["年度汇总", "零租设备月报表", "包月外租设备月报"]
        if any(name not in wb.sheetnames for name in names):
            raise MustReportError("请上传含年度汇总、零租设备月报表和包月外租设备月报的机械报表")
        annual = wb["年度汇总"]
        columns = _month_columns(annual)
        missing = selected - set(columns.values())
        if missing:
            raise MustReportError("年度汇总没有所选月份：" + "、".join(sorted(missing)))
        records = {}
        for row in range(2, annual.max_row + 1):
            name = str(annual.cell(row, 1).value or "").strip()
            key = SUMMARY_FIELDS.get(str(annual.cell(row, 2).value or "").strip())
            if not key or not name or name in ("总计", "合计"):
                continue
            for column, ym in columns.items():
                if ym not in selected:
                    continue
                record = records.setdefault((ym, name), {"核算月": ym, "报表项目名": name})
                value = _value(annual, formulas[annual.title], row, column)
                if value is not None:
                    label = LABELS[key]
                    if label in record:
                        raise MustReportError(f"年度汇总「{name}」同月同项出现重复行，请核对")
                    record[label] = str(to_money(value))
        for sheet_name, column, key in [(names[1], 9, "mech_4"), (names[2], 15, "mech_5")]:
            sheet = wb[sheet_name]
            ym = _report_month(sheet)
            if ym not in selected:
                continue
            current = ""
            sums = {}
            for row in range(5, sheet.max_row + 1):
                serial = sheet.cell(row, 1).value
                name = str(sheet.cell(row, 2).value or "").strip()
                if isinstance(serial, str) and re.fullmatch(r"[一二三四五六七八九十百]+", serial.strip()):
                    current = name
                    records.setdefault((ym, current), {"核算月": ym, "报表项目名": current})
                    continue  # 项目分组的汇总值不与明细重复相加。
                if any(word in str(label or "") for label in (serial, name)
                       for word in ("合计", "总计", "小计")):
                    continue
                if current:
                    value = _value(sheet, formulas[sheet_name], row, column)
                    if value is not None:
                        sums[current] = sums.get(current, Decimal(0)) + value
            for name, value in sums.items():
                records[(ym, name)][LABELS[key]] = str(to_money(value))
        return list(records.values())
    finally:
        wb.close()
        formulas.close()


def available_months(data):
    """历史操作脚本取年度汇总实际月份；累计列不作为月份。"""
    wb = load_workbook(BytesIO(data), data_only=True)
    try:
        return sorted(set(_month_columns(wb["年度汇总"]).values()))
    finally:
        wb.close()


def read_file(data, ym):
    """向通用导入返回该月规范表；源表映射固定，界面仍复用通用预览与结果。"""
    records = read_months(data, [ym])
    return {"sheet": "机械费报表", "headers": HEADERS,
            "grid_rows": [[r.get(label) for label in HEADERS] for r in records],
            "first_row": 1, "sheets": ["年度汇总", "零租设备月报表", "包月外租设备月报"]}
