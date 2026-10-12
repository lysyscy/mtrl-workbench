"""Excel 工作簿机械读取：两个引擎的遍历、合并坐标与资源释放归一处。

格式选择、单元格转换、选表与异常解释由调用者保留；这里不识别业务字段。
xlrd 沿用默认读取模式，不顺便开启 formatting_info（会改变现有网格）。
"""

from __future__ import annotations

import io


class _Workbook:
    """内部工作簿；仅由 open_workbook 构造，并作为上下文使用。"""

    def __init__(self, book, kind: str):
        self._book = book
        self._kind = kind

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        if self._kind == "xls":
            self._book.release_resources()
        else:
            self._book.close()

    @property
    def names(self) -> list[str]:
        """工作表名按文件顺序返回，不作选表。"""
        if self._kind == "xls":
            return [self._book.sheet_by_index(i).name for i in range(self._book.nsheets)]
        return list(self._book.sheetnames)

    def grid(self, sheet: str | None, *, xls_cell, xlsx_value) -> tuple[list[list], list[tuple], str]:
        """网格、0基闭区间合并块、表名；两种格式的原有值转换由调用者提供。

        xls_cell 接收 xlrd 单元格与日期模式，xlsx_value 接收原生值；不替调用者统一
        日期、布尔、数字或空值。缺表的引擎异常原样抛出，空表原样返回。
        """
        if self._kind == "xls":
            sh = self._book.sheet_by_name(sheet) if sheet is not None else self._book.sheet_by_index(0)
            grid = [[xls_cell(sh.cell(r, c), self._book.datemode) for c in range(sh.ncols)]
                    for r in range(sh.nrows)]
            merges = [(r1, r2 - 1, c1, c2 - 1) for r1, r2, c1, c2 in sh.merged_cells]
            return grid, merges, sh.name
        ws = self._book[sheet] if sheet is not None else self._book.worksheets[0]
        grid = [[xlsx_value(v) for v in row] for row in ws.iter_rows(values_only=True)]
        merges = ([(r.min_row - 1, r.max_row - 1, r.min_col - 1, r.max_col - 1)
                   for r in ws.merged_cells.ranges] if getattr(ws, "merged_cells", None) else [])
        return grid, merges, ws.title


def open_workbook(data: bytes, kind: str, *, read_only: bool = False) -> _Workbook:
    """只开书；kind/read_only由调用者决定，加载异常在调用者原出口解释。"""
    if kind == "xls":
        import xlrd

        return _Workbook(xlrd.open_workbook(file_contents=data), kind)
    if kind == "xlsx":
        from openpyxl import load_workbook

        return _Workbook(load_workbook(io.BytesIO(data), data_only=True, read_only=read_only), kind)
    raise ValueError(f"Unknown workbook kind: {kind}")


def fill_merges(grid: list[list], merges: list[tuple], *, rows: int | None = None) -> list[list]:
    """仅铺真实合并块；原地更新，rows为展开行上限，None表示整表。"""
    limit = len(grid) if rows is None else min(rows, len(grid))
    for r1, r2, c1, c2 in merges:
        if r1 >= limit or c1 >= len(grid[r1]):
            continue
        value = grid[r1][c1]
        if value is None or str(value).strip() == "":
            continue
        for r in range(r1, min(r2 + 1, limit)):
            for c in range(c1, min(c2 + 1, len(grid[r]))):
                grid[r][c] = value
    return grid
