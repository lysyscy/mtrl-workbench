"""历史付款线索逐行只读抽取；口径见 TAIL-285 设计稿 §二、§4.2。"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'artifacts/hist_supplement'
MAIN = ('1月份', '2月份', '3月份', '2026年4月份', '2026年5月份', '2026年6月', '2026年7月', '2026年8月')
BILL = re.compile(r'^[A-Za-z]{2,}\d{4,}$|^[A-Za-z]{2,}-\d{4}')


def dump(path, value):
    """保存稳定 JSON。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str, sort_keys=True) + '\n', encoding='utf-8')


def csv_out(path, rows, fields):
    """保存含原文的 CSV。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction='ignore')
        w.writeheader()
        for row in rows:
            w.writerow({k: json.dumps(v, ensure_ascii=False, default=str) if isinstance(v, (dict, list)) else v for k, v in row.items()})


def money(v):
    """只接受明确数值，按分定量。"""
    if v is None or isinstance(v, (bool, datetime, date)):
        return None
    try:
        return Decimal(str(v).replace(',', '').strip()).quantize(Decimal('.01'))
    except InvalidOperation:
        return None


def txt(v):
    """保留源值的可读形态。"""
    return '' if v is None else v.isoformat() if isinstance(v, (date, datetime)) else str(v).strip()


def project_key(s):
    """列明项目简称的有限映射，原名始终保留在证据中。"""
    for token, key in (('灵宝','灵宝'), ('珠江','珠江'), ('慧能','慧能'), ('宜化','宜化'), ('普康','普康'), ('晋江','晋江'), ('龙源','龙源'), ('南平','邵武'), ('邵武','邵武'), ('泗阳','泗阳'), ('大唐','大唐'), ('雷州','大唐'), ('压储','内蒙储能'), ('内蒙储能','内蒙储能'), ('国风','国风'), ('泰石','泰石'), ('安彩','安彩')):
        if token in s:
            return key
    return re.sub(r'\s+', '', s)


def source_date(v):
    """费用/付款时间列标准化，保留Excel序列值的1900日期系统解释。"""
    s = txt(v)
    if not s:
        return ''
    try:
        if re.fullmatch(r'\d{5}(?:\.0+)?', s):
            return from_excel(float(s)).date().isoformat()
        if re.fullmatch(r'20\d{6}', s):
            return datetime.strptime(s, '%Y%m%d').date().isoformat()
        if re.match(r'^20\d{2}-\d{2}-\d{2}', s):
            return date.fromisoformat(s[:10]).isoformat()
        m = re.fullmatch(r'(20\d{2})[./年-](\d{1,2})[./月-](\d{1,2})日?', s)
        return date(*map(int, m.groups())).isoformat() if m else ''
    except (ValueError, OverflowError):
        return ''


def bare_date_evidence(records):
    """裸M.D仅在I列/同项目汇总日期有同日证据时确证，跨年仍待确认。"""
    summary_columns = {'4月汇总': ('B','D'), '汇总5': ('D','C'), '汇总6': ('C','B'), '汇总7': ('D','C'), '汇总8': ('E','D')}
    summary_dates = defaultdict(list)
    for r in records:
        if r['sheet'] in summary_columns:
            pc, dc = summary_columns[r['sheet']]
            day = source_date(r['raw'].get(dc))
            if day:
                summary_dates[(project_key(r['raw'].get(pc,'')), day)].append(dict(sheet=r['sheet'], line=r['line'], cell=dc, value=r['raw'].get(dc)))
    for r in records:
        if not r['month'] or r['pay_date']:
            continue
        short = re.fullmatch(r'(\d{1,2})[./月](\d{1,2})(?:日|号)?', r['annotation'])
        if not short:
            continue
        m, d = map(int, short.groups())
        if m > r['month']:
            continue
        try:
            day = date(2026,m,d).isoformat()
        except ValueError:
            continue
        evidence = list(summary_dates.get((project_key(r['project']),day),[]))
        i_date = source_date(r['raw'].get('I')) if r['annotation_col'] != 'I' else ''
        if i_date == day:
            evidence.append(dict(sheet=r['sheet'],line=r['line'],cell='I',value=r['raw']['I']))
        if evidence:
            r.update(pay_date=day,date_basis=r['annotation_col']+':裸M.D经I/汇总同日交叉确证')
            r['evidence']['date_crosscheck']=evidence
            if r['classification']=='待确认' and r['依据'][0].startswith(r['annotation_col']+':无可确证日'):
                if r['block'] and r['block']['suspected_shift']:
                    r['依据']=['区块E−F与G不等超过30%，疑似错位隔离']
                elif money(r['amount']) is not None and money(r['amount'])>0:
                    r.update(classification='候选',依据=['付款批注明确金额；日期交叉确证'])
                else:
                    r['依据']=['G须本行与前后月多源一致，待第二阶段核验']


def payment_date(v, month, col):
    """仅接受明确年月日或带付款动词的 M.D；裸 7.2 等留待确认。"""
    s = txt(v)
    if isinstance(v, (date, datetime)):
        d = v.date() if isinstance(v, datetime) else v
    else:
        full = re.search(r'(20\d{2})[./年-](\d{1,2})[./月-](\d{1,2})日?', s)
        compact = re.fullmatch(r'(20\d{2})(\d{2})(\d{2})', s)
        short = re.search(r'(?<!\d)(\d{1,2})[.月/](\d{1,2})日?\s*(?:付|付款)', s)
        chinese_day = re.fullmatch(r'(\d{1,2})月(\d{1,2})(?:日|号)', s)
        try:
            if compact:
                d = date(*map(int, compact.groups()))
            elif full:
                d = date(*map(int, full.groups()))
            elif short or chinese_day:
                m, day = map(int, (short or chinese_day).groups())
                if m > month:
                    return '', col + ':跨年/后月日期待确认'
                d = date(2026, m, day)
            else:
                return '', col + ':无可确证日/裸日期歧义'
        except ValueError:
            return '', col + ':无效日期'
    if d > date(2026, 10, 11):
        return '', col + ':未来日期'
    return d.isoformat(), col + ':' + s


def explicit_amount(v):
    """从付款批注明确的付金额取数，不把日期当金额。"""
    m = re.search(r'付(?:款)?\s*([\d,]+(?:\.\d+)?)\s*(万|元)', txt(v))
    if not m:
        return None
    return money(Decimal(m.group(1).replace(',', '')) * (10000 if m.group(2) == '万' else 1))


def main():
    """抽取全部源行并保存互斥分类与证据。"""
    p = argparse.ArgumentParser()
    p.add_argument('--source', type=Path, required=True)
    args = p.parse_args()
    source = args.source.resolve()
    wb = load_workbook(source, read_only=True, data_only=True)
    formulas = load_workbook(source, read_only=True, data_only=False)
    records, blocks, mappings = [], [], []
    for sh in wb.sheetnames:
        ws = wb[sh]
        values = list(ws.iter_rows(values_only=True))
        forms = list(formulas[sh].iter_rows(values_only=True))
        month = MAIN.index(sh) + 1 if sh in MAIN else 0
        titles = [(i, txt(r[0])) for i, r in enumerate(values, 1) if '付款计划' in ' '.join(txt(v) for v in r) and '外租机械' not in txt(r[0])]
        if month and (not titles or titles[0][0] > 1):
            titles.insert(0, (1, '无标题区间（项目待确认）'))
        local_blocks = []
        for k, (start, title) in enumerate(titles):
            end = titles[k + 1][0] - 1 if k + 1 < len(titles) else len(values)
            project = title.split('项目')[0].strip()
            eligible = [r for r in values[start:end] if BILL.match(txt(r[2])) and money(r[4]) is not None and money(r[6]) is not None]
            shifted = sum(money(r[4]) - (money(r[5]) or Decimal(0)) != money(r[6]) for r in eligible)
            block = dict(sheet=sh, start=start, end=end, title=title, project=project, numeric_rows=len(eligible), inconsistent_rows=shifted, suspected_shift=bool(eligible and shifted / len(eligible) > .30))
            local_blocks.append(block)
            blocks.append(block)
        if month:
            mappings.append(dict(sheet=sh, bill_no='C', payee='D', expense_date='B', flow_amount='E', historical_paid='F', planned_amount='G', payment_time='I', method='J', remark='K', annotation='I' if month == 3 else 'L', anomaly='无表头' if month == 1 else '首区间无标题' if month == 3 else '标题多格' if month == 6 else ''))
        for i, r in enumerate(values, 1):
            raw = {chr(65 + c): txt(v) for c, v in enumerate(r)}
            raw_form = {chr(65 + c): txt(v) for c, v in enumerate(forms[i-1]) if txt(v).startswith('=') or txt(v).startswith('#')}
            block = next((b for b in local_blocks if b['start'] <= i <= b['end']), None)
            col = 'I' if month == 3 else 'L'
            annotation = raw.get(col, '')
            rec = dict(sheet=sh, line=i, month=month, block=block, project=block['project'] if block else '', bill_no=raw.get('C', ''), payee=raw.get('D', ''), expense_date=source_date(raw.get('B')), raw=raw, formulas=raw_form, annotation=annotation, annotation_col=col, classification='佐证', 依据=[], evidence={'F': raw.get('F', ''), 'expense_date_raw':raw.get('B','')})
            rec['pay_date'], rec['date_basis'] = payment_date(r[8] if month == 3 else r[11] if len(r) > 11 else None, month, col)
            explicit = explicit_amount(annotation)
            rec['amount'] = str(explicit) if explicit is not None else ''
            rec['amount_basis'] = col + ':显式付款金额' if explicit is not None else ''
            structural = ' '.join(raw.values())
            if not month:
                cls, why = '排除', '汇总表仅作佐证'
            elif month == 8 and block and any(x in block['title'] for x in ('新疆慧能', '内蒙古宜化')):
                cls, why = '排除', '用户指定8月已手录区块硬排除'
            elif not any(raw.values()):
                cls, why = '排除', '空白行'
            elif i in [x[0] for x in titles] and '付款计划' in structural:
                cls, why = '排除', '区块标题'
            elif '小计' in raw.get('C', '') or '合计' in raw.get('C', '') or '#REF!' in structural or '#REF!' in ' '.join(raw_form.values()):
                cls, why = '排除', '小计/合计/损坏公式'
            elif '外租机械' in raw.get('C', '') or ('外租机械费' == raw.get('A', '') and not raw.get('C', '') and not raw.get('D', '')):
                cls, why = '排除', '机械段说明行'
            elif '科目' == raw.get('A', '') or '收款单位' == raw.get('D', ''):
                cls, why = '排除', '表头'
            elif '退回' in structural or any((money(r[c]) is not None and money(r[c]) < 0) for c in (4, 5, 6)):
                cls, why = '排除', '负数/退回'
            elif not annotation:
                cls, why = '佐证', '无实际付款批注；F仅佐证/未付计划'
            elif not BILL.match(rec['bill_no']):
                cls, why = '待确认', '无可识别编号'
            elif not rec['pay_date']:
                cls, why = '待确认', rec['date_basis']
            elif block and block['suspected_shift']:
                cls, why = '待确认', '区块E−F与G不等超过30%，疑似错位隔离'
            elif explicit is not None and explicit > 0:
                cls, why = '候选', '付款批注明确日期及金额'
            else:
                cls, why = '待确认', 'G须本行与前后月多源一致，待第二阶段核验'
            rec['classification'], rec['依据'] = cls, [why]
            records.append(rec)
    wb.close()
    formulas.close()
    bare_date_evidence(records)
    # 同编号+费用日期，取邻月同项；只在至少一个邻月演变能确证且其余证据无冲突时放行。
    for rec in records:
        if rec['依据'] != ['G须本行与前后月多源一致，待第二阶段核验']:
            continue
        e, f, g = (money(rec['raw'].get(c)) for c in ('E', 'F', 'G'))
        f = f or Decimal(0)
        adjacent = [x for x in records if x['bill_no'] == rec['bill_no'] and x['expense_date'] == rec['expense_date'] and x['payee']==rec['payee'] and project_key(x['project'])==project_key(rec['project']) and abs(x['month'] - rec['month']) == 1]
        proof = []
        for x in adjacent:
            xe, xf, xg = (money(x['raw'].get(c)) for c in ('E', 'F', 'G'))
            xf = xf or Decimal(0)
            if xe == e and xg is not None and e is not None:
                delta = f - xf if x['month'] < rec['month'] else xf - f
                expected = xg if x['month'] < rec['month'] else g
                consistent = (delta == 0 and xg == g) or (delta == expected and delta >= 0)
                proof.append(dict(sheet=x['sheet'], line=x['line'], E=str(xe), F=str(xf), G=str(xg), consistent=consistent))
            else:
                proof.append(dict(sheet=x['sheet'],line=x['line'], E=str(xe),F=str(xf),G=str(xg),consistent=False,reason='邻月同项金额缺失或E变动，不能自动确证'))
        rec['evidence']['adjacent'] = proof
        rec['evidence']['blank_F_as_excel_zero'] = not rec['raw'].get('F')
        if e is not None and g is not None and g > 0 and e - f == g and proof and all(x['consistent'] for x in proof):
            rec.update(classification='候选', amount=str(g), amount_basis='G:E−F相等且邻月同编号同费用日期演变一致', 依据=['G多源一致；邻月引用见evidence.adjacent'])
        else:
            rec['依据'] = ['金额无可确证多源依据（本行不等/邻月缺失/演变不一致）']
    out = OUT / 'extract'
    out.mkdir(parents=True, exist_ok=True)
    (out / '抽取明细.jsonl').write_text(''.join(json.dumps(x, ensure_ascii=False, default=str) + '\n' for x in records), encoding='utf-8')
    fields = ['sheet', 'line', 'project', 'bill_no', 'classification', '依据', 'raw', 'formulas']
    for cls in ('排除', '待确认'):
        csv_out(out / (cls + '清单.csv'), [x for x in records if x['classification'] == cls], fields)
    csv_out(out / '列映射表.csv', mappings, list(mappings[0]))
    csv_out(out / '区块边界表.csv', blocks, list(blocks[0]))
    stats = dict(source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(), total_rows=len(records), classes=dict(Counter(x['classification'] for x in records)), sheets={sh: dict(Counter(x['classification'] for x in records if x['sheet'] == sh)) for sh in wb.sheetnames})
    assert sum(stats['classes'].values()) == len(records)
    dump(out / '抽取统计.json', stats)
    print(json.dumps(stats, ensure_ascii=False))


if __name__ == '__main__':
    main()
