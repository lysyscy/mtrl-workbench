# -*- coding: utf-8 -*-
"""`MOD-07` 提取成品「结余」列 —— **真机端到端**（连服务 + 连库）。

跑什么（**写 → 导出 → 读成品 → 清空 → 再导 → 复原**）：
  ① 读 `/api/m07/extract` 找**真实有数据的月 / 仓库**（不写死样本，库里换数据也不用改脚本）；
  ② `POST /api/m07/extract/balance` 录一个结余 ⇒ 回执 `ok=True`；
  ③ 非法输入（`abc` / 负数）⇒ **被拒且原值不变**；
  ④ 导出成品 → 用 `openpyxl` 读回该格 = 刚录的数（**成品里真的带上了**）；
  ⑤ 锁行清理本次结余 ⇒ 再导出 ⇒ 该格**如实留空**（⛔ 不是 0）；
  ⑥ **直查库**：目标结余键回到基线（跑完不留痕迹）。

⚠️ 本脚本**会写库**，只能在独占维护窗口运行；异常时尝试恢复，恢复失败必须报红；`PASS/FAIL` 逐条打印，末尾给 rc（有 FAIL ⇒ 1）。

用法： backend/.venv/Scripts/python.exe tools/probe/_e2e_m07bal.py
"""
from __future__ import annotations

import io
import json
import sys
import urllib.error
import urllib.parse
import urllib.request
from decimal import Decimal
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
BACKEND = ROOT / "backend"
BASE = "http://127.0.0.1:8000"

# ⚠️ 必须在**导入 `app.*` 之前**把后端目录加进 `sys.path`（本脚本在仓根、不在 backend 下跑）。
sys.path.insert(0, str(BACKEND))

_ok = _bad = 0


def ck(name: str, got, want=True) -> None:
    global _ok, _bad
    good = (got == want)
    print(("PASS  " if good else "FAIL  ") + name + ("" if good else f"（实得 {got!r}，期望 {want!r}）"))
    _ok, _bad = _ok + good, _bad + (not good)


def get_json(path: str) -> dict:
    with urllib.request.urlopen(BASE + path, timeout=30) as r:
        return json.loads(r.read().decode("utf-8"))


def post_json(path: str, body: dict) -> dict:
    req = urllib.request.Request(
        BASE + path, data=json.dumps(body).encode("utf-8"),
        headers={"Content-Type": "application/json"}, method="POST")
    with urllib.request.urlopen(req, timeout=30) as r:
        return json.loads(r.read().decode("utf-8"))


def get_bytes(path: str) -> tuple[bytes, str]:
    with urllib.request.urlopen(BASE + path, timeout=120) as r:
        return r.read(), r.headers.get("content-type", "")


def pick_sample() -> tuple[str, str]:
    """**真实有数据的**（月, 焊材仓库）—— 不写死样本。"""
    ex = get_json("/api/m07/extract")
    by_kind = (ex.get("amounts") or {}).get("焊材") or {}
    if not by_kind:
        raise SystemExit("库里没有任何焊材提取数据 ⇒ 无法真机验证（先抓一期收发结存）")
    months = sorted({ym for m in by_kind.values() for ym in m})
    ym = months[-1]
    wh = sorted(by_kind)[0]
    return ym, wh


def sheet_rows(content: bytes, sheet: str) -> list[list]:
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    try:
        ws = wb[sheet]
        return [[ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
                for r in range(1, ws.max_row + 1)]
    finally:
        wb.close()


def balance_of(content: bytes, sheet: str, warehouse: str):
    """成品里该仓库那一行的「结余」格（B 列）。"""
    for row in sheet_rows(content, sheet)[1:]:
        if str(row[0] or "") == warehouse:
            return row[1]
    raise AssertionError(f"成品 [{sheet}] 里找不到仓库 `{warehouse}`")


_MISSING = object()


def balance_value(kind: str, warehouse: str):
    """读取一个精确的「类别 × 仓库」结余键；不存在与 0 分开。"""
    from sqlalchemy import select

    from app.core.db import session_scope
    from app.mod07_setl.models import WeldGasBalance

    with session_scope() as s:
        row = s.execute(
            select(WeldGasBalance).where(
                WeldGasBalance.kind == str(kind),
                WeldGasBalance.warehouse_name == str(warehouse),
            )
        ).scalars().first()
        return _MISSING if row is None else Decimal(str(row.amount))


def restore_balance_if_owned(*, kind: str, warehouse: str, expected) -> bool:
    """锁行比较后只删除本次新建结余，不覆盖并发值。"""
    from sqlalchemy import select
    from app.core.db import session_scope
    from app.mod07_setl.models import WeldGasBalance

    with session_scope() as s:
        row = s.execute(select(WeldGasBalance).where(
            WeldGasBalance.kind == kind, WeldGasBalance.warehouse_name == warehouse,
        ).with_for_update()).scalars().first()
        if row is None:
            return False
        if Decimal(str(row.amount)) != expected:
            raise RuntimeError(f"结余键 {kind}/{warehouse} 已被其他值改动，停止恢复以免覆盖")
        s.delete(row)
    return True


def main() -> int:
    from app.report_export.weldgas_export import SHEET_OF_KIND  # noqa: E402

    ym, wh = pick_sample()
    sheet = SHEET_OF_KIND["焊材"]
    url = "/api/m07/extract/export?" + urllib.parse.urlencode({"status": "全部", "period": ym})
    print(f"样本（现读现取）：核算月 {ym} · 仓库 {wh} · 工作表 {sheet}\n")

    baseline = balance_value("焊材", wh)
    ck("起点：目标结余键没有记录（干净）", baseline is _MISSING, True)
    if baseline is not _MISSING:
        print("  ⛔ 目标键已有真实结余，立即停止，不覆盖")
        return 1
    if balance_value("焊材", wh) != baseline:
        ck("写入前目标结余键仍保持基线", False)
        return 1

    try:
        # ① 录一个
        r = post_json("/api/m07/extract/balance", {"kind": "焊材", "warehouse": wh, "raw": "1234.56"})
        ck("录入回执 ok", r.get("ok"), True)
        ck("录入回执 value", str(r.get("value")), "1234.56")
        ck("目标键写入", balance_value("焊材", wh), Decimal("1234.56"))

        # ② 读端点带上它
        ex = get_json("/api/m07/extract")
        ck("读端点带上该值", (ex.get("balances") or {}).get("焊材", {}).get(wh), "1234.56")

        # ③ 非法输入被拒、原值不变
        for bad in ("abc", "-1"):
            rb = post_json("/api/m07/extract/balance", {"kind": "焊材", "warehouse": wh, "raw": bad})
            ck(f"非法输入 {bad!r} 被拒", rb.get("ok"), False)
            ck(f"非法输入 {bad!r} 带中文原因", bool(rb.get("message")), True)
        ex = get_json("/api/m07/extract")
        ck("非法输入后原值不变", (ex.get("balances") or {}).get("焊材", {}).get(wh), "1234.56")

        # ④ 导出 → 成品里真的带上了
        content, ctype = get_bytes(url)
        ck("导出回的是 xlsx", "spreadsheetml" in ctype, True)
        got = balance_of(content, sheet, wh)
        ck("成品「结余」格 = 刚录的数", got is not None and abs(Decimal(str(got)) - Decimal("1234.56")) < Decimal("0.005"), True)

        # ⑤ 清空 → 再导 ⇒ 如实留空
        restored = restore_balance_if_owned(
            kind="焊材", warehouse=wh, expected=Decimal("1234.56")
        )
        ck("清理本次结余成功", restored, True)
        ck("清空后目标键不存在", balance_value("焊材", wh) is _MISSING, True)
        content2, _ = get_bytes(url)
        ck("清空后成品该格如实留空（⛔ 不是 0）", balance_of(content2, sheet, wh) in (None, ""), True)
    finally:
        try:
            restored = restore_balance_if_owned(
                kind="焊材", warehouse=wh, expected=Decimal("1234.56")
            )
            if restored:
                print("  ✅ 异常/中断清理：目标结余键已回到基线")
        except Exception as exc:  # noqa: BLE001
            ck("异常/中断清理成功", False, str(exc))

    print(f"\n===== {_ok} 通过 / {_bad} 失败 =====")
    return 1 if _bad else 0


if __name__ == "__main__":
    sys.exit(main())
