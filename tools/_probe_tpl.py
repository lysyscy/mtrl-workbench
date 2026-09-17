# -*- coding: utf-8 -*-
"""就地取证：用素材包的模板法管线跑我们**真机抓取**的两份导出件。

只读真机件、产物落临时目录；**不动工程仓、不动交付面**。
"""
import os
import sys
import traceback

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.stderr.reconfigure(encoding="utf-8", errors="replace")

import openpyxl

PKG = r"D:\工作\AI\PRD编写夹 - Workbuddy专用\PRD素材包-20260911"
INPUT = os.path.join(PKG, "input")
TPL_DIR = os.path.join(PKG, "data", "templates")
DIAG = r"D:\dev\mtrl-workbench\artifacts\_diag"
OUT = os.path.join(os.environ.get("TEMP", "."), "tpl-probe")
sys.path.insert(0, INPUT)


def peek(path, label):
    print("=" * 70)
    print("[%s] %s" % (label, os.path.basename(path)))
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    print("  sheets:", wb.sheetnames)
    ws = wb[wb.sheetnames[0]]
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        cells = [str(x)[:14] if x is not None else "" for x in row[:10]]
        print("  r%-2d %s" % (i, cells))
        if i >= 3:
            break
    wb.close()


def main():
    os.makedirs(OUT, exist_ok=True)

    for f, lab in (("inv_2026-09.xlsx", "真机·库存单据"), ("setl_2026-09.xlsx", "真机·收发结存")):
        peek(os.path.join(DIAG, f), lab)
    for f in ("模板_库存单据明细表.xlsx", "模板_收发结存表.xlsx"):
        peek(os.path.join(TPL_DIR, f), "模板")

    print("=" * 70)
    print("① 库存单据 · kcdj_process.main()")
    try:
        import kcdj_process

        out = kcdj_process.main(os.path.join(DIAG, "inv_2026-09.xlsx"),
                                os.path.join(OUT, "库存单据明细表"))
        print(">>> 产出:", out, os.path.getsize(out), "B")
    except Exception:
        print(">>> 失败")
        traceback.print_exc()

    print("=" * 70)
    print("② 收发结存 · sfjc_process.main()")
    try:
        import sfjc_process

        rc = sfjc_process.main(os.path.join(DIAG, "setl_2026-09.xlsx"),
                               os.path.join(OUT, "收发结存"), "2026年9月")
        print(">>> rc =", rc)
        for root, _, files in os.walk(os.path.join(OUT, "收发结存")):
            for x in files:
                p = os.path.join(root, x)
                print("    产出:", x, os.path.getsize(p), "B")
    except Exception:
        print(">>> 失败")
        traceback.print_exc()


if __name__ == "__main__":
    main()
