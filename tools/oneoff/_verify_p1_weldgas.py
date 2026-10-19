# -*- coding: utf-8 -*-
"""片 1（焊材 / 气体提取 · 路径① 模板填充）的**可复跑验证**。

做四件事：
1. **两条独立路径**对账：① 从 `setl_detail` 明细**自己**重算（本文件内的朴素实现）；
   ② 从**产物的显示区**读回 —— 两条必须逐格相同；
2. 结构性检查：数据区行数 = 所筛选项目数、尾块紧贴、**尾块以下一格不许有**；
3. **反向自检**（尺子有没有区分力）：把产物**故意弄坏**四种（残留 / 该空没空 / 留公式 / 少一行），
   守卫必须**逐条报红**；
4. 打印成品形态，供人肉复核（与 Excel 打开看到的对齐）。

⚠️ 用**库里的真数据**（读，不写）；⛔ 不连 i8。
"""
import sys
import tempfile
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, "backend")

import openpyxl

from app.core.db import session_scope
from app.mod07_setl import extract as X
from app.mod07_setl import service as M7
from app.mod07_setl.repository import load_rows
from app.report_export import tpl_fill, weldgas_export as W

OK, BAD = [], []


def chk(cond, msg):
    (OK if cond else BAD).append(msg)
    print("   %s %s" % ("✓" if cond else "✗", msg))


def _num(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    try:
        return Decimal(str(v))
    except Exception:
        return None


def naive_grid(details, kind, warehouses, year, report_ym):
    """① 朴素重算（**不用域内核**）：各月 = 收金额之和；`结余` 本轮**留空**（手填列，无来源）。"""
    months = ["%s-%02d" % (year, m) for m in range(1, 13)]
    cells = {}
    for d in details:
        if not X.match_kind(str(d.get("resource_name") or ""), kind, X.default_rules()):
            continue
        wh = str(d.get("warehouse") or "")
        if str(d.get("ym")) in months:
            key = (wh, str(d.get("ym")))
            cells[key] = cells.get(key, Decimal("0")) + Decimal(str(d.get("in") or 0))
    out = []
    for w in warehouses:
        # ⚠️ 「算出来是 0.00 ⇒ 该格留空」是**显示约定**（壳里同款：`v ? M(v) : '—'`）
        vals = [cells.get((w, m)) for m in months]
        out.append({"wh": w,
                    "cells": [(v or None) for v in vals],
                    "raw": vals,
                    "total": sum([v for v in vals if v is not None], Decimal("0"))})
    return {"months": months, "rows": out}


def read_display(path, sheet, first_row):
    """② 从产物读回显示区（数据行 + 尾块，读到尾块下一行为止）。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        ws = wb[sheet]
        out = []
        for r in range(first_row, ws.max_row + 1):
            out.append([ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)])
        return out, ws.max_row
    finally:
        wb.close()


def main():
    print("=" * 78)
    print("① 取数（库里的真数据）")
    print("=" * 78)
    with session_scope() as s:
        details = load_rows(s)
        state = {"status": X.STATUS_DEFAULT, "ext_period": "2026-09"}
        plan = M7.build_extract_export(s, M7.parse_extract_state(state))
        whs = {k: [r["warehouse"] for r in g["rows"]] for k, g in plan["grids"].items()}
    print("   报告月 =", plan["report_ym"], "| 年 =", plan["year"])
    for k, v in plan["grids"].items():
        print("   [%s] 行数 = %d  项目 = %s" % (k, len(v["rows"]), whs[k]))
    chk(plan["report_ym"] == "2026-09", "报告月 = 2026-09（单月 ⇒ 取该月）")
    chk(all(len(g["months"]) == 12 for g in plan["grids"].values()), "月份列 = 12 列（1月…12月）")

    print("\n" + "=" * 78)
    print("② 生成成品（落 scratch 目录）")
    print("=" * 78)
    scratch = Path(tempfile.mkdtemp(prefix="p1-"))
    out = W.build_xlsx(plan["grids"], scratch, plan["report_ym"])
    print("   成品 =", out)
    chk(Path(out).exists(), "成品已产出：%s" % Path(out).name)

    print("\n" + "=" * 78)
    print("③ 两条独立路径对账（朴素重算 vs 产物显示区）")
    print("=" * 78)
    for kind, sheet in W.SHEET_OF_KIND.items():
        grid = plan["grids"][kind]
        mine = naive_grid(details, kind, whs[kind], plan["year"], plan["report_ym"])
        disp, max_row = read_display(out, sheet, 2)
        # 数据行 + 尾块；尾块 = 最后一行
        data_rows, tail_row = disp[:-1], disp[-1]
        chk(len(data_rows) == len(mine["rows"]),
            "[%s] 数据行数 %d = 项目数 %d" % (sheet, len(data_rows), len(mine["rows"])))
        bad = 0
        for i, (dr, mr) in enumerate(zip(data_rows, mine["rows"])):
            if str(dr[0] or "") != mr["wh"]:
                bad += 1
                print("       行%d 名称不符：%r vs %r" % (i, dr[0], mr["wh"]))
            if dr[1] is not None and str(dr[1]).strip() != "":
                bad += 1
                print("       行%d 结余应当是空（手填列，本轮无来源）：%r" % (i, dr[1]))
            for j, mv in enumerate(mr["cells"]):
                got = _num(dr[2 + j])
                if not _eq(got, mv):
                    bad += 1
                    print("       行%d %d月 不符：%r vs %r" % (i, j + 1, dr[2 + j], mv))
            if not _eq(_num(dr[14]), mr["total"]):
                bad += 1
                print("       行%d 合计不符：%r vs %r" % (i, dr[14], mr["total"]))
        chk(bad == 0, "[%s] 逐格对账 0 差异（名称/12 月/合计）＋ 结余如实为空" % sheet)
        # 尾块 = 各月列之和 ＋ 合计之和（自己算一遍；⚠️ bad 必须归零，否则上面那条会连坐）
        bad = 0
        for j, mv in enumerate(mine["months"]):
            want = sum([r["raw"][j] for r in mine["rows"] if r["raw"][j] is not None],
                       Decimal("0"))
            if not _eq(_num(tail_row[2 + j]), want):
                bad += 1
                print("       尾块 %d月 不符：%r vs %r" % (j + 1, tail_row[2 + j], want))
        chk(bad == 0, "[%s] 尾块合计行 = 各月列之和（自己算一遍）" % sheet)
        zeros = [(i, j) for i, mr in enumerate(mine["rows"])
                 for j, v in enumerate(mr["raw"]) if v == 0]
        if zeros:
            chk(all(_num(data_rows[i][2 + j]) is None for i, j in zeros),
                "[%s] 算出来是 0.00 的 %d 格**留空**（与壳同款：`v ? M(v) : '—'`，⛔ 不写 0）"
                % (sheet, len(zeros)))
        chk(max_row == 2 + len(mine["rows"]), "[%s] 末行 = %d（尾块紧贴数据区，下面一格没有）"
            % (sheet, max_row))
        for c in range(1, 16):
            v = openpyxl.load_workbook(out, data_only=True)[sheet].cell(row=max_row, column=c).value
            if v is not None and isinstance(v, str) and v.startswith("="):
                chk(False, "[%s] 还有公式：%s%d" % (sheet, chr(64 + c), max_row))
        print("   [%s] 成品逐行：" % sheet)
        for i, dr in enumerate(data_rows):
            print("      R%-2d %-22s 结余=%-12s 1..12月=%s 合计=%s"
                  % (2 + i, dr[0], dr[1], [str(x) if x is not None else "—" for x in dr[2:14]],
                     dr[14]))
        print("      R%-2d %-22s %s" % (max_row, "(合计行)", [str(x) if x is not None else "—"
                                                          for x in tail_row[1:]]))

    print("\n" + "=" * 78)
    print("④ 反向自检：把成品故意弄坏 4 种 ⇒ 守卫必须逐条报红")
    print("=" * 78)
    _corrupt_cases(out, plan)

    print("\n" + "=" * 78)
    print("⑤ 真机产物（若已导出）与朴素重算对账")
    print("=" * 78)
    _check_real(details, plan)

    print("\n" + "=" * 78)
    print("结果：%d 绿 / %d 红" % (len(OK), len(BAD)))
    for m in BAD:
        print("   ✗", m)
    return 1 if BAD else 0


def _check_real(details, plan):
    """真机产物（`exports/焊材气体提取表/<期间标签>/`）逐格对账。"""
    ym = plan["report_ym"]
    real = (Path("exports/焊材气体提取表")
            / W.period_label(ym) / W.file_name(ym))
    if not real.exists():
        print("   ⏭ 真机产物还没导出（先跑 GET /api/m07/extract/export）：", real)
        return
    print("   真机产物 =", real, "(%d B)" % real.stat().st_size)
    for kind, sheet in W.SHEET_OF_KIND.items():
        whs = [r["warehouse"] for r in plan["grids"][kind]["rows"]]
        mine = naive_grid(details, kind, whs, plan["year"], plan["report_ym"])
        disp, max_row = read_display(real, sheet, 2)
        data_rows, tail_row = disp[:-1], disp[-1]
        bad = 0
        for i, (dr, mr) in enumerate(zip(data_rows, mine["rows"])):
            if str(dr[0] or "") != mr["wh"]:
                bad += 1
            for j, mv in enumerate(mr["cells"]):
                if not _eq(_num(dr[2 + j]), mv):
                    bad += 1
                    print("       %s 行%d %d月：%r vs %r" % (sheet, i, j + 1, dr[2 + j], mv))
            if dr[1] is not None and str(dr[1]).strip() != "":
                bad += 1
                print("       %s 结余应为空：%r" % (sheet, dr[1]))
        chk(bad == 0, "[%s] 真机产物：显示区 = 朴素重算（%d 行，末行 %d）"
            % (sheet, len(data_rows), max_row))
        for j in range(12):
            want = sum([r["raw"][j] for r in mine["rows"] if r["raw"][j] is not None],
                       Decimal("0"))
            if not _eq(_num(tail_row[2 + j]), want):
                bad += 1
                print("       %s 合计行 %d月：%r vs %r" % (sheet, j + 1, tail_row[2 + j], want))
        chk(bad == 0, "[%s] 真机产物：合计行 = 各月列之和" % sheet)


def _eq(a, b):
    if a is None or b is None:
        return a is None and b is None
    return abs(float(a) - float(b)) < 0.005


def _corrupt_cases(out, plan):
    scratch = Path(tempfile.mkdtemp(prefix="p1-bad-"))
    wb = openpyxl.Workbook()
    template = None
    from app.report_export import resources
    template = resources.WELD_TEMPLATE

    def rebuild(mutate):
        """用同一份 spec 重新生成，再按 `mutate` 破坏产物 ⇒ 返回 `(path, problems)`。"""
        p = scratch / "m.xlsx"
        specs = W.build_specs(plan["grids"])
        pl = tpl_fill.fill_workbook(template, p, specs)
        mutate(p)
        return p, tpl_fill.verify_filled(p, pl)

    def _mut_residue(p):                       # ① 尾块以下塞一个旧值
        w = openpyxl.load_workbook(p)
        ws = w["焊材"]
        for r in range(ws.max_row + 1, ws.max_row + 4):
            ws.cell(row=r, column=1).value = "上一版的仓库"
            ws.cell(row=r, column=2).value = 999.0
        w.save(p)
        w.close()

    def _mut_notblank(p):                      # ② 该空的月份格写上 0
        w = openpyxl.load_workbook(p)
        ws = w["焊材"]
        ws.cell(row=2, column=13).value = 0.0   # M2（11月）本该空
        w.save(p)
        w.close()

    def _mut_formula(p):                       # ③ 数据格里留公式（值覆盖不彻底）
        w = openpyxl.load_workbook(p)
        ws = w["焊材"]
        ws.cell(row=2, column=15).value = "=SUM(B2:N2)"
        w.save(p)
        w.close()

    def _mut_short(p):                         # ④ 少一行（旧实现"只写前 n 行"的形态）
        w = openpyxl.load_workbook(p)
        ws = w["焊材"]
        ws.delete_rows(2, 1)
        w.save(p)
        w.close()

    for name, fn, want in (
            ("残留（尾块以下有旧值）", _mut_residue, "尾块以下"),
            ("该空没空（11月写 0）", _mut_notblank, "与期望不符"),
            ("值覆盖不彻底（留公式）", _mut_formula, "还是公式"),
            ("少一行（只写前 n 行的旧形态）", _mut_short, "与期望不符")):
        p, problems = rebuild(fn)
        hit = any(want in x for x in problems)
        chk(bool(problems) and hit,
            "反向自检：%s ⇒ 守卫报红（%d 条，命中「%s」）" % (name, len(problems), want))
        for x in problems[:2]:
            print("       ·", x)


if __name__ == "__main__":
    sys.exit(main())
