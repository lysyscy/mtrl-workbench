# -*- coding: utf-8 -*-
"""`TAIL-04` 收发结存成品的**独立对账**（可复跑）。

它回答一个问题：**导出的成品文件里的数字，与数据库里的明细对得上吗？**

⚠️ 为什么要有它：导出链路内部有一层「缓存级对账」（`setl_export.verify` —— 从新文件的
   透视缓存重算 vs 从同一份 `df` 重算），但那是**同源比对**：它只证明"没有分叉"，
   **证明不了 `df` 本身是对的**（`df` 是从库里映射来的，映射错一列照样过）。
⇒ 本脚本走**第三条独立路径**：直接用 SQL 在**库**上算四个金额口径，与成品文件里
   `Sheet1` 的「总计」行逐分比对；并与**分仓库文件**的行数/合计对账。

用法：`cd D:\\dev\\mtrl-workbench && backend/.venv/Scripts/python.exe tools/_verify_tail04.py [期间]`
默认期间 = `2026-09`。只读（不写库、不抓 i8）。
"""

import pathlib
import sys
import zipfile
from decimal import Decimal

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1] / "backend"))
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

import openpyxl  # noqa: E402
from sqlalchemy import text  # noqa: E402

from app.core.db import session_scope  # noqa: E402
from app.report_export.service import EXPORT_ROOT  # noqa: E402

PERIOD = sys.argv[1] if len(sys.argv) > 1 else "2026-09"
MONEY = ("begin_amt", "in_amt", "out_amt", "end_amt")
LABEL = ("期初金额", "收金额", "发金额", "结存金额")
problems: list[str] = []


def ok(cond, msg):
    print(("  ✅ " if cond else "  ⛔ ") + msg)
    if not cond:
        problems.append(msg)


def main() -> int:
    ym_dir = EXPORT_ROOT / "收发结存" / PERIOD
    zips = sorted(ym_dir.glob("*.zip"))
    print("成品目录 =", ym_dir)
    if not zips:
        print("⛔ 找不到成品 zip —— 先跑一次导出")
        return 1
    zpath = zips[-1]
    print("成品 =", zpath.name, zpath.stat().st_size, "B")

    # ---- ① 库侧独立取数（SQL，**与导出链路不共用任何代码**）
    cols = ", ".join(f"sum({c})" for c in MONEY)
    with session_scope() as s:
        db_total = s.execute(text(f"SELECT {cols} FROM setl_detail WHERE ym = :y"),
                             {"y": PERIOD}).one()
        db_rows = s.execute(text("SELECT count(*) FROM setl_detail WHERE ym = :y"),
                            {"y": PERIOD}).scalar()
        db_whs = [r[0] for r in s.execute(text(
            "SELECT warehouse_name FROM setl_detail WHERE ym = :y "
            "GROUP BY warehouse_name ORDER BY warehouse_name"), {"y": PERIOD})]
        db_per_wh = {r[0]: (int(r[1]), r[2]) for r in s.execute(text(
            "SELECT warehouse_name, count(*), sum(in_amt) FROM setl_detail "
            "WHERE ym = :y GROUP BY warehouse_name"), {"y": PERIOD})}
    db_total = [Decimal(str(x or 0)) for x in db_total]
    print(f"\n库侧：{db_rows} 行 · {len(db_whs)} 仓 · 总计 " +
          " ".join(f"{l}={v:.2f}" for l, v in zip(LABEL, db_total)))

    z = zipfile.ZipFile(zpath)
    names = z.namelist()
    print("\n=== ① 包内文件 ===")
    ok(len(names) == len(db_whs) + 1,
       f"包内 {len(names)} 个文件 == 分仓库 {len(db_whs)} ＋ 总表 1")
    for w in db_whs:
        ok(f"{w}.xlsx" in names, f"分仓库文件在：{w}.xlsx")

    main_name = [n for n in names if n.endswith("收发结存表.xlsx")]
    ok(len(main_name) == 1, f"总表唯一：{main_name}")
    if not main_name:
        return 1

    # ---- ② 总表：报表1 行数 + Sheet1 总计
    import tempfile
    tmp = pathlib.Path(tempfile.mkdtemp(prefix="t04verify-"))
    for n in (main_name[0],):
        (tmp / n).write_bytes(z.read(n))
    wb = openpyxl.load_workbook(tmp / main_name[0], data_only=True)
    print("\n=== ② 总表 ===")
    print("  sheets =", wb.sheetnames)
    r1 = list(wb["报表1"].iter_rows(values_only=True))
    ok(len(r1) == db_rows + 1, f"报表1 行数 = 表头 1 ＋ 明细 {db_rows}（实得 {len(r1)}）")

    ws1 = wb["Sheet1"]
    grand_row = None
    for r in ws1.iter_rows(min_row=1, max_row=ws1.max_row, values_only=False):
        if r[0].value == "总计":
            grand_row = r
            break
    ok(grand_row is not None, "Sheet1 有「总计」行")
    if grand_row is not None:
        for k, (lbl, want) in enumerate(zip(LABEL, db_total)):
            got = Decimal(str(grand_row[1 + k].value or 0))
            ok(abs(got - want) < Decimal("0.005"),
               f"⭐ 总计 {lbl}：文件 {got:.2f} vs 库 SQL {want:.2f}（**第三条独立路径**）")

    # ---- ③ 分仓库文件：行数与「合计」口径
    print("\n=== ③ 分仓库文件 ===")
    bad = []
    for w in db_whs:
        fn = f"{w}.xlsx"
        if fn not in names:
            continue
        (tmp / fn).write_bytes(z.read(fn))
        wbw = openpyxl.load_workbook(tmp / fn, data_only=True)
        rws = [r for r in wbw["Sheet1"].iter_rows(values_only=True)
               if any(x is not None and str(x).strip() != "" for x in r)]
        nrows, sum_in = db_per_wh[w]
        got_rows = len(rws) - 1                      # 去掉表头
        got_sum = sum(Decimal(str(r[11])) for r in rws[1:] if r[11] is not None)
        if got_rows != nrows or abs(got_sum - Decimal(str(sum_in or 0))) > Decimal("0.005"):
            bad.append(f"{w}: 行 {got_rows}vs{nrows} · 收金额 {got_sum:.2f}vs{sum_in:.2f}")
        wbw.close()
    ok(not bad, f"{len(db_whs)} 个分仓库文件的行数与「收金额」逐仓对平" + (("；不符：" + "；".join(bad)) if bad else ""))

    # ---- ④ 透视表：打开即刷新 + location
    #      ⚠️ 透视在**总表 xlsx 内部**（总表本身又是一个 zip）—— 别去外层 zip 里找（实测踩过）
    #      ⚠️⚠️ `refreshOnLoad` 的**唯一合法归属** = `pivotCacheDefinition1.xml` 的
    #          `<pivotCacheDefinition>`；放 `<pivotTableDefinition>` 或 `workbook.xml` 的
    #          `<pivotCache>` 都是**不合 schema** 的（openpyxl 读工作簿当场 `TypeError`）。
    #          ② 那次 `load_workbook`（非 read_only、会解析透视）正是能抓出它的那一步。
    print("\n=== ④ 透视表（原样留在模板里、只换缓存） ===")
    import re
    inner = zipfile.ZipFile(tmp / main_name[0])
    pt = inner.read("xl/pivotTables/pivotTable1.xml").decode("utf-8")
    ok("refreshOnLoad" not in pt,
       "pivotTableDefinition 上**没有** refreshOnLoad（该属性不在 pivotTable schema 里）")
    cd = inner.read("xl/pivotCache/pivotCacheDefinition1.xml").decode("utf-8")
    ok('refreshOnLoad="1"' in cd,
       "pivotCacheDefinition 上已置 refreshOnLoad=\"1\"（Excel 打开即刷新）")
    loc = re.search(r'<location ref="([^"]*)"', pt)
    ok(loc is not None and loc.group(1).endswith(str(4 + len(db_whs))),
       f"pivotTable1 location = {loc.group(1) if loc else '?'}（行数随仓库集合派生，末行 = 4+{len(db_whs)}）")
    cr = inner.read("xl/pivotCache/pivotCacheRecords1.xml").decode("utf-8")
    ok(f'count="{db_rows}"' in cr, f"缓存记录数 == 明细行数 {db_rows}")
    inner.close()
    z.close()

    print("\n" + ("=" * 60))
    if problems:
        print(f"⛔ 未通过 {len(problems)} 条：")
        for p in problems:
            print("   -", p)
        return 1
    print("✅ 独立对账全部通过（库 SQL / 成品透视 / 分仓库文件 三条路径一致）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
