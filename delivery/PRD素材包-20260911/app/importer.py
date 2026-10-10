# -*- coding: utf-8 -*-
"""
导入层：识别文件类型 → 解析 → 应用规则 → 入库（upsert）。
类型：sfjc(收发结存) / kcdj(库存单据) / pay_notice(付款通知书) / pay_plan(付款计划)
规则（剔除/分类/项目映射/小微）全部来自 config 表，程序不写死。
"""
import json
import os
import sys

import openpyxl

BASE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(BASE)

# 表头/解析工具统一来源（S7-1）：input/schema_cols
sys.path.insert(0, os.path.join(ROOT, 'input'))
from schema_cols import (SFJC_HEADER, KCDJ_COLS, KCDJ_RAW_FIELDS,
                         to_number, read_sfjc_grid)

import db  # noqa
import warehouse_sync  # noqa


def _ensure_project(name, is_small='否', conn=None, full_name=None, wh_cache=None):
    """ensure_project 包装：三维匹配 + 仓库列表缓存兜底补 full_name。
    wh_cache = warehouse_sync.load_cache() 的结果（None 表示不兜底）。"""
    pid = db.ensure_project(name, is_small, conn=conn, full_name=full_name)
    if wh_cache and conn is not None:
        it = (warehouse_sync.lookup_by_name(name, wh_cache)
              or (warehouse_sync.lookup_by_full(full_name, wh_cache) if full_name else None))
        if it and it.get('full_name'):
            row = conn.execute('SELECT full_name FROM projects WHERE id=?', (pid,)).fetchone()
            if row and not (row['full_name'] or ''):
                conn.execute('UPDATE projects SET full_name=? WHERE id=?', (it['full_name'], pid))
    return pid


# ================= 文件类型识别 =================
def _first_rows(path, n=8):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = []
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i >= n:
            break
        rows.append([str(x) if x is not None else '' for x in r])
    wb.close()
    return rows


def detect_type(path):
    rows = _first_rows(path)
    head = [cell for r in rows for cell in r]
    joined = '|'.join(head)
    if all(k in joined for k in ('仓库', '资源编码', '期初数量')):
        return 'sfjc'
    if all(k in joined for k in ('核算年', '出入库方向', '仓库')):
        return 'kcdj'
    if all(k in joined for k in ('单据编码', '单据名称', '单据日期', '付款项目')):
        return 'pay_notice'
    if all(k in joined for k in ('单据号', '预付款金额')):
        return 'pay_plan'
    return None


# ================= 收发结存 =================
# 表头 SFJC_HEADER 与解析工具 read_sfjc_grid / to_number 统一来自 schema_cols（S7-1）

def _read_sfjc_rows(path):
    """读原始导出：拆合并填充、去空行（迁用 schema_cols.read_sfjc_grid，与 sfjc_process 共用）"""
    return read_sfjc_grid(path)


# ================= 付款通知书 规则（S7-2 规则下沉，纯函数便于单测） =================
def is_excluded(payproj, pname, recv, lrr, btype, excl):
    """是否命中剔除规则（按原顺序：付款项目名/收款单位/录入人/业务类型/单据名称关键字）"""
    if payproj in excl.get('付款项目名', []):
        return True
    if recv in excl.get('收款单位', []):
        return True
    if lrr in excl.get('录入人', []):
        return True
    if btype in excl.get('业务类型', []):
        return True
    if any(k in pname for k in excl.get('单据名称关键字', [])):
        return True
    return False


def initial_biz_class(btype, ptype, cls_rules):
    """非剔除行的初判分类：机械锁定业务类型优先，否则按付款类型初判，未知→待确认"""
    if btype in cls_rules.get('机械锁定业务类型', []):
        return '机械'
    return cls_rules.get('付款类型初判', {}).get(ptype, '待确认')


def decide_biz_class(payproj, pname, recv, lrr, btype, ptype, excl, cls_rules):
    """组合规则（剔除优先，再初判）。导入流程因需先做项目解析、保留 None 中间态而不直接调用。"""
    if is_excluded(payproj, pname, recv, lrr, btype, excl):
        return '剔除'
    return initial_biz_class(btype, ptype, cls_rules)


def is_small_project(name, small_map):
    """判断名字是否小微：name 是 small_map 的 key（小微壳）或 value 的 key（具体小微项目）。
    不再用『仓库名含"小微"』（实际小微仓库名如"大唐雷州发电项目"不含"小微"，旧判据失效）。"""
    if not name:
        return False
    if name in small_map:
        return True
    for sms in small_map.values():
        if isinstance(sms, dict) and name in sms:
            return True
    return False


def build_kcdj_raw(d):
    """inv_detail.raw_json 仅保留未独立成列的非结构化字段（S7-4 瘦身）"""
    return json.dumps({k: str(d.get(k)) for k in KCDJ_RAW_FIELDS}, ensure_ascii=False)


def import_sfjc(path, month):
    rows = _read_sfjc_rows(path)
    conn = db.get_conn()
    small_map = db.get_config('small_map', {})
    wh_cache = warehouse_sync.load_cache()
    inserted = updated = 0
    projects_seen = {}
    try:
        # 活跃业务不自动建档：陌生仓库报错（须先同步项目，见总纲 §0.6 建档入口收敛）
        _existing = {r['name'] for r in conn.execute('SELECT name FROM projects')}
        _unmatched = set()
        for raw in rows:
            _wh = str(dict(zip(SFJC_HEADER, raw))['仓库']).strip()
            if _wh and _wh != '合计' and _wh not in _existing:
                _unmatched.add(_wh)
        if _unmatched:
            raise RuntimeError('以下仓库未在项目档案中，请先在配置页「同步项目」：%s'
                               % '、'.join(sorted(_unmatched)))
        # 幂等重跑：先清当月自动行。否则当月数据变少/换月时会残留旧行
        # （2026-09-08 实测：9 月库里混进 636 行 8 月残留，跨月衔接对不上）
        conn.execute("DELETE FROM stock_monthly WHERE month=? "
                     "AND (source IS NULL OR source='auto')", (month,))
        for raw in rows:
            d = dict(zip(SFJC_HEADER, raw))
            wh = str(d['仓库']).strip()
            if not wh or wh == '合计':
                continue
            is_small = '是' if is_small_project(wh, small_map) else '否'
            if wh not in projects_seen:
                projects_seen[wh] = _ensure_project(wh, is_small, conn=conn, wh_cache=wh_cache)
            pid = projects_seen[wh]

            num = to_number  # S7-1：复用共享数字解析

            cur = conn.execute(
                '''INSERT INTO stock_monthly(month, project_id, 仓库, 资源分类, 资源编码, 资源名称,
                   规格型号, 计量单位, 物料性质, 期初数量, 期初金额, 收数量, 收单价, 收金额,
                   发数量, 发单价, 发金额, 结存数量, 结存金额)
                   VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                   ON CONFLICT(month, 仓库, 资源编码) DO UPDATE SET
                   资源名称=excluded.资源名称, 规格型号=excluded.规格型号, 计量单位=excluded.计量单位,
                   物料性质=excluded.物料性质, 期初数量=excluded.期初数量, 期初金额=excluded.期初金额,
                   收数量=excluded.收数量, 收单价=excluded.收单价, 收金额=excluded.收金额,
                   发数量=excluded.发数量, 发单价=excluded.发单价, 发金额=excluded.发金额,
                   结存数量=excluded.结存数量, 结存金额=excluded.结存金额, source='auto' ''',
                (month, pid, wh, d['资源分类'], d['资源编码'], d['资源名称'],
                 d['规格型号'], d['计量单位'], d['物料性质'],
                 num(d['期初数量']), num(d['期初金额']), num(d['收数量']), num(d['收单价']),
                 num(d['收金额']), num(d['发数量']), num(d['发单价']), num(d['发金额']),
                 num(d['结存数量']), num(d['结存金额'])))
            if cur.rowcount == 1:
                inserted += 1
            else:
                updated += 1
        conn.commit()
    finally:
        conn.close()
    return {'type': 'sfjc', 'month': month, 'inserted': inserted, 'updated': updated,
            'rows': inserted + updated, 'projects': list(projects_seen.keys())}


# ================= 库存单据明细 =================
# 表头 KCDJ_COLS / KCDJ_RAW_FIELDS 统一来自 schema_cols（S7-1/4）


def import_kcdj(path, month=None):
    sys.path.insert(0, os.path.join(ROOT, 'input'))
    from common_tpl import read_data
    hdr, data = read_data(path, col_count=28)
    conn = db.get_conn()
    small_map = db.get_config('small_map', {})
    wh_cache = warehouse_sync.load_cache()
    inserted = updated = 0
    projects_seen = {}
    whs = set()
    try:
        # 活跃业务不自动建档：陌生仓库报错（须先同步项目，见总纲 §0.6 建档入口收敛）
        _existing = {r['name'] for r in conn.execute('SELECT name FROM projects')}
        _unmatched = set()
        for row in data:
            _wh = str(dict(zip(KCDJ_COLS.values(), row)).get('仓库', '')).strip()
            if _wh and _wh not in _existing:
                _unmatched.add(_wh)
        if _unmatched:
            raise RuntimeError('以下仓库未在项目档案中，请先在配置页「同步项目」：%s'
                               % '、'.join(sorted(_unmatched)))
        # 幂等重跑：先清本次涉及月份的自动行（同 import_sfjc，防残留）
        months = {month} if month else set()
        if not month:
            for row in data:
                d = dict(zip(KCDJ_COLS.values(), row))
                try:
                    months.add('%s-%02d' % (int(float(d['核算年'])), int(float(d['核算月']))))
                except (TypeError, ValueError):
                    pass
        for m in months:
            conn.execute("DELETE FROM inv_detail WHERE month=? "
                         "AND (source IS NULL OR source='auto')", (m,))
        for row in data:
            d = dict(zip(KCDJ_COLS.values(), row))
            wh = str(d.get('仓库', '')).strip()
            if not wh:
                continue
            whs.add(wh)
            proj_name = str(d.get('项目名称', '')).strip()   # 全名（库存单据第4列）
            m = month or ('%s-%02d' % (int(float(d['核算年'])), int(float(d['核算月']))))
            is_small = '是' if is_small_project(wh, small_map) else '否'
            if wh not in projects_seen:
                projects_seen[wh] = _ensure_project(wh, is_small, conn=conn,
                                                    full_name=proj_name, wh_cache=wh_cache)
            pid = projects_seen[wh]

            num = to_number  # S7-1

            raw = build_kcdj_raw(d)  # S7-4：raw_json 仅保留非结构化字段
            cur = conn.execute(
                '''INSERT INTO inv_detail(month, project_id, 序号, 仓库, 出入库方向, 物料性质,
                   物料名称, 规格型号, 计量单位, 业务类型, 项目名称, 单据编号, 供应商,
                   班组, 部门, 单据日期, 审核日期, 是否作废, cbs名称, wbs名称,
                   数量, 不含税单价, 含税单价, 金额, 价税合计, raw_json)
                   VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                   ON CONFLICT(month, 单据编号, 序号) DO UPDATE SET
                   仓库=excluded.仓库, 出入库方向=excluded.出入库方向, 物料性质=excluded.物料性质,
                   物料名称=excluded.物料名称, 规格型号=excluded.规格型号, 计量单位=excluded.计量单位,
                   业务类型=excluded.业务类型, 项目名称=excluded.项目名称, 供应商=excluded.供应商,
                   班组=excluded.班组, 部门=excluded.部门, 单据日期=excluded.单据日期,
                   审核日期=excluded.审核日期, 是否作废=excluded.是否作废, cbs名称=excluded.cbs名称,
                   wbs名称=excluded.wbs名称, 数量=excluded.数量, 不含税单价=excluded.不含税单价,
                   含税单价=excluded.含税单价, 金额=excluded.金额, 价税合计=excluded.价税合计,
                   raw_json=excluded.raw_json, source='auto' ''',
                (m, pid, d.get('序号'), wh, d.get('出入库方向'), d.get('物料性质'),
                 d.get('物料名称'), d.get('规格型号'), d.get('计量单位'), d.get('业务类型'),
                 d.get('项目名称'), d.get('单据编号'), d.get('供应商'), d.get('班组'),
                 d.get('部门'), d.get('单据日期'), d.get('审核日期'), d.get('是否作废'),
                 d.get('cbs名称'), d.get('wbs名称'),
                 num(d.get('数量')), num(d.get('不含税单价')), num(d.get('含税单价')),
                 num(d.get('金额')), num(d.get('价税合计')), raw))
            if cur.rowcount == 1:
                inserted += 1
            else:
                updated += 1
        conn.commit()
    finally:
        conn.close()
    return {'type': 'kcdj', 'month': month or 'auto', 'inserted': inserted, 'updated': updated,
            'rows': inserted + updated, 'projects': sorted(whs)}


# ================= 付款通知书 =================
def _load_projects_index():
    """加载 项目名/全名/别名 → 项目行 索引（付款导入匹配依据）"""
    conn = db.get_conn()
    try:
        rows = conn.execute('SELECT id, name, full_name, is_small, aliases FROM projects').fetchall()
    finally:
        conn.close()
    idx = {}
    for r in rows:
        idx[r['name']] = r
        if r['full_name']:
            idx[r['full_name']] = r
        for a in (r['aliases'] or '').split('|'):
            if a:
                idx[a] = r
    return idx


PAY_HDR_ALIAS = {
    '单据编码': '单据编码', '单据名称': '单据名称', '单据日期': '单据日期',
    '付款项目': '付款项目', '申请人': '申请人', '申请部门': '申请部门',
    '付款类型': '付款类型', '业务类型': '业务类型', '付款属性': '付款属性',
    '付款组织': '付款组织', '收款单位': '收款单位', '申请金额': '申请金额',
    '审批金额': '审批金额', '累计实付': '累计实付金额', '支付状态': '支付状态',
    '冲红状态': '冲红状态', '录入人': '录入人', '审批状态': '审批状态',
    '打印状态': '打印状态', '备注': '备注', '记账状态': '记账状态',
}
MONEY_F = {'申请金额', '审批金额', '累计实付金额'}


def _parse_pay_header(rows):
    """找含'单据编码'的表头行 → (表头列表, 表头行号)"""
    for i, r in enumerate(rows):
        strs = [str(x) if x is not None else '' for x in r]
        if '单据编码' in strs and '单据名称' in strs:
            return list(r), i
    return None, None


def import_pay_notice(path, source_month=None):
    rows = _first_rows(path, n=6)
    hdr, hi = _parse_pay_header(rows)
    if hdr is None:
        raise RuntimeError('找不到付款通知书表头（含 单据编码/单据名称）')

    col = {}
    for j, name in enumerate(hdr):
        s = str(name).strip() if name is not None else ''
        for alias, std in PAY_HDR_ALIAS.items():
            if alias in s:
                col[std] = j
                break

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    small_map = db.get_config('small_map', {})
    excl = db.get_config('exclude_rules', {})
    cls_rules = db.get_config('classify_rules', {})
    proj_idx = _load_projects_index()  # 项目名/全名/别名 → 项目行（project_map 已退役）
    wh_cache = warehouse_sync.load_cache()

    conn = db.get_conn()
    inserted = updated = skipped = 0
    pending_projects, pending_smalls = set(), set()
    biz_class_stat = {}
    file_codes = set()
    try:
        for r in ws.iter_rows(min_row=hi + 2, values_only=True):
            def gv(name):
                j = col.get(name)
                return r[j] if j is not None and j < len(r) else None

            code = gv('单据编码')
            if code is None or str(code).strip() == '':
                continue
            code = str(code).strip()
            file_codes.add(code)

            def money(name):
                return to_number(gv(name))  # S7-1：复用共享数字解析

            pname = str(gv('单据名称') or '').strip()
            payproj = str(gv('付款项目') or '').strip()
            recv = str(gv('收款单位') or '').strip()
            lrr = str(gv('录入人') or '').strip()
            btype = str(gv('业务类型') or '').strip()
            ptype = str(gv('付款类型') or '').strip()

            # ---- 剔除规则（顺序执行，规则下沉到 importer.is_excluded，便于单测）----
            excluded = is_excluded(payproj, pname, recv, lrr, btype, excl)
            biz_class = '剔除' if excluded else None
            if excluded:
                skipped += 1

            # ---- 项目解析（2.4）----
            # 剔除/忽略行也解析统计名（用于按项目筛选和显示），但不建档
            stat_name, small_name = None, None
            if biz_class in ('剔除', '忽略'):
                if payproj in small_map:
                    for sm, kws in small_map[payproj].items():
                        if any(k in pname for k in kws):
                            small_name, stat_name = sm, sm
                            break
                else:
                    prow = proj_idx.get(payproj)
                    if prow:
                        stat_name = prow['name']
            elif biz_class is None:
                # ① 小微合并项目：按单据名称关键词拆具体小微
                if payproj in small_map:
                    for sm, kws in small_map[payproj].items():
                        if any(k in pname for k in kws):
                            small_name, stat_name = sm, sm
                            break
                    if stat_name is None:
                        pending_smalls.add(payproj + ' :: ' + (pname[:30] if pname else '?'))
                        biz_class = '待确认'
                else:
                    # ② 普通项目：按 项目名/别名 匹配（project_map 已退役）
                    prow = proj_idx.get(payproj)
                    if prow:
                        stat_name = prow['name']
                    else:
                        pending_projects.add(payproj or '(空)')
                        biz_class = '待确认'

            # ---- 分类（2.3，规则下沉到 importer.initial_biz_class）----
            if biz_class is None:
                biz_class = initial_biz_class(btype, ptype, cls_rules)
                if biz_class == '待确认':
                    pending_projects.add('【分类待确认】' + (payproj or '(空)'))
            if biz_class:
                biz_class_stat[biz_class] = biz_class_stat.get(biz_class, 0) + 1

            # ---- 建档 + upsert ----
            # 只对 辅材/机械 建档；剔除/忽略 行仅保留统计名用于筛选显示
            pid = None
            if biz_class in ('辅材', '机械') and stat_name:
                # full_name=付款项目（全名）：普通项目=自己的全名，小微拆解行=壳全名（付款项目即壳全名）
                pid = _ensure_project(stat_name, '是' if small_name else '否', conn=conn,
                                      full_name=payproj, wh_cache=wh_cache)

            cur = conn.execute(
                '''INSERT INTO pay_notice(单据编码, 单据名称, 单据日期, project_id, stat_name,
                   small_name, 付款项目, 申请人, 申请部门, 录入人, 付款类型, 业务类型,
                   付款属性, 付款组织, 收款单位, 申请金额, 审批金额, 累计实付金额, 支付状态,
                   冲红状态, 审批状态, 打印状态, 记账状态, 备注, biz_class, manual, confirmed,
                   source_month, source)
                   VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,0,0,?,?)
                   ON CONFLICT(单据编码) DO UPDATE SET
                   单据名称=excluded.单据名称, 单据日期=excluded.单据日期,
                   project_id=excluded.project_id, stat_name=excluded.stat_name,
                   small_name=excluded.small_name, 付款项目=excluded.付款项目,
                   申请人=excluded.申请人, 申请部门=excluded.申请部门, 录入人=excluded.录入人,
                   付款类型=excluded.付款类型, 业务类型=excluded.业务类型,
                   付款组织=excluded.付款组织, 收款单位=excluded.收款单位, 申请金额=excluded.申请金额,
                   审批金额=excluded.审批金额, 累计实付金额=excluded.累计实付金额,
                   支付状态=excluded.支付状态, 冲红状态=excluded.冲红状态,
                   审批状态=excluded.审批状态, 打印状态=excluded.打印状态,
                   记账状态=excluded.记账状态, 备注=excluded.备注,
                   biz_class=CASE WHEN pay_notice.manual=1 THEN pay_notice.biz_class ELSE excluded.biz_class END,
                   source_month=excluded.source_month, source='auto' ''',
                (code, pname, str(gv('单据日期') or ''), pid, stat_name, small_name,
                 payproj, str(gv('申请人') or ''), str(gv('申请部门') or ''), lrr,
                 ptype, btype, str(gv('付款属性') or ''), str(gv('付款组织') or ''), recv,
                 money('申请金额'), money('审批金额'), money('累计实付金额'),
                 str(gv('支付状态') or ''), str(gv('冲红状态') or ''),
                 str(gv('审批状态') or ''), str(gv('打印状态') or ''),
                 str(gv('记账状态') or ''), str(gv('备注') or ''), biz_class,
                 source_month or (str(gv('单据日期') or '')[:7]), 'auto'))
            if cur.rowcount == 1:
                inserted += 1
            else:
                updated += 1
        conn.commit()

        # ---- 消失行检测（二次导入：库中有、但本次文件没有的单据，可能被信息化删除）----
        if source_month:
            existing = conn.execute(
                "SELECT id, 单据编码, 单据名称, 单据日期, biz_class FROM pay_notice "
                "WHERE deleted=0 AND source_month=?", (source_month,)).fetchall()
        else:
            existing = conn.execute(
                "SELECT id, 单据编码, 单据名称, 单据日期, biz_class FROM pay_notice "
                "WHERE deleted=0").fetchall()
        removed = [dict(r) for r in existing if r['单据编码'] not in file_codes]
    finally:
        conn.close()
        wb.close()
    return {'type': 'pay_notice', 'inserted': inserted, 'updated': updated,
            'rows': inserted + updated, 'skipped': skipped,
            'biz_class_stat': biz_class_stat,
            'pending_projects': sorted(pending_projects),
            'pending_smalls': sorted(pending_smalls),
            'removed': removed}


# ================= 付款计划 =================
def import_pay_plan(path, month):
    rows = _first_rows(path, n=6)
    hdr = None
    for r in rows:
        strs = [str(x) if x is not None else '' for x in r]
        if '单据号' in strs:
            hdr = list(r)
            break
    if hdr is None:
        raise RuntimeError('找不到付款计划表头（含 单据号）')
    col = {}
    for j, name in enumerate(hdr):
        s = str(name).strip() if name is not None else ''
        for kw, std in {'单据号': '单据号', '供应商': '供应商', '单据金额': '单据金额',
                        '预付款金额': '预付款金额', '项目': '项目', '备注': '备注'}.items():
            if kw in s:
                col[std] = j
                break
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    conn = db.get_conn()
    inserted = 0
    try:
        for r in ws.iter_rows(min_row=2, values_only=True):
            code = r[col['单据号']] if '单据号' in col and col['单据号'] < len(r) else None
            code = str(code).strip() if code is not None else ''

            def num(j):
                return to_number(r[j] if j is not None else None)  # S7-1

            proj = str(r[col['项目']]).strip() if '项目' in col else ''
            pid = db.ensure_project(proj, conn=conn) if proj else None
            conn.execute(
                '''INSERT INTO pay_plan(month, project_id, 单据号, 供应商, 单据金额, 预付款金额, 备注)
                   VALUES(?,?,?,?,?,?,?)
                   ON CONFLICT(month, 单据号, 供应商) DO UPDATE SET
                   供应商=excluded.供应商, 单据金额=excluded.单据金额,
                   预付款金额=excluded.预付款金额, 备注=excluded.备注 ''',
                (month, pid, str(code).strip(),
                 str(r[col['供应商']]).strip() if '供应商' in col else '',
                 num(col.get('单据金额')), num(col.get('预付款金额')),
                 str(r[col['备注']]).strip() if '备注' in col else ''))
            inserted += 1
        conn.commit()
    finally:
        conn.close()
    return {'type': 'pay_plan', 'month': month, 'inserted': inserted, 'rows': inserted}


TYPE_NAMES = {'sfjc': '收发结存表', 'kcdj': '库存单据明细表',
              'pay_notice': '付款通知书', 'pay_plan': '付款计划'}


def import_file(path, month=None, forced_type=None):
    """统一入口：识别类型（可强制）→ 分发导入 → 返回结果 dict"""
    ftype = forced_type or detect_type(path)
    if ftype is None:
        raise RuntimeError('无法识别文件类型，请确认是收发结存/库存单据/付款通知书/付款计划之一')
    if ftype == 'sfjc':
        return import_sfjc(path, month)
    if ftype == 'kcdj':
        return import_kcdj(path, month)
    if ftype == 'pay_notice':
        return import_pay_notice(path, month)
    if ftype == 'pay_plan':
        return import_pay_plan(path, month)
    raise RuntimeError('不支持的类型: %s' % ftype)


if __name__ == '__main__':
    db.init_db()
    print('db ok')
