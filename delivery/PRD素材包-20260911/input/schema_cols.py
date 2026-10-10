# -*- coding: utf-8 -*-
"""
收发结存 / 库存单据 列定义与共享解析工具（表头收敛，S7-1）。

此前列定义散落在 importer.py（SFJC_HEADER/KCDJ_COLS）与 sfjc_process.py（HEADER_COLS
等），信息化加一列要改两处、漏改即错位。统一收口到这里，所有导入/处理/测试共用。

同时提供两个共享解析工具，消除四处重复的 num()/money() 本地函数：
- to_number(v)          千分位逗号 / null / 空 容错的数字解析
- read_sfjc_grid(path)  收发结存原始导出 → list[list]（拆合并填充、去空行，未做数字转换/剔合计）
"""
import openpyxl

# ================= 收发结存（报表1，17 列，1-based 1..17） =================
SFJC_HEADER = ['仓库', '资源分类', '资源编码', '资源名称', '规格型号', '计量单位', '物料性质',
               '期初数量', '期初金额', '收数量', '收单价', '收金额',
               '发数量', '发单价', '发金额', '结存数量', '结存金额']
SFJC_NUM_FIELDS = {'期初数量', '期初金额', '收数量', '收单价', '收金额',
                   '发数量', '发单价', '发金额', '结存数量', '结存金额'}
SFJC_TEXT_FIELDS = [c for c in SFJC_HEADER if c not in SFJC_NUM_FIELDS]
# 列宽（xlsx 输出布局相关，顺序与 SFJC_HEADER 绑定）
SFJC_COL_WIDTHS = {'A': 46.9, 'B': 25.7, 'C': 30.1, 'D': 32.9, 'E': 23.6, 'F': 19.3,
                   'G': 12.0, 'H': 12.0, 'I': 14.0, 'J': 12.0, 'K': 12.0, 'L': 14.0,
                   'M': 12.0, 'N': 12.0, 'O': 14.0, 'P': 12.0, 'Q': 14.0}

# ================= 库存单据明细（28 列，0-based 0..27） =================
KCDJ_COLS = {
    0: '序号', 1: '核算年', 2: '核算月', 3: '业务类型', 4: '项目名称', 5: '单据编号',
    6: '供应商', 7: '安全费用', 8: '物料名称', 9: '计量单位', 10: '规格型号',
    11: '物料性质', 12: '数量', 13: '不含税单价', 14: '含税单价', 15: '金额',
    16: '价税合计', 17: '出入库方向', 18: '仓库', 19: '班组', 20: '部门',
    21: '单据日期', 22: '审核日期', 23: '是否作废', 24: 'wbs编号', 25: 'wbs名称',
    26: 'cbs编号', 27: 'cbs名称',
}
KCDJ_NUM_FIELDS = {'数量', '不含税单价', '含税单价', '金额', '价税合计'}
# inv_detail 已独立成列的字段之外，raw_json 只保留以下"非结构化"字段（S7-4 瘦身依据）
KCDJ_RAW_FIELDS = ['核算年', '核算月', '安全费用', 'wbs编号', 'cbs编号']


def to_number(v):
    """通用数字解析：去千分位逗号、容错 null/空/None。失败或空返回 None。
    统一替代 importer 内四处 num()/money() 本地定义，行为完全一致。"""
    try:
        s = str(v).replace(',', '').strip()
        return float(s) if s and s.lower() != 'null' else None
    except (ValueError, TypeError):
        return None


def read_sfjc_grid(path):
    """读取收发结存原始导出为 list[list]（17 列，1-based 1..17）。

    - 拆合并单元格（合并区取左上角值填充）
    - 第 3 行起为数据，跳过整行全空
    - 不做数字转换、不剔除"合计"行（由调用方按仓库名决定），保持最小公约数
    importer 与 sfjc_process 共用，避免同一份 xlsx 两套解析逻辑。
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['报表1']
    fill = {}
    for mr in ws.merged_cells.ranges:
        v = ws.cell(row=mr.min_row, column=mr.min_col).value
        for r in range(mr.min_row, mr.max_row + 1):
            for c in range(mr.min_col, mr.max_col + 1):
                fill[(r, c)] = v
    rows = []
    for r in range(3, ws.max_row + 1):
        row = [fill.get((r, c), ws.cell(row=r, column=c).value) for c in range(1, 18)]
        if any(x is not None and str(x).strip() != '' for x in row):
            rows.append(row)
    wb.close()
    return rows
