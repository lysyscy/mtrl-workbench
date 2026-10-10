"""
库存单据明细表 处理（模板法）→ 输出 1 个 xlsx：
  报表1（原始数据）+ 过程资料（真透视）+ 安全文明施工费（真透视）
透视配置（对齐用户8月手作）：
  过程资料：行=仓库, 列=出入库方向(入库左/出库右) × 物料性质2(主材+辅材组合为"数据组1"),
            值=求和项:金额, 带入库汇总/出库汇总/总计列
  安全文明施工费：页筛选=cbs名称="安全文明施工费", 行=仓库, 值=求和项:金额
用法: python kcdj_process.py <导出的xlsx> [输出目录]
"""
import os
import re
import sys
import zipfile

BASE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(BASE)
sys.path.insert(0, BASE)
from common_tpl import rebuild, read_data  # noqa
import openpyxl  # noqa

TEMPLATE = os.path.join(ROOT, 'data', 'templates', '模板_库存单据明细表.xlsx')  # 用户手作=模板基底（勿动）

COL_WAREHOUSE = 18   # 仓库
COL_CBS = 27         # cbs名称
# 数字列（0-based，同时是缓存字段索引）
NUM_COLS = {0, 1, 2, 12, 13, 14, 15, 16, 23}   # 序号/核算年/核算月/数量/不含税单价/含税单价/金额/价税合计/是否作废
EMPTY_FIELDS = (28,)                            # 物料性质2 = 主材+辅材分组字段（缓存空，分组由透视定义处理）


def verify_output(out, tpl=None):
    """交付前结构验证：关键文件/透视location+refreshOnLoad/追加字符串/读回检查"""
    tpl = tpl or TEMPLATE
    z = zipfile.ZipFile(out)
    print('关键文件:')
    for n in ['xl/workbook.xml', 'xl/worksheets/sheet1.xml', 'xl/worksheets/sheet2.xml', 'xl/worksheets/sheet3.xml',
              'xl/pivotCache/pivotCacheDefinition1.xml', 'xl/pivotCache/pivotCacheRecords1.xml',
              'xl/pivotTables/pivotTable1.xml', 'xl/pivotTables/pivotTable2.xml', 'xl/sharedStrings.xml']:
        print('  %s: %s' % (n, 'OK' if n in z.namelist() else '缺失!'))
    for n in ('xl/pivotTables/pivotTable1.xml', 'xl/pivotTables/pivotTable2.xml'):
        s = z.read(n).decode('utf-8')
        loc = re.search(r'<location ref="([^"]*)"', s)
        ro = 'refreshOnLoad="1"' in s
        print('%s location=%s refreshOnLoad=%s' % (n.split('/')[-1], loc.group(1) if loc else '?', ro))

    def sslist(zz):
        s = zz.read('xl/sharedStrings.xml').decode('utf-8')
        return re.findall(r'<t[^>]*>([^<]*)</t>', s)
    a, b = sslist(zipfile.ZipFile(tpl)), sslist(z)
    adds = [x for x in b if x not in a]
    print('sharedStrings 追加: %d 条' % len(adds) + ((' %s...' % adds[:8]) if adds else ''))
    z.close()

    wb = openpyxl.load_workbook(out, read_only=True, data_only=True)
    print('sheets:', wb.sheetnames)
    rows = list(wb['报表1'].iter_rows(values_only=True))
    print('报表1: 总行=%d, 表头=%s, 最后序号=%s' % (len(rows), rows[0][:4], rows[-1][0]))
    for sn in ('过程资料', '安全文明施工费'):
        r2 = [r for r in wb[sn].iter_rows(values_only=True) if any(x is not None for x in r)]
        print('%s 行数: %d' % (sn, len(r2)))
        for r in r2[:3]:
            print('  ', [str(x)[:14] if x is not None else '' for x in r[:4]])


def main(src, out_dir=None):
    if out_dir is None:
        out_dir = os.path.join(ROOT, 'output', '库存单据明细表')
    os.makedirs(out_dir, exist_ok=True)
    # 命名：库存单据明细表_2026年8月.xlsx（取输出目录月份标签）
    label = os.path.basename(os.path.normpath(out_dir))
    out = os.path.join(out_dir, '库存单据明细表_%s.xlsx' % label)

    # 预读数据（计算安全文明施工费的仓库子集）
    hdr, data = read_data(src, col_count=28)
    safe_wh = {str(r[COL_WAREHOUSE]).strip() for r in data
               if str(r[COL_CBS]).strip() == '安全文明施工费'}

    cfg = {
        'template': TEMPLATE,
        'src': src,
        'out': out,
        'data': data,   # 预读数据，引擎不再重读
        'data_sheet': {'names': ('报表1', 'Sheet1'), 'header_keywords': ('序号', '金额'),
                       'col_count': 28, 'num_cols': NUM_COLS},
        'shared': {'sheet_file': 'xl/worksheets/sheet3.xml'},   # 报表1 = sheet3
        'cache': {'field_count': 29, 'num_cols': NUM_COLS, 'empty_fields': EMPTY_FIELDS,
                  'wh_field': COL_WAREHOUSE,
                  'wh_order_pivot_file': 'xl/pivotTables/pivotTable1.xml'},
        'pivots': [
            # 过程资料：全部仓库
            {'file': 'xl/pivotTables/pivotTable1.xml',
             'location_ref': lambda n: 'A3:N%d' % (7 + n),
             'rowitems_mode': 'leading_empty'},
            # 安全文明施工费：仅有该费用数据的仓库
            {'file': 'xl/pivotTables/pivotTable2.xml',
             'location_ref': lambda n: 'A3:B%d' % (4 + n),
             'rowitems_mode': 'plain',
             'warehouses': safe_wh},
        ],
        'verify': [
            {'name': '过程资料', 'row': 18, 'cols': (17, 11), 'value': 15,
             'group_map': {'主材': '主辅材', '辅材': '主辅材'}},
            {'name': '安全文明', 'row': 18, 'value': 15, 'filter': (27, '安全文明施工费')},
        ],
    }
    rebuild(cfg)
    print('输出:', out)
    verify_output(out)
    return out


if __name__ == '__main__':
    src = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, 'data', 'templates', '模板_库存单据明细表.xlsx')
    out_dir = sys.argv[2] if len(sys.argv) > 2 else None
    main(src, out_dir)
