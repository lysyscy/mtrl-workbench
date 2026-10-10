"""
6月交叉验证处理：模板法 + 透视表动态适配新仓库集合
- 分仓库文件（9个）
- 总表：Sheet1透视显示 + 报表1全量 + WPS原生透视表（基于8月模板，动态更新仓库相关部分）
"""
import os
import re
import sys
import zipfile
import warnings
from copy import copy
from html import unescape
from xml.sax.saxutils import escape

import openpyxl
import pandas as pd

warnings.filterwarnings('ignore')

BASE = os.path.dirname(os.path.abspath(__file__))  # 脚本所在目录=input
ROOT = os.path.dirname(BASE)
sys.path.insert(0, BASE)
import xlsxwriter  # noqa
# S7-1 表头收敛：列定义与解析工具统一来自 schema_cols
from schema_cols import (SFJC_HEADER, SFJC_NUM_FIELDS, SFJC_TEXT_FIELDS,
                         SFJC_COL_WIDTHS, read_sfjc_grid)

# ===== 读取规则（收发结存表特有：拆合并/剔合计/去千分位逗号） =====
# 列定义统一来自 schema_cols（S7-1 收敛）
HEADER_COLS = SFJC_HEADER
NUM_COLS = SFJC_NUM_FIELDS
TEXT_COLS = SFJC_TEXT_FIELDS
COL_WIDTHS = SFJC_COL_WIDTHS


def read_raw(path):
    """读取原始导出：构造单行表头、拆合并填充、剔除合计、转数字（复用 schema_cols.read_sfjc_grid）"""
    rows = read_sfjc_grid(path)
    df = pd.DataFrame(rows, columns=SFJC_HEADER)
    df = df[df['仓库'] != '合计'].copy()
    for c in SFJC_NUM_FIELDS:
        df[c] = pd.to_numeric(df[c].astype(str).str.replace(',', ''), errors='coerce')
    return df


TEMPLATE = os.path.join(ROOT, 'data', 'templates', '模板_收发结存表.xlsx')
NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
MONEY_FIELDS = ['期初金额', '收金额', '发金额', '结存金额']
SRC = None
OUT_DIR = os.path.join(ROOT, 'output')
LABEL = ''


def fmt_num(x):
    s = ('%.10f' % float(x)).rstrip('0').rstrip('.')
    return s if s not in ('', '-0') else '0'


def write_warehouse_file(df, path):
    wb = xlsxwriter.Workbook(path)
    ws = wb.add_worksheet('Sheet1')
    header_fmt = wb.add_format({'bold': True, 'bg_color': '#ED7D31', 'font_color': '#1A1B1C',
                                'border': 1, 'align': 'center', 'valign': 'vcenter', 'font_size': 11})
    text_fmt = wb.add_format({'border': 1, 'valign': 'vcenter', 'align': 'left', 'font_size': 11})
    num_fmt = wb.add_format({'border': 1, 'valign': 'vcenter', 'align': 'right',
                             'num_format': '#,##0.00', 'font_size': 11})
    for j, col in enumerate(HEADER_COLS):
        ws.write(0, j, col, header_fmt)
    for i, (_, row) in enumerate(df.iterrows(), start=1):
        for j, col in enumerate(HEADER_COLS):
            fmt = text_fmt if col in TEXT_COLS else num_fmt
            v = row[col]
            if col in NUM_COLS and pd.isna(v):
                v = None
            ws.write(i, j, v, fmt)
    for letter, width in COL_WIDTHS.items():
        ws.set_column(f'{letter}:{letter}', width)
    ws.set_row(0, 33)
    for r in range(1, len(df) + 1):
        ws.set_row(r, 28)
    wb.close()


def update_sheet_data(df, out_path, show_order):
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb['报表1']
    data = df.values.tolist()
    n = len(data)
    for i in range(n):
        r = 2 + i
        for c in range(1, 18):
            cell = ws.cell(row=r, column=c)
            if r > ws.max_row:
                tpl = ws.cell(row=ws.max_row, column=c)
                cell._style = copy(tpl._style)
            cell.value = data[i][c - 1]
    for r in range(2 + n, ws.max_row + 1):
        for c in range(1, 18):
            ws.cell(row=r, column=c).value = None
    # 删除尾部多余行（模板674数据行 vs 实际n行）
    if n < ws.max_row - 1:
        ws.delete_rows(2 + n, ws.max_row - (1 + n))
    ws1 = wb['Sheet1']
    pv = df.groupby('仓库')[MONEY_FIELDS].sum()
    grand = pv.sum()
    for i, wh in enumerate(show_order):
        r = 4 + i
        ws1.cell(row=r, column=1).value = wh
        for j, m in enumerate(MONEY_FIELDS):
            ws1.cell(row=r, column=2 + j).value = float(pv.loc[wh, m])
    r_grand = 4 + len(show_order)
    ws1.cell(row=r_grand, column=1).value = '总计'
    for j, m in enumerate(MONEY_FIELDS):
        ws1.cell(row=r_grand, column=2 + j).value = float(grand[m])
    wb.save(out_path)
    return pv, grand


def patch_pivot(in_path, out_path, df, show_order):
    """基于8月模板patch：cacheDefinition + records + pivotTable1(动态适配仓库)"""
    zt = zipfile.ZipFile(TEMPLATE)
    cd = zt.read('xl/pivotCache/pivotCacheDefinition1.xml').decode('utf-8')
    pt_orig = zt.read('xl/pivotTables/pivotTable1.xml').decode('utf-8')
    zt.close()

    # ===== 1) cacheDefinition: sharedItems追加 + recordCount + sourceReference =====
    field_maps = []
    new_parts = []
    for col in df.columns:
        m = re.search(r'<cacheField name="' + re.escape(col) + r'"[^>]*>(.*?)</cacheField>', cd, re.S)
        seg = m.group(1)
        is_num = col in NUM_COLS
        if is_num:
            items = re.findall(r'<n v="([^"]*)"', seg)
            seen = [float(x) for x in items]
        else:
            items = re.findall(r'<s v="([^"]*)"', seg)
            seen = [unescape(x) for x in items]
        idx = {v: i for i, v in enumerate(seen)}
        add_nodes = []
        for v in df[col]:
            if v is None or (isinstance(v, float) and pd.isna(v)):
                v = ''
            key = float(v) if is_num else str(v)
            if key not in idx:
                idx[key] = len(seen)
                seen.append(key)
                add_nodes.append(key)
        extra = ''
        for k in add_nodes:
            if is_num:
                extra += f'<n v="{fmt_num(k)}"/>'
            else:
                extra += f'<s v="{escape(str(k), {chr(34): "&quot;", chr(39): "&apos;"})}"/>'
        if is_num:
            nums = [float(x) for x in seen]
            new_seg = re.sub(r'minValue="[^"]*"', f'minValue="{fmt_num(min(nums))}"', seg)
            new_seg = re.sub(r'maxValue="[^"]*"', f'maxValue="{fmt_num(max(nums))}"', new_seg)
            new_seg = re.sub(r'count="\d+"', f'count="{len(nums)}"', new_seg)
        else:
            new_seg = re.sub(r'count="\d+"', f'count="{len(seen)}"', seg)
        new_seg = new_seg[:-len('</sharedItems>')] + extra + '</sharedItems>'
        new_parts.append(new_seg)
        field_maps.append(idx)
    cd_new = re.sub(r'(<cacheFields count="17">).*?(</cacheFields>)',
                    lambda mo: mo.group(1) + ''.join(
                        f'<cacheField name="{escape(col)}" numFmtId="0">{seg}</cacheField>'
                        for col, seg in zip(df.columns, new_parts)) + mo.group(2),
                    cd, flags=re.S)
    cd_new = re.sub(r'recordCount="\d+"', f'recordCount="{len(df)}"', cd_new)
    cd_new = re.sub(r'ref="A1:Q\d+"', f'ref="A1:Q{len(df) + 1}"', cd_new)

    # ===== 2) records重建 =====
    rows = []
    for _, row in df.iterrows():
        xs = []
        for j, col in enumerate(df.columns):
            v = row[col]
            if v is None or (isinstance(v, float) and pd.isna(v)):
                v = ''
            key = float(v) if col in NUM_COLS else str(v)
            xs.append(f'<x v="{field_maps[j][key]}"/>')
        rows.append('<r>' + ''.join(xs) + '</r>')
    cr_new = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
              f'<pivotCacheRecords xmlns="{NS}" count="{len(df)}">'
              + ''.join(rows) + '</pivotCacheRecords>')

    # ===== 3) pivotTable1: 动态适配新仓库 =====
    # 仓库缓存索引
    wh_cache_idx = {w: field_maps[0][w] for w in df['仓库'].unique()}
    axis_items = [wh_cache_idx[w] for w in show_order]
    pt = pt_orig
    # location
    r_grand = 4 + len(show_order)
    loc_ref = f'A3:E{r_grand}'
    pt = re.sub(r'<location ref="[^"]*"', f'<location ref="{loc_ref}"', pt)
    # axisRow pivotField items
    pat_axis = re.compile(r'<pivotField axis="axisRow"[^>]*><items count="\d+">.*?</items></pivotField>', re.S)
    m = pat_axis.search(pt)
    if m:
        head = m.group(0)[:m.group(0).find('<items')]
        items_xml = (f'<items count="{len(axis_items) + 1}">'
                     + ''.join(f'<item x="{i}"/>' for i in axis_items)
                     + '<item t="default"/></items></pivotField>')
        pt = pt[:m.start()] + head + items_xml + pt[m.end():]
    # dataField pivotField items（4个金额字段）
    pat_df = re.compile(r'<pivotField dataField="1"[^>]*><items count="\d+">.*?</items></pivotField>', re.S)
    dmatches = list(pat_df.finditer(pt))
    for mm, col in zip(reversed(dmatches), reversed(MONEY_FIELDS)):
        uniq_n = len(field_maps[list(df.columns).index(col)])
        head = mm.group(0)[:mm.group(0).find('<items')]
        items_xml = (f'<items count="{uniq_n + 1}">'
                     + ''.join(f'<item x="{i}"/>' for i in range(uniq_n))
                     + '<item t="default"/></items></pivotField>')
        pt = pt[:mm.start()] + head + items_xml + pt[mm.end():]
    # rowItems
    pat_rows = re.compile(r'<rowItems count="\d+">.*?</rowItems>', re.S)
    m = pat_rows.search(pt)
    if m:
        row_items = ('<rowItems count="' + str(len(show_order) + 1) + '">'
                     + '<i><x/></i>'
                     + ''.join(f'<i><x v="{k}"/></i>' for k in range(1, len(show_order)))
                     + '<i t="grand"><x/></i></rowItems>')
        pt = pt[:m.start()] + row_items + pt[m.end():]

    # ===== 4) 写回 =====
    zin = zipfile.ZipFile(in_path)
    zout = zipfile.ZipFile(out_path, 'w', zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == 'xl/pivotCache/pivotCacheDefinition1.xml':
            data = cd_new.encode('utf-8')
        elif item.filename == 'xl/pivotCache/pivotCacheRecords1.xml':
            data = cr_new.encode('utf-8')
        elif item.filename == 'xl/pivotTables/pivotTable1.xml':
            data = pt.encode('utf-8')
        zout.writestr(item, data)
    zout.close()
    zin.close()


def verify(out_path, df, pv, grand, show_order):
    import xml.etree.ElementTree as ET
    problems = []
    zf = zipfile.ZipFile(out_path)
    for name in ['xl/pivotCache/pivotCacheDefinition1.xml', 'xl/pivotCache/pivotCacheRecords1.xml',
                 'xl/pivotTables/pivotTable1.xml']:
        try:
            ET.fromstring(zf.read(name))
        except Exception as e:
            problems.append(f'{name}: {e}')
    cd = ET.fromstring(zf.read('xl/pivotCache/pivotCacheDefinition1.xml'))
    cr = ET.fromstring(zf.read('xl/pivotCache/pivotCacheRecords1.xml'))
    fields = cd.findall(f'{{{NS}}}cacheFields/{{{NS}}}cacheField')
    fvals = []
    for f in fields:
        si = f.find(f'{{{NS}}}sharedItems')
        vals = [float(ch.get('v')) if ch.tag.endswith('n') else str(ch.get('v')) for ch in si]
        fvals.append(vals)
    money_idx = [list(df.columns).index(m) for m in MONEY_FIELDS]
    sums = {m: {} for m in MONEY_FIELDS}
    for r in cr.findall(f'{{{NS}}}r'):
        xs = [int(x.get('v', 0)) for x in r.findall(f'{{{NS}}}x')]
        wh = fvals[0][xs[0]]
        for k, m in enumerate(MONEY_FIELDS):
            sums[m][wh] = sums[m].get(wh, 0.0) + fvals[money_idx[k]][xs[money_idx[k]]]
    for m in MONEY_FIELDS:
        for wh in show_order:
            got = round(sums[m].get(wh, 0), 2)
            exp = round(pv.loc[wh, m], 2)
            if abs(got - exp) > 0.01:
                problems.append(f'对账失败 {m}/{wh}: {got} vs {exp}')
        if abs(round(sum(sums[m].values()), 2) - round(grand[m], 2)) > 0.01:
            problems.append(f'总计对账失败 {m}')
    zf.close()
    return problems


def main(src=None, out_dir=None, label=None):
    global SRC, OUT_DIR, LABEL
    if src:
        SRC = src
    if out_dir:
        OUT_DIR = out_dir
    if label:
        LABEL = label
    # 清空输出目录（防止月与月之间仓库集合变化时旧文件残留）
    if os.path.isdir(OUT_DIR):
        for f in os.listdir(OUT_DIR):
            try:
                os.remove(os.path.join(OUT_DIR, f))
            except OSError:
                pass
    os.makedirs(OUT_DIR, exist_ok=True)
    df = read_raw(SRC)
    show_order = [w for w in df['仓库'].unique() if w == w]
    os.makedirs(OUT_DIR, exist_ok=True)
    print('数据:', len(df), '行 | 仓库(%d):' % len(show_order), show_order)

    files = []
    for wh in show_order:
        sub = df[df['仓库'] == wh].copy()
        fn = os.path.join(OUT_DIR, f'{wh}.xlsx')
        write_warehouse_file(sub, fn)
        files.append(fn)
        print('  分仓库:', os.path.basename(fn), len(sub), '行')

    work = os.path.join(OUT_DIR, '_work.xlsx')
    main_fn = os.path.join(OUT_DIR, f'第四工程公司{LABEL}收发结存表.xlsx')
    pv, grand = update_sheet_data(df, work, show_order)
    patch_pivot(work, main_fn, df, show_order)
    os.remove(work)
    files.append(main_fn)

    problems = verify(main_fn, df, pv, grand, show_order)
    if problems:
        print('!! 验证问题:')
        for p in problems:
            print('  -', p)
    else:
        print('✓ 透视表对账全部通过 | 总计: 期初%.2f 收%.2f 发%.2f 结存%.2f' % (
            grand['期初金额'], grand['收金额'], grand['发金额'], grand['结存金额']))

    zip_fn = os.path.join(OUT_DIR, f'第四工程公司{LABEL}收发结存表.zip')
    with zipfile.ZipFile(zip_fn, 'w', zipfile.ZIP_DEFLATED) as zf:
        for fn in files:
            zf.write(fn, os.path.basename(fn))
    print('✓ 打包:', os.path.basename(zip_fn))
    # 输出目录只保留zip（分仓库/总表已进zip，删除中间xlsx）
    for fn in files:
        try:
            os.remove(fn)
            print('  清理中间文件:', os.path.basename(fn))
        except OSError:
            pass
    return 0


if __name__ == '__main__':
    main()
